#!/usr/bin/env python3
"""
drawio_toolkit.py — single-file bundle of the draw.io academic-figure toolkit
(modes: image -> drawio, text -> drawio, enhance for publication). See SKILL.md / REFERENCE.md.

USAGE
  python drawio_toolkit.py <tool> [args...]     run a tool (auto-unpacks once to a cache dir)
  python drawio_toolkit.py unpack [DIR] [--force]  write scripts/ assets/ examples/ to DIR
                                                 (default ./tidyfigs) — a normal skill folder
  python drawio_toolkit.py path                  print the unpacked scripts/ dir (for sys.path)
  python drawio_toolkit.py list                  list tools and bundled files

  From Python:  import drawio_toolkit; drawio_toolkit.setup()   # unpacks + adds scripts/ to sys.path
                from drawio_kit import Diagram;  from figlib import *;  from drawio_doc import Doc

TOOLS (python drawio_toolkit.py <tool> --help for options)
  image mode : new_build  inspect_image  iterate  compare
  text mode  : analyze_text  new_figures  prototype  finalize
  enhance    : start_enhance  brief  rebuild  enhance  audit_drawio  report
  shared     : render_preview  lint_drawio  faithful
  libraries  : drawio_kit (builder)  figlib (figure components)  drawio_doc (edit existing files)

Cache dir: $DRAWIO_TOOLKIT_HOME, else ~/.cache/drawio_toolkit/<version>-<hash>, else a temp dir.
Requires Python 3.9+, numpy, pillow; matplotlib (math), pyyaml (YAML themes / enhance).
Optional: scipy or opencv-python, pytesseract + tesseract (OCR), python-docx / pypdf, draw.io desktop CLI.
"""
import hashlib
import os
import subprocess
import sys
import tempfile

VERSION = "2.0"
FILES = {}

# ================================================================================================
FILES['scripts/analyze_text.py'] = r'''#!/usr/bin/env python3
"""
analyze_text.py — read academic prose and suggest which figures it supports.

    python analyze_text.py paper_section.md            # .md .txt .tex (.docx/.pdf if libs present)
    python analyze_text.py --text "We propose ..."     # inline text
    python analyze_text.py file.tex --split --json out.json

For each passage it reports: word count, display/inline equations, structure cues per figure
family (process, feedback, causal, comparison, hierarchy, temporal, quantitative, architecture,
equation), enumerations, acronyms, numbers, and the three suggested figure types (always from
different layouts). The agent still reads the passage itself — this is a checklist, not a parser.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import re
import subprocess
import sys

CUES = {
    "process": [r"\bthen\b", r"\bnext\b", r"\bfinally\b", r"\bfirst(ly)?\b", r"\bsecond(ly)?\b", r"\bsubsequently\b",
                r"\bafter\b", r"\bbefore\b", r"\bstages?\b", r"\bsteps?\b", r"\bpipeline\b", r"\bworkflow\b",
                r"\b(converted|passed|fed|inserted|encoded|segmented|stored|retrieved|re-?ranked|embedded|generated|reviewed|logged|extracted|filtered|aggregated|trained|evaluated)\b",
                r"\bat (query|run|inference|training) time\b", r"\b(offline|online)\b",
                r"\binputs?\b", r"\boutputs?\b", r"\bprocedure\b"],
    "feedback": [r"\bloop\b", r"\bfeedback\b", r"\biterat", r"\brefine", r"\btune[sd]?\b", r"\bupdate[sd]?\b", r"\bcycle\b"],
    "causal": [r"\bdriv(e|es|en)\b", r"\bleads? to\b", r"\bresults? in\b", r"\bcaus", r"\baffect", r"\binfluenc",
               r"\bmoderat", r"\bmediat", r"\bmanifest", r"\bcontribut", r"\bdetermin", r"\binforms?\b", r"\bdue to\b",
               r"\bbecause\b", r"\bin turn\b", r"\boutcome\b"],
    "comparison": [r"\bwhereas\b", r"\bwhile\b", r"\bcompared\b", r"\bin contrast\b", r"\bunlike\b", r"\btrade-?offs?\b",
                   r"\bversus\b", r"\bvs\.?", r"\boutperform", r"\bbetter\b", r"\bworse\b", r"\badvantages?\b",
                   r"\blimitations?\b", r"\b(two|three|four|several) (families|approaches|methods|models|types)\b",
                   r"\bbest\b", r"\bleast\b", r"\bhighest\b", r"\blowest\b", r"\bdominates?\b"],
    "hierarchy": [r"\bconsists? of\b", r"\bcompris", r"\bcomposed of\b", r"\bcategor", r"\btypes? of\b", r"\bclassif",
                  r"\btaxonom", r"\bincludes?\b", r"\bsubsets?\b", r"\bsub-?(types|classes|categories)\b", r"\bhierarch"],
    "temporal": [r"\b(1[89]|20)\d\d\b", r"\bdecades?\b", r"\bsince\b", r"\bhistor", r"\btimeline\b", r"\beras?\b",
                 r"\bphases?\b", r"\bearly\b", r"\blater\b", r"\brecently\b", r"\bevolv"],
    "quantitative": [r"\d+(\.\d+)?\s?%", r"\bpercent", r"\bincreas", r"\bdecreas", r"\baccuracy\b", r"\bF1\b", r"\bmean\b",
                     r"\bmedian\b", r"\bstatistically\b", r"\bp\s?[<=]", r"\bimprov", r"\breduc"],
    "architecture": [r"\blayers?\b", r"\bmodules?\b", r"\bcomponents?\b", r"\bsystem\b", r"\barchitecture\b", r"\bframework\b",
                     r"\binterfaces?\b"],
    "matrix_computation": [r"\b(matrix|matrices|tensor|factoriz|decompos|singular value|svd|eigenvector|eigenvalue|spectral|inner product|dot product|outer product|hadamard|rank|subspace|low-rank|projection)\b"],
    "geometric_metric_space": [r"\b(metric|distance|euclidean|cosine|manifold|neighborhood|geodesic|coordinate|vector space|angle|unit sphere|hypersphere|chord|dimension|embedding|topology|topological|phase space|trajectory|centroid|clustering)\b"],
    "mechanistic_causal": [r"\b(pathway|reaction|stoichiometry|enzyme|kinetics|concentration|substrate|flux|sir model|compartmental|dag\b|directed acyclic|causal graph|instrumental variable|treatment effect|endogeneity|structural equation)\b"],
    "probabilistic_graphical": [r"\b(bayesian|dirichlet|prior|posterior|likelihood|generative|latent|plate notation|markov|mcmc|stochastic|sampling|discrete distribution|bernoulli|poisson)\b"],
}

# figure families per structure: (key, description). The first entries are preferred.
FAMILIES = {
    "process": [("swimlane_pipeline", "stages as lanes (two phases / actors) with routed connectors"),
                ("layered_architecture", "system layers stacked top-to-bottom with components"),
                ("stage_chain", "single left-to-right chain of numbered stage cards"),
                ("timeline", "stages on a time axis")],
    "feedback": [("numbered_cycle", "stages around a hub; closing arrow = feedback")],
    "causal": [("causal_columns", "drivers -> mechanisms -> indicators -> decisions columns"),
               ("concept_map", "hub entity with labeled relations"),
               ("stepped_funnel", "stages narrowing top-to-bottom (compact overview)"),
               ("nested_containers", "factors grouped inside containers")],
    "comparison": [("rating_matrix", "items x criteria with ordinal dots + source wording"),
                   ("method_cards", "one card per item: how / strengths / limitations / best fit"),
                   ("positioning_chart", "two criteria as axes, third as bubble size")],
    "hierarchy": [("tree", "top-down taxonomy tree"), ("nested_boxes", "containment boxes"),
                  ("outline_columns", "indented columns per level")],
    "temporal": [("timeline", "milestones on a horizontal axis"), ("phase_bands", "eras as colored bands"),
                 ("swimlane_pipeline", "actors as lanes over time")],
    "quantitative": [("dot_or_bar_chart", "only values stated in the text"), ("results_table", "values with units"),
                     ("before_after", "slope chart for two conditions")],
    "architecture": [("layered_architecture", "system layers"), ("block_diagram", "components and interfaces")],
    "equation": [("annotated_equation", "each term color-coded with bracketed explanations"),
                 ("computation_graph", "how the quantity is assembled from inputs"),
                 ("intuition_glossary", "geometric/continuum intuition + symbol table")],
    "matrix_computation": [("matrix_factorization_grid", "populated 2D matrix grids with cell reconstruction dot-product arithmetic"),
                           ("comparative_state_spaces", "orthogonal high-D failure vs compressed low-rank semantic subspace"),
                           ("spectral_decomposition", "eigenvector projection and singular value decomposition axes")],
    "geometric_metric_space": [("metric_geometry_clusters", "continuous 2D metric space with unit sphere arc, chord vs cosine, and cluster centroid attraction"),
                               ("divisive_density_clustering", "hyperplane cutting vs density cores and outlier noise segregation"),
                               ("manifold_projection_spaces", "high-D to 2D projection showing local neighborhood vs global geometry preservation")],
    "mechanistic_causal": [("causal_dag", "directed acyclic graph with signed paths and confounder conditioning"),
                           ("reaction_pathway", "reaction network with stoichiometric fluxes, activation and inhibition"),
                           ("compartmental_flow", "state transition boxes with differential rates and feedback")],
    "probabilistic_graphical": [("hierarchical_generative_plate", "Bayesian plate notation with hyperparameter priors and discrete sampling"),
                                ("markov_state_graph", "transition probability states with emission nodes"),
                                ("probabilistic_simplex", "Dirichlet topic mixture simplices and distribution boundaries")],
}

DISPLAY_EQ = [r"\$\$(.+?)\$\$", r"\\\[(.+?)\\\]", r"\\begin\{(equation|align|gather|multline)\*?\}(.+?)\\end\{\1\*?\}"]
INLINE_EQ = r"(?<![\\$])\$(?!\$)(.+?)(?<!\\)\$"


def read_input(path=None, text=None):
    if text is not None:
        return text
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            import docx  # python-docx
            return "\n\n".join(p.text for p in docx.Document(path).paragraphs)
        except ImportError:
            sys.exit("reading .docx needs python-docx (pip install python-docx), or paste the text")
    if ext == ".pdf":
        try:
            return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout
        except Exception:
            try:
                import pypdf
                return "\n\n".join(p.extract_text() or "" for p in pypdf.PdfReader(path).pages)
            except ImportError:
                sys.exit("reading .pdf needs pdftotext or pypdf, or paste the text")
    return open(path, encoding="utf-8", errors="replace").read()


def clean_tex(t):
    t = re.sub(r"(?<!\\)%.*", "", t)                                   # comments
    t = re.sub(r"\\(cite|ref|eqref|label)\{[^}]*\}", "", t)
    t = re.sub(r"\\(emph|textit|textbf|underline)\{([^}]*)\}", r"\2", t)
    return t


def split_passages(text):
    """Split into passages at headings / blank lines; keep display equations with their paragraph."""
    text = text.replace("\r\n", "\n")
    blocks, cur, heading = [], [], None
    lines = text.split("\n")
    in_eq = False
    for ln in lines:
        s = ln.strip()
        if re.match(r"^(#{1,6}\s|\\(sub)*section\*?\{)", s):
            if cur:
                blocks.append((heading, "\n".join(cur).strip())); cur = []
            heading = re.sub(r"^#+\s*|\\(sub)*section\*?\{|\}$", "", s)
            continue
        if s.count("$$") % 2 == 1 or re.match(r"\\begin\{(equation|align)", s) or re.match(r"\\end\{(equation|align)", s):
            in_eq = not in_eq if "$$" in s or "\\begin" in s else False
        if not s and not in_eq:
            if cur:
                blocks.append((heading, "\n".join(cur).strip())); cur = []
            continue
        cur.append(ln)
    if cur:
        blocks.append((heading, "\n".join(cur).strip()))
    # merge a block that is only a display equation into the previous paragraph
    merged = []
    for h, b in blocks:
        if merged and re.fullmatch(r"\s*(\$\$.+?\$\$|\\\[.+?\\\]|\\begin\{.+?\\end\{[a-z*]+\})\s*[.,]?\s*", b, re.S):
            ph, pb = merged[-1]; merged[-1] = (ph, pb + "\n" + b)
        else:
            merged.append((h, b))
    return merged


def analyze_passage(t):
    plain = re.sub(r"\$\$.+?\$\$|\\\[.+?\\\]", " EQUATION ", t, flags=re.S)
    words = re.findall(r"[A-Za-z][A-Za-z\-']*", plain)
    disp = []
    for pat in DISPLAY_EQ:
        for m in re.finditer(pat, t, re.S):
            disp.append(m.group(m.lastindex).strip())
    inline = [m.group(1).strip() for m in re.finditer(INLINE_EQ, re.sub(r"\$\$.+?\$\$", "", t, flags=re.S))]
    low = plain.lower()
    cues = {k: sum(len(re.findall(p, low)) for p in pats) for k, pats in CUES.items()}
    cues["architecture"] = cues["architecture"] // 2          # generic words; weak evidence on its own
    if disp or len(inline) >= 3:
        cues["equation"] = 6 + 3 * len(disp) + len(inline) // 2
    enums = re.findall(r"\((?:i{1,3}|iv|v|[a-e]|\d)\)|\b(?:first|second|third|finally)\b", low)
    acr = sorted(set(re.findall(r"\b[A-Z][A-Z0-9\-]{1,}[a-z]?\b", plain)) - {"EQUATION"})
    nums = re.findall(r"\b\d[\d,]*(?:\.\d+)?\s?(?:%|-dimensional|D\b|[a-zA-Z]{1,3}\b)?", plain)
    sentences = [s.strip() for s in re.split(r"(?<=[.;:])\s+(?=[A-Z(])", plain) if len(s.split()) > 3]
    return dict(words=len(words), sentences=len(sentences), display_equations=disp, inline_math=inline,
                cues=cues, enumerations=len(enums), acronyms=acr, numbers=[n.strip() for n in nums][:30])


def suggest(cues):
    """Three distinct figure types. Default: three layouts of the dominant structure; if a second
    structure is strong (>= half the dominant score) use two + one alternative lens from it."""
    ranked = [k for v, k in sorted(((v, k) for k, v in cues.items() if v > 0), reverse=True)]
    if not ranked:
        return [("concept_map", "fallback"), ("stage_chain", "fallback"), ("method_cards", "fallback")]
    # Active scientific/computational families take precedence over generic cards/timelines
    sci_fams = ["geometric_metric_space", "matrix_computation", "mechanistic_causal", "probabilistic_graphical", "equation"]
    sci_candidates = [(cues[k], k) for k in sci_fams if k in cues and cues[k] >= 2]
    if sci_candidates:
        sci_candidates.sort(reverse=True)
        best_sci = sci_candidates[0][1]
        if best_sci in ranked:
            ranked.remove(best_sci)
            ranked.insert(0, best_sci)
    dom = ranked[0]
    second = next((k for k in ranked[1:] if cues[k] >= 0.5 * cues[dom] and k not in ("architecture",)), None)
    if dom in sci_fams and second not in sci_fams:
        second = None
    picks, seen = [], set()

    def take(fam_key, n):
        for fam, desc in FAMILIES.get(fam_key, []):
            if len(picks) < n and fam not in seen:
                picks.append((fam, f"{fam_key}: {desc}")); seen.add(fam)
    take(dom, 2 if second else 3)
    if second:
        take(second, 3)
    for k in ranked[1:]:
        take(k, 3)
    if dom in ("process", "architecture") and cues.get("feedback", 0) >= 2 and "numbered_cycle" not in seen:
        picks[-1] = ("numbered_cycle", "feedback: stages around a hub; closing arrow = feedback")
    return picks[:3]


def slugify(t, n=4):
    w = [x.lower() for x in re.findall(r"[A-Za-z0-9]+", t) if x.lower() not in STOP][:n]
    return "_".join(w) or "passage"


STOP = set("a an the of and or to in on for with by from as at is are was were be been this that these those we our "
           "it its each such their which into using use used than then also both can may".split())


def analyze(text, split=None, min_words=40):
    text = clean_tex(text)
    blocks = split_passages(text)
    total = sum(len(b.split()) for _, b in blocks)
    if split is None:   # long input, or several headed sections -> one passage per paragraph/section
        split = total > 350 or len({h for h, _ in blocks}) >= 2
    if not split:
        blocks = [(blocks[0][0] if blocks else None, "\n\n".join(b for _, b in blocks))]
    out = []
    for i, (h, b) in enumerate(blocks):
        a = analyze_passage(b)
        if a["words"] < min_words and split:
            continue
        a.update(index=len(out) + 1, heading=h, text=b, slug=slugify(h or b), suggestions=suggest(a["cues"]))
        out.append(a)
    return out


def to_markdown(passages):
    md = []
    for p in passages:
        md += [f"## Passage {p['index']}: {p['heading'] or p['slug']}", "",
               f"- words: {p['words']}, sentences: {p['sentences']}, enumerations: {p['enumerations']}",
               f"- display equations: {len(p['display_equations'])}" + "".join(f"\n  - `{e}`" for e in p["display_equations"]),
               f"- inline math: {', '.join('`' + m + '`' for m in p['inline_math'][:12]) or '—'}",
               f"- acronyms: {', '.join(p['acronyms']) or '—'}",
               f"- numbers: {', '.join(p['numbers'][:15]) or '—'}",
               "- structure cues: " + ", ".join(f"{k} {v}" for k, v in sorted(p["cues"].items(), key=lambda kv: -kv[1]) if v),
               "- suggested prototypes:"] + [f"  {j + 1}. **{f}** — {why}" for j, (f, why) in enumerate(p["suggestions"])] + [""]
    return "\n".join(md)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?"); ap.add_argument("--text")
    ap.add_argument("--split", action="store_true", help="force one passage per paragraph/section")
    ap.add_argument("--whole", action="store_true", help="force the whole input as one passage")
    ap.add_argument("--min-words", type=int, default=40); ap.add_argument("--json")
    a = ap.parse_args()
    if not a.path and a.text is None:
        a.text = sys.stdin.read()
    t = read_input(a.path, a.text)
    ps = analyze(t, True if a.split else False if a.whole else None, a.min_words)
    print(to_markdown(ps))
    if a.json:
        json.dump(ps, open(a.json, "w"), indent=1)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/audit_drawio.py'] = r'''#!/usr/bin/env python3
"""
audit_drawio.py — is this figure ready for an academic article?

    python audit_drawio.py fig.drawio [--brief brief.yaml] [--theme theme.yaml] [--out-dir DIR] [--tag before]

Checks (all at the TARGET PRINT SIZE from the brief: width_in, min_font_pt, grayscale, colorblind):
  legibility   font sizes in pt at print size, hairline strokes, low-resolution embedded images
  color        palette size, off-theme colors, text contrast (WCAG), grayscale and color-vision-
               deficiency (deuteranopia / protanopia) separability of category fills
  typography   number of font families / sizes, sizes not in theme
  math         math flag, LaTeX that fails to parse, Unicode math in plain labels (inconsistent notation)
  layout       lint (overlaps, overflow, near-miss alignment, uneven gaps), margins, excess whitespace,
               aspect ratio vs column, in-figure titles, long labels, density
  connectors   dangling edge ends next to shapes, stroke / arrow-head variety
Writes <tag>_audit.md, <tag>_audit.json and <tag>_print_preview.png (actual print size at 150 dpi).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import base64
import io
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.dont_write_bytecode = True
from drawio_doc import Doc  # noqa: E402

DEFAULT_BRIEF = dict(width_in=7.0, min_font_pt=7.0, grayscale=False, colorblind=True)
CVD = {"deuteranopia": [[0.367322, 0.860646, -0.227968], [0.280085, 0.672501, 0.047413], [-0.011820, 0.042940, 0.968881]],
       "protanopia": [[0.152286, 1.052583, -0.204868], [0.114503, 0.786281, 0.099216], [-0.003882, -0.048116, 1.051998]]}
UNICODE_MATH = "λθαβγδεσμπρτφωΣΔΩ∑∏∫≤≥≠≈×÷∈∉∞√∂∇±→⇒"


def load_yaml(p):
    if not p or not os.path.exists(p):
        return {}
    if p.endswith((".yaml", ".yml")):
        import yaml
        return yaml.safe_load(open(p, encoding="utf-8")) or {}
    if p.endswith(".json"):
        return json.load(open(p, encoding="utf-8"))
    import tomllib
    return tomllib.loads(open(p, encoding="utf-8").read())


# ---------------- color helpers
def hex2rgb(h):
    h = h.strip().lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def lin(c):
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def delin(c):
    c = max(0.0, min(1.0, c))
    return 12.92 * c if c <= 0.0031308 else 1.055 * c ** (1 / 2.4) - 0.055


def luminance(rgb):
    r, g, b = (lin(c) for c in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b):
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def to_lab(rgb):
    r, g, b = (lin(c) for c in rgb)
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116  # noqa: E731
    return (116 * f(y) - 16, 500 * (f(x) - f(y)), 200 * (f(y) - f(z)))


def de(a, b):
    return math.dist(to_lab(a), to_lab(b))


def simulate(rgb, kind):
    m = CVD[kind]
    l = [lin(c) for c in rgb]
    return tuple(delin(sum(m[i][j] * l[j] for j in range(3))) for i in range(3))


def blend(rgb, alpha, bg=(1, 1, 1)):
    return tuple(alpha * c + (1 - alpha) * b for c, b in zip(rgb, bg))


def is_hex(v):
    return bool(v) and re.fullmatch(r"#[0-9a-fA-F]{3}([0-9a-fA-F]{3})?", v.strip()) is not None


# ---------------- audit
def category_count(path):
    """Distinct effective fill colors (opacity applied, dE >= 8 apart) — a proxy for color-coded categories."""
    doc = Doc.load(path)
    page_bg = hex2rgb(doc.model.get("background")) if is_hex(doc.model.get("background")) else (1, 1, 1)
    reps = []
    for c in doc.vertices():
        f = c.get("fillColor")
        if not is_hex(f) or c.is_text:
            continue
        a = float(c.get("fillOpacity", 100)) / 100 * float(c.get("opacity", 100)) / 100
        eff = blend(hex2rgb(f), a, page_bg)
        if contrast(eff, page_bg) < 1.04:
            continue
        if all(de(eff, r) >= 8 for r in reps):
            reps.append(eff)
    return len(reps)


def audit(path, brief=None, theme_path=None, accepted=(), reference=None):
    b = dict(DEFAULT_BRIEF); b.update({k: v for k, v in (brief or {}).items() if v is not None})
    doc = Doc.load(path)
    pw, ph = doc.page
    cb = doc.content_bbox()
    fig_w = max(cb[2], 1)
    pt_per_px = b["width_in"] * 72 / max(pw, 1)
    issues = {k: [] for k in ("legibility", "color", "typography", "math", "layout", "connectors")}
    stats = {}
    page_bg = hex2rgb(doc.model.get("background")) if is_hex(doc.model.get("background")) else (1, 1, 1)

    # --- legibility
    sizes, fams = {}, {}
    for c in doc.vertices():
        if not c.plain or c.is_group:
            continue
        fs = float(c.get("fontSize", 11))
        sizes[fs] = sizes.get(fs, 0) + 1
        if not c.is_math:
            fams[c.get("fontFamily", "Helvetica")] = fams.get(c.get("fontFamily", "Helvetica"), 0) + 1
        pt = fs * pt_per_px * (1.0 if not c.is_math else 1.0)
        if pt < b["min_font_pt"] - 0.05:
            issues["legibility"].append(f"{c.id}: '{c.plain[:30]}' is {pt:.1f} pt at {b['width_in']} in wide (min {b['min_font_pt']} pt)")
    for c in doc.cells:
        sw = float(c.get("strokeWidth", 1)) if (c.is_edge or (c.is_vertex and c.get("strokeColor") not in (None, "none"))) else None
        if sw is not None and sw * pt_per_px < 0.45 and not c.is_text:
            issues["legibility"].append(f"{c.id}: stroke {sw:g}px = {sw * pt_per_px:.2f} pt (< 0.5 pt hairline)")
            break
    for c in doc.vertices():
        if c.is_image:
            m = re.match(r"data:image/(png|jpeg|jpg|webp),(.+)", c.get("image", ""), re.S)
            if m:
                try:
                    from PIL import Image
                    im = Image.open(io.BytesIO(base64.b64decode(m.group(2))))
                    g = c.abs_geo()
                    printed_in = g[2] * pt_per_px / 72
                    dpi = im.width / max(printed_in, 1e-6)
                    if dpi < 300:
                        issues["legibility"].append(f"{c.id}: embedded raster is {dpi:.0f} dpi at print size (journals want >= 300)")
                except Exception:
                    pass

    # --- color
    th = load_yaml(theme_path) if theme_path else {}
    palette = {v.lower() for v in (th.get("colors") or {}).values()}
    used = {}
    for c in doc.cells:
        for k in ("fillColor", "strokeColor", "fontColor", "gradientColor"):
            v = c.get(k)
            if is_hex(v):
                used[v.lower()] = used.get(v.lower(), 0) + 1
        for m in re.finditer(r'color="?(#[0-9a-fA-F]{3,6})', c.value):
            used[m.group(1).lower()] = used.get(m.group(1).lower(), 0) + 1
    stats["colors_used"] = used
    chromatic = [h for h in used if max(hex2rgb(h)) - min(hex2rgb(h)) > 0.08]
    if len(chromatic) > 4:
        issues["color"].append(f"{len(chromatic)} distinct hues ({', '.join(sorted(chromatic))}); academic figures read best with <= 3 hues + tints")
    if palette:
        off = sorted(h for h in used if h not in palette)
        if off:
            issues["color"].append(f"off-theme colors: {', '.join(off)}")
    # category fills (effective color after opacity)
    fills = {}
    for c in doc.vertices():
        f = c.get("fillColor")
        if not is_hex(f) or c.is_text:
            continue
        a = float(c.get("fillOpacity", 100)) / 100 * float(c.get("opacity", 100)) / 100
        eff = blend(hex2rgb(f), a, page_bg)
        if contrast(eff, page_bg) < 1.08:
            continue
        key = "#%02x%02x%02x" % tuple(round(x * 255) for x in eff)
        fills.setdefault(key, []).append(c.id)
    keys = sorted(fills)
    stats["category_fills"] = {k: len(v) for k, v in fills.items()}
    for i in range(len(keys)):
        for j in range(i + 1, len(keys)):
            A, B = hex2rgb(keys[i]), hex2rgb(keys[j])
            if de(A, B) < 8:
                continue
            if b.get("grayscale") and abs(to_lab(A)[0] - to_lab(B)[0]) < 12:
                issues["color"].append(f"fills {keys[i]} and {keys[j]} differ by only {abs(to_lab(A)[0] - to_lab(B)[0]):.0f} L* in grayscale "
                                       f"(e.g. {fills[keys[i]][0]} vs {fills[keys[j]][0]}) — add pattern/outline or change lightness")
            if b.get("colorblind"):
                for kind in CVD:
                    d2 = de(simulate(A, kind), simulate(B, kind))
                    if d2 < 8:
                        issues["color"].append(f"fills {keys[i]} vs {keys[j]} nearly identical under {kind} (dE {d2:.1f})")
                        break
    if reference and os.path.exists(reference):
        before, after = category_count(reference), category_count(path)
        stats["category_colors"] = {"before": before, "after": after}
        if after < before:
            issues["color"].append(f"color-coded categories dropped from {before} to {after} distinct fills vs the original — "
                                   "a color pass or edit erased a distinction (restore it or confirm it was intended)")
    # text contrast
    verts = [c for c in doc.vertices() if c.abs_geo()]
    for c in verts:
        if not c.plain or c.is_math and False:
            continue
        fc = c.get("fontColor", "#000000")
        if not is_hex(fc):
            continue
        g = c.abs_geo()
        bgc = None
        if is_hex(c.get("fillColor")) and not c.is_text:
            a = float(c.get("fillOpacity", 100)) / 100
            bgc = blend(hex2rgb(c.get("fillColor")), a, page_bg)
        else:
            cx, cy = g[0] + g[2] / 2, g[1] + g[3] / 2
            best = None
            for o in verts:
                if o is c or not is_hex(o.get("fillColor")) or o.is_text or o.is_image:
                    continue
                og = o.abs_geo()
                if og[0] <= cx <= og[0] + og[2] and og[1] <= cy <= og[1] + og[3]:
                    if best is None or og[2] * og[3] < best[0]:
                        best = (og[2] * og[3], o)
            if best:
                o = best[1]
                bgc = blend(hex2rgb(o.get("fillColor")), float(o.get("fillOpacity", 100)) / 100, page_bg)
        bgc = bgc or page_bg
        ratio = contrast(hex2rgb(fc), bgc)
        big = float(c.get("fontSize", 11)) * pt_per_px >= 14 or (int(float(c.get("fontStyle", 0))) & 1)
        if ratio < (3.0 if big else 4.5):
            issues["color"].append(f"{c.id}: text contrast {ratio:.1f}:1 ('{c.plain[:25]}', {fc} on {'#%02x%02x%02x' % tuple(round(x * 255) for x in bgc)})")

    # --- typography
    stats["font_families"] = fams; stats["font_sizes"] = sizes
    if len(fams) > 2:
        issues["typography"].append(f"{len(fams)} font families: {', '.join(fams)} (use 1, at most 2)")
    if len(sizes) > 6:
        issues["typography"].append(f"{len(sizes)} distinct font sizes {sorted(sizes)} (use <= 5-6 roles)")
    if th.get("sizes"):
        roles = {float(v) for v in th["sizes"].values()}
        odd = sorted(s for s in sizes if s not in roles)
        if odd:
            issues["typography"].append(f"font sizes not in theme: {odd}")

    # --- math
    mlabels = [c for c in doc.vertices() if c.is_math]
    inline = [c for c in doc.vertices() if not c.is_math and re.search(r"\\\(.+?\\\)", c.plain)]
    stats["math_labels"] = len(mlabels); stats["inline_math_labels"] = len(inline)
    if (mlabels or inline) and not doc.math:
        issues["math"].append("labels contain $$…$$ but the diagram math flag is off (draw.io shows raw LaTeX)")
    try:
        from drawio_kit import render_math
        for c in mlabels:
            src = re.fullmatch(r"\s*\$\$(.+)\$\$\s*", c.plain, re.S).group(1)
            if render_math(src, 20) is None:
                issues["math"].append(f"{c.id}: LaTeX may not render: {src[:50]}")
    except Exception:
        pass
    uni = [c for c in doc.vertices() if c.plain and not c.is_math and any(ch in UNICODE_MATH for ch in c.plain)]
    if uni and (mlabels or inline):
        issues["math"].append(f"{len(uni)} plain labels use Unicode math ({', '.join(sorted({ch for c in uni for ch in c.plain if ch in UNICODE_MATH}))}) "
                              f"while {len(mlabels)} use LaTeX — make notation consistent (e.g. {uni[0].id})")
    elif uni:
        issues["math"].append(f"{len(uni)} labels write math in Unicode ({uni[0].id}: '{uni[0].plain[:25]}'); consider $$LaTeX$$ labels matching the paper")

    # --- layout
    try:
        from lint_drawio import lint
        L = lint(path, theme_path)
        # off-theme values are already reported under color / typography
        issues["layout"] += [f"lint: {w}" for w in L["errors"] + L["warnings"] if not w.startswith("off-theme")]
    except Exception as e:
        issues["layout"].append(f"lint failed: {e}")
    m = [cb[0], cb[1], pw - cb[0] - cb[2], ph - cb[1] - cb[3]]
    stats["margins"] = [round(v, 1) for v in m]
    if min(m) < 6:
        issues["layout"].append(f"content touches the page edge (margins l,t,r,b = {[round(v) for v in m]})")
    if m[0] + m[2] > 0.2 * pw or m[1] + m[3] > 0.2 * ph:
        issues["layout"].append(f"excess page whitespace (margins {[round(v) for v in m]}) — fit page to content")
    aspect = cb[3] / max(cb[2], 1)
    stats["aspect_h_over_w"] = round(aspect, 2)
    if b["width_in"] <= 3.6 and aspect > 1.3:
        issues["layout"].append(f"tall figure (h/w = {aspect:.2f}) for a single column; consider a wider layout")
    if b["width_in"] >= 6 and aspect < 0.25:
        issues["layout"].append(f"very flat figure (h/w = {aspect:.2f}); text will be small when scaled to width")
    texts = [c for c in doc.vertices() if c.plain and not c.is_math]
    stats["text_items"] = len(texts)
    if len(texts) > 35:
        issues["layout"].append(f"{len(texts)} text items — dense; move explanations to the caption")
    for c in texts:
        if len(c.plain.split()) > 16:
            issues["layout"].append(f"{c.id}: long label ({len(c.plain.split())} words) — shorten or move to caption")
    if texts:
        top = [c for c in texts if c.abs_geo()[1] < cb[1] + 0.08 * cb[3]]
        mx = max(float(c.get("fontSize", 11)) for c in texts)
        for c in top:
            if float(c.get("fontSize", 11)) >= mx and float(c.get("fontSize", 11)) >= 1.3 * sorted(sizes)[len(sizes) // 2] \
                    and len(c.plain.split()) >= 3 and c.abs_geo()[2] > 0.5 * cb[2]:
                issues["layout"].append(f"{c.id}: in-figure title '{c.plain[:40]}' — journals put titles in the caption")

    # --- connectors
    sws, heads = {}, {}
    for e in doc.edges():
        sws[e.get("strokeWidth", "1")] = 1
        if e.get("endArrow", "classic") != "none":
            heads[e.get("endSize", "6")] = 1
        for side, ptag in (("source", "sourcePoint"), ("target", "targetPoint")):
            if e.el.get(side):
                continue
            p = next((p for p in e.el.iter("mxPoint") if p.get("as") == ptag), None)
            if p is None:
                continue
            x, y = float(p.get("x", 0)), float(p.get("y", 0))
            for v in verts:
                if v.is_text or v.is_group or v.is_image:
                    continue
                g = v.abs_geo()
                dx = max(g[0] - x, 0, x - g[0] - g[2]); dy = max(g[1] - y, 0, y - g[1] - g[3])
                inside = g[0] < x < g[0] + g[2] and g[1] < y < g[1] + g[3]
                if (math.hypot(dx, dy) <= 8 and not inside and g[2] > 1 and g[3] > 1):
                    issues["connectors"].append(f"{e.id}: free {side} end ({x:.0f},{y:.0f}) is {math.hypot(dx, dy):.0f}px from {v.id} — attach it")
                    break
    if len(sws) > 3:
        issues["connectors"].append(f"{len(sws)} different edge stroke widths {sorted(sws, key=float)} (use 2-3 roles)")
    if len(heads) > 2:
        issues["connectors"].append(f"{len(heads)} arrow-head sizes {sorted(heads, key=float)}")

    accepted = [x for x in (accepted or []) if x]
    ignored = 0
    if accepted:
        for k in issues:
            keep = [x for x in issues[k] if not any(s in x for s in accepted)]
            ignored += len(issues[k]) - len(keep); issues[k] = keep
    stats["accepted_ignored"] = ignored
    status = {k: ("pass" if not v else "warn" if len(v) <= 3 else "fail") for k, v in issues.items()}
    n = sum(len(v) for v in issues.values())
    score = max(0, 100 - sum(min(len(v), 8) * w for (k, v), w in zip(issues.items(), (4, 3, 2, 4, 2, 2))))
    return dict(file=path, brief=b, pt_per_px=round(pt_per_px, 3), issues=issues, status=status, stats=stats,
                score=score, n_issues=n)


def write_report(r, out_dir, tag):
    os.makedirs(out_dir, exist_ok=True)
    b = r["brief"]
    md = [f"# Audit ({tag}): `{os.path.basename(r['file'])}`", "",
          f"Target: **{b['width_in']} in** wide, min font **{b['min_font_pt']} pt**, grayscale-safe: {b.get('grayscale')}, "
          f"colorblind-safe: {b.get('colorblind')}  (1 px = {r['pt_per_px']} pt)", "",
          f"**Readiness score: {r['score']}/100**  ({r['n_issues']} issues)", "",
          "| area | status | issues |", "|---|---|---|"]
    for k, v in r["issues"].items():
        md.append(f"| {k} | {r['status'][k]} | {len(v)} |")
    for k, v in r["issues"].items():
        if v:
            md += ["", f"## {k}", ""] + [f"- {x}" for x in v[:40]] + ([f"- … {len(v) - 40} more"] if len(v) > 40 else [])
    s = r["stats"]
    md += ["", "## stats", "", f"- font families: {s.get('font_families')}", f"- font sizes: {s.get('font_sizes')}",
           f"- colors used: {s.get('colors_used')}", f"- margins l,t,r,b: {s.get('margins')}; h/w = {s.get('aspect_h_over_w')}",
           f"- text items: {s.get('text_items')}; math labels: {s.get('math_labels')} (+{s.get('inline_math_labels', 0)} with inline math)",
           f"- accepted warnings ignored (brief review.accepted_warnings): {s.get('accepted_ignored', 0)}"]
    p = os.path.join(out_dir, f"{tag}_audit.md")
    open(p, "w", encoding="utf-8").write("\n".join(md) + "\n")
    json.dump(r, open(os.path.join(out_dir, f"{tag}_audit.json"), "w", encoding="utf-8"), indent=1, default=str)
    return p


def print_preview(path, out_dir, tag, width_in, dpi=150):
    from render_preview import render_kit
    out = os.path.join(out_dir, f"{tag}_print_preview.png")
    render_kit(path, out, scale=2.0, width=int(width_in * dpi))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("drawio"); ap.add_argument("--brief"); ap.add_argument("--theme")
    ap.add_argument("--out-dir", default="."); ap.add_argument("--tag", default="audit")
    ap.add_argument("--width-in", type=float); ap.add_argument("--min-pt", type=float)
    ap.add_argument("--reference", help="original .drawio: report lost color categories")
    a = ap.parse_args()
    br = load_yaml(a.brief) if a.brief else {}
    br = br.get("target", br) if isinstance(br, dict) else {}
    if a.width_in: br["width_in"] = a.width_in
    if a.min_pt: br["min_font_pt"] = a.min_pt
    full = load_yaml(a.brief) if a.brief else {}
    r = audit(a.drawio, br, a.theme, (full.get("review") or {}).get("accepted_warnings", []), a.reference)
    p = write_report(r, a.out_dir, a.tag)
    pp = print_preview(a.drawio, a.out_dir, a.tag, r["brief"]["width_in"])
    print(f"score {r['score']}/100, {r['n_issues']} issues: " + ", ".join(f"{k} {len(v)}" for k, v in r["issues"].items()))
    for k, v in r["issues"].items():
        for x in v[:6]:
            print(f"  [{k}] {x}")
    print("wrote", p, "and", pp)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/brief.py'] = r'''#!/usr/bin/env python3
"""
brief.py — record interview answers and derive the concrete targets.

    python brief.py WORK --set target.venue=elsevier --set target.column=one_and_half \
                         --set target.print_mode=grayscale_safe --set style.font="Times New Roman" \
                         --set style.palette=okabe_ito --set intent.takeaway="..."  --resolve
    python brief.py WORK --show

--resolve derives: target.width_in / min_font_pt from assets/venues.yaml (unless set explicitly),
target.grayscale from print_mode, and applies style.palette (okabe_ito | grayscale | navy_default |
keep | a path to a theme file) and style.font (family name | keep) to WORK/theme.yaml, keeping the
theme's other roles, sizes and strokes. Values are parsed as YAML (numbers, booleans, lists).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import os
import sys

import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
COLUMN_ALIASES = {"single column": "single", "single": "single", "1.5 column": "one_and_half", "1.5": "one_and_half",
                  "one_and_half": "one_and_half", "double": "double", "full text width (double column)": "double",
                  "full": "double", "double column": "double"}
VENUE_ALIASES = {"ieee journal / conference": "ieee", "elsevier journal": "elsevier", "springer (lncs)": "springer_lncs",
                 "thesis / dissertation": "thesis", "lncs": "springer_lncs", "springer": "springer_lncs"}
PALETTES = {"okabe_ito": "okabe_ito.yaml", "color-blind-safe preset (okabe-ito)": "okabe_ito.yaml",
            "grayscale": "grayscale.yaml", "grayscale preset": "grayscale.yaml", "navy_default": "navy_default.yaml"}


def parse_value(v):
    """'a | b' -> list of strings; true/false/numbers -> typed; everything else stays a raw string."""
    v = v.strip()
    if not v:
        return ""
    if " | " in v:
        return [x.strip() for x in v.split(" | ") if x.strip()]
    try:
        x = yaml.safe_load(v)
        return x if isinstance(x, (bool, int, float)) else v
    except yaml.YAMLError:
        return v


def set_key(d, dotted, value):
    ks = dotted.split(".")
    for k in ks[:-1]:
        d = d.setdefault(k, {}) if isinstance(d.get(k), dict) else d.setdefault(k, {}) if k not in d else d[k]
    d[ks[-1]] = value


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work"); ap.add_argument("--set", action="append", default=[]); ap.add_argument("--resolve", action="store_true")
    ap.add_argument("--show", action="store_true")
    a = ap.parse_args()
    bp = os.path.join(a.work, "brief.yaml")
    b = yaml.safe_load(open(bp, encoding="utf-8")) or {}
    explicit = set()
    for kv in a.set:
        k, v = kv.split("=", 1)
        set_key(b, k.strip(), parse_value(v))
        explicit.add(k.strip())
    if a.resolve:
        venues = yaml.safe_load(open(os.path.join(SKILL, "assets", "venues.yaml"), encoding="utf-8"))
        t = b.setdefault("target", {})
        venue = VENUE_ALIASES.get(str(t.get("venue", "generic")).lower(), str(t.get("venue", "generic")).lower())
        col = COLUMN_ALIASES.get(str(t.get("column", "double")).lower(), str(t.get("column", "double")).lower())
        preset = venues.get(venue, venues["generic"])
        cols = preset["columns"]
        if "target.width_in" not in explicit and col in cols:
            t["width_in"] = cols[col]
        elif "target.width_in" not in explicit and col not in cols:
            t["width_in"] = cols.get("double", 7.0)
        if "target.min_font_pt" not in explicit:
            t["min_font_pt"] = preset["min_font_pt"]
        t["venue"], t["column"] = venue, col
        pm = str(t.get("print_mode", "color")).lower()
        t["grayscale"] = "gray" in pm
        if isinstance(t.get("colorblind"), str):
            t["colorblind"] = t["colorblind"].lower().startswith("y")
        # theme edits
        tp = b.get("style", {}).get("theme") or os.path.join(a.work, "theme.yaml")
        theme = yaml.safe_load(open(tp, encoding="utf-8"))
        pal = str(b.get("style", {}).get("palette", "keep"))
        src = PALETTES.get(pal.lower()) if pal.lower() in PALETTES else (pal if os.path.exists(pal) else None)
        if src:
            p = src if os.path.exists(src) else os.path.join(SKILL, "assets", "themes", src)
            new = yaml.safe_load(open(p, encoding="utf-8"))["colors"]
            theme.setdefault("colors", {}).update(new)
            print(f"palette: applied {os.path.basename(p)} to {tp}")
        font = str(b.get("style", {}).get("font", "keep"))
        if font.lower() not in ("keep", "keep current", ""):
            fam = {"times new roman (times-based journals)": "Times New Roman", "helvetica / arial (sans)": "Helvetica",
                   "latin modern / computer modern (latex look)": "Latin Modern Roman"}.get(font.lower(), font)
            theme.setdefault("fonts", {})["paper"] = {"family": fam}
            theme["default_font"] = "paper"
            print(f"font: default label font -> {fam}")
        yaml.safe_dump(theme, open(tp, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)
    yaml.safe_dump(b, open(bp, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)
    if a.show or a.resolve:
        t = b.get("target", {})
        print(f"target: {t.get('venue')} {t.get('column')} -> {t.get('width_in')} in, min {t.get('min_font_pt')} pt, "
              f"grayscale {t.get('grayscale')}, colorblind {t.get('colorblind')}")
        print(f"intent: {b.get('intent')}")
        print(f"content: {b.get('content')}")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/compare.py'] = r'''#!/usr/bin/env python3
"""
compare.py — quantify and visualize how closely a rebuilt diagram matches the source image.

    python compare.py SOURCE.png PREVIEW.png --out-dir work/compare --tag iter_03
    python compare.py SOURCE.png fig.drawio  --out-dir work/compare          # renders the .drawio first

Writes (prefix = tag):
  <tag>_side.png      source | preview side by side, shared coordinate grid
  <tag>_overlay.png   50/50 blend (ghosting = misalignment)
  <tag>_edges.png     source edges RED, preview edges CYAN, matches WHITE/GRAY  <- best for alignment work
  <tag>_heat.png      per-cell mismatch heatmap with cell coordinates (worst cells outlined)
  <tag>_metrics.json  scores + worst cells (in source px) — use the boxes with `inspect_image.py zoom`

Scores (0..1, higher is better):
  pixel_similarity  1 - mean |gray difference| (after slight blur) — color/fill agreement
  edge_f1           F1 of edge pixels within a 3 px tolerance — geometry/text placement agreement
  score             0.4 * pixel_similarity + 0.6 * edge_f1
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import sys

import numpy as np
from PIL import Image, ImageDraw, ImageFilter, ImageFont

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from drawio_kit import find_font_file  # noqa: E402


def _font(px):
    p = find_font_file("Helvetica")
    return ImageFont.truetype(p, max(8, int(px))) if p else ImageFont.load_default()


def edges(gray):
    try:
        import cv2
        return cv2.Canny(gray, 60, 160) > 0
    except ImportError:
        g = gray.astype(float)
        gx = np.zeros_like(g); gy = np.zeros_like(g)
        gx[:, 1:-1] = g[:, 2:] - g[:, :-2]
        gy[1:-1, :] = g[2:, :] - g[:-2, :]
        mag = np.hypot(gx, gy)
        return mag > max(40, np.percentile(mag, 90))


def dilate(mask, r):
    try:
        from scipy import ndimage
        return ndimage.binary_dilation(mask, iterations=r)
    except ImportError:
        out = mask.copy()
        for dy in range(-r, r + 1):
            for dx in range(-r, r + 1):
                out |= np.roll(np.roll(mask, dy, 0), dx, 1)
        return out


def compare(src_path, prev_path, out_dir, tag="cmp", cell=None, tol=3, quiet=False):
    os.makedirs(out_dir, exist_ok=True)
    if prev_path.endswith(".drawio"):
        from render_preview import render_kit
        rendered = os.path.join(out_dir, f"{tag}_preview.png")
        with Image.open(src_path) as s:
            render_kit(prev_path, rendered, width=s.width)
        prev_path = rendered
    A = Image.open(src_path).convert("RGB")
    B = Image.open(prev_path).convert("RGB")
    if B.size != A.size:
        B = B.resize(A.size, Image.LANCZOS)
    W, H = A.size
    ga = np.asarray(A.convert("L").filter(ImageFilter.GaussianBlur(1.2))).astype(float)
    gb = np.asarray(B.convert("L").filter(ImageFilter.GaussianBlur(1.2))).astype(float)
    diff = np.abs(ga - gb) / 255.0
    pix = 1 - diff.mean()
    ea, eb = edges(np.asarray(A.convert("L"))), edges(np.asarray(B.convert("L")))
    da, db = dilate(ea, tol), dilate(eb, tol)
    prec = (eb & da).sum() / max(1, eb.sum())
    rec = (ea & db).sum() / max(1, ea.sum())
    f1 = 2 * prec * rec / max(1e-9, prec + rec)
    score = 0.4 * pix + 0.6 * f1

    # per-cell scores
    cell = cell or max(40, int(round(max(W, H) / 14 / 10) * 10))
    cells = []
    for y in range(0, H, cell):
        for x in range(0, W, cell):
            sl = (slice(y, min(H, y + cell)), slice(x, min(W, x + cell)))
            ca, cb = ea[sl], eb[sl]
            if ca.sum() + cb.sum() < 30 and diff[sl].mean() < 0.03:
                continue
            p = (cb & da[sl]).sum() / max(1, cb.sum()); r = (ca & db[sl]).sum() / max(1, ca.sum())
            cf1 = 2 * p * r / max(1e-9, p + r)
            cpix = 1 - diff[sl].mean()
            cells.append(dict(box=[x, y, min(W, x + cell), min(H, y + cell)], edge_f1=round(float(cf1), 3),
                              pixel=round(float(cpix), 3), bad=round(float(1 - (0.4 * cpix + 0.6 * cf1)), 3),
                              missing_edges=int((ca & ~db[sl]).sum()), extra_edges=int((cb & ~da[sl]).sum())))
    worst = sorted(cells, key=lambda c: -c["bad"])[:8]

    # visuals
    f = _font(max(11, W / 90))
    side = Image.new("RGB", (W * 2 + 10, H), "white")
    side.paste(A, (0, 0)); side.paste(B, (W + 10, 0))
    d = ImageDraw.Draw(side, "RGBA")
    step = 50 if max(W, H) <= 1600 else 100
    for off in (0, W + 10):
        for x in range(0, W, step):
            d.line([(off + x, 0), (off + x, H)], fill=(255, 0, 0, 40))
            if (x // step) % 2 == 0:
                d.text((off + x + 2, 2), str(x), fill=(200, 0, 0, 200), font=f)
        for y in range(0, H, step):
            d.line([(off, y), (off + W, y)], fill=(0, 0, 255, 40))
            if (y // step) % 2 == 0:
                d.text((off + 2, y + 1), str(y), fill=(0, 0, 200, 200), font=f)
    d.text((W - 90, H - 24), "SOURCE", fill=(200, 0, 0), font=f)
    d.text((2 * W - 80, H - 24), "PREVIEW", fill=(0, 120, 200), font=f)
    side_p = os.path.join(out_dir, f"{tag}_side.png"); side.save(side_p)
    over_p = os.path.join(out_dir, f"{tag}_overlay.png"); Image.blend(A, B, 0.5).save(over_p)
    e = np.full((H, W, 3), 255, np.uint8)
    e[ea & ~db] = (230, 0, 0)          # in source only (missing in preview)
    e[eb & ~da] = (0, 170, 220)        # in preview only (extra / misplaced)
    e[ea & db] = (90, 90, 90)          # matched
    edges_p = os.path.join(out_dir, f"{tag}_edges.png"); Image.fromarray(e).save(edges_p)
    heat = A.convert("L").convert("RGB")
    hd = ImageDraw.Draw(heat, "RGBA")
    for c in cells:
        a = int(200 * min(1, c["bad"] * 1.6))
        hd.rectangle(c["box"], fill=(255, 0, 0, a // 2))
    for i, c in enumerate(worst):
        hd.rectangle(c["box"], outline=(255, 0, 0, 255), width=3)
        hd.text((c["box"][0] + 3, c["box"][1] + 2), f"#{i + 1} {c['box'][0]},{c['box'][1]}", fill=(120, 0, 0), font=f)
    heat_p = os.path.join(out_dir, f"{tag}_heat.png"); heat.save(heat_p)

    metrics = dict(tag=tag, source=src_path, preview=prev_path, pixel_similarity=round(float(pix), 4),
                   edge_precision=round(float(prec), 4), edge_recall=round(float(rec), 4),
                   edge_f1=round(float(f1), 4), score=round(float(score), 4), cell_size=cell, worst_cells=worst,
                   images=dict(side=side_p, overlay=over_p, edges=edges_p, heat=heat_p))
    json.dump(metrics, open(os.path.join(out_dir, f"{tag}_metrics.json"), "w", encoding="utf-8"), indent=1)
    if not quiet:
        print(f"[{tag}] score={score:.3f}  pixel={pix:.3f}  edge_f1={f1:.3f} (precision {prec:.3f}, recall {rec:.3f})")
        print("  worst cells (x0,y0,x1,y1): " + "; ".join(f"{c['box']} bad={c['bad']}" for c in worst[:5]))
        print("  images:", side_p, edges_p, heat_p)
    return metrics


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source"); ap.add_argument("preview", help="PNG preview or .drawio file")
    ap.add_argument("--out-dir", default="compare"); ap.add_argument("--tag", default="cmp")
    ap.add_argument("--cell", type=int, default=None); ap.add_argument("--tol", type=int, default=3)
    a = ap.parse_args()
    compare(a.source, a.preview, a.out_dir, a.tag, a.cell, a.tol)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/drawio_doc.py'] = r'''"""
drawio_doc — load, inspect and edit an EXISTING .drawio file (any origin: image-to-drawio,
text-to-drawio, hand-made). Complements drawio_kit (which builds new diagrams).

    from drawio_doc import Doc
    doc = Doc.load("fig.drawio")
    for c in doc.vertices(): print(c.id, c.plain, c.style.get("fillColor"))
    doc.find(text="Encoder")[0].set_style(fillColor="#19355d", fontColor="#ffffff")
    doc.move("box1", dx=4); doc.align(["a", "b", "c"], "top"); doc.distribute(["a", "b", "c"], "x")
    doc.save("fig#ENHANCED.drawio")

Geometry helpers work in ABSOLUTE page coordinates (children of groups are converted for you).
Styles are dicts; base tokens (e.g. "text", "ellipse", "group") are kept in Cell.tokens.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import base64
import copy
import html
import re
import urllib.parse
import xml.etree.ElementTree as ET
import zlib

__all__ = ["Doc", "Cell", "parse_style", "style_to_str", "plain_text"]


def parse_style(s):
    tokens, kv = [], {}
    for part in (s or "").split(";"):
        if not part:
            continue
        if "=" in part:
            k, v = part.split("=", 1)
            kv[k] = v
        else:
            tokens.append(part)
    return tokens, kv


def style_to_str(tokens, kv):
    parts = list(tokens) + [f"{k}={v}" for k, v in kv.items()]
    return ";".join(parts) + (";" if parts else "")


def plain_text(value):
    v = re.sub(r"<br\s*/?>", "\n", value or "")
    v = re.sub(r"</(div|p)>", "\n", v)
    return html.unescape(re.sub(r"<[^>]+>", "", v)).strip()


class Cell:
    def __init__(self, doc, el, wrapper=None):
        self.doc, self.el, self.wrapper = doc, el, wrapper
        self.tokens, self.style = parse_style(el.get("style"))

    # identity / kind
    @property
    def id(self):
        return (self.wrapper if self.wrapper is not None else self.el).get("id")

    @property
    def is_vertex(self):
        return self.el.get("vertex") == "1"

    @property
    def is_edge(self):
        return self.el.get("edge") == "1"

    @property
    def is_group(self):
        return "group" in self.tokens

    @property
    def is_text(self):
        return "text" in self.tokens

    @property
    def is_image(self):
        return self.style.get("shape") == "image"

    @property
    def is_math(self):
        return bool(re.fullmatch(r"\s*\$\$.+\$\$\s*", self.plain, re.S))

    @property
    def parent(self):
        return self.el.get("parent")

    @property
    def value(self):
        w = self.wrapper
        return (w.get("label") if w is not None else self.el.get("value")) or ""

    @value.setter
    def value(self, v):
        (self.wrapper if self.wrapper is not None else self.el).set("label" if self.wrapper is not None else "value", v)

    @property
    def plain(self):
        return plain_text(self.value)

    # style
    def set_style(self, **kv):
        for k, v in kv.items():
            if v is None:
                self.style.pop(k, None)
            else:
                self.style[k] = str(v)
        self._flush()
        return self

    def get(self, k, default=None):
        return self.style.get(k, default)

    def _flush(self):
        self.el.set("style", style_to_str(self.tokens, self.style))

    # geometry
    @property
    def geo_el(self):
        return self.el.find("mxGeometry")

    def rel_geo(self):
        g = self.geo_el
        if g is None:
            return None
        return [float(g.get(k, 0) or 0) for k in ("x", "y", "width", "height")]

    def abs_geo(self):
        g = self.rel_geo()
        if g is None:
            return None
        p = self.doc.get(self.parent) if self.parent not in (None, "0", "1") else None
        if p is not None and p.is_vertex:
            pg = p.abs_geo()
            if pg:
                g[0] += pg[0]; g[1] += pg[1]
        return g

    def set_abs_geo(self, x=None, y=None, w=None, h=None):
        g = self.geo_el
        cur = self.abs_geo()
        rel = self.rel_geo()
        ox, oy = cur[0] - rel[0], cur[1] - rel[1]
        if x is not None: g.set("x", _fmt(x - ox))
        if y is not None: g.set("y", _fmt(y - oy))
        if w is not None: g.set("width", _fmt(w))
        if h is not None: g.set("height", _fmt(h))
        return self

    def __repr__(self):
        kind = "edge" if self.is_edge else "group" if self.is_group else "text" if self.is_text else "vertex"
        return f"<{kind} {self.id} {self.plain[:30]!r}>"


def _fmt(v):
    v = round(float(v), 2)
    return str(int(v)) if v.is_integer() else str(v)


class Doc:
    def __init__(self, tree, root_el, model, diagram, compressed):
        self.tree, self.root_el, self.model, self.diagram, self.compressed = tree, root_el, model, diagram, compressed
        self._index()

    # ---------- io
    @classmethod
    def load(cls, path):
        tree = ET.parse(path)
        top = tree.getroot()
        if top.tag == "mxGraphModel":
            return cls(tree, top, top, None, False)
        diag = top.find("diagram")
        model = diag.find("mxGraphModel")
        compressed = False
        if model is None:
            data = base64.b64decode(diag.text.strip())
            model = ET.fromstring(urllib.parse.unquote(zlib.decompress(data, -15).decode()))
            diag.text = None
            diag.append(model)
            compressed = True
        return cls(tree, top, model, diag, compressed)

    def save(self, path):
        self._index()
        self.tree.write(path, encoding="utf-8", xml_declaration=False)
        return path

    def _index(self):
        self.root = self.model.find("root")
        self.cells, self._by = [], {}
        for o in list(self.root):
            if o.tag == "mxCell":
                c = Cell(self, o)
            else:
                inner = o.find("mxCell")
                if inner is None:
                    continue
                c = Cell(self, inner, wrapper=o)
            self.cells.append(c)
            self._by[c.id] = c

    # ---------- query
    def get(self, cid):
        return self._by.get(cid)

    def vertices(self, include_groups=False):
        return [c for c in self.cells if c.is_vertex and (include_groups or not c.is_group)]

    def edges(self):
        return [c for c in self.cells if c.is_edge]

    def find(self, text=None, style=None, kind=None):
        """text: regex on plain label; style: dict of key->value (or regex string) that must match."""
        out = []
        for c in self.cells:
            if kind == "vertex" and not c.is_vertex: continue
            if kind == "edge" and not c.is_edge: continue
            if kind == "text" and not c.is_text: continue
            if text and not re.search(text, c.plain, re.I): continue
            if style and not all(re.fullmatch(str(v), c.style.get(k, "")) for k, v in style.items()): continue
            out.append(c)
        return out

    @property
    def page(self):
        return float(self.model.get("pageWidth", 850)), float(self.model.get("pageHeight", 1100))

    def set_page(self, w, h):
        self.model.set("pageWidth", _fmt(w)); self.model.set("pageHeight", _fmt(h))
        self.model.set("dx", _fmt(w)); self.model.set("dy", _fmt(h))

    @property
    def math(self):
        return self.model.get("math") == "1"

    @math.setter
    def math(self, on):
        self.model.set("math", "1" if on else "0")

    def content_bbox(self):
        boxes = [c.abs_geo() for c in self.vertices() if c.abs_geo()]
        for e in self.edges():
            for p in e.el.iter("mxPoint"):
                if p.get("as") in ("sourcePoint", "targetPoint", None):
                    x, y = float(p.get("x", 0)), float(p.get("y", 0))
                    boxes.append([x, y, 0, 0])
        if not boxes:
            return [0, 0, 0, 0]
        x0 = min(b[0] for b in boxes); y0 = min(b[1] for b in boxes)
        x1 = max(b[0] + b[2] for b in boxes); y1 = max(b[1] + b[3] for b in boxes)
        return [x0, y0, x1 - x0, y1 - y0]

    # ---------- edit
    def move(self, cid, dx=0, dy=0):
        c = self.get(cid)
        g = c.abs_geo()
        c.set_abs_geo(x=g[0] + dx, y=g[1] + dy)
        return c

    def translate_all(self, dx, dy):
        """Shift every top-level vertex and every free edge point / waypoint."""
        for c in self.cells:
            if c.is_vertex and c.parent in ("1", None):
                g = c.geo_el
                g.set("x", _fmt(float(g.get("x", 0)) + dx)); g.set("y", _fmt(float(g.get("y", 0)) + dy))
            elif c.is_edge and c.parent in ("1", None):
                for p in c.el.iter("mxPoint"):
                    p.set("x", _fmt(float(p.get("x", 0)) + dx)); p.set("y", _fmt(float(p.get("y", 0)) + dy))

    def align(self, ids, edge="top"):
        """Align vertices on left|right|top|bottom|center-x|center-y to their mean."""
        cs = [self.get(i) for i in ids]
        gs = [c.abs_geo() for c in cs]
        f = {"left": lambda g: g[0], "right": lambda g: g[0] + g[2], "center-x": lambda g: g[0] + g[2] / 2,
             "top": lambda g: g[1], "bottom": lambda g: g[1] + g[3], "center-y": lambda g: g[1] + g[3] / 2}[edge]
        target = sum(f(g) for g in gs) / len(gs)
        for c, g in zip(cs, gs):
            d = target - f(g)
            if edge in ("left", "right", "center-x"):
                c.set_abs_geo(x=g[0] + d)
            else:
                c.set_abs_geo(y=g[1] + d)

    def distribute(self, ids, axis="x"):
        """Equal gaps between vertices (sorted by position) keeping the first and last fixed."""
        cs = sorted((self.get(i) for i in ids), key=lambda c: c.abs_geo()[0 if axis == "x" else 1])
        gs = [c.abs_geo() for c in cs]
        k = 0 if axis == "x" else 1
        span = gs[-1][k] + gs[-1][k + 2] - gs[0][k]
        total = sum(g[k + 2] for g in gs)
        gap = (span - total) / (len(gs) - 1) if len(gs) > 1 else 0
        pos = gs[0][k]
        for c, g in zip(cs, gs):
            c.set_abs_geo(**({"x": pos} if axis == "x" else {"y": pos}))
            pos += g[k + 2] + gap

    def equalize(self, ids, width=True, height=True, anchor="center"):
        cs = [self.get(i) for i in ids]
        gs = [c.abs_geo() for c in cs]
        W = max(g[2] for g in gs); H = max(g[3] for g in gs)
        for c, g in zip(cs, gs):
            nw, nh = (W if width else g[2]), (H if height else g[3])
            nx = g[0] - (nw - g[2]) / 2 if anchor == "center" else g[0]
            ny = g[1] - (nh - g[3]) / 2 if anchor == "center" else g[1]
            c.set_abs_geo(nx, ny, nw, nh)

    def remove(self, cid):
        c = self.get(cid)
        if c is None:
            return
        kids = [k.id for k in self.cells if k.parent == cid]
        for k in kids:
            self.remove(k)
        self.root.remove(c.wrapper if c.wrapper is not None else c.el)
        for e in self.edges():  # detach edges pointing to it
            for side in ("source", "target"):
                if e.el.get(side) == cid:
                    del e.el.attrib[side]
        self._index()

    def add_vertex(self, cid, value, style, x, y, w, h, parent="1"):
        el = ET.SubElement(self.root, "mxCell", id=cid, value=value, style=style, vertex="1", parent=parent)
        ET.SubElement(el, "mxGeometry", x=_fmt(x), y=_fmt(y), width=_fmt(w), height=_fmt(h), **{"as": "geometry"})
        self._index()
        return self.get(cid)

    def add_edge(self, cid, style, src=None, tgt=None, points=(), source_point=None, target_point=None, value=""):
        attrs = dict(id=cid, value=value, style=style, edge="1", parent="1")
        if src: attrs["source"] = src
        if tgt: attrs["target"] = tgt
        el = ET.SubElement(self.root, "mxCell", **attrs)
        g = ET.SubElement(el, "mxGeometry", relative="1", **{"as": "geometry"})
        if source_point:
            ET.SubElement(g, "mxPoint", x=_fmt(source_point[0]), y=_fmt(source_point[1]), **{"as": "sourcePoint"})
        if target_point:
            ET.SubElement(g, "mxPoint", x=_fmt(target_point[0]), y=_fmt(target_point[1]), **{"as": "targetPoint"})
        if points:
            arr = ET.SubElement(g, "Array", **{"as": "points"})
            for px, py in points:
                ET.SubElement(arr, "mxPoint", x=_fmt(px), y=_fmt(py))
        self._index()
        return self.get(cid)

    def new_id(self, base="enh"):
        i = 1
        while f"{base}_{i}" in self._by:
            i += 1
        return f"{base}_{i}"

    def clone(self):
        return Doc(copy.deepcopy(self.tree), None, None, None, self.compressed)._reattach()

    def _reattach(self):
        top = self.tree.getroot()
        self.root_el = top
        if top.tag == "mxGraphModel":
            self.model, self.diagram = top, None
        else:
            self.diagram = top.find("diagram"); self.model = self.diagram.find("mxGraphModel")
        self._index()
        return self
'''

# ================================================================================================
FILES['scripts/drawio_kit.py'] = r'''#!/usr/bin/env python3
r"""
drawio_kit — build *editable* draw.io (.drawio) diagrams that replicate a source image.

Design rules baked into this library
  * Coordinates are given in SOURCE-IMAGE pixels (measure them on the inspection grid).
    `scale` multiplies geometry, font sizes and page size at save time.
  * Colors, fonts, font sizes, stroke widths and opacities are referenced by THEME ROLE
    ("primary", "heading", "thick", "tint" ...). Raw hex/numbers still work but are reported
    as off-theme so lint can flag them.
  * Everything that draw.io can represent natively is emitted natively (boxes, text, edges,
    ellipses, stars, groups). Free-form geometry (polygons, blobs, meshes) becomes an embedded
    SVG image; photos are cropped from the source and embedded as JPEG/PNG.

Minimal example
    import sys; sys.path.insert(0, "<skill>/scripts")
    from drawio_kit import Diagram
    d = Diagram.from_source("images/fig.png", theme="theme.yaml")
    d.rect(40, 40, 200, 60, "**Encoder**", fill="primary", color="background", size="heading", rounded=True)
    d.text("Bidirectional context.", 140, 130, size="body", bold=True)
    d.edge([(140, 100), (140, 120)], arrow="end")
    d.save("images/fig.drawio")

Text markup (used by every `label`/`text` argument unless raw=True):
    **bold**   *italic*   __underline__   {role|colored text}   {#ff0000|hex color}   \n = line break
    \* = literal asterisk
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import base64
import contextlib
import html
import io
import json
import math
import os
import random
import re
import sys
import urllib.parse
from html.parser import HTMLParser
from xml.sax.saxutils import quoteattr

__all__ = ["Theme", "Diagram", "html_lines", "measure_line", "find_font_file", "catmull_rom", "render_math", "measure_math"]

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
DEFAULT_THEME_PATH = os.path.join(SKILL_DIR, "assets", "theme.yaml")

# --------------------------------------------------------------------------------------
# Theme
# --------------------------------------------------------------------------------------
_BUILTIN_THEME = {
    "colors": {"primary": "#19355d", "secondary": "#4276b0", "neutral": "#a4a4a4",
               "text": "#000000", "background": "#ffffff"},
    "fonts": {"sans": {"family": "Helvetica"}, "serif": {"family": "Times New Roman"},
              "mono": {"family": "Courier New"}},
    "default_font": "sans",
    "sizes": {"title": 28, "heading": 20, "subheading": 16, "body": 14, "label": 12,
              "small": 10, "caption": 18},
    "line_height": 1.2,
    "strokes": {"hairline": 0.75, "thin": 1, "normal": 1.5, "thick": 2.5, "heavy": 4},
    "opacity": {"tint": 25, "soft": 45, "faded": 35},
    "shape": {"corner_radius": 8, "shadow": False},
    "spacing": {"padding": 10, "gap": 12},
}


def _deep_merge(base, over):
    out = dict(base)
    for k, v in (over or {}).items():
        out[k] = _deep_merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def _norm_hex(c):
    c = c.strip().lower()
    if re.fullmatch(r"#[0-9a-f]{3}", c):
        c = "#" + "".join(ch * 2 for ch in c[1:])
    return c


class Theme:
    """Role-based theme loaded from YAML / TOML / JSON (merged over built-in defaults)."""

    def __init__(self, data=None, path=None):
        self.data = _deep_merge(_BUILTIN_THEME, data or {})
        self.path = path
        self.colors = {k: _norm_hex(v) for k, v in self.data["colors"].items()}
        self.fonts = self.data["fonts"]
        self.sizes = self.data["sizes"]
        self.strokes = self.data["strokes"]
        self.opacities = self.data["opacity"]
        self.line_height = float(self.data.get("line_height", 1.2))
        self.off_theme = set()   # raw values used that are not part of the theme

    # ---- loading -------------------------------------------------------------------
    @classmethod
    def load(cls, path=None, overrides=None):
        if isinstance(path, Theme):
            return path
        if isinstance(path, dict):
            return cls(_deep_merge(path, overrides or {}))
        path = path or (DEFAULT_THEME_PATH if os.path.exists(DEFAULT_THEME_PATH) else None)
        data = {}
        if path:
            ext = os.path.splitext(path)[1].lower()
            with open(path, "rb") as fh:
                raw = fh.read()
            if ext in (".yaml", ".yml"):
                try:
                    import yaml
                except ImportError as e:  # pragma: no cover
                    raise SystemExit("PyYAML not installed: pip install pyyaml, or use a .toml/.json theme") from e
                data = yaml.safe_load(raw) or {}
            elif ext == ".toml":
                try:
                    import tomllib
                except ImportError:  # Python < 3.11
                    import tomli as tomllib  # type: ignore
                data = tomllib.loads(raw.decode())
            elif ext == ".json":
                data = json.loads(raw)
            else:
                raise ValueError(f"Unsupported theme format: {path}")
        return cls(_deep_merge(data, overrides or {}), path)

    def to_dict(self):
        return self.data

    # ---- resolvers -----------------------------------------------------------------
    def color(self, c):
        if c is None or c == "none":
            return "none"
        if c in self.colors:
            return self.colors[c]
        if isinstance(c, str) and c.startswith("#"):
            c = _norm_hex(c)
            if c not in self.colors.values():
                self.off_theme.add(f"color {c}")
            return c
        raise KeyError(f"Unknown color role '{c}'. Roles: {sorted(self.colors)}")

    def size(self, s):
        if isinstance(s, str):
            if s not in self.sizes:
                raise KeyError(f"Unknown size role '{s}'. Roles: {sorted(self.sizes)}")
            return float(self.sizes[s])
        if float(s) not in [float(v) for v in self.sizes.values()]:
            self.off_theme.add(f"font size {s}")
        return float(s)

    def stroke(self, s):
        if isinstance(s, str):
            return float(self.strokes[s])
        return float(s)

    def opacity(self, o):
        if o is None:
            return None
        if isinstance(o, str):
            return int(self.opacities[o])
        return int(o)

    def font_key(self, f=None):
        f = f or self.data.get("default_font", "sans")
        if f not in self.fonts:
            self.off_theme.add(f"font {f}")
            return None
        return f

    def font_family(self, f=None):
        k = self.font_key(f)
        return self.fonts[k]["family"] if k else f

    def font_source(self, f=None):
        k = self.font_key(f)
        if k and self.fonts[k].get("google"):
            fam = self.fonts[k]["family"].replace(" ", "+")
            return urllib.parse.quote(f"https://fonts.googleapis.com/css?family={fam}", safe="")
        return None


# --------------------------------------------------------------------------------------
# Fonts + text measurement (PIL), shared with render_preview.py and lint_drawio.py
# --------------------------------------------------------------------------------------
_FONT_DIRS = ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts"),
              os.path.expanduser("~/.local/share/fonts"), "/Library/Fonts", "/System/Library/Fonts",
              os.path.expanduser("~/Library/Fonts"), "C:/Windows/Fonts"]
_FONT_INDEX = None
_FONT_ALIASES = {  # family keyword -> preferred file stems (metric-compatible first)
    "sans": ["LiberationSans", "Arial", "Helvetica", "arial", "DejaVuSans"],
    "serif": ["LiberationSerif", "TimesNewRoman", "Times New Roman", "times", "DejaVuSerif"],
    "mono": ["LiberationMono", "CourierNew", "cour", "DejaVuSansMono"],
}


def _font_index():
    global _FONT_INDEX
    if _FONT_INDEX is None:
        _FONT_INDEX = {}
        for d in _FONT_DIRS:
            if os.path.isdir(d):
                for root, _, files in os.walk(d):
                    for f in files:
                        if f.lower().endswith((".ttf", ".otf", ".ttc")):
                            _FONT_INDEX.setdefault(os.path.splitext(f)[0].lower(), os.path.join(root, f))
    return _FONT_INDEX


def _family_kind(family):
    f = (family or "").lower()
    if any(k in f for k in ("times", "serif", "georgia", "garamond", "cambria", "palatino")) and "sans" not in f:
        return "serif"
    if any(k in f for k in ("courier", "mono", "consol", "menlo")):
        return "mono"
    return "sans"


def find_font_file(family="Helvetica", bold=False, italic=False):
    """Best local TTF for a family (used for measuring/previewing; draw.io uses the real family)."""
    idx = _font_index()
    suffixes = (["-bolditalic", "bi", "-boldoblique", "z"] if bold and italic else
                ["-bold", "bd", "b"] if bold else ["-italic", "-oblique", "i"] if italic else
                ["-regular", "", "-book"])
    stems = [family.replace(" ", ""), family] + _FONT_ALIASES[_family_kind(family)]
    for stem in stems:
        for suf in suffixes:
            p = idx.get((stem + suf).lower())
            if p:
                return p
    for stem in _FONT_ALIASES[_family_kind(family)]:
        p = idx.get(stem.lower())
        if p:
            return p
    return next(iter(idx.values()), None)


_PIL_FONTS = {}


def _pil_font(family, size, bold=False, italic=False):
    key = (family, round(size * 4) / 4, bold, italic)
    if key not in _PIL_FONTS:
        try:
            from PIL import ImageFont
            path = find_font_file(family, bold, italic)
            _PIL_FONTS[key] = ImageFont.truetype(path, max(1, int(round(size)))) if path else None
        except Exception:
            _PIL_FONTS[key] = None
    return _PIL_FONTS[key]


INLINE_MATH = re.compile(r"\\\((.+?)\\\)")


def measure_line(segs, family="Helvetica", size=14.0):
    """Width in px of one line given segments [{'text','bold','italic','size'}].
    Inline LaTeX written as \\( ... \\) is measured as rendered math."""
    w = 0.0
    expanded = []
    for s in segs:
        parts = INLINE_MATH.split(s["text"])
        for i, p in enumerate(parts):
            if not p:
                continue
            if i % 2:
                w += measure_math(p, s.get("size") or size)[0]
            else:
                expanded.append(dict(s, text=p))
    for s in expanded:
        sz = s.get("size") or size
        f = _pil_font(family, sz, s.get("bold"), s.get("italic"))
        if f is None:
            w += len(s["text"]) * sz * (0.6 if s.get("bold") else 0.55)
        else:
            w += f.getlength(s["text"]) * (sz / max(1, int(round(sz))))
    return w


class _SegParser(HTMLParser):
    """Parse draw.io HTML labels into lines of styled segments."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines = [[]]
        self.stack = [dict(bold=False, italic=False, underline=False, color=None, size=None)]

    def _push(self, **kw):
        st = dict(self.stack[-1]); st.update({k: v for k, v in kw.items() if v is not None})
        self.stack.append(st)

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "br":
            self.lines.append([]); return
        if tag in ("div", "p") and self.lines[-1]:
            self.lines.append([])
        kw = {}
        if tag in ("b", "strong"): kw["bold"] = True
        if tag in ("i", "em"): kw["italic"] = True
        if tag == "u": kw["underline"] = True
        if tag == "font" and a.get("color"): kw["color"] = a["color"]
        css = a.get("style") or ""
        m = re.search(r"color:\s*(#[0-9a-fA-F]{3,6})", css)
        if m: kw["color"] = m.group(1)
        if re.search(r"font-weight:\s*(bold|[6-9]00)", css): kw["bold"] = True
        if "font-style: italic" in css or "font-style:italic" in css: kw["italic"] = True
        m = re.search(r"font-size:\s*([\d.]+)px", css)
        if m: kw["size"] = float(m.group(1))
        self._push(**kw)

    def handle_endtag(self, tag):
        if tag in ("br",):
            return
        if len(self.stack) > 1:
            self.stack.pop()
        if tag in ("div", "p"):
            self.lines.append([])

    def handle_startendtag(self, tag, attrs):
        if tag == "br":
            self.lines.append([])

    def handle_data(self, data):
        if data:
            st = self.stack[-1]
            self.lines[-1].append(dict(text=data, **st))


def html_lines(value, raw_text=False):
    """draw.io label value -> list of lines, each a list of segment dicts."""
    if raw_text or value is None:
        return [[dict(text=l, bold=False, italic=False, underline=False, color=None, size=None)]
                for l in (value or "").split("\n")]
    p = _SegParser()
    p.feed(value)
    lines = p.lines
    while len(lines) > 1 and not lines[-1]:
        lines.pop()
    return lines


# --------------------------------------------------------------------------------------
# Markup -> draw.io HTML
# --------------------------------------------------------------------------------------
def markup_to_html(s, theme):
    s = s.replace("\\*", "\x00")
    s = html.escape(s, quote=False)

    def col(m):
        return f'<font color="{theme.color(m.group(1))}">{m.group(2)}</font>'
    s = re.sub(r"\{(#?[\w]+)\|(.+?)\}", col, s, flags=re.S)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s, flags=re.S)
    s = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"<i>\1</i>", s, flags=re.S)
    s = re.sub(r"__(.+?)__", r"<u>\1</u>", s, flags=re.S)
    return s.replace("\x00", "*").replace("\n", "<br>")


# --------------------------------------------------------------------------------------
# Geometry helpers
# --------------------------------------------------------------------------------------
def catmull_rom(points, samples_per_seg=12, closed=False):
    """Smooth polyline through points (centripetal-ish uniform Catmull-Rom)."""
    pts = [tuple(map(float, p)) for p in points]
    if len(pts) < 3:
        return pts
    if closed:
        pts = pts[-1:] + pts + pts[:2]
    else:
        pts = [pts[0]] + pts + [pts[-1]]
    out = []
    for i in range(1, len(pts) - 2):
        p0, p1, p2, p3 = pts[i - 1], pts[i], pts[i + 1], pts[i + 2]
        for k in range(samples_per_seg):
            t = k / samples_per_seg
            t2, t3 = t * t, t * t * t
            out.append(tuple(0.5 * ((2 * p1[j]) + (-p0[j] + p2[j]) * t +
                                    (2 * p0[j] - 5 * p1[j] + 4 * p2[j] - p3[j]) * t2 +
                                    (-p0[j] + 3 * p1[j] - 3 * p2[j] + p3[j]) * t3) for j in (0, 1)))
    out.append(pts[-2])
    return out


def _resample(points, n):
    """Resample a polyline to n points evenly by arc length."""
    pts = [tuple(map(float, p)) for p in points]
    d = [0.0]
    for a, b in zip(pts, pts[1:]):
        d.append(d[-1] + math.dist(a, b))
    L = d[-1] or 1.0
    out, j = [], 0
    for i in range(n):
        t = L * i / (n - 1)
        while j < len(d) - 2 and d[j + 1] < t:
            j += 1
        seg = (d[j + 1] - d[j]) or 1.0
        f = (t - d[j]) / seg
        out.append((pts[j][0] + f * (pts[j + 1][0] - pts[j][0]), pts[j][1] + f * (pts[j + 1][1] - pts[j][1])))
    return out


def point_in_polygon(px, py, poly):
    inside = False
    for (x1, y1), (x2, y2) in zip(poly, poly[1:] + poly[:1]):
        if (y1 > py) != (y2 > py) and px < (x2 - x1) * (py - y1) / (y2 - y1) + x1:
            inside = not inside
    return inside


def latex_for_preview(latex):
    """MathJax-only constructs -> matplotlib mathtext equivalents (used for measuring/previewing)."""
    t = latex
    t = re.sub(r"\\(?:textcolor|color)\{[^}]*\}\{((?:[^{}]|\{[^{}]*\})*)\}", r"\1", t)
    t = re.sub(r"\\color\{[^}]*\}", "", t)
    t = t.replace(r"\text{", r"\mathrm{").replace(r"\operatorname{", r"\mathrm{")
    t = t.replace(r"\lVert", r"\Vert").replace(r"\rVert", r"\Vert").replace(r"\displaystyle", "")
    # $$...$$ is display style in MathJax: show fractions full size like MathJax does
    t = re.sub(r"\\(?<![dt])frac", r"\\dfrac", t.replace(r"\tfrac", r"\frac").replace(r"\dfrac", r"\frac"))
    return t


def render_math(latex, size_px, color="#000000"):
    """LaTeX -> transparent RGBA PIL image via matplotlib mathtext (Computer Modern), or None."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        matplotlib.rcParams["mathtext.fontset"] = "cm"
        from matplotlib.figure import Figure
        from PIL import Image
        fig = Figure(figsize=(0.01, 0.01))
        fig.text(0, 0, "$" + latex_for_preview(latex) + "$", fontsize=size_px, color=color)
        buf = io.BytesIO()
        fig.savefig(buf, dpi=72, transparent=True, bbox_inches="tight", pad_inches=0.02, format="png")
        buf.seek(0)
        return Image.open(buf).convert("RGBA")
    except Exception:
        return None


def measure_math(latex, size_px):
    im = render_math(latex, size_px)
    if im is None:
        return len(latex) * size_px * 0.45, size_px * 1.4
    return float(im.width), float(im.height)


SHAPES = {  # friendly name -> draw.io base style
    "rect": "", "ellipse": "ellipse", "circle": "ellipse;aspect=fixed",
    "star": "shape=mxgraph.basic.star", "triangle": "triangle;direction=north",
    "triangle_down": "triangle;direction=south", "triangle_right": "triangle",
    "rhombus": "rhombus", "diamond": "rhombus", "hexagon": "shape=hexagon;perimeter=hexagonPerimeter2;size=0.25",
    "parallelogram": "shape=parallelogram;perimeter=parallelogramPerimeter;size=0.2",
    "trapezoid": "shape=trapezoid;perimeter=trapezoidPerimeter;size=0.2",
    "cylinder": "shape=cylinder3;boundedLbl=1;backgroundOutline=1;size=10",
    "document": "shape=document;boundedLbl=1;size=0.25", "cloud": "ellipse;shape=cloud",
    "process": "shape=process;backgroundOutline=1;size=0.1", "note": "shape=note;size=14",
    "card": "shape=card;size=14", "actor": "shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top",
    "arrow_right": "shape=singleArrow;arrowWidth=0.5;arrowSize=0.3",
    "double_arrow": "shape=doubleArrow;arrowWidth=0.5;arrowSize=0.3",
    "cross": "shape=cross;size=0.25", "plus": "shape=cross;size=0.25",
    "donut": "ellipse;shape=doubleEllipse", "step": "shape=step;perimeter=stepPerimeter;size=0.2",
    "callout_box": "shape=callout;perimeter=calloutPerimeter;position2=0.5;size=12;position=0.4;base=16",
}


# --------------------------------------------------------------------------------------
# Diagram
# --------------------------------------------------------------------------------------
class Diagram:
    """Collects elements in source-pixel coordinates and serializes them as .drawio XML."""

    def __init__(self, width, height, theme=None, scale=1.0, source=None, background="background",
                 name="Page-1", grid=True, coord_scale=(1.0, 1.0), math=False):
        self.theme = Theme.load(theme)
        self.width, self.height, self.S = float(width), float(height), float(scale)
        # coord_scale: multiply every measured x/y (not font sizes) — use when coordinates were read
        # off a resized view of the source, e.g. coord_scale=(src_w / view_w, src_h / view_h)
        self.cs = tuple(map(float, coord_scale))
        self.math_enabled = math   # draw.io renders LaTeX labels ($$...$$); equation() turns it on
        self.source = source
        self._src_img = None
        self.background = background
        self.name = name
        self.grid = grid
        self.els = []
        self._ids = set()
        self._groups = []
        self._counter = 0

    @classmethod
    def from_source(cls, image_path, theme=None, scale=1.0, **kw):
        from PIL import Image
        with Image.open(image_path) as im:
            w, h = im.size
        return cls(w, h, theme=theme, scale=scale, source=image_path, **kw)

    # ---- ids / groups ---------------------------------------------------------------
    def _new_id(self, want=None):
        base = re.sub(r"[^\w\-]", "_", want) if want else None
        if base and base not in self._ids:
            self._ids.add(base); return base
        while True:
            self._counter += 1
            cand = f"{base or 'el'}_{self._counter}"
            if cand not in self._ids:
                self._ids.add(cand); return cand

    def _parent(self):
        return self._groups[-1] if self._groups else None

    @contextlib.contextmanager
    def group(self, name=None):
        """Everything created inside is grouped (moves together in draw.io). Edges stay top-level."""
        gid = self._new_id(name or "group")
        self.els.append(dict(kind="group", id=gid, parent=self._parent()))
        self._groups.append(gid)
        try:
            yield gid
        finally:
            self._groups.pop()

    def get(self, id):
        return next(e for e in self.els if e["id"] == id)

    def bbox(self, id):
        e = self.get(id)
        return e["geo"]

    # ---- style helpers --------------------------------------------------------------
    def _font_style(self, font=None, size="body", color="text", bold=False, italic=False,
                    underline=False, align="center", valign="middle", wrap=True, spacing=None):
        t = self.theme
        st = {"html": "1", "fontFamily": t.font_family(font), "fontSize": t.size(size),
              "fontColor": t.color(color), "align": align, "verticalAlign": valign,
              "whiteSpace": "wrap" if wrap else "nowrap"}
        fs = (1 if bold else 0) | (2 if italic else 0) | (4 if underline else 0)
        if fs:
            st["fontStyle"] = fs
        src = t.font_source(font)
        if src:
            st["fontSource"] = src
        if spacing is not None:
            st["spacing"] = spacing
        return st

    def _shape_style(self, fill="background", stroke="text", sw="normal", rounded=False, dashed=False,
                     dash=None, opacity=None, fill_opacity=None, stroke_opacity=None, shadow=None,
                     gradient=None, gradient_dir="south"):
        t = self.theme
        st = {"fillColor": t.color(fill), "strokeColor": t.color(stroke), "strokeWidth": t.stroke(sw)}
        if gradient:
            st.update(gradientColor=t.color(gradient), gradientDirection=gradient_dir)
        if rounded:
            r = t.data["shape"]["corner_radius"] if rounded is True else rounded
            st.update(rounded=1, absoluteArcSize=1, arcSize=round(2 * r, 1))
        else:
            st["rounded"] = 0
        if dashed:
            st.update(dashed=1, dashPattern=dash or "6 4")
        for k, v in (("opacity", opacity), ("fillOpacity", fill_opacity), ("strokeOpacity", stroke_opacity)):
            if v is not None:
                st[k] = t.opacity(v)
        if shadow if shadow is not None else t.data["shape"].get("shadow"):
            st["shadow"] = 1
        return st

    def _label(self, s, raw):
        if not s:
            return ""
        return s if raw else markup_to_html(s, self.theme)

    def _add_vertex(self, id, x, y, w, h, tokens, style, value=""):
        el = dict(kind="vertex", id=self._new_id(id), parent=self._parent(),
                  geo=[float(x), float(y), float(w), float(h)], tokens=tokens, style=style, value=value)
        self.els.append(el)
        return el["id"]

    # ---- primitives -----------------------------------------------------------------
    def rect(self, x, y, w, h, label="", *, fill="background", stroke="text", sw="normal", rounded=False,
             dashed=False, dash=None, opacity=None, fill_opacity=None, stroke_opacity=None, shadow=None,
             id=None, raw=False, extra=None, font=None, size="body", color="text", bold=False, italic=False,
             underline=False, align="center", valign="middle", wrap=True, spacing=None, shape="rect",
             gradient=None, gradient_dir="south"):
        """Box (or any SHAPES entry via shape=) with an optional centered/aligned label.
        rounded: True (theme corner_radius) or a radius in px. gradient: second color role (top->bottom)."""
        st = self._shape_style(fill, stroke, sw, rounded, dashed, dash, opacity, fill_opacity, stroke_opacity, shadow,
                               gradient, gradient_dir)
        st.update(self._font_style(font, size, color, bold, italic, underline, align, valign, wrap, spacing))
        st.update(extra or {})
        base = SHAPES.get(shape, shape)
        return self._add_vertex(id, x, y, w, h, base, st, self._label(label, raw))

    def shape(self, kind, x, y, w, h, label="", **kw):
        """Any named shape: star, triangle, rhombus, hexagon, cylinder, document, cloud, ... (see SHAPES)
        or a raw draw.io base style such as 'shape=mxgraph.aws4.lambda'."""
        return self.rect(x, y, w, h, label, shape=kind, **kw)

    def ellipse(self, x, y, w, h, label="", **kw):
        return self.rect(x, y, w, h, label, shape="ellipse", **kw)

    def circle(self, cx, cy, r, label="", **kw):
        return self.rect(cx - r, cy - r, 2 * r, 2 * r, label, shape="circle", **kw)

    def dot(self, cx, cy, r=4, fill="neutral", stroke="none", **kw):
        """Small marker (scatter point)."""
        kw.setdefault("sw", "thin")
        return self.circle(cx, cy, r, fill=fill, stroke=stroke, **kw)

    def star(self, cx, cy, size, fill="secondary", stroke="primary", **kw):
        return self.shape("star", cx - size / 2, cy - size / 2, size, size, fill=fill, stroke=stroke, **kw)

    def text(self, s, x, y, *, w=None, h=None, anchor="center", size="body", font=None, color="text",
             bold=False, italic=False, underline=False, align=None, valign="middle", wrap=False, rotation=0,
             opacity=None, id=None, raw=False, pad=4, extra=None):
        """Free text. (x, y) is the anchor point: center|n|s|e|w|ne|nw|se|sw. Auto-sizes if w/h omitted.
        bold/italic apply to the whole label; use **..** / *..* for partial styling."""
        value = self._label(s, raw)
        fam, sz = self.theme.font_family(font), self.theme.size(size)
        lines = html_lines(value)
        if bold:
            for l in lines:
                for sg in l: sg["bold"] = True
        if w is None:
            w = max((measure_line(l, fam, sz) for l in lines), default=0) + pad
        if h is None:
            h = len(lines) * sz * self.theme.line_height + 2
        ax = {"w": 0, "nw": 0, "sw": 0, "e": 1, "ne": 1, "se": 1}.get(anchor, 0.5)
        ay = {"n": 0, "nw": 0, "ne": 0, "s": 1, "sw": 1, "se": 1}.get(anchor, 0.5)
        align = align or {0: "left", 1: "right"}.get(ax, "center")
        st = self._font_style(font, size, color, bold, italic, underline, align, valign, wrap, spacing=0)
        st.update(strokeColor="none", fillColor="none")
        if rotation:
            st["rotation"] = rotation
        if opacity is not None:
            st["textOpacity"] = self.theme.opacity(opacity)
        st.update(extra or {})
        return self._add_vertex(id, x - ax * w, y - ay * h, w, h, "text", st, value)

    def icon(self, glyph, cx, cy, size=24, **kw):
        """Emoji / unicode glyph used as an icon placeholder (draw.io renders it as text)."""
        kw.setdefault("raw", True)
        return self.text(glyph, cx, cy, size=size, w=size * 1.4, h=size * 1.4, **kw)

    def equation(self, latex, x, y, *, w=None, h=None, anchor="center", size="body", color="text",
                 align=None, opacity=None, id=None, pad=8, extra=None):
        """LaTeX equation as a native draw.io math label (MathJax; editable in draw.io).
        Enables math for the diagram. Auto-sizes with matplotlib mathtext when w/h are omitted.
        Write plain LaTeX without delimiters, e.g. r"s(q,d)=\\lambda\\,\\mathrm{sim}(\\mathbf{q},\\mathbf{d})".
        To annotate individual terms, place each term as its own equation() (positions are then known)."""
        self.math_enabled = True
        sz = self.theme.size(size)
        if w is None or h is None:
            mw, mh = measure_math(latex, sz)
            w = w or mw * 1.08 + pad
            h = h or mh + pad
        ax = {"w": 0, "nw": 0, "sw": 0, "e": 1, "ne": 1, "se": 1}.get(anchor, 0.5)
        ay = {"n": 0, "nw": 0, "ne": 0, "s": 1, "sw": 1, "se": 1}.get(anchor, 0.5)
        align = align or {0: "left", 1: "right"}.get(ax, "center")
        st = self._font_style(None, size, color, align=align, valign="middle", wrap=False, spacing=0)
        st.update(strokeColor="none", fillColor="none")
        if opacity is not None:
            st["textOpacity"] = self.theme.opacity(opacity)
        st.update(extra or {})
        value = "$$" + html.escape(latex, quote=False) + "$$"
        return self._add_vertex(id or "eq", x - ax * w, y - ay * h, w, h, "text", st, value)

    # ---- edges ----------------------------------------------------------------------
    _HEADS = {"block": ("block", 1), "classic": ("classic", 1), "open": ("open", 0), "oval": ("oval", 1),
              "diamond": ("diamond", 1), "openblock": ("block", 0), "dash": ("dash", 0), "none": ("none", 0)}

    def edge(self, pts=(), src=None, tgt=None, *, color="text", sw="normal", arrow="end", head="block",
             head_size=6, dashed=False, dash=None, route="straight", exit=None, entry=None, label=None,
             label_size="label", opacity=None, id=None, extra=None):
        """Connector/line. pts lists the points draw.io cannot compute, in source px:
            [start (only if no src), *waypoints, end (only if no tgt)].
        e.g. edge([(0,0), (50,0), (50,40)]) free polyline; edge([(50,0)], src="a", tgt="b") one
        waypoint between attached shapes; edge([(10,10)], tgt="b") free start, attached end.
        route: straight | orthogonal | rounded | curved (spline through waypoints).
        exit/entry: (fx, fy) fractions on the source/target bounds, e.g. (1, 0.5) = right middle."""
        t = self.theme
        pts = [tuple(map(float, p)) for p in pts]
        st = {"html": "1", "strokeColor": t.color(color), "strokeWidth": t.stroke(sw)}
        st.update({"straight": {"edgeStyle": "none", "rounded": 0},
                   "orthogonal": {"edgeStyle": "orthogonalEdgeStyle", "rounded": 0},
                   "rounded": {"edgeStyle": "orthogonalEdgeStyle", "rounded": 1},
                   "curved": {"edgeStyle": "none", "curved": 1}}[route])
        end_h, start_h = (head if arrow in ("end", "both") else "none"), (head if arrow in ("start", "both") else "none")
        for side, h in (("end", end_h), ("start", start_h)):
            name, fill = self._HEADS[h]
            st[f"{side}Arrow"] = name
            st[f"{side}Fill"] = fill
            st[f"{side}Size"] = head_size
        if dashed:
            st.update(dashed=1, dashPattern=dash or "6 4")
        if opacity is not None:
            st["opacity"] = t.opacity(opacity)
        if exit:
            st.update(exitX=exit[0], exitY=exit[1], exitDx=0, exitDy=0)
        if entry:
            st.update(entryX=entry[0], entryY=entry[1], entryDx=0, entryDy=0)
        value = ""
        if label:
            st.update(self._font_style(None, label_size, "text", wrap=False))
            st["labelBackgroundColor"] = t.color("background")
            value = self._label(label, False)
        st.update(extra or {})
        el = dict(kind="edge", id=self._new_id(id), parent=None, style=st, tokens="", value=value,
                  src=src, tgt=tgt, pts=pts)
        self.els.append(el)
        return el["id"]

    def line(self, pts, **kw):
        kw.setdefault("arrow", None)
        return self.edge(pts, **kw)

    def arrow(self, pts, **kw):
        return self.edge(pts, arrow="end", **kw)

    def flex_arrow(self, pts, *, width=14, end_width=None, end_size=7, fill="primary", stroke=None,
                   curved=True, opacity=None, id=None):
        """Wide 'block' arrow (draw.io flexArrow). With curved=True and waypoints it becomes a wavy band."""
        t = self.theme
        st = {"html": "1", "shape": "flexArrow", "width": width, "endWidth": end_width or width,
              "endSize": end_size, "fillColor": t.color(fill), "strokeColor": t.color(stroke or fill)}
        st.update({"curved": 1, "edgeStyle": "none"} if curved else {"edgeStyle": "none"})
        if opacity is not None:
            st["opacity"] = t.opacity(opacity)
        el = dict(kind="edge", id=self._new_id(id), parent=None, style=st, tokens="", value="",
                  src=None, tgt=None, pts=[tuple(map(float, p)) for p in pts], scale_keys=("width", "endWidth", "endSize"))
        self.els.append(el)
        return el["id"]

    # ---- free-form geometry (embedded SVG) -------------------------------------------
    def _svg_paint(self, v, default="none"):
        return default if v in (None, "none") else self.theme.color(v)

    def svg_layer(self, polys, id=None, pad=2):
        """Embed free-form geometry as one SVG image element.
        polys: list of dicts with keys
            pts (list of (x,y)) | circle=(cx,cy,r)
            closed=True, smooth=False (Catmull-Rom through pts)
            fill (role/hex), fill_opacity (percent or role), stroke, stroke_opacity, sw, dash ("6 4")."""
        xs, ys = [], []
        for p in polys:
            if "circle" in p:
                cx, cy, r = p["circle"]; xs += [cx - r, cx + r]; ys += [cy - r, cy + r]
            else:
                pts = catmull_rom(p["pts"], closed=p.get("closed", True)) if p.get("smooth") else p["pts"]
                p["_pts"] = pts
                xs += [q[0] for q in pts]; ys += [q[1] for q in pts]
        x0, y0, x1, y1 = min(xs) - pad, min(ys) - pad, max(xs) + pad, max(ys) + pad
        body = []
        for p in polys:
            fill = self._svg_paint(p.get("fill"))
            stroke = self._svg_paint(p.get("stroke"))
            fo = (self.theme.opacity(p["fill_opacity"]) / 100) if p.get("fill_opacity") is not None else 1
            so = (self.theme.opacity(p["stroke_opacity"]) / 100) if p.get("stroke_opacity") is not None else 1
            sw = self.theme.stroke(p.get("sw", "normal"))
            dash = f' stroke-dasharray="{p["dash"]}"' if p.get("dash") else ""
            attrs = (f'fill="{fill}" fill-opacity="{fo:.3f}" stroke="{stroke}" stroke-opacity="{so:.3f}" '
                     f'stroke-width="{sw}" stroke-linejoin="round" stroke-linecap="round"{dash}')
            if "circle" in p:
                cx, cy, r = p["circle"]
                body.append(f'<circle cx="{cx - x0:.2f}" cy="{cy - y0:.2f}" r="{r:.2f}" {attrs}/>')
            else:
                d = " ".join(f"{q[0] - x0:.2f},{q[1] - y0:.2f}" for q in p["_pts"])
                tag = "polygon" if p.get("closed", True) else "polyline"
                if tag == "polyline":
                    attrs = attrs.replace(f'fill="{fill}"', 'fill="none"') if not p.get("fill") else attrs
                body.append(f'<{tag} points="{d}" {attrs}/>')
        W, H = x1 - x0, y1 - y0
        svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="{W:.1f}" height="{H:.1f}" '
               f'viewBox="0 0 {W:.2f} {H:.2f}">' + "".join(body) + "</svg>")
        b64 = base64.b64encode(svg.encode()).decode()
        st = {"shape": "image", "html": "1", "imageAspect": 0, "aspect": "fixed",
              "image": f"data:image/svg+xml,{b64}"}  # NOTE: no ';base64' — draw.io style syntax
        return self._add_vertex(id or "svg", x0, y0, W, H, "", st)

    def polygon(self, pts, *, fill="neutral", stroke="none", fill_opacity=None, stroke_opacity=None,
                sw="normal", smooth=False, closed=True, dash=None, id=None):
        return self.svg_layer([dict(pts=pts, fill=fill, stroke=stroke, fill_opacity=fill_opacity,
                                    stroke_opacity=stroke_opacity, sw=sw, smooth=smooth, closed=closed,
                                    dash=dash)], id=id)

    def surface_mesh(self, back, front, *, nu=30, nv=14, bulge=0.0, fill="secondary", fill_opacity="tint",
                     line="primary", line_opacity=55, line_width="hairline", outline="primary", id=None):
        """Pseudo-3-D surface patch (e.g. a manifold) between two edge curves, drawn as a mesh.
        back/front: control points from the left tip to the right tip (same direction).
        bulge: px to lift the middle of the patch (negative y) for a 3-D look."""
        N = 80
        B = _resample(catmull_rom(back), N + 1)
        F = _resample(catmull_rom(front), N + 1)

        def P(i, v):
            u = i / N
            x = (1 - v) * B[i][0] + v * F[i][0]
            y = (1 - v) * B[i][1] + v * F[i][1] - bulge * math.sin(math.pi * v) * math.sin(math.pi * u)
            return (x, y)
        outline_pts = [P(i, 0) for i in range(N + 1)] + [P(N - i, 1) for i in range(N + 1)]
        polys = [dict(pts=outline_pts, fill=fill, fill_opacity=fill_opacity)]
        for k in range(1, nv):
            polys.append(dict(pts=[P(i, k / nv) for i in range(N + 1)], closed=False, stroke=line,
                              stroke_opacity=line_opacity, sw=line_width))
        for k in range(1, nu):
            i = round(N * k / nu)
            polys.append(dict(pts=[P(i, j / 40) for j in range(41)], closed=False, stroke=line,
                              stroke_opacity=line_opacity, sw=line_width))
        if outline:
            polys.append(dict(pts=outline_pts, stroke=outline, stroke_opacity=80, sw="thin"))
        return self.svg_layer(polys, id=id or "surface")

    # ---- raster content ---------------------------------------------------------------
    def _source(self):
        if self._src_img is None:
            from PIL import Image
            if not self.source:
                raise ValueError("Diagram has no source image (use Diagram.from_source)")
            self._src_img = Image.open(self.source).convert("RGB")
        return self._src_img

    def crop(self, bbox, upscale=2):
        from PIL import Image
        x0, y0, x1, y1 = map(int, bbox)
        im = self._source().crop((x0, y0, x1, y1))
        if upscale and upscale != 1:
            im = im.resize((max(1, int((x1 - x0) * upscale)), max(1, int((y1 - y0) * upscale))), Image.LANCZOS)
        return im

    def save_crop(self, bbox, path, upscale=2):
        self.crop(bbox, upscale).save(path)
        return path

    def image_crop(self, bbox, x=None, y=None, w=None, h=None, *, inset=1, upscale=2, fmt="jpeg",
                   quality=85, border=None, id=None):
        """Embed a region of the source image (photo, logo, complex illustration).
        bbox=(x0,y0,x1,y1) in source px; placed at the same spot unless x/y/w/h given.
        inset trims a few px so the source's own border isn't duplicated."""
        x0, y0, x1, y1 = bbox
        im = self.crop((x0 + inset, y0 + inset, x1 - inset, y1 - inset), upscale)
        return self._embed_pil(im, x0 if x is None else x, y0 if y is None else y,
                               (x1 - x0) if w is None else w, (y1 - y0) if h is None else h,
                               fmt, quality, border, id or "image")

    def image_file(self, path, x, y, w, h, *, border=None, id=None):
        from PIL import Image
        im = Image.open(path)
        fmt = "png" if im.mode in ("RGBA", "LA", "P") else "jpeg"
        return self._embed_pil(im.convert("RGBA" if fmt == "png" else "RGB"), x, y, w, h, fmt, 90, border, id)

    def _embed_pil(self, im, x, y, w, h, fmt, quality, border, id):
        buf = io.BytesIO()
        if fmt == "jpeg":
            im.convert("RGB").save(buf, "JPEG", quality=quality)
        else:
            im.save(buf, "PNG", optimize=True)
        b64 = base64.b64encode(buf.getvalue()).decode()
        st = {"shape": "image", "html": "1", "imageAspect": 0, "image": f"data:image/{fmt},{b64}"}
        if border:
            st["imageBorder"] = self.theme.color(border)
        return self._add_vertex(id, x, y, w, h, "", st)

    # ---- composites -------------------------------------------------------------------
    def callout(self, x, y, w, h, label, tip, *, base_width=40, base_at=None, fill="neutral",
                fill_opacity="soft", stroke="text", sw="normal", rounded=True, id=None, **font):
        """Speech-bubble box with a pointer toward tip=(tx,ty). The pointer leaves from the box side
        facing the tip; base_at = position along that side in px (default: nearest to tip)."""
        tx, ty = tip
        side = ("top" if ty < y else "bottom" if ty > y + h else "left" if tx < x else "right")
        horiz = side in ("top", "bottom")
        lo, hi = (x, x + w) if horiz else (y, y + h)
        c = base_at if base_at is not None else min(max(tx if horiz else ty, lo + base_width / 2 + 8), hi - base_width / 2 - 8)
        a, b = c - base_width / 2, c + base_width / 2
        edge_coord = {"top": y, "bottom": y + h, "left": x, "right": x + w}[side]
        inward = {"top": 1, "bottom": -1, "left": 1, "right": -1}[side] * (self.theme.stroke(sw) + 0.5)
        if horiz:
            p0, p1 = (a, edge_coord + inward), (b, edge_coord + inward)
        else:
            p0, p1 = (edge_coord + inward, a), (edge_coord + inward, b)
        font.setdefault("size", "body")
        with self.group(id or "callout") as gid:
            self.rect(x, y, w, h, label, fill=fill, fill_opacity=fill_opacity, stroke=stroke, sw=sw,
                      rounded=rounded, id=f"{gid}_box", spacing=6, **font)
            self.svg_layer([dict(pts=[p0, tip, p1], fill="background"),
                            dict(pts=[p0, tip, p1], fill=fill, fill_opacity=fill_opacity),
                            dict(pts=[p0, tip, p1], closed=False, stroke=stroke, sw=sw)], id=f"{gid}_pointer")
        return gid

    def panel(self, x, y, w, h, title=None, *, header_h=None, header_fill="primary", header_color="background",
              title_size="heading", title_align="center", fill="background", stroke="text", sw="normal",
              rounded=False, header_band=True, header_opacity=None, shadow=None, id=None, **font):
        """Container with an optional header band (or a plain title line if header_band=False).
        Returns (panel_id, (cx, cy, cw, ch)) where the tuple is the content area below the header."""
        pid = self.rect(x, y, w, h, fill=fill, stroke=stroke, sw=sw, rounded=rounded, shadow=shadow, id=id or "panel")
        hh = header_h or (self.theme.size(title_size) * 2.0 if title else 0)
        if title:
            if header_band:
                self.rect(x, y, w, hh, title, fill=header_fill, stroke=header_fill if stroke != "none" else "none",
                          sw=sw, fill_opacity=header_opacity, color=header_color, size=title_size, bold=True,
                          align=title_align, id=f"{pid}_header", spacing=8, **font)
            else:
                pad = self.theme.data["spacing"]["padding"]
                self.rect(x + pad, y, w - 2 * pad, hh, title, fill="none", stroke="none", color=header_color,
                          size=title_size, bold=True, align=title_align, id=f"{pid}_title", **font)
        return pid, (x, y + hh, w, h - hh)

    # ---- layout helpers -------------------------------------------------------------
    def gap(self):
        return self.theme.data["spacing"]["gap"]

    def row(self, x, y, n, w, h, gap=None):
        g = self.gap() if gap is None else gap
        return [(x + i * (w + g), y, w, h) for i in range(n)]

    def row_widths(self, x, y, widths, h, gap=None):
        """Boxes of varying widths left-to-right with one constant gap."""
        g = self.gap() if gap is None else gap
        out, cx = [], x
        for w in widths:
            out.append((cx, y, w, h)); cx += w + g
        return out

    def column(self, x, y, n, w, h, gap=None):
        g = self.gap() if gap is None else gap
        return [(x, y + i * (h + g), w, h) for i in range(n)]

    def fit_row(self, x0, x1, y, n, h, gap=None):
        """n equal boxes spanning [x0, x1] with equal gaps."""
        g = self.gap() if gap is None else gap
        w = (x1 - x0 - g * (n - 1)) / n
        return self.row(x0, y, n, w, h, g)

    def grid_boxes(self, x, y, cols, rows, w, h, gx=None, gy=None):
        gx = self.gap() if gx is None else gx
        gy = gx if gy is None else gy
        return [(x + c * (w + gx), y + r * (h + gy), w, h) for r in range(rows) for c in range(cols)]

    @staticmethod
    def snap(v, step=2):
        return round(v / step) * step

    # ---- scatter helpers --------------------------------------------------------------
    @staticmethod
    def scatter_in_polygon(poly, n, keepout=(), min_dist=8, seed=0, max_tries=200000):
        """n random points inside poly, outside keepout rects (x0,y0,x1,y1), >= min_dist apart."""
        rnd = random.Random(seed)
        xs, ys = [p[0] for p in poly], [p[1] for p in poly]
        pts, tries = [], 0
        while len(pts) < n and tries < max_tries:
            tries += 1
            px, py = rnd.uniform(min(xs), max(xs)), rnd.uniform(min(ys), max(ys))
            if not point_in_polygon(px, py, poly):
                continue
            if any(a <= px <= c and b <= py <= d for a, b, c, d in keepout):
                continue
            if all((px - qx) ** 2 + (py - qy) ** 2 >= min_dist ** 2 for qx, qy in pts):
                pts.append((px, py))
        return pts

    @staticmethod
    def scatter_along(fn, x0, x1, n, spread=10, min_dist=8, seed=0, max_tries=100000):
        """n points near the curve y=fn(x), x in [x0, x1], gaussian spread in y."""
        rnd = random.Random(seed)
        pts, tries = [], 0
        while len(pts) < n and tries < max_tries:
            tries += 1
            x = rnd.uniform(x0, x1); y = fn(x) + rnd.gauss(0, spread)
            if all((x - a) ** 2 + (y - b) ** 2 >= min_dist ** 2 for a, b in pts):
                pts.append((x, y))
        return pts

    @staticmethod
    def scatter_blob(cx, cy, n, sx=20, sy=None, min_dist=8, seed=0, max_tries=100000):
        rnd = random.Random(seed)
        sy = sx if sy is None else sy
        pts, tries = [], 0
        while len(pts) < n and tries < max_tries:
            tries += 1
            x, y = rnd.gauss(cx, sx), rnd.gauss(cy, sy)
            if all((x - a) ** 2 + (y - b) ** 2 >= min_dist ** 2 for a, b in pts):
                pts.append((x, y))
        return pts

    # ---- serialization ----------------------------------------------------------------
    def _style_str(self, tokens, st, scale_keys=()):
        parts = [t for t in tokens.split(";") if t] if tokens else []
        for k, v in st.items():
            if k in ("fontSize", "arcSize") or k in scale_keys:
                v = round(float(v) * self.S, 2)
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            parts.append(f"{k}={v}")
        return ";".join(parts) + ";"

    def _compute_group_bboxes(self):
        kids = {}
        for e in self.els:
            if e.get("parent"):
                kids.setdefault(e["parent"], []).append(e)
        groups = [e for e in self.els if e["kind"] == "group"]

        def bb(g):
            if "geo" in g:
                return g["geo"]
            boxes = []
            for c in kids.get(g["id"], []):
                if c["kind"] == "group":
                    if bb(c): boxes.append(bb(c))
                elif c["kind"] == "vertex":
                    boxes.append(c["geo"])
            if not boxes:
                g["geo"] = None; return None
            x0 = min(b[0] for b in boxes); y0 = min(b[1] for b in boxes)
            x1 = max(b[0] + b[2] for b in boxes); y1 = max(b[1] + b[3] for b in boxes)
            g["geo"] = [x0, y0, x1 - x0, y1 - y0]
            return g["geo"]
        for g in groups:
            bb(g)

    def to_xml(self):
        self._compute_group_bboxes()
        S = self.S
        by_id = {e["id"]: e for e in self.els}
        cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']

        def _r(v):
            v = round(v, 2)
            return int(v) if float(v).is_integer() else v

        def fx(v):
            return _r(v * S * self.cs[0])

        def fy(v):
            return _r(v * S * self.cs[1])
        f = fx
        for e in self.els:
            if e["kind"] == "group" and not e.get("geo"):
                continue
            parent = e.get("parent") or "1"
            if e["kind"] in ("vertex", "group"):
                x, y, w, h = e["geo"]
                if e.get("parent"):
                    px, py = by_id[e["parent"]]["geo"][:2]
                    x, y = x - px, y - py
                if e["kind"] == "group":
                    style = "group;"
                    extra = ' connectable="0"'
                    value = ""
                else:
                    style = self._style_str(e["tokens"], e["style"])
                    extra = ""
                    value = e["value"]
                cells.append(f'<mxCell id="{e["id"]}" value={quoteattr(value)} style={quoteattr(style)} '
                             f'vertex="1"{extra} parent="{parent}"><mxGeometry x="{fx(x)}" y="{fy(y)}" '
                             f'width="{fx(w)}" height="{fy(h)}" as="geometry"/></mxCell>')
            else:
                style = self._style_str("", e["style"], e.get("scale_keys", ()))
                a = f'id="{e["id"]}" value={quoteattr(e["value"])} style={quoteattr(style)} edge="1" parent="1"'
                if e["src"]: a += f' source="{e["src"]}"'
                if e["tgt"]: a += f' target="{e["tgt"]}"'
                pts = list(e["pts"])
                g = ""
                # pts = only the points draw.io cannot compute: a free start (no src), waypoints,
                # and a free end (no tgt). With src/tgt given, the remaining points are waypoints.
                if pts and not e["src"]:
                    g += f'<mxPoint x="{fx(pts[0][0])}" y="{fy(pts[0][1])}" as="sourcePoint"/>'
                    pts = pts[1:]
                if pts and not e["tgt"]:
                    g += f'<mxPoint x="{fx(pts[-1][0])}" y="{fy(pts[-1][1])}" as="targetPoint"/>'
                    pts = pts[:-1]
                way = pts
                if way:
                    g += '<Array as="points">' + "".join(f'<mxPoint x="{fx(x)}" y="{fy(y)}"/>' for x, y in way) + "</Array>"
                cells.append(f'<mxCell {a}><mxGeometry relative="1" as="geometry">{g}</mxGeometry></mxCell>')
        bg = self.theme.color(self.background) if self.background else "none"
        did = re.sub(r"[^\w]", "_", self.name)
        return ('<mxfile host="app.diagrams.net" agent="image-to-drawio">'
                f'<diagram id="{did}" name={quoteattr(self.name)}>'
                f'<mxGraphModel dx="{_r(self.width * S)}" dy="{_r(self.height * S)}" grid="{1 if self.grid else 0}" gridSize="10" '
                'guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" '
                f'pageWidth="{_r(self.width * S)}" pageHeight="{_r(self.height * S)}" background="{bg}" math="{1 if self.math_enabled else 0}" shadow="0">'
                "<root>" + "".join(cells) + "</root></mxGraphModel></diagram></mxfile>")

    def save(self, path, quiet=False):
        xml = self.to_xml()
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(xml)
        if not quiet:
            n_v = sum(e["kind"] == "vertex" for e in self.els)
            n_e = sum(e["kind"] == "edge" for e in self.els)
            print(f"wrote {path}  ({n_v} vertices, {n_e} edges, scale={self.S})")
            if self.theme.off_theme:
                print("  off-theme values used:", ", ".join(sorted(self.theme.off_theme)))
        return path


if __name__ == "__main__":
    print(__doc__)
'''

# ================================================================================================
FILES['scripts/enhance.py'] = r'''#!/usr/bin/env python3
"""
enhance.py — deterministic clean-up passes on ANY .drawio (run after structural edits).

    python enhance.py fig.drawio [-o "fig#ENHANCED.drawio"] --theme theme.yaml --brief brief.yaml \
        --colors --fonts --sizes --strokes --align --attach --fit-page 16 --math [--map "#d6d6d6=neutral"]

Passes (each logs what it changed to stdout and, with --log, appends to a CHANGES.md):
  --colors        snap near-identical colors (CIELAB dE <= --max-de, default 10) to theme colors,
                  express tints as theme role + fillOpacity, keep distinct colors (never erased);
                  --map "#hex=role" or "#hex=role@25" (fill at 25 % opacity) forces a mapping — use it to
                  consolidate a palette deliberately; --add-roles writes kept colors into the theme
  --fonts         set every non-math label to the theme's default font family (or --family NAME)
  --sizes         merge near-identical font sizes into consistent levels (--to-theme: snap levels to
                  theme roles), then raise anything below the brief's minimum print size
  --strokes       snap stroke widths to theme stroke roles; unify arrow-head sizes (--head 6)
  --align         snap near-miss alignments (<= --tol px, default 3) among comparable shapes
  --attach        attach free edge ends lying within --attach-tol px (default 8) of a shape border
  --fit-page M    move content to (M, M) and set the page to content + 2M
  --math          turn the diagram math flag on if any label is $$LaTeX$$
Default output: "<stem>#ENHANCED.drawio" next to the input (suffix configurable with --suffix).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.dont_write_bytecode = True
from audit_drawio import hex2rgb, is_hex, load_yaml, to_lab  # noqa: E402
from drawio_doc import Doc  # noqa: E402

COLOR_KEYS = ("fillColor", "strokeColor", "fontColor", "gradientColor", "imageBorder", "labelBackgroundColor")


def enhanced_path(path, suffix="#ENHANCED"):
    stem, ext = os.path.splitext(path)
    if stem.endswith(suffix):
        return path
    return stem + suffix + (ext or ".drawio")


class Log:
    def __init__(self):
        self.lines = []

    def __call__(self, msg):
        self.lines.append(msg)
        print("  " + msg)


def pass_colors(doc, theme, log, close_de=10.0, forced=None, add_roles=False, theme_path=None):
    """Bring colors onto the theme WITHOUT erasing meaning.
    1. explicit --map wins;  2. a color within `close_de` (CIELAB dE) of a theme color -> that color;
    3. a fill that is a tint of a theme color (theme color at 5-70 % opacity over the page) -> that
       role + fillOpacity;  4. anything else is kept and reported; with --add-roles it is written to
       the theme as a named role so the figure is explicitly on-theme. Never maps a distinct color
       to the background."""
    pal = {r: c.lower() for r, c in (theme.get("colors") or {}).items()}
    fop = {}                                      # forced opacity for fills: "#hex=role@25"
    fx = {}
    for k, v in (forced or {}).items():
        role, _, op = v.partition("@")
        fx[k.lower()] = pal.get(role, role).lower()
        if op:
            fop[k.lower()] = float(op)
    forced = fx
    labs = {c: to_lab(hex2rgb(c)) for c in set(pal.values())}
    page_bg = hex2rgb(doc.model.get("background")) if is_hex(doc.model.get("background")) else (1, 1, 1)

    def near(h):
        L = to_lab(hex2rgb(h))
        best = min(labs, key=lambda c: math.dist(L, labs[c])) if labs else None
        return best, (math.dist(L, labs[best]) if best else 1e9)

    tint_cache = {}

    def as_tint(h):
        if h in tint_cache:
            return tint_cache[h]
        tint_cache[h] = _as_tint(h)
        return tint_cache[h]

    def _as_tint(h):
        rgb = hex2rgb(h)
        best = None
        for role_hex in labs:
            base = hex2rgb(role_hex)
            if max(abs(x - y) for x, y in zip(base, page_bg)) < 0.05:
                continue
            for a100 in range(5, 71):
                a = a100 / 100
                mix = tuple(a * x + (1 - a) * y for x, y in zip(base, page_bg))
                d = math.dist(to_lab(rgb), to_lab(mix))
                if best is None or d < best[0]:
                    best = (d, role_hex, a100)
        return best if best and best[0] <= close_de else None
    mapped = tinted = 0
    kept = {}
    for c in doc.cells:
        for k in COLOR_KEYS:
            v = c.get(k)
            if not is_hex(v):
                continue
            v = v.lower()
            if v in forced:
                c.set_style(**{k: forced[v]}); mapped += 1
                if k == "fillColor" and v in fop:
                    c.set_style(fillOpacity=round(fop[v] * float(c.get("fillOpacity", 100)) / 100))
                continue
            if v in labs:
                continue
            best, dist = near(v)
            if best and dist <= close_de:
                c.set_style(**{k: best}); mapped += 1; continue
            if k == "fillColor":
                t = as_tint(v)
                if t:
                    prev = float(c.get("fillOpacity", 100)) / 100
                    c.set_style(fillColor=t[1], fillOpacity=round(t[2] * prev)); tinted += 1; continue
            kept[v] = kept.get(v, 0) + 1
        val = c.value
        nv = re.sub(r'(color="?)(#[0-9a-fA-F]{3,6})',
                    lambda m: m.group(1) + (forced.get(m.group(2).lower()) or (near(m.group(2))[0] if near(m.group(2))[1] <= close_de else m.group(2))), val)
        if nv != val:
            c.value = nv; mapped += 1
    bg = (doc.model.get("background") or "").lower()
    if is_hex(bg) and bg not in labs:
        nb = forced.get(bg) or (near(bg)[0] if near(bg)[1] <= close_de else None)
        if nb:
            doc.model.set("background", nb); mapped += 1
        else:
            kept[bg] = kept.get(bg, 0) + 1
    if add_roles and kept and theme_path and theme_path.endswith((".yaml", ".yml")):
        import yaml
        t = yaml.safe_load(open(theme_path, encoding="utf-8"))
        cols = t.setdefault("colors", {})
        existing = {v.lower() for v in cols.values()}
        n = 1
        for h in sorted(kept):
            if h in existing:
                continue
            while f"extra_{n}" in cols:
                n += 1
            cols[f"extra_{n}"] = h; n += 1
        yaml.safe_dump(t, open(theme_path, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)
        log(f"colors: added {len(kept)} figure-specific colors to the theme as extra_* roles (rename them semantically)")
    log(f"colors: {mapped} snapped to theme colors, {tinted} fills expressed as theme role + opacity"
        + (f"; kept {len(kept)} distinct colors not in the theme: {', '.join(sorted(kept))}" if kept else ""))


def pass_fonts(doc, theme, log, family=None):
    fonts = theme.get("fonts") or {}
    key = theme.get("default_font", "sans")
    fam = family or (fonts.get(key) or {}).get("family", "Helvetica")
    n = 0
    for c in doc.vertices():
        if c.plain and not c.is_math and c.get("fontFamily") != fam:
            c.set_style(fontFamily=fam); n += 1
        if c.value and "font-family" in c.value:
            c.value = re.sub(r"font-family:[^;\"]+;?", "", c.value)
    for e in doc.edges():
        if e.plain and e.get("fontFamily") != fam:
            e.set_style(fontFamily=fam); n += 1
    log(f"fonts: {n} labels set to '{fam}'")


def pass_sizes(doc, theme, brief, log, to_theme=False, merge=0.12, theme_path=None):
    """Harmonize font sizes: merge sizes within `merge` (12 %) into one level (the most used size),
    optionally snap levels to theme size roles (--to-theme: only when the figure was built with this
    theme's scale), then raise anything below the brief's minimum print size."""
    cells = [c for c in doc.cells if c.plain and "fontSize" in c.style]
    count = {}
    for c in cells:
        fs = float(c.get("fontSize")); count[fs] = count.get(fs, 0) + 1
    levels, mapping = [], {}
    for fs in sorted(count):
        if levels and fs <= levels[-1][0] * (1 + merge):
            levels[-1][1].append(fs)
        else:
            levels.append([fs, [fs]])
    for _, members in levels:
        rep = max(members, key=lambda m: (count[m], m))
        for m in members:
            mapping[m] = rep
    roles = sorted({float(v) for v in (theme.get("sizes") or {}).values()})
    if to_theme and roles:
        mapping = {k: min(roles, key=lambda r: abs(r - v)) for k, v in mapping.items()}
    pt_per_px = brief.get("width_in", 7.0) * 72 / max(doc.page[0], 1)
    min_px = brief.get("min_font_pt", 7.0) / pt_per_px
    changed = raised = 0
    bumped = {}
    for c in cells:
        fs = float(c.get("fontSize"))
        nfs = mapping.get(fs, fs)
        if nfs < min_px:
            bumped[nfs] = math.ceil(min_px * 2) / 2
            nfs = bumped[nfs]; raised += 1
        if abs(nfs - fs) > 0.01:
            c.set_style(fontSize=round(nfs, 2)); changed += 1
    if bumped and theme_path and theme_path.endswith((".yaml", ".yml")):
        # keep the theme consistent with the figure: raise the affected size roles too
        import yaml
        t = yaml.safe_load(open(theme_path, encoding="utf-8"))
        for role, val in list((t.get("sizes") or {}).items()):
            if float(val) in bumped:
                t["sizes"][role] = bumped[float(val)]
        yaml.safe_dump(t, open(theme_path, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)
        log(f"sizes: theme roles updated for the print minimum: {bumped}")
    lv = sorted({mapping[k] for k in mapping})
    log(f"sizes: {len(count)} sizes -> {len(lv)} levels {lv}; {changed} labels changed, {raised} raised to the "
        f"{brief.get('min_font_pt', 7)} pt minimum ({min_px:.1f} px at {brief.get('width_in', 7.0)} in)")


def pass_strokes(doc, theme, log, head=None):
    roles = sorted({float(v) for v in (theme.get("strokes") or {}).values()})
    n = h = 0
    for c in doc.cells:
        if "strokeWidth" in c.style and roles:
            sw = float(c.get("strokeWidth"))
            ns = min(roles, key=lambda r: abs(r - sw))
            if abs(ns - sw) > 0.01:
                c.set_style(strokeWidth=ns); n += 1
        if c.is_edge and head and c.get("endArrow", "classic") not in ("none",) and c.get("shape") != "flexArrow":
            if c.get("endSize") != str(head):
                c.set_style(endSize=head); h += 1
    log(f"strokes: {n} widths snapped to theme roles" + (f", {h} arrow heads set to {head}" if head else ""))


def pass_align(doc, log, tol=3.0):
    shapes = [c for c in doc.vertices() if not c.is_text and not c.is_image and c.abs_geo() and c.abs_geo()[2] >= 14 and c.abs_geo()[3] >= 14]
    moved = 0
    for key in ("center-y", "center-x", "top", "left"):
        f = {"left": lambda g: g[0], "top": lambda g: g[1], "center-x": lambda g: g[0] + g[2] / 2,
             "center-y": lambda g: g[1] + g[3] / 2}[key]
        horiz = key in ("left", "center-x")
        parent = {c.id: c.id for c in shapes}

        def find(i):
            while parent[i] != i:
                parent[i] = parent[parent[i]]; i = parent[i]
            return i
        for i in range(len(shapes)):
            for j in range(i + 1, len(shapes)):
                a, b = shapes[i].abs_geo(), shapes[j].abs_geo()
                if not (0.01 < abs(f(a) - f(b)) <= tol):
                    continue
                ratio = max(a[2], a[3]) / max(b[2], b[3])
                gap = (max(0, max(a[1], b[1]) - min(a[1] + a[3], b[1] + b[3])) if horiz
                       else max(0, max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2])))
                contains = lambda p, q: p[0] <= q[0] and p[1] <= q[1] and p[0] + p[2] >= q[0] + q[2] and p[1] + p[3] >= q[1] + q[3]  # noqa: E731
                if 1 / 3 <= ratio <= 3 and gap < 160 and not contains(a, b) and not contains(b, a):
                    parent[find(shapes[i].id)] = find(shapes[j].id)
        groups = {}
        for c in shapes:
            groups.setdefault(find(c.id), []).append(c)
        for members in groups.values():
            if len(members) < 2:
                continue
            vals = sorted(f(m.abs_geo()) for m in members)
            target = vals[len(vals) // 2]
            for m in members:
                g = m.abs_geo(); d = target - f(g)
                if abs(d) > 0.01:
                    m.set_abs_geo(**({"x": g[0] + d} if horiz else {"y": g[1] + d})); moved += 1
    log(f"align: {moved} near-miss alignments snapped (tol {tol}px)")


def pass_attach(doc, log, tol=8.0):
    verts = [v for v in doc.vertices() if not v.is_text and not v.is_image and v.abs_geo() and v.abs_geo()[2] > 1]
    n = 0
    for e in doc.edges():
        for side, ptag in (("source", "sourcePoint"), ("target", "targetPoint")):
            if e.el.get(side):
                continue
            g_el = e.geo_el
            p = next((p for p in (g_el.findall("mxPoint") if g_el is not None else []) if p.get("as") == ptag), None)
            if p is None:
                continue
            x, y = float(p.get("x", 0)), float(p.get("y", 0))
            best = None
            for v in verts:
                g = v.abs_geo()
                inside = g[0] < x < g[0] + g[2] and g[1] < y < g[1] + g[3]
                dist = math.hypot(max(g[0] - x, 0, x - g[0] - g[2]), max(g[1] - y, 0, y - g[1] - g[3]))
                if not inside and dist <= tol and (best is None or dist < best[0]):
                    best = (dist, v)
            if best:
                e.el.set(side, best[1].id)
                g_el.remove(p)
                e.set_style(**{("exitX" if side == "source" else "entryX"): None, ("exitY" if side == "source" else "entryY"): None})
                n += 1
    log(f"attach: {n} free edge ends attached to the shape they touch")


def pass_fit(doc, log, margin=16.0):
    x, y, w, h = doc.content_bbox()
    dx, dy = margin - x, margin - y
    if abs(dx) > 0.01 or abs(dy) > 0.01:
        doc.translate_all(dx, dy)
    doc.set_page(w + 2 * margin, h + 2 * margin)
    log(f"fit-page: content moved by ({dx:.0f}, {dy:.0f}); page = {w + 2 * margin:.0f} x {h + 2 * margin:.0f}")


def pass_math(doc, log):
    has = any(c.is_math for c in doc.vertices())
    if has and not doc.math:
        doc.math = True
        log("math: diagram math flag turned on")
    else:
        log(f"math: {'on' if doc.math else 'off'} ({sum(c.is_math for c in doc.vertices())} LaTeX labels)")


def run(src, out=None, theme=None, brief=None, **opts):
    doc = Doc.load(src)
    th = load_yaml(theme) if isinstance(theme, str) else (theme or {})
    br = load_yaml(brief) if isinstance(brief, str) else (brief or {})
    br = br.get("target", br)
    log = Log()
    if opts.get("colors"): pass_colors(doc, th, log, opts.get("max_de", 10.0), opts.get("map"), opts.get("add_roles"),
                                       theme if isinstance(theme, str) else None)
    if opts.get("fonts"): pass_fonts(doc, th, log, opts.get("family"))
    if opts.get("sizes"): pass_sizes(doc, th, br, log, opts.get("to_theme", False),
                                     theme_path=theme if isinstance(theme, str) and opts.get("update_theme") else None)
    if opts.get("strokes"): pass_strokes(doc, th, log, opts.get("head"))
    if opts.get("attach"): pass_attach(doc, log, opts.get("attach_tol", 8.0))
    if opts.get("align"): pass_align(doc, log, opts.get("tol", 3.0))
    if opts.get("fit_page") is not None: pass_fit(doc, log, opts["fit_page"])
    if opts.get("math"): pass_math(doc, log)
    out = out or enhanced_path(src, opts.get("suffix", "#ENHANCED"))
    doc.save(out)
    return out, log.lines


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("drawio"); ap.add_argument("-o", "--out")
    ap.add_argument("--theme"); ap.add_argument("--brief")
    ap.add_argument("--suffix", default="#ENHANCED")
    for f in ("colors", "fonts", "sizes", "strokes", "align", "attach", "math", "all"):
        ap.add_argument("--" + f, action="store_true")
    ap.add_argument("--fit-page", type=float, default=None)
    ap.add_argument("--to-theme", action="store_true", help="with --sizes: snap size levels to theme roles")
    ap.add_argument("--update-theme", action="store_true", help="with --sizes: write raised sizes back into the theme")
    ap.add_argument("--add-roles", action="store_true", help="with --colors: add kept colors to the theme as extra_* roles")
    ap.add_argument("--max-de", type=float, default=10.0, help="CIELAB distance treated as 'same color'"); ap.add_argument("--map", action="append", default=[])
    ap.add_argument("--family"); ap.add_argument("--head", type=float); ap.add_argument("--tol", type=float, default=3.0)
    ap.add_argument("--attach-tol", type=float, default=8.0); ap.add_argument("--log", help="append changes to this markdown file")
    a = ap.parse_args()
    if a.all:
        a.colors = a.fonts = a.sizes = a.strokes = a.align = a.attach = a.math = True
        a.fit_page = 16 if a.fit_page is None else a.fit_page
    forced = dict(m.split("=", 1) for m in a.map)
    print("enhance passes:")
    out, lines = run(a.drawio, a.out, a.theme, a.brief, colors=a.colors, fonts=a.fonts, sizes=a.sizes, strokes=a.strokes,
                     align=a.align, attach=a.attach, math=a.math, fit_page=a.fit_page, max_de=a.max_de, map=forced,
                     family=a.family, head=a.head, to_theme=a.to_theme, update_theme=a.update_theme, add_roles=a.add_roles, tol=a.tol, attach_tol=a.attach_tol, suffix=a.suffix)
    if a.log:
        with open(a.log, "a", encoding="utf-8") as fh:
            fh.write(f"\n### enhance.py on `{os.path.basename(a.drawio)}`\n" + "".join(f"- {l}\n" for l in lines))
    print("wrote", out)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/faithful.py'] = r'''"""faithful — label words that do not occur in the source text (shared with text-to-drawio).

    from faithful import faithfulness
    missing = faithfulness("fig.drawio", open("passage.md", encoding="utf-8").read(), allowed=["layer"])
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import html
import re
import xml.etree.ElementTree as ET

FUNCTION_WORDS = set("""a an the and or of to in on for with by from as at is are be via per vs not no all each
its it this that these those than then into over under between within across e.g i.e etc""".split())


def tokens(s):
    s = s.replace("\u2011", "-").replace("\u2013", "-")
    out = []
    for t in re.findall(r"[A-Za-z][A-Za-z'\-]*|\d[\d,]*(?:\.\d+)?", s):
        for part in t.split("-"):
            part = part.strip("'").lower().replace(",", "") if part[:1].isdigit() else part.strip("'").lower()
            if part:
                out.append(part)
    return out


def latex_words(latex):
    inner = " ".join(re.findall(r"\\(?:mathrm|text|operatorname|mathit|textbf|mathbf)\{([^}]*)\}", latex))
    rest = re.sub(r"\\[A-Za-z]+", " ", latex)
    return inner + " " + rest


def label_texts(drawio_path):
    root = ET.parse(drawio_path).getroot()
    out = []
    for c in root.iter("mxCell"):
        v = c.get("value") or ""
        if not v:
            continue
        plain = html.unescape(re.sub(r"<br\s*/?>", "\n", v))
        plain = re.sub(r"<[^>]+>", "", plain)
        m = re.fullmatch(r"\s*\$\$(.+)\$\$\s*", plain, re.S)
        text = latex_words(m.group(1)) if m else re.sub(r"\\\((.+?)\\\)", lambda mm: latex_words(mm.group(1)), plain)
        out.append((c.get("id"), text, plain))
    return out


def stem(w):
    for suf, rep in (("ations", ""), ("ation", ""), ("ings", ""), ("ing", ""), ("ions", ""), ("ion", ""),
                     ("ies", "y"), ("ied", "y"), ("es", ""), ("ed", ""), ("s", ""), ("ly", ""), ("al", "")):
        if w.endswith(suf) and len(w) - len(suf) >= 3:
            return w[: len(w) - len(suf)] + rep
    return w


def faithfulness(drawio_path, passage_text, allowed):
    """Label words that do not occur in the passage (light stemming; acronyms match their expansion)."""
    raw = passage_text + " " + latex_words(passage_text)
    src_list = tokens(raw)
    src = set(src_list)
    stems = {stem(t) for t in src}
    initials = "".join(t[0] for t in src_list if t[0].isalpha())
    allowed = {a.lower() for a in allowed} | FUNCTION_WORDS

    def ok(w):
        if w in src or w in allowed or len(w) <= 2:
            return True
        if w[0].isdigit():
            return w in src
        if stem(w) in stems:
            return True
        if w.isalpha() and 2 <= len(w) <= 6 and w in initials:     # acronym, e.g. llm <- large language model
            return True
        k = min(len(w), 6)
        return any(t[:k] == w[:k] for t in src if len(t) >= k)
    missing = {}
    for cid, text, plain in label_texts(drawio_path):
        for w in tokens(text):
            if not ok(w):
                missing.setdefault(w, set()).add(cid)
    return missing


'''

# ================================================================================================
FILES['scripts/figlib.py'] = r'''"""
figlib — reusable academic-figure components built on drawio_kit.

Every function draws onto a Diagram `d` (source px coordinates, theme roles for colors/sizes) and
returns ids / geometry so you can connect things. Import in build.py:  `from figlib import *`

Style presets (dicts for d.rect(**NODE) etc.):
    NODE   white box, primary outline          — default step / entity
    SOFT   secondary tint, secondary outline   — second category (e.g. online, mechanisms)
    DARK   primary fill, white text            — emphasis (outputs, key results, expert step)
    MUTED  neutral tint, dashed outline        — context / moderators / optional
    GHOST  no fill, dashed neutral outline     — group frames
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import math

from drawio_kit import measure_math

NODE = dict(fill="background", stroke="primary", sw="normal", rounded=True, size="body")
SOFT = dict(fill="secondary", fill_opacity="tint", stroke="secondary", sw="normal", rounded=True, size="body")
DARK = dict(fill="primary", stroke="primary", color="background", rounded=True, size="body")
MUTED = dict(fill="neutral", fill_opacity="tint", stroke="neutral", sw="normal", rounded=True, size="body", dashed=True)
GHOST = dict(fill="none", stroke="neutral", sw="thin", rounded=10, dashed=True)

__all__ = ["NODE", "SOFT", "DARK", "MUTED", "GHOST", "style", "pill", "lane", "chain", "cycle", "layers",
           "frame", "bracket", "spread", "annotated_equation", "rating_dots", "rating_legend", "card",
           "continuum", "legend", "badge", "axes", "table", "connect", "eq"]


def style(base, **over):
    """Copy a preset with overrides: style(NODE, size="label")."""
    s = dict(base); s.update(over); return s


def pill(d, cx, y, text, w=200, h=34, fill="primary", color="background", size="subheading", id=None):
    """Rounded header pill (column / stage heading)."""
    return d.rect(cx - w / 2, y, w, h, f"**{text}**", fill=fill, stroke=fill, color=color, rounded=h / 2, size=size, id=id)


def badge(d, cx, cy, n, r=13, fill="primary", color="background", id=None):
    """Numbered circle (step numbers)."""
    return d.circle(cx, cy, r, f"**{n}**", fill=fill, stroke="background", color=color, size="small", id=id)


def lane(d, x, y, w, h, title, band_w=38, band="primary", wash="neutral", id=None):
    """Swimlane: washed background + rotated title band on the left. Returns content box."""
    d.rect(x, y, w, h, fill=wash, fill_opacity="wash", stroke="neutral", sw="thin", rounded=6, id=id)
    d.rect(x, y, band_w, h, fill=band, stroke=band, rounded=6)
    d.text(f"**{title}**", x + band_w / 2, y + h / 2, size="subheading", color="background", rotation=-90)
    return (x + band_w, y, w - band_w, h)


def connect(d, a, b, side="h", color="primary", sw="thick", **kw):
    """Attach an arrow a->b leaving/entering mid-sides: side 'h' (right->left) or 'v' (bottom->top)."""
    ex, en = ((1, .5), (0, .5)) if side == "h" else ((.5, 1), (.5, 0))
    return d.edge(kw.pop("pts", []), src=a, tgt=b, color=color, sw=sw, exit=ex, entry=en, **kw)


def chain(d, labels, x, y, w, h, gap=30, styles=None, ids=None, direction="h", arrows=True,
          arrow_color="primary", arrow_sw="thick", shapes=None):
    """Row (direction='h') or column ('v') of equal boxes joined by arrows.
    styles: one preset for all or a list per box; shapes: optional list of SHAPES names per box.
    Returns list of ids (in order)."""
    n = len(labels)
    boxes = d.row(x, y, n, w, h, gap) if direction == "h" else d.column(x, y, n, w, h, gap)
    styles = styles if isinstance(styles, list) else [styles or NODE] * n
    ids = ids or [None] * n
    out = []
    for i, ((bx, by, bw, bh), t) in enumerate(zip(boxes, labels)):
        kw = dict(styles[i])
        if shapes and shapes[i]:
            kw["shape"] = shapes[i]
        out.append(d.rect(bx, by, bw, bh, t, id=ids[i], **kw))
    if arrows:
        for a, b in zip(out, out[1:]):
            connect(d, a, b, "h" if direction == "h" else "v", color=arrow_color, sw=arrow_sw)
    return out


def cycle(d, labels, cx, cy, rx, ry, w=156, h=60, styles=None, badges=True, start_deg=-90, close=True,
          close_color="secondary", close_dashed=True, bulge=1.22, arrow_color="primary"):
    """Boxes evenly placed on an ellipse with curved arrows between neighbours (clockwise).
    The closing arrow (last -> first) can be styled separately (feedback). Returns (ids, centers)."""
    n = len(labels)
    styles = styles if isinstance(styles, list) else [styles or NODE] * n
    ids, pos = [], []
    for i, t in enumerate(labels):
        a = math.radians(start_deg) + i * 2 * math.pi / n
        x, y = cx + rx * math.cos(a), cy + ry * math.sin(a)
        pos.append((x, y))
        ids.append(d.rect(x - w / 2, y - h / 2, w, h, t, id=f"cyc{i}", **styles[i]))
        if badges:
            badge(d, x - w / 2, y - h / 2, i + 1)
    for i in range(n if close else n - 1):
        j = (i + 1) % n
        (x1, y1), (x2, y2) = pos[i], pos[j]
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        ctrl = (cx + (mx - cx) * bulge, cy + (my - cy) * bulge)   # control point outside the chord
        last = close and i == n - 1
        d.edge([ctrl], src=ids[i], tgt=ids[j], route="curved", color=close_color if last else arrow_color,
               sw="thick", dashed=last and close_dashed)
    return ids, pos


def layers(d, x, y, w, spec, layer_h=92, gap=26, tab_w=150, comp_h=56, comp_gap=40, styles=None,
           flow_arrows=True, last_highlight=True):
    """Stacked architecture layers. spec = [(layer_name, [component labels...]), ...].
    Components in a layer are joined left->right; a down arrow links consecutive layers.
    styles: optional {layer_index: preset}. Returns {(li, ci): id} and the list of layer boxes."""
    ids, boxes_out = {}, []
    for li, (name, comps) in enumerate(spec):
        ly = y + li * (layer_h + gap)
        last = last_highlight and li == len(spec) - 1
        d.rect(x, ly, w, layer_h, fill="secondary" if last else "neutral", fill_opacity="wash", stroke="neutral",
               sw="thin", rounded=6)
        tab = "secondary" if last else "primary"
        d.rect(x, ly, tab_w, layer_h, f"**{name}**", fill=tab, stroke=tab, color="background", size="subheading", rounded=6)
        cb = d.fit_row(x + tab_w + 25, x + w - 25, ly + (layer_h - comp_h) / 2, len(comps), comp_h, gap=comp_gap)
        st = (styles or {}).get(li, NODE)
        for ci, ((bx, by, bw, bh), t) in enumerate(zip(cb, comps)):
            ids[(li, ci)] = d.rect(bx, by, bw, bh, t, id=f"L{li}c{ci}", **st)
        for ci in range(len(comps) - 1):
            connect(d, ids[(li, ci)], ids[(li, ci + 1)])
        if flow_arrows and li < len(spec) - 1:
            mx = cb[len(cb) // 2][0] + cb[len(cb) // 2][2] / 2
            d.arrow([(mx, ly + layer_h + 2), (mx, ly + layer_h + gap - 2)], color="primary", sw="heavy", head_size=5)
        boxes_out.append((x, ly, w, layer_h))
    return ids, boxes_out


def frame(d, x, y, w, h, color="neutral", id=None, label=None):
    """Dashed group frame (no fill) — e.g. around a column of related items; connect arrows to it."""
    fid = d.rect(x, y, w, h, fill="none", stroke=color, sw="thin", dashed=True, rounded=10, id=id)
    if label:
        d.text(f"*{label}*", x + 8, y - 10, anchor="w", size="small", color=color)
    return fid


def bracket(d, x0, x1, y, color="primary", up=False, tick=8, sw="thick"):
    """Horizontal bracket under (or over, up=True) a span; returns its midpoint x."""
    t = -tick if up else tick
    d.line([(x0, y - t), (x0, y), (x1, y), (x1, y - t)], color=color, sw=sw)
    return (x0 + x1) / 2


def spread(centers, width, gap, lo, hi):
    """Shift box centers left->right so boxes of `width` don't overlap and stay inside [lo, hi]."""
    out = []
    for c in centers:
        c = max(c, lo + width / 2)
        if out and c - out[-1] < width + gap:
            c = out[-1] + width + gap
        out.append(c)
    over = (out[-1] + width / 2 - hi) if out else 0
    return [c - max(0, over) for c in out]


def eq(d, latex, x, y, **kw):
    """Shorthand for d.equation (LaTeX label)."""
    return d.equation(latex, x, y, **kw)


def annotated_equation(d, terms, cx, cy, notes_above=(), notes_below=(), size="equation", term_gap=10,
                       box_w=230, box_h=78, top_y=30, bottom_y=320, lo=30, hi=None, bracket_offset=34):
    """Equation split into terms, each its own LaTeX label, with bracketed callouts.
    terms: [(latex, color_role), ...] in reading order.
    notes_above / notes_below: [(term_index, color_role, markup_text), ...].
    Callout boxes are spread so they never overlap. Returns list of (x0, x1) spans per term."""
    hi = hi or d.width - 30
    sz = d.theme.size(size)
    widths = [measure_math(t, sz)[0] * 1.08 + 6 for t, _ in terms]
    total = sum(widths) + term_gap * (len(terms) - 1)
    x = cx - total / 2
    spans = []
    for (t, col), w in zip(terms, widths):
        d.equation(t, x, cy, w=w, h=sz * 1.9, anchor="w", size=size, color=col)
        spans.append((x, x + w)); x += w + term_gap
    for row, up in ((list(notes_above), True), (list(notes_below), False)):
        if not row:
            continue
        mids = [(spans[i][0] + spans[i][1]) / 2 for i, _, _ in row]
        centers = spread(mids, box_w, 24, lo, hi)
        for (i, col, txt), mx, bx in zip(row, mids, centers):
            x0, x1 = spans[i]
            by = cy - bracket_offset if up else cy + bracket_offset
            bracket(d, x0 + 2, x1 - 2, by, col, up=up)
            top = top_y if up else bottom_y
            yconn = top + box_h if up else top
            ym = (by + yconn) / 2
            d.line([(mx, by), (mx, ym), (bx, ym), (bx, yconn)], color=col, sw="thin")
            d.rect(bx - box_w / 2, top, box_w, box_h, txt, fill="background", stroke=col, sw="normal",
                   rounded=True, size="label")
    return spans


def rating_dots(d, cx, cy, n, of=3, r=7, gap=22, color="primary"):
    """n of `of` filled dots centered at (cx, cy) — Harvey-ball style ordinal rating."""
    for k in range(of):
        d.circle(cx + (k - (of - 1) / 2) * gap, cy, r, fill=color if k < n else "background", stroke=color, sw="normal")


def rating_legend(d, x, y, of=3, best="most favourable", worst="least favourable"):
    rating_dots(d, x + 14, y, of, of, r=5, gap=14)
    d.text(best, x + 14 + of * 9 + 10, y, anchor="w", size="small")
    x2 = x + 190
    rating_dots(d, x2 + 14, y, 1, of, r=5, gap=14)
    d.text(worst, x2 + 14 + of * 9 + 10, y, anchor="w", size="small")


def card(d, x, y, w, h, title, intro=None, sections=(), footer=None, header_fill="primary", id=None):
    """Card with a header band, optional italic intro, bullet sections and a tinted footer.
    sections: [(heading, [items], color_role, marker), ...] e.g. ("Strengths", [...], "secondary", "+")."""
    with d.group(id or f"card_{title.split()[0].lower()}") as gid:
        d.rect(x, y, w, h, fill="background", stroke="neutral", sw="thin", rounded=12, shadow=True)
        d.rect(x, y, w, 54, f"**{title}**", fill=header_fill, stroke=header_fill, color="background", rounded=12, size="heading")
        cy = y + 84
        if intro:
            d.text(f"*{intro}*", x + 18, cy, anchor="w", w=w - 36, h=40, wrap=True, size="label")
            cy += 46
        for head, items, col, mark in sections:
            d.text(f"**{head}**", x + 18, cy, anchor="w", size="body", color=col)
            cy += 28
            for it in items:
                d.text(f"{{{col}|**{mark}**}}  {it}", x + 18, cy, anchor="w", size="label")
                cy += 26
            cy += 20
        if footer:
            d.rect(x + 14, y + h - 84, w - 28, 70, footer, fill="secondary", fill_opacity="tint", stroke="none",
                   rounded=8, size="label")
    return gid


def continuum(d, x0, x1, y, left, right, left_eq=None, right_eq=None, marker=None, marker_label=None,
              colors=("neutral", "secondary")):
    """Horizontal gradient bar with end labels (optionally LaTeX) and an optional marker at fraction `marker`."""
    d.rect(x0, y - 8, x1 - x0, 16, fill=colors[0], gradient=colors[1], gradient_dir="east", stroke="none", rounded=8)
    for x, lab, eq in ((x0, left, left_eq), (x1, right, right_eq)):
        if eq:
            d.equation(eq, x, y + 34, size="eq_small", color="primary")
        d.text(lab, x, y + (62 if eq else 30), size="small")
    if marker is not None:
        xb = x0 + marker * (x1 - x0)
        d.shape("triangle_down", xb - 11, y - 36, 22, 18, fill="primary", stroke="primary")
        if marker_label:
            d.text(marker_label, xb, y - 50, size="small", color="primary")


def legend(d, x, y, items, spacing=190, sw=22, sh=16):
    """items: [(preset_dict, label), ...] drawn left->right."""
    for k, (st, lab) in enumerate(items):
        xx = x + k * spacing
        d.rect(xx, y - sh / 2, sw, sh, **{kk: v for kk, v in st.items() if kk not in ("size", "bold", "color")})
        d.text(lab, xx + sw + 8, y, anchor="w", size="label")


def axes(d, ox, oy, x1, y1, xlabel=None, ylabel=None, xticks=(), yticks=(), color="text", wash=True):
    """Qualitative/quantitative axes with arrowheads. ticks: [(position_px, label), ...]."""
    if wash:
        d.rect(ox, y1, x1 - ox, oy - y1, fill="neutral", fill_opacity="wash", stroke="none")
    d.arrow([(ox, oy), (x1 + 20, oy)], color=color, sw="normal")
    d.arrow([(ox, oy), (ox, y1 - 20)], color=color, sw="normal")
    if xlabel:
        d.text(xlabel, (ox + x1) / 2, oy + 48, size="body")
    if ylabel:
        d.text(ylabel, ox - 62, (oy + y1) / 2, rotation=-90, size="body")
    for x, t in xticks:
        d.text(t, x, oy + 18, size="small", color="neutral")
    for y, t in yticks:
        d.text(t, ox - 30, y, size="small", color="neutral")


def table(d, x, y, col_widths, headers, rows, row_h=52, header_h=36, cell=None, zebra=True):
    """Simple table. rows: list of lists of markup strings; cell(d, ri, ci, box) may draw custom content
    (return True to skip the default text). Returns list of row boxes."""
    cx = x
    for w, h in zip(col_widths, headers):
        d.rect(cx, y, w, header_h, f"**{h}**", fill="primary", stroke="background", color="background", size="label")
        cx += w
    out = []
    total = sum(col_widths)
    for ri, row in enumerate(rows):
        ry = y + header_h + ri * row_h
        odd = zebra and ri % 2 == 0
        d.rect(x, ry, total, row_h, fill="secondary" if odd else "background", fill_opacity="wash" if odd else None,
               stroke="neutral", sw="thin")
        cx = x
        for ci, (w, val) in enumerate(zip(col_widths, row)):
            box = (cx, ry, w, row_h)
            if not (cell and cell(d, ri, ci, box)) and val:
                d.text(val, cx + 14, ry + row_h / 2, anchor="w", size="body")
            cx += w
        out.append((x, ry, total, row_h))
    return out
'''

# ================================================================================================
FILES['scripts/finalize.py'] = r'''#!/usr/bin/env python3
"""
finalize.py — keep the prototype the user picked.

    python finalize.py PASSAGE_DIR --pick v2 [--name rag_pipeline] [--clean]

Copies <pick>_<style>.drawio/.png to <name>.drawio/.png in the same folder (default name: the
passage folder name, or "figure" for a flat folder named *_figures), writes CHOSEN.md with the
draft caption and inferred-content list, and with --clean deletes the other prototypes'
.drawio/.png files (the _work/ snapshots are kept for regeneration).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import glob
import importlib.util
import os
import re
import shutil
import sys

sys.dont_write_bytecode = True


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("passage_dir"); ap.add_argument("--pick", required=True)
    ap.add_argument("--name"); ap.add_argument("--clean", action="store_true")
    a = ap.parse_args()
    pdir = os.path.abspath(a.passage_dir)
    src = sorted(glob.glob(os.path.join(pdir, f"{a.pick}_*.drawio")))
    if not src:
        raise SystemExit(f"no {a.pick}_*.drawio in {pdir} (run prototype.py first)")
    folder = os.path.basename(pdir)
    name = a.name or re.sub(r"^p\d+_", "", folder.replace("_figures", "")) or "figure"
    dst = os.path.join(pdir, name + ".drawio")
    shutil.copy(src[0], dst)
    png = src[0][:-7] + ".png"
    if os.path.exists(png):
        shutil.copy(png, dst[:-7] + ".png")
    caption, inferred = "", []
    bp = os.path.join(pdir, "build.py")
    if os.path.exists(bp):
        try:
            sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
            spec = importlib.util.spec_from_file_location("b", bp); m = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(m)
            v = next((v for v in m.VARIANTS if v["key"] == a.pick), {})
            caption = v.get("caption", "")
            inferred = [x for x in getattr(m, "INFERRED", []) if x.lower().startswith(a.pick + ":") or ":" not in x[:4]]
        except Exception:
            pass
    open(os.path.join(pdir, "CHOSEN.md"), "w", encoding="utf-8").write(
        f"# Chosen figure\n\n- prototype: `{os.path.basename(src[0])}`\n- final: `{os.path.basename(dst)}`\n"
        f"- draft caption: {caption or '—'}\n- inferred content: {'; '.join(inferred) or 'none'}\n")
    if a.clean:
        for f in glob.glob(os.path.join(pdir, "v[0-9]_*.drawio")) + glob.glob(os.path.join(pdir, "v[0-9]_*.png")):
            os.remove(f)
    print("final:", dst)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/inspect_image.py'] = r'''#!/usr/bin/env python3
"""
inspect_image.py — measure a source figure before rebuilding it in draw.io.

Subcommands (all coordinates are SOURCE pixels):
  info    SRC                         size, background color, a few dominant colors
  grid    SRC [--step 50] [-o]        coordinate grid overlay (labels every major line)
  zoom    SRC --box X0 Y0 X1 Y1 [--factor 3] [--step 10] [-o]
                                      magnified crop with a fine grid labeled in source coords
  palette SRC [-k 8] [--theme T]      dominant colors (k-means) + swatch; maps to theme roles if given
  boxes   SRC [--min-area 150] [-o]   connected components / rectangles -> JSON + numbered overlay
  ocr     SRC [--upscale 2] [-o]      text lines with bounding boxes + estimated font px (tesseract)
  lines   SRC [--min-len 60]          long straight lines (axes, dividers, borders), needs OpenCV
  all     SRC --out-dir DIR           run everything, write DIR/summary.md (start here)

Typical use:
  python inspect_image.py all images/fig.png --out-dir images/fig_drawio_work/inspect
  python inspect_image.py zoom images/fig.png --box 40 180 200 330 --factor 4
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import math
import os
import sys

import numpy as np
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from drawio_kit import Theme, find_font_file, measure_line  # noqa: E402


def load(src):
    return Image.open(src).convert("RGB")


def hexc(rgb):
    return "#%02x%02x%02x" % tuple(int(v) for v in rgb[:3])


def _font(px):
    p = find_font_file("Helvetica")
    return ImageFont.truetype(p, max(8, int(px))) if p else ImageFont.load_default()


def background_color(im):
    a = np.asarray(im).reshape(-1, 3)
    border = np.concatenate([np.asarray(im)[0], np.asarray(im)[-1], np.asarray(im)[:, 0], np.asarray(im)[:, -1]])
    q = (border // 8).astype(np.int32)
    keys, counts = np.unique(q[:, 0] * 10000 + q[:, 1] * 100 + q[:, 2], return_counts=True)
    k = keys[counts.argmax()]
    sel = border[(q[:, 0] * 10000 + q[:, 1] * 100 + q[:, 2]) == k]
    return tuple(int(v) for v in np.median(sel, axis=0))


# ------------------------------------------------------------------ grid / zoom
def cmd_grid(src, step=50, out=None):
    im = load(src)
    W, H = im.size
    ov = im.copy()
    d = ImageDraw.Draw(ov, "RGBA")
    f = _font(max(10, min(W, H) / 70))
    for x in range(0, W, step):
        major = (x // step) % 2 == 0
        d.line([(x, 0), (x, H)], fill=(255, 0, 0, 110 if major else 55), width=1)
        if major:
            d.text((x + 2, 2), str(x), fill=(200, 0, 0, 255), font=f)
            d.text((x + 2, H - f.size - 4), str(x), fill=(200, 0, 0, 255), font=f)
    for y in range(0, H, step):
        major = (y // step) % 2 == 0
        d.line([(0, y), (W, y)], fill=(0, 0, 255, 110 if major else 55), width=1)
        if major:
            d.text((2, y + 1), str(y), fill=(0, 0, 200, 255), font=f)
            d.text((W - 4 * f.size, y + 1), str(y), fill=(0, 0, 200, 255), font=f)
    out = out or os.path.splitext(src)[0] + f".grid{step}.png"
    ov.save(out)
    print("wrote", out)
    return out


def cmd_zoom(src, box, factor=3, step=10, out=None):
    im = load(src)
    x0, y0, x1, y1 = [int(v) for v in box]
    x0, y0 = max(0, x0), max(0, y0)
    x1, y1 = min(im.width, x1), min(im.height, y1)
    crop = im.crop((x0, y0, x1, y1)).resize(((x1 - x0) * factor, (y1 - y0) * factor), Image.NEAREST)
    margin = 34
    canvas = Image.new("RGB", (crop.width + margin, crop.height + margin), "white")
    canvas.paste(crop, (margin, margin))
    d = ImageDraw.Draw(canvas, "RGBA")
    f = _font(11)
    start_x = (x0 // step + (x0 % step > 0)) * step
    for x in range(start_x, x1 + 1, step):
        X = margin + (x - x0) * factor
        major = x % (step * 5) == 0
        d.line([(X, margin), (X, canvas.height)], fill=(255, 0, 0, 140 if major else 60))
        if major or step * factor >= 40:
            d.text((X + 1, 2), str(x), fill=(200, 0, 0), font=f)
    start_y = (y0 // step + (y0 % step > 0)) * step
    for y in range(start_y, y1 + 1, step):
        Y = margin + (y - y0) * factor
        major = y % (step * 5) == 0
        d.line([(margin, Y), (canvas.width, Y)], fill=(0, 0, 255, 140 if major else 60))
        if major or step * factor >= 40:
            d.text((1, Y + 1), str(y), fill=(0, 0, 200), font=f)
    out = out or os.path.splitext(src)[0] + f".zoom_{x0}_{y0}_{x1}_{y1}.png"
    canvas.save(out)
    print("wrote", out)
    return out


# ------------------------------------------------------------------ palette
def _srgb_to_lab(rgb):
    c = np.asarray(rgb, dtype=float) / 255.0
    c = np.where(c > 0.04045, ((c + 0.055) / 1.055) ** 2.4, c / 12.92)
    M = np.array([[0.4124, 0.3576, 0.1805], [0.2126, 0.7152, 0.0722], [0.0193, 0.1192, 0.9505]])
    xyz = c @ M.T / np.array([0.95047, 1.0, 1.08883])
    f = np.where(xyz > 0.008856, np.cbrt(xyz), 7.787 * xyz + 16 / 116)
    return np.stack([116 * f[..., 1] - 16, 500 * (f[..., 0] - f[..., 1]), 200 * (f[..., 1] - f[..., 2])], -1)


def kmeans(X, k, iters=25, seed=0):
    rng = np.random.default_rng(seed)
    C = X[rng.choice(len(X), size=min(k, len(X)), replace=False)].astype(float)
    for _ in range(iters):
        d = ((X[:, None, :] - C[None]) ** 2).sum(-1)
        lab = d.argmin(1)
        newC = np.array([X[lab == i].mean(0) if (lab == i).any() else C[i] for i in range(len(C))])
        if np.allclose(newC, C, atol=0.5):
            break
        C = newC
    return C, lab


def cmd_palette(src, k=8, theme=None, out=None, quiet=False, write=True):
    im = load(src)
    im.thumbnail((400, 400))
    X = np.asarray(im).reshape(-1, 3).astype(float)
    bg = np.array(background_color(im))
    C, lab = kmeans(X, k)
    counts = np.bincount(lab, minlength=len(C))
    order = counts.argsort()[::-1]
    rows = []
    for i in order:
        rgb = C[i]
        mx, mn = rgb.max(), rgb.min()
        sat = 0 if mx == 0 else (mx - mn) / mx
        rows.append(dict(hex=hexc(rgb), share=round(float(counts[i] / counts.sum()), 4),
                         lightness=round(float(_srgb_to_lab(rgb)[0]), 1), saturation=round(float(sat), 2),
                         is_background=bool(np.abs(rgb - bg).max() < 18)))
    th = Theme.load(theme) if theme else None
    if th:
        names = list(th.colors)
        T = _srgb_to_lab(np.array([[int(th.colors[n][j:j + 2], 16) for j in (1, 3, 5)] for n in names]))
        for r in rows:
            L = _srgb_to_lab(np.array([int(r["hex"][j:j + 2], 16) for j in (1, 3, 5)]))
            dist = np.sqrt(((T - L) ** 2).sum(-1))
            r["nearest_role"], r["delta_e"] = names[int(dist.argmin())], round(float(dist.min()), 1)
    # swatch
    sw = Image.new("RGB", (520, 34 * len(rows) + 10), "white")
    d = ImageDraw.Draw(sw)
    f = _font(14)
    for j, r in enumerate(rows):
        d.rectangle([10, 8 + 34 * j, 60, 36 + 34 * j], fill=r["hex"], outline="black")
        txt = f'{r["hex"]}  {100 * r["share"]:5.1f}%  L={r["lightness"]:5.1f}  sat={r["saturation"]:.2f}'
        if "nearest_role" in r:
            txt += f'  ~{r["nearest_role"]} (dE {r["delta_e"]})'
        if r["is_background"]:
            txt += "  [bg]"
        d.text((70, 14 + 34 * j), txt, fill="black", font=f)
    out = out or os.path.splitext(src)[0] + ".palette.png"
    if write:
        sw.save(out)
    if not quiet:
        print(f"background ~ {hexc(bg)}")
        for r in rows:
            print(" ", r)
        if write:
            print("wrote", out)
    return dict(background=hexc(bg), colors=rows, swatch=out if write else None)


# ------------------------------------------------------------------ components
def _label(mask):
    try:
        from scipy import ndimage
        lab, n = ndimage.label(mask)
        return lab, n
    except ImportError:
        pass
    try:
        import cv2
        n, lab = cv2.connectedComponents(mask.astype(np.uint8), connectivity=8)
        return lab, n - 1
    except ImportError:
        raise SystemExit("boxes needs scipy or opencv-python")


def cmd_boxes(src, min_area=150, thresh=40, out=None, quiet=False):
    """Components that differ from the background. Classifies each as filled box / outline / text-ish / mark."""
    im = load(src)
    a = np.asarray(im).astype(int)
    bg = np.array(background_color(im))
    diff = np.abs(a - bg).max(-1)
    mask = diff > thresh
    lab, n = _label(mask)
    comps = []
    if n:
        try:
            from scipy import ndimage
            slices = ndimage.find_objects(lab)
        except ImportError:
            slices = []
            for i in range(1, n + 1):
                ys, xs = np.nonzero(lab == i)
                slices.append((slice(ys.min(), ys.max() + 1), slice(xs.min(), xs.max() + 1)) if len(ys) else None)
        for i, sl in enumerate(slices, start=1):
            if sl is None:
                continue
            ys, xs = sl
            h, w = ys.stop - ys.start, xs.stop - xs.start
            region = lab[sl] == i
            area = int(region.sum())
            if w * h < min_area:
                continue
            fill_ratio = area / float(w * h)
            px = a[sl][region]
            med = np.median(px, axis=0)
            kind = ("filled box" if fill_ratio > 0.85 and min(w, h) > 12 else
                    "outline/frame" if fill_ratio < 0.25 and min(w, h) > 20 else
                    "line" if min(w, h) <= 4 else
                    "glyph/text-ish" if max(w, h) < 40 else "mark/shape")
            comps.append(dict(id=len(comps), bbox=[int(xs.start), int(ys.start), int(xs.stop), int(ys.stop)],
                              w=int(w), h=int(h), area=area, fill_ratio=round(fill_ratio, 2),
                              color=hexc(med), kind=kind))
    comps.sort(key=lambda c: -c["w"] * c["h"])
    for j, c in enumerate(comps):
        c["id"] = j
    ov = im.copy()
    d = ImageDraw.Draw(ov, "RGBA")
    f = _font(12)
    for c in comps[:300]:
        if c["kind"] == "glyph/text-ish":
            continue
        x0, y0, x1, y1 = c["bbox"]
        col = {"filled box": (0, 160, 0, 200), "outline/frame": (220, 0, 0, 200), "line": (0, 0, 220, 200)}.get(c["kind"], (200, 120, 0, 200))
        d.rectangle([x0, y0, x1, y1], outline=col, width=2)
        d.text((x0 + 2, y0 + 1), str(c["id"]), fill=col, font=f)
    out = out or os.path.splitext(src)[0] + ".boxes.png"
    ov.save(out)
    js = os.path.splitext(out)[0] + ".json"
    json.dump(dict(background=hexc(bg), components=comps), open(js, "w"), indent=1)
    if not quiet:
        print(f"{len(comps)} components (background {hexc(bg)}); wrote {out} and {js}")
        for c in comps[:40]:
            if c["kind"] != "glyph/text-ish":
                print(" ", c)
    return comps, out, js


def cmd_lines(src, min_len=60, quiet=False):
    try:
        import cv2
    except ImportError:
        print("lines needs opencv-python; skipping")
        return []
    g = cv2.cvtColor(np.asarray(load(src)), cv2.COLOR_RGB2GRAY)
    edges = cv2.Canny(g, 50, 150)
    L = cv2.HoughLinesP(edges, 1, np.pi / 180, threshold=60, minLineLength=min_len, maxLineGap=4)
    out = []
    for l in (L if L is not None else []):
        x1, y1, x2, y2 = map(int, l[0])
        ang = math.degrees(math.atan2(y2 - y1, x2 - x1)) % 180
        kind = "horizontal" if ang < 2 or ang > 178 else "vertical" if abs(ang - 90) < 2 else "diagonal"
        out.append(dict(p1=[x1, y1], p2=[x2, y2], length=round(math.dist((x1, y1), (x2, y2))), kind=kind))
    out.sort(key=lambda r: -r["length"])
    if not quiet:
        for r in out[:40]:
            print(" ", r)
    return out


# ------------------------------------------------------------------ OCR
def cmd_ocr(src, upscale=2, out=None, quiet=False, psm=11, family="Helvetica"):
    try:
        import pytesseract
    except ImportError:
        print("ocr needs pytesseract + tesseract; skipping (read the text visually instead)")
        return []
    im = load(src)
    big = im.resize((im.width * upscale, im.height * upscale), Image.LANCZOS)
    data = pytesseract.image_to_data(big, config=f"--psm {psm}", output_type=pytesseract.Output.DICT)
    words = []
    for i, t in enumerate(data["text"]):
        if t.strip() and float(data["conf"][i]) > 30:
            words.append(dict(text=t, x=data["left"][i] / upscale, y=data["top"][i] / upscale,
                              w=data["width"][i] / upscale, h=data["height"][i] / upscale,
                              key=(data["block_num"][i], data["par_num"][i], data["line_num"][i])))
    # Also run inverted pass for white/light text on dark backgrounds
    try:
        from PIL import ImageOps
        inv = ImageOps.invert(big.convert("L"))
        data_inv = pytesseract.image_to_data(inv, config=f"--psm {psm}", output_type=pytesseract.Output.DICT)
        for i, t in enumerate(data_inv["text"]):
            if t.strip() and float(data_inv["conf"][i]) > 30:
                ix = data_inv["left"][i] / upscale
                iy = data_inv["top"][i] / upscale
                iw = data_inv["width"][i] / upscale
                ih = data_inv["height"][i] / upscale
                overlap = any(abs(ix - w["x"]) < max(iw, w["w"]) * 0.7 and abs(iy - w["y"]) < max(ih, w["h"]) * 0.7 for w in words)
                if not overlap:
                    words.append(dict(text=t, x=ix, y=iy, w=iw, h=ih,
                                      key=(100 + data_inv["block_num"][i], data_inv["par_num"][i], data_inv["line_num"][i])))
    except Exception:
        pass
    # group by tesseract line, then merge horizontally adjacent words at the same baseline
    lines = {}
    for w in words:
        lines.setdefault(w["key"], []).append(w)
    rows = []
    for ws in lines.values():
        ws.sort(key=lambda w: w["x"])
        cur = [ws[0]]
        for w in ws[1:]:
            prev = cur[-1]
            if w["x"] - (prev["x"] + prev["w"]) < 1.2 * max(prev["h"], w["h"]) and abs(w["y"] - prev["y"]) < 0.6 * prev["h"]:
                cur.append(w)
            else:
                rows.append(cur); cur = [w]
        rows.append(cur)
    res = []
    for r in rows:
        x0 = min(w["x"] for w in r); y0 = min(w["y"] for w in r)
        x1 = max(w["x"] + w["w"] for w in r); y1 = max(w["y"] + w["h"] for w in r)
        text = " ".join(w["text"] for w in r)
        has_desc = any(ch in text for ch in "gjpqy(),;")
        caps = text.upper() == text
        est = (y1 - y0) / (1.04 if has_desc else 0.74 if caps else 0.84)   # calibrated on Helvetica/Arial
        # width-based estimate (more robust when OCR text is right): size where measured width == bbox width
        w100 = measure_line([dict(text=text)], family, 100) or 1
        b100 = measure_line([dict(text=text, bold=True)], family, 100) or 1
        res.append(dict(text=text, bbox=[round(x0), round(y0), round(x1), round(y1)],
                        center=[round((x0 + x1) / 2), round((y0 + y1) / 2)], est_font_px=round(est, 1),
                        est_px_if_regular=round(100 * (x1 - x0) / w100, 1),
                        est_px_if_bold=round(100 * (x1 - x0) / b100, 1)))
    res.sort(key=lambda r: (r["bbox"][1] // 8, r["bbox"][0]))
    ov = im.copy()
    d = ImageDraw.Draw(ov, "RGBA")
    f = _font(10)
    for j, r in enumerate(res):
        d.rectangle(r["bbox"], outline=(255, 0, 255, 200), width=1)
        d.text((r["bbox"][0], r["bbox"][3]), f"{j}:{r['est_font_px']:.0f}", fill=(200, 0, 200), font=f)
    out = out or os.path.splitext(src)[0] + ".ocr.png"
    ov.save(out)
    js = os.path.splitext(out)[0] + ".json"
    json.dump(res, open(js, "w"), indent=1)
    if not quiet:
        for j, r in enumerate(res):
            print(f"  [{j}] {r['bbox']} ~{r['est_font_px']}px (width fit: {r['est_px_if_regular']} reg / {r['est_px_if_bold']} bold)  {r['text']}")
        print("wrote", out, js)
    return res


# ------------------------------------------------------------------ all
def cmd_all(src, out_dir, theme=None, step=None):
    os.makedirs(out_dir, exist_ok=True)
    im = load(src)
    W, H = im.size
    step = step or (25 if max(W, H) <= 700 else 50 if max(W, H) <= 1600 else 100)
    stem = os.path.join(out_dir, "source")
    im.save(stem + ".png")
    grid = cmd_grid(stem + ".png", step, stem + f".grid{step}.png")
    fine = cmd_grid(stem + ".png", max(10, step // 2), stem + f".grid{max(10, step // 2)}.png")
    pal = cmd_palette(stem + ".png", 10, theme, stem + ".palette.png", quiet=True)
    comps, boxes_png, boxes_js = cmd_boxes(stem + ".png", out=stem + ".boxes.png", quiet=True)
    ocr = cmd_ocr(stem + ".png", out=stem + ".ocr.png", quiet=True)
    lines = cmd_lines(stem + ".png", quiet=True)
    md = [f"# Inspection summary for `{src}`", "",
          f"- size: **{W} x {H}** px, background ~ `{pal['background']}`",
          f"- grids: `{os.path.basename(grid)}` (step {step}), `{os.path.basename(fine)}`",
          f"- palette swatch: `{os.path.basename(pal['swatch'])}`  |  boxes: `{os.path.basename(boxes_png)}`"
          f"  |  OCR: `source.ocr.png`", "",
          "## Palette (largest share first)", "", "| hex | share | L | sat | note |", "|---|---|---|---|---|"]
    for c in pal["colors"]:
        note = ("background" if c["is_background"] else "") + (f" ~{c['nearest_role']} dE={c['delta_e']}" if "nearest_role" in c else "")
        md.append(f"| `{c['hex']}` | {100 * c['share']:.1f}% | {c['lightness']} | {c['saturation']} | {note} |")
    md += ["", "## Large components (candidate boxes/panels/lines)", "", "| id | bbox x0,y0,x1,y1 | w x h | kind | color |", "|---|---|---|---|---|"]
    for c in [c for c in comps if c["kind"] != "glyph/text-ish"][:60]:
        md.append(f"| {c['id']} | {c['bbox']} | {c['w']}x{c['h']} | {c['kind']} | `{c['color']}` |")
    if lines:
        md += ["", "## Long straight lines", ""] + [f"- {l['kind']} {l['p1']} -> {l['p2']} (len {l['length']})" for l in lines[:30]]
    if ocr:
        sizes = sorted({round(r["est_font_px"]) for r in ocr})
        md += ["", f"## OCR text lines (est. font px; distinct ~ {sizes})", "",
               "| # | bbox | center | ~px (height) | px if regular / bold (width fit) | text |", "|---|---|---|---|---|---|"]
        for j, r in enumerate(ocr):
            md.append(f"| {j} | {r['bbox']} | {r['center']} | {r['est_font_px']} | {r['est_px_if_regular']} / {r['est_px_if_bold']} | {r['text'].replace('|', '/')} |")
        md += ["", "> OCR can misread; always confirm text visually. Prefer the width-fit size when the OCR text is",
               "> correct (pick the regular or bold column by looking); the height estimate is +/-15%."]
    else:
        md += ["", "_OCR unavailable — transcribe text by viewing the image._"]
    path = os.path.join(out_dir, "summary.md")
    open(path, "w").write("\n".join(md) + "\n")
    print("wrote", path)
    return path


def cmd_probe(src, box, ascii_map=True):
    """Deep inspect a region (x0, y0, x1, y1): dominant colors, shapes, text, and ASCII map."""
    im = load(src)
    x0, y0, x1, y1 = [int(v) for v in box]
    x0, y0 = max(0, x0), max(0, y0)
    x1, y1 = min(im.width, x1), min(im.height, y1)
    crop = im.crop((x0, y0, x1, y1))
    arr = np.asarray(crop)
    h, w = arr.shape[:2]
    print(f"=== Probe region [{x0}, {y0}, {x1}, {y1}] ({w}x{h} px) ===")

    # 1. Colors
    pixels = arr.reshape(-1, 3)
    u_colors, counts = np.unique(pixels, axis=0, return_counts=True)
    idx = np.argsort(-counts)
    print("Top colors (RGB):")
    for i in range(min(5, len(idx))):
        c = u_colors[idx[i]]
        print(f"  #{c[0]:02x}{c[1]:02x}{c[2]:02x} ({c}): {counts[idx[i]]} px ({counts[idx[i]]/(w*h)*100:.1f}%)")

    # 2. Text in this box (normal and inverted)
    try:
        import pytesseract
        from PIL import ImageOps
        txt_norm = pytesseract.image_to_string(crop).strip()
        txt_inv = pytesseract.image_to_string(ImageOps.invert(crop.convert("L"))).strip()
        if txt_norm:
            print(f"Text (normal): '{txt_norm}'")
        if txt_inv and txt_inv != txt_norm:
            print(f"Text (inverted): '{txt_inv}'")
    except Exception:
        pass

    # 3. Geometric features (OpenCV if available)
    try:
        import cv2
        gray = cv2.cvtColor(arr, cv2.COLOR_RGB2GRAY)
        circles = cv2.HoughCircles(gray, cv2.HOUGH_GRADIENT, 1, 10, param1=50, param2=20, minRadius=5, maxRadius=max(w, h))
        if circles is not None:
            print(f"Circles detected: {len(circles[0])}")
            for c in circles[0][:5]:
                print(f"  circle: center=({c[0]+x0:.1f}, {c[1]+y0:.1f}), r={c[2]:.1f}")
        edges = cv2.Canny(gray, 50, 150)
        lines = cv2.HoughLinesP(edges, 1, np.pi/180, 15, minLineLength=15, maxLineGap=4)
        if lines is not None:
            print(f"Lines detected: {len(lines)}")
            for l in lines[:5]:
                lx1, ly1, lx2, ly2 = l[0]
                print(f"  line: ({lx1+x0}, {ly1+y0}) -> ({lx2+x0}, {ly2+y0}) len={math.hypot(lx2-lx1, ly2-ly1):.1f}")
    except Exception:
        pass

    # 4. ASCII Density Map
    if ascii_map:
        grid_h = min(25, max(5, h // 2))
        grid_w = min(60, max(10, w // 2))
        small = crop.convert("L").resize((grid_w, grid_h), Image.BILINEAR)
        sm_arr = np.asarray(small)
        print("ASCII density map:")
        for r in range(grid_h):
            line = "".join("#" if sm_arr[r, c] < 80 else "+" if sm_arr[r, c] < 160 else "." if sm_arr[r, c] < 230 else " " for c in range(grid_w))
            print(f"  {line}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("info"); p.add_argument("src")
    p = sub.add_parser("grid"); p.add_argument("src"); p.add_argument("--step", type=int, default=50); p.add_argument("-o", "--out")
    p = sub.add_parser("zoom"); p.add_argument("src"); p.add_argument("--box", type=float, nargs=4, required=True)
    p.add_argument("--factor", type=int, default=3); p.add_argument("--step", type=int, default=10); p.add_argument("-o", "--out")
    p = sub.add_parser("palette"); p.add_argument("src"); p.add_argument("-k", type=int, default=8); p.add_argument("--theme"); p.add_argument("-o", "--out")
    p = sub.add_parser("boxes"); p.add_argument("src"); p.add_argument("--min-area", type=int, default=150)
    p.add_argument("--thresh", type=int, default=40); p.add_argument("-o", "--out")
    p = sub.add_parser("ocr"); p.add_argument("src"); p.add_argument("--upscale", type=int, default=2); p.add_argument("--psm", type=int, default=11); p.add_argument("-o", "--out")
    p = sub.add_parser("lines"); p.add_argument("src"); p.add_argument("--min-len", type=int, default=60)
    p = sub.add_parser("probe"); p.add_argument("src"); p.add_argument("--box", type=float, nargs=4, required=True)
    p.add_argument("--no-ascii", action="store_true")
    p = sub.add_parser("all"); p.add_argument("src"); p.add_argument("--out-dir", required=True); p.add_argument("--theme"); p.add_argument("--step", type=int)
    a = ap.parse_args()
    if a.cmd == "info":
        im = load(a.src)
        print(f"size {im.width}x{im.height}, background ~ {hexc(background_color(im))}")
        cmd_palette(a.src, 6, write=False)
    elif a.cmd == "grid":
        cmd_grid(a.src, a.step, a.out)
    elif a.cmd == "zoom":
        cmd_zoom(a.src, a.box, a.factor, a.step, a.out)
    elif a.cmd == "palette":
        cmd_palette(a.src, a.k, a.theme, a.out)
    elif a.cmd == "boxes":
        cmd_boxes(a.src, a.min_area, a.thresh, a.out)
    elif a.cmd == "ocr":
        cmd_ocr(a.src, a.upscale, a.out, psm=a.psm)
    elif a.cmd == "lines":
        cmd_lines(a.src, a.min_len)
    elif a.cmd == "probe":
        cmd_probe(a.src, a.box, ascii_map=not a.no_ascii)
    elif a.cmd == "all":
        cmd_all(a.src, a.out_dir, a.theme, a.step)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/iterate.py'] = r'''#!/usr/bin/env python3
"""
iterate.py — one build/verify cycle: run build.py -> render preview -> compare -> lint -> log.

    python iterate.py WORKDIR [--note "moved header text up 4px"] [--no-snapshot]

Outputs in WORKDIR/iterations/:
    iter_NN.drawio            snapshot of the output (roll back by copying it to the output path)
    iter_NN_preview.png       page-aligned render of the .drawio at source resolution
    iter_NN_side.png          source | preview with a coordinate grid       <- look at this
    iter_NN_edges.png         red = only in source, cyan = only in preview  <- alignment errors
    iter_NN_heat.png          worst cells outlined with their coordinates
    iter_NN_metrics.json, iter_NN_lint.txt
and appends a row to WORKDIR/history.json (printed as a table).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from compare import compare  # noqa: E402
from lint_drawio import lint  # noqa: E402
from render_preview import render_kit  # noqa: E402


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workdir"); ap.add_argument("--note", default="")
    ap.add_argument("--no-snapshot", action="store_true")
    a = ap.parse_args()
    work = os.path.abspath(a.workdir)
    cfg = json.load(open(os.path.join(work, "config.json")))
    hist_p = os.path.join(work, "history.json")
    hist = json.load(open(hist_p)) if os.path.exists(hist_p) else []
    n = len(hist) + 1
    tag = f"iter_{n:02d}"
    it = os.path.join(work, "iterations"); os.makedirs(it, exist_ok=True)

    r = subprocess.run([sys.executable, os.path.join(work, "build.py")], capture_output=True, text=True)
    print(r.stdout.strip())
    if r.returncode != 0:
        print(r.stderr)
        sys.exit("build.py failed — fix the error above (no iteration recorded)")
    if r.stderr.strip():
        print(r.stderr.strip())
    out = cfg["out"]
    if not a.no_snapshot:
        shutil.copy(out, os.path.join(it, f"{tag}.drawio"))
    from PIL import Image
    with Image.open(cfg["source"]) as s:
        sw = s.width
    prev = os.path.join(it, f"{tag}_preview.png")
    render_kit(out, prev, width=sw)
    m = compare(cfg["source"], prev, it, tag, quiet=True)
    L = lint(out, cfg["theme"], cfg.get("scale", 1.0))
    with open(os.path.join(it, f"{tag}_lint.txt"), "w") as fh:
        for k in ("errors", "warnings", "info"):
            for line in L[k]:
                fh.write(f"{k[:-1].upper() if k != 'info' else 'INFO'}  {line}\n")
    hist.append(dict(iter=n, score=m["score"], pixel=m["pixel_similarity"], edge_f1=m["edge_f1"],
                     errors=len(L["errors"]), warnings=len(L["warnings"]), note=a.note))
    json.dump(hist, open(hist_p, "w"), indent=1)

    print(f"\n{'iter':>4} {'score':>6} {'pixel':>6} {'edgeF1':>6} {'err':>4} {'warn':>5}  note")
    for h in hist:
        print(f"{h['iter']:>4} {h['score']:>6.3f} {h['pixel']:>6.3f} {h['edge_f1']:>6.3f} {h['errors']:>4} {h['warnings']:>5}  {h['note']}")
    if len(hist) > 1 and hist[-1]["score"] < hist[-2]["score"] - 0.005:
        print(f"\n!! score dropped vs iter {n - 1}; consider reverting (iterations/iter_{n - 1:02d}.drawio)")
    print("\nworst cells (x0,y0,x1,y1):")
    for c in m["worst_cells"][:6]:
        print(f"  {c['box']}  bad={c['bad']}  missing_edges={c['missing_edges']} extra_edges={c['extra_edges']}")
    for e in L["errors"][:10]:
        print("ERROR", e)
    for w in L["warnings"][:15]:
        print("WARN ", w)
    if len(L["warnings"]) > 15:
        print(f"  ... {len(L['warnings']) - 15} more in {tag}_lint.txt")
    print(f"\nVIEW: {m['images']['side']}\n      {m['images']['edges']}\n      {m['images']['heat']}")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/lint_drawio.py'] = r'''#!/usr/bin/env python3
"""
lint_drawio.py — catch structural errors, theme drift, text overflow, overlaps and sloppy spacing.

    python lint_drawio.py fig.drawio --theme theme.yaml [--scale 1.0] [--json report.json]

Checks
  ERROR  invalid XML, duplicate ids, missing parent/source/target, ';base64' in image data URIs
         (breaks draw.io style parsing), missing geometry
  WARN   off-theme colors / font families / font sizes, text overflowing its cell, overlapping
         text, elements outside the page, near-miss alignments (1-3 px), unequal sizes or gaps
         among siblings that look like a row/column
  INFO   counts per element type
Exit status 1 when any ERROR is found.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from drawio_kit import Theme, html_lines, measure_line, measure_math  # noqa: E402
import html as _html
from render_preview import load_model, parse_style, fnum  # noqa: E402

COLOR_KEYS = ("fillColor", "strokeColor", "fontColor", "imageBorder", "labelBackgroundColor", "gradientColor")


def lint(path, theme_path=None, scale=1.0, max_items=25):
    errors, warns, infos = [], [], []
    try:
        model = load_model(path)
    except Exception as e:
        return dict(errors=[f"cannot parse: {e}"], warnings=[], info=[])
    root = model.find("root")
    cells = []
    for o in root:
        c = o if o.tag == "mxCell" else o.find("mxCell")
        if c is None:
            continue
        if o.tag != "mxCell":
            c.set("id", o.get("id")); c.set("value", o.get("label", ""))
        cells.append(c)
    ids = Counter(c.get("id") for c in cells)
    for i, n in ids.items():
        if n > 1:
            errors.append(f"duplicate id '{i}' ({n}x)")
    by_id = {c.get("id"): c for c in cells}
    for c in cells:
        p = c.get("parent")
        if c.get("id") not in ("0",) and p and p not in by_id:
            errors.append(f"{c.get('id')}: parent '{p}' missing")
        for k in ("source", "target"):
            if c.get(k) and c.get(k) not in by_id:
                errors.append(f"{c.get('id')}: {k} '{c.get(k)}' missing")
        st = c.get("style") or ""
        if re.search(r"image=data:image/[\w+.-]+;base64", st):
            errors.append(f"{c.get('id')}: image data URI uses ';base64,' — draw.io needs 'data:image/png,<b64>'")
        if (c.get("vertex") == "1" or c.get("edge") == "1") and c.find("mxGeometry") is None:
            errors.append(f"{c.get('id')}: no mxGeometry")

    # absolute geometry
    absg = {}

    def ag(cid):
        if cid in absg:
            return absg[cid]
        c = by_id[cid]
        g = c.find("mxGeometry")
        if g is None:
            absg[cid] = None; return None
        x, y, w, h = (fnum(g.attrib, k, 0) for k in ("x", "y", "width", "height"))
        p = c.get("parent")
        if p in by_id and by_id[p].get("vertex") == "1":
            pg = ag(p)
            if pg:
                x, y = x + pg[0], y + pg[1]
        absg[cid] = (x, y, w, h)
        return absg[cid]

    verts = [c for c in cells if c.get("vertex") == "1"]
    edges = [c for c in cells if c.get("edge") == "1"]
    pw, ph = fnum(model.attrib, "pageWidth", 0), fnum(model.attrib, "pageHeight", 0)

    # theme
    th = Theme.load(theme_path) if theme_path else None
    off_colors, off_fonts, off_sizes = Counter(), Counter(), Counter()
    if th:
        palette = set(th.colors.values()) | {"none", "default", "#ffffff00"}
        families = {v["family"] for v in th.fonts.values()}
        sizes = {round(float(v) * scale, 2) for v in th.sizes.values()}
    kinds = Counter()
    texts = []
    for c in cells:
        tokens, kv = parse_style(c.get("style"))
        if th:
            for k in COLOR_KEYS:
                v = kv.get(k)
                if v and v.lower() not in palette:
                    off_colors[v.lower()] += 1
            for m in re.finditer(r'color="?(#[0-9a-fA-F]{3,6})', c.get("value") or ""):
                if m.group(1).lower() not in palette:
                    off_colors[m.group(1).lower()] += 1
            if c.get("value") and kv.get("fontFamily") and kv["fontFamily"] not in families:
                off_fonts[kv["fontFamily"]] += 1
            if c.get("value") and "fontSize" in kv and round(float(kv["fontSize"]), 2) not in sizes:
                if not re.fullmatch(r"[\W\s]*", re.sub(r"<[^>]+>", "", c.get("value") or "")):  # ignore icon glyphs
                    off_sizes[kv["fontSize"]] += 1
        if c.get("vertex") != "1":
            continue
        g = ag(c.get("id"))
        if g is None:
            continue
        x, y, w, h = g
        kind = ("group" if "group" in tokens else "text" if "text" in tokens else
                "image" if kv.get("shape") == "image" else kv.get("shape") or next((t for t in tokens if t in ("ellipse", "rhombus", "triangle")), "rect"))
        kinds[kind] += 1
        rot = abs(fnum(kv, "rotation", 0)) % 180
        bx, by, bw, bh = (x + w / 2 - h / 2, y + h / 2 - w / 2, h, w) if 45 < rot < 135 else (x, y, w, h)
        if pw and (bx < -1 or by < -1 or bx + bw > pw + 1 or by + bh > ph + 1):
            warns.append(f"{c.get('id')}: outside page ({x:.0f},{y:.0f},{w:.0f}x{h:.0f}) vs page {pw:.0f}x{ph:.0f}")
        val = c.get("value") or ""
        if val and kind != "group":
            lines = html_lines(val) if kv.get("html") == "1" else html_lines(val, raw_text=True)
            plain = " / ".join("".join(s["text"] for s in l) for l in lines).strip(" /")
            if not plain:
                continue
            fam = kv.get("fontFamily", "Helvetica"); size = fnum(kv, "fontSize", 11)
            mm = re.fullmatch(r"\s*\$\$(.+)\$\$\s*", _html.unescape(plain), re.S)
            if mm:   # LaTeX label: measure the rendered math, not the source
                mw, mh = measure_math(mm.group(1), size)
                al = kv.get("align", "center")
                tx = x if al == "left" else x + w - mw if al == "right" else x + (w - mw) / 2
                if mw > w * 1.05 + 4:
                    warns.append(f"{c.get('id')}: equation is {mw:.0f}px wide but its box is {w:.0f}px")
                texts.append((c.get("id"), plain, (tx, y + (h - mh) / 2, mw, mh)))
                continue
            bold = int(fnum(kv, "fontStyle", 0)) & 1
            if bold:
                for l in lines:
                    for s in l: s["bold"] = True
            spacing = fnum(kv, "spacing", 2)
            avail = w - 2 * spacing - fnum(kv, "spacingLeft", 0) - fnum(kv, "spacingRight", 0)
            widths = [measure_line(l, fam, size) for l in lines]
            nlines = len(lines)
            if kv.get("whiteSpace") == "wrap" and avail > 0:
                nlines = sum(max(1, int(-(-wd // avail))) for wd in widths)
                text_w = min(max(widths), avail)
            else:
                text_w = max(widths)
                if text_w > w * 1.02 + 2 and kind != "text":
                    warns.append(f"{c.get('id')}: label '{plain[:40]}' is {text_w:.0f}px wide but the box is {w:.0f}px")
            text_h = nlines * size * 1.2
            if text_h > h * 1.05 + 2 and kv.get("verticalLabelPosition") != "bottom" and kind != "text":
                warns.append(f"{c.get('id')}: label '{plain[:40]}' needs ~{text_h:.0f}px height, box is {h:.0f}px ({nlines} lines)")
            al = kv.get("align", "center")
            tx = x + spacing if al == "left" else x + w - spacing - text_w if al == "right" else x + (w - text_w) / 2
            ty = y + (h - text_h) / 2
            if "rotation" not in kv:
                texts.append((c.get("id"), plain, (tx, ty, text_w, text_h)))
    # overlapping text
    for i in range(len(texts)):
        for j in range(i + 1, len(texts)):
            a, b = texts[i][2], texts[j][2]
            ix = max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0]))
            iy = max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1]))
            small = min(a[2] * a[3], b[2] * b[3]) or 1
            if ix * iy / small > 0.15:
                warns.append(f"text overlap: '{texts[i][1][:25]}' ({texts[i][0]}) and '{texts[j][1][:25]}' ({texts[j][0]})")
    # partial overlaps: shapes that intersect without one containing the other, and text that
    # crosses a shape border (both usually mean a crowded or mis-placed element)
    boxes = []
    for c in verts:
        tokens, kv = parse_style(c.get("style"))
        g = ag(c.get("id"))
        if not g or "group" in tokens or "text" in tokens or kv.get("shape") == "image":
            continue
        if g[2] >= 14 and g[3] >= 14 and kv.get("fillColor", "") != "none" or (g[2] >= 14 and g[3] >= 14 and kv.get("strokeColor") not in (None, "none")):
            boxes.append((c.get("id"), g))

    def inter(a, b):
        return (max(0, min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])),
                max(0, min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])))

    def inside(a, b, tol=1):
        return a[0] >= b[0] - tol and a[1] >= b[1] - tol and a[0] + a[2] <= b[0] + b[2] + tol and a[1] + a[3] <= b[1] + b[3] + tol
    ov = []
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            (ia, a), (ib, b) = boxes[i], boxes[j]
            ix, iy = inter(a, b)
            small = min(a[2] * a[3], b[2] * b[3]) <= 30 * 30   # badges / markers on corners are intentional
            if ix > 2 and iy > 2 and not small and not inside(a, b) and not inside(b, a):
                ov.append(f"shapes overlap: {ia} and {ib} ({ix:.0f}x{iy:.0f}px)")
    for tid, plain, t in texts:
        for bid, b in boxes:
            ix, iy = inter(t, b)
            if ix > 3 and iy > 3 and not inside(t, b, tol=2) and not inside(b, t):
                ov.append(f"text crosses a shape border: '{plain[:30]}' ({tid}) vs {bid}")
    warns += ov[:max_items]
    # alignment / spacing among sized, non-text shapes
    shapes = []
    for c in verts:
        tokens, kv = parse_style(c.get("style"))
        g = ag(c.get("id"))
        if not g or "group" in tokens or "text" in tokens or kv.get("shape") == "image":
            continue
        if g[2] >= 14 and g[3] >= 14:
            shapes.append((c.get("id"), g))
    near = []
    keys = {"left": (lambda g: g[0], 1), "right": (lambda g: g[0] + g[2], 1), "center-x": (lambda g: g[0] + g[2] / 2, 1),
            "top": (lambda g: g[1], 0), "bottom": (lambda g: g[1] + g[3], 0), "center-y": (lambda g: g[1] + g[3] / 2, 0)}

    def contains(a, b):
        return a[0] <= b[0] and a[1] <= b[1] and a[0] + a[2] >= b[0] + b[2] and a[1] + a[3] >= b[1] + b[3]

    def axis_gap(a, b, vertical):  # gap between boxes along the other axis
        if vertical:
            return max(0, max(a[1], b[1]) - min(a[1] + a[3], b[1] + b[3]))
        return max(0, max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2]))
    for name, (fn, vertical) in keys.items():
        for i in range(len(shapes)):
            for j in range(i + 1, len(shapes)):
                (ia, a), (ib, b) = shapes[i], shapes[j]
                dv = abs(fn(a) - fn(b))
                ratio = max(a[2], a[3]) / max(1, max(b[2], b[3]))
                if (0.6 < dv <= 3 and 1 / 3 <= ratio <= 3 and not contains(a, b) and not contains(b, a)
                        and axis_gap(a, b, vertical) < 120):
                    near.append(f"near-miss {name}: {ia}={fn(a):.1f} vs {ib}={fn(b):.1f} (snap to one value)")
    warns += near[:max_items]
    # rows: same top & height
    rows = defaultdict(list)
    for i, g in shapes:
        rows[(round(g[1] / 4), round(g[3] / 4))].append((i, g))
    for key, items in rows.items():
        if len(items) < 3:
            continue
        items.sort(key=lambda t: t[1][0])
        ws = [g[2] for _, g in items]
        gaps = [b[1][0] - (a[1][0] + a[1][2]) for a, b in zip(items, items[1:])]
        if max(ws) - min(ws) > 1 and max(ws) - min(ws) <= 0.08 * max(ws):
            warns.append(f"row {[i for i, _ in items]}: widths {['%.0f' % w for w in ws]} differ slightly — make equal")
        if all(g > 0 for g in gaps) and max(gaps) - min(gaps) > 3 and max(gaps) - min(gaps) <= 0.6 * max(gaps):
            warns.append(f"row {[i for i, _ in items]}: uneven gaps {['%.0f' % g for g in gaps]} — distribute evenly")
    if th:
        for v, n in off_colors.most_common(max_items):
            warns.append(f"off-theme color {v} used {n}x")
        for v, n in off_fonts.most_common():
            warns.append(f"off-theme font family '{v}' used {n}x")
        for v, n in off_sizes.most_common(max_items):
            warns.append(f"off-theme font size {v} used {n}x (theme sizes x scale: {sorted(sizes)})")
    infos.append(f"{len(verts)} vertices ({dict(kinds)}), {len(edges)} edges")
    return dict(errors=errors, warnings=warns, info=infos)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("drawio"); ap.add_argument("--theme"); ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--json"); ap.add_argument("--max", type=int, default=25)
    a = ap.parse_args()
    r = lint(a.drawio, a.theme, a.scale, a.max)
    for e in r["errors"]:
        print("ERROR ", e)
    for w in r["warnings"]:
        print("WARN  ", w)
    for i in r["info"]:
        print("INFO  ", i)
    print(f"-- {len(r['errors'])} errors, {len(r['warnings'])} warnings")
    if a.json:
        json.dump(r, open(a.json, "w", encoding="utf-8"), indent=1)
    sys.exit(1 if r["errors"] else 0)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/new_build.py'] = r'''#!/usr/bin/env python3
"""
new_build.py — scaffold a work directory for converting one image to .drawio.

    python new_build.py ./images/test_fig1.png [--theme my_theme.yaml] [--scale 1.0] [--no-inspect]

Creates (next to the image):
    images/test_fig1_drawio_work/
        config.json     source/out/theme/scale paths (read by build.py and iterate.py)
        theme.yaml      copy of --theme (or the skill default) — edit roles/fonts here
        build.py        the build script you edit (from assets/build_template_image.py)
        NOTES.md        element inventory / plan / decisions (fill it in)
        inspect/        grids, zoom crops, palette, boxes, OCR, summary.md
        iterations/     previews, comparisons, lint reports, .drawio snapshots per iteration
Final output: images/test_fig1.drawio  (same path, .drawio extension)
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)


def scaffold_layout(work):
    boxes_path = os.path.join(work, "inspect", "source.boxes.json")
    ocr_path = os.path.join(work, "inspect", "source.ocr.json")
    if not os.path.exists(boxes_path):
        return ""
    try:
        boxes_data = json.load(open(boxes_path, encoding="utf-8"))
        comps = boxes_data.get("components", [])
    except Exception:
        return ""
    ocr_items = []
    if os.path.exists(ocr_path):
        try:
            ocr_items = json.load(open(ocr_path, encoding="utf-8"))
        except Exception:
            pass

    valid = [c for c in comps if c.get("kind") in ("filled box", "outline/frame", "mark/shape") and c["w"] >= 24 and c["h"] >= 16]
    if not valid:
        return ""

    valid.sort(key=lambda c: -c["w"] * c["h"])
    lines = [
        "    # ------------------------------------------------------------------ DETECTED COMPONENTS (Auto-scaffolded)",
        "    # Initial layout anchors detected from source inspection (coordinates in source px).",
        "    # Review/refine against inspect/source.boxes.png and inspect/source.grid*.png:",
    ]

    for c in valid[:35]:
        cid = c["id"]
        x0, y0, x1, y1 = c["bbox"]
        w, h = c["w"], c["h"]
        matched_text = []
        for o in ocr_items:
            ox0, oy0, ox1, oy1 = o["bbox"]
            cx, cy = (ox0 + ox1) / 2, (oy0 + oy1) / 2
            if x0 <= cx <= x1 and y0 <= cy <= y1:
                matched_text.append(o["text"])
        label = "\\n".join(matched_text).replace('"', '\\"') if matched_text else ""
        fill = "background"  # scaffold uses only built-in roles (was: undefined panel_bg)
        stroke = "neutral" if w * h > 30000 else "primary"
        role_label = f'"{label}"' if label else '""'
        lines.append(f'    d.rect({x0}, {y0}, {w}, {h}, label={role_label}, fill="{fill}", stroke="{stroke}", id="comp_{cid}")')

    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("image")
    ap.add_argument("--theme", help="YAML/TOML/JSON theme to start from (default: assets/theme_image.yaml)")
    ap.add_argument("--scale", type=float, default=1.0, help="output scale vs source px (1.0 = same size)")
    ap.add_argument("--workdir", help="override work dir location")
    ap.add_argument("--force", action="store_true", help="overwrite an existing build.py")
    ap.add_argument("--no-inspect", action="store_true")
    a = ap.parse_args()

    image = os.path.abspath(a.image)
    if not os.path.exists(image):
        sys.exit(f"image not found: {image}")
    stem = os.path.splitext(image)[0]
    out = stem + ".drawio"
    work = os.path.abspath(a.workdir or stem + "_drawio_work")
    for sub in ("", "inspect", "iterations", "crops", "scratch"):
        os.makedirs(os.path.join(work, sub), exist_ok=True)

    theme_src = a.theme or os.path.join(SKILL, "assets", "theme_image.yaml")
    theme_dst = os.path.join(work, "theme" + os.path.splitext(theme_src)[1])
    if not os.path.exists(theme_dst) or a.theme:
        shutil.copy(theme_src, theme_dst)

    cfg = dict(source=image, out=out, theme=theme_dst, scale=a.scale, scripts=HERE, workdir=work)
    json.dump(cfg, open(os.path.join(work, "config.json"), "w", encoding="utf-8"), indent=1)

    if not a.no_inspect:
        subprocess.run([sys.executable, os.path.join(HERE, "inspect_image.py"), "all", image,
                        "--out-dir", os.path.join(work, "inspect"), "--theme", theme_dst], check=False)

    build = os.path.join(work, "build.py")
    if not os.path.exists(build) or a.force:
        t = open(os.path.join(SKILL, "assets", "build_template_image.py"), encoding="utf-8").read()
        scaffold = scaffold_layout(work)
        if scaffold:
            example_marker = "    # Examples (delete when you start):"
            if example_marker in t:
                parts = t.split(example_marker)
                # Keep up to the example marker, insert scaffold, then skip down to 'return d'
                tail = parts[1].split("    return d")[-1]
                t = parts[0] + scaffold + "\n\n    return d" + tail
        t = (t.replace("__IMAGE_NAME__", os.path.basename(image)).replace("__OUT_NAME__", os.path.basename(out))
              .replace("__SCRIPTS__", HERE).replace("__WORKDIR__", work))
        open(build, "w", encoding="utf-8").write(t)
    notes = os.path.join(work, "NOTES.md")
    if not os.path.exists(notes):
        open(notes, "w", encoding="utf-8").write(
            f"# Conversion notes: {os.path.basename(image)}\n\n"
            "## Complexity tier\n(1 simple boxes | 2 tile grid/dashboard | 3 plot/geometric | 4 complex/3-D/photos)\n\n"
            "## Theme mapping (source color -> role)\n\n"
            "## Font sizes (measured px -> size role)\n\n"
            "## Element inventory (region by region, z-order)\n- [ ] \n\n"
            "## Native vs embedded decisions\n\n"
            "## Iteration log\n")
    print(f"\nwork dir : {work}\nbuild    : {build}\ntheme    : {theme_dst}\noutput   : {out}")
    print(f"next     : read {os.path.join(work, 'inspect', 'summary.md')}, view the grid images, fill NOTES.md,\n"
          f"           edit build.py, then: python {os.path.join(HERE, 'iterate.py')} {work} --note 'first pass'")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/new_figures.py'] = r'''#!/usr/bin/env python3
"""
new_figures.py — set up output folders for turning prose into three figure prototypes per passage.

    python new_figures.py ./drafts/methods.md                     # -> ./drafts/methods_figures/
    python new_figures.py --text "We propose ..." [--name rag]    # -> ./figures/<name or slug>/
    python new_figures.py paper.tex --split --passages 2,4        # several passages -> subfolders
    python new_figures.py notes.md --out-dir ./paper/figs --theme my_theme.yaml

Output location rules (the user can always override with --out-dir):
    input file  <dir>/<stem>.<ext>   ->  <dir>/<stem>_figures/
    pasted text                      ->  ./figures/<name or slug-of-first-words>/
    one passage   -> files directly in that folder
    several       -> one subfolder per passage: p1_<slug>/, p2_<slug>/, ...

Each passage folder gets:
    passage.md      the exact source text (the only allowed source of figure content)
    NOTES.md        claims inventory + chosen figure families (fill in before building)
    build.py        v1 / v2 / v3 builders (from assets/build_template_text.py, families pre-filled)
    _work/          config.json, analysis, iteration snapshots, lint/faithfulness reports
and the output root gets theme.yaml (shared by all passages) and index.md.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import re
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from analyze_text import analyze, read_input, slugify, to_markdown  # noqa: E402


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", nargs="?", help=".md .txt .tex (.docx/.pdf with optional libs)")
    ap.add_argument("--text", help="inline text instead of a file ('-' = stdin)")
    ap.add_argument("--name", help="folder name for pasted text (default: slug of the first words)")
    ap.add_argument("--out-dir", help="explicit output folder (overrides the defaults)")
    ap.add_argument("--theme", help="theme file to copy (default: assets/theme.yaml); kept if already present")
    ap.add_argument("--split", action="store_true", help="one passage per paragraph/section")
    ap.add_argument("--whole", action="store_true", help="treat the whole input as one passage")
    ap.add_argument("--passages", default="all", help="'all' or comma list of passage numbers, e.g. 1,3")
    ap.add_argument("--min-words", type=int, default=40)
    ap.add_argument("--force", action="store_true", help="overwrite existing build.py / NOTES.md")
    a = ap.parse_args()

    if a.text == "-":
        a.text = sys.stdin.read()
    if not a.path and a.text is None:
        sys.exit("give a file path or --text")
    text = read_input(a.path, a.text)
    passages = analyze(text, True if a.split else False if a.whole else None, a.min_words)
    if a.passages != "all":
        keep = {int(x) for x in a.passages.split(",")}
        passages = [p for p in passages if p["index"] in keep]
    if not passages:
        sys.exit("no passages selected (try --whole or a lower --min-words)")

    if a.out_dir:
        root = os.path.abspath(a.out_dir)
    elif a.path:
        stem = os.path.splitext(os.path.abspath(a.path))[0]
        root = stem + "_figures"
    else:
        root = os.path.abspath(os.path.join("figures", a.name or slugify(text, 5)))
    os.makedirs(root, exist_ok=True)

    theme_dst = os.path.join(root, "theme" + os.path.splitext(a.theme or "x.yaml")[1])
    if a.theme or not os.path.exists(theme_dst):
        shutil.copy(a.theme or os.path.join(SKILL, "assets", "theme.yaml"), theme_dst)

    template = open(os.path.join(SKILL, "assets", "build_template_text.py"), encoding="utf-8").read()
    index = [f"# Figure prototypes", "", f"source: `{a.path or '(pasted text)'}`  ", f"theme: `{theme_dst}`", ""]
    flat = len(passages) == 1
    for p in passages:
        pdir = root if flat else os.path.join(root, f"p{p['index']}_{p['slug']}")
        work = os.path.join(pdir, "_work")
        os.makedirs(work, exist_ok=True)
        open(os.path.join(pdir, "passage.md"), "w", encoding="utf-8").write((f"# {p['heading']}\n\n" if p["heading"] else "") + p["text"] + "\n")
        fams = [f for f, _ in p["suggestions"]] + ["figure"] * 3
        cfg = dict(scripts=HERE, theme=theme_dst, passage=os.path.join(pdir, "passage.md"), passage_dir=pdir,
                   suggestions=p["suggestions"], preview_scale=1.5)
        json.dump(cfg, open(os.path.join(work, "config.json"), "w", encoding="utf-8"), indent=1)
        json.dump({k: v for k, v in p.items() if k != "text"}, open(os.path.join(work, "analysis.json"), "w", encoding="utf-8"), indent=1)
        b = os.path.join(pdir, "build.py")
        if a.force or not os.path.exists(b):
            t = (template.replace("__TITLE__", p["heading"] or p["slug"]).replace("__SCRIPTS__", HERE)
                 .replace("__PASSAGE_DIR__", pdir).replace("__FAMILY1__", fams[0]).replace("__FAMILY2__", fams[1])
                 .replace("__FAMILY3__", fams[2]))
            open(b, "w", encoding="utf-8").write(t)
        n = os.path.join(pdir, "NOTES.md")
        if a.force or not os.path.exists(n):
            eqs = "\n".join(f"- `{e}`" for e in p["display_equations"]) or "- (none)"
            open(n, "w", encoding="utf-8").write(
                f"# Notes — {p['heading'] or p['slug']}\n\n"
                "## Claims inventory (every figure label must trace to one of these)\n"
                "| # | element (entity / step / relation / quantity / conclusion) | quote from passage |\n|---|---|---|\n| 1 | | |\n\n"
                f"## Equations\n{eqs}\n\n"
                "## Prototype plan (three DIFFERENT figure families)\n"
                + "".join(f"- v{i + 1}: **{f}** — {why}\n" for i, (f, why) in enumerate(p["suggestions"]))
                + "\n## Inferred / added content (must also be listed in build.py INFERRED)\n- \n\n## Iteration log\n")
        index.append(f"- **{os.path.relpath(pdir, root) if not flat else '.'}** — {p['words']} words; suggested: "
                     + ", ".join(f for f, _ in p["suggestions"]))
    open(os.path.join(root, "index.md"), "w", encoding="utf-8").write("\n".join(index) + "\n\n" + to_markdown(passages))
    print(to_markdown(passages))
    print(f"output root : {root}\ntheme       : {theme_dst}")
    for p in passages:
        pdir = root if flat else os.path.join(root, f"p{p['index']}_{p['slug']}")
        print(f"passage {p['index']:>2}  : {pdir}")
    print(f"next        : fill NOTES.md, write v1-v3 in build.py, then python {os.path.join(HERE, 'prototype.py')} <passage dir>")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/prototype.py'] = r'''#!/usr/bin/env python3
"""
prototype.py — build all three prototypes of a passage and check them.

    python prototype.py PASSAGE_DIR [--only v2] [--note "what changed"]

Steps: import PASSAGE_DIR/build.py -> for each VARIANTS entry create a Diagram of its canvas size,
call its builder, save  <key>_<name>.drawio  and a preview  <key>_<name>.png  -> lint (theme,
text fit, overlaps, alignment) -> faithfulness check (label words not found in passage.md) ->
prototypes.png contact sheet + prototypes.md summary -> snapshot in _work/iter_NN/ + history.json.

Exit status 1 if a builder raised or any lint ERROR occurred.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import html
import importlib.util
import json
import os
import re
import shutil
import sys
import traceback
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
sys.dont_write_bytecode = True   # keep passage folders free of __pycache__
SKILL = os.path.dirname(HERE)

FUNCTION_WORDS = set("""a an the and or of to in on for with by from as at is are be via per vs not no all each
its it this that these those than then into over under between within across e.g i.e etc""".split())


def load_cfg(pdir):
    p = os.path.join(pdir, "_work", "config.json")
    if os.path.exists(p):
        return json.load(open(p))
    theme = next((t for t in (os.path.join(pdir, "theme.yaml"), os.path.join(os.path.dirname(pdir), "theme.yaml"),
                              os.path.join(SKILL, "assets", "theme.yaml")) if os.path.exists(t)), None)
    return dict(scripts=HERE, theme=theme, passage=os.path.join(pdir, "passage.md"), passage_dir=pdir, preview_scale=1.5)


def tokens(s):
    s = s.replace("\u2011", "-").replace("\u2013", "-")
    out = []
    for t in re.findall(r"[A-Za-z][A-Za-z'\-]*|\d[\d,]*(?:\.\d+)?", s):
        for part in t.split("-"):
            part = part.strip("'").lower().replace(",", "") if part[:1].isdigit() else part.strip("'").lower()
            if part:
                out.append(part)
    return out


def latex_words(latex):
    inner = " ".join(re.findall(r"\\(?:mathrm|text|operatorname|mathit|textbf|mathbf)\{([^}]*)\}", latex))
    rest = re.sub(r"\\[A-Za-z]+", " ", latex)
    return inner + " " + rest


def label_texts(drawio_path):
    root = ET.parse(drawio_path).getroot()
    out = []
    for c in root.iter("mxCell"):
        v = c.get("value") or ""
        if not v:
            continue
        plain = html.unescape(re.sub(r"<br\s*/?>", "\n", v))
        plain = re.sub(r"<[^>]+>", "", plain)
        m = re.fullmatch(r"\s*\$\$(.+)\$\$\s*", plain, re.S)
        out.append((c.get("id"), latex_words(m.group(1)) if m else plain, plain))
    return out


def stem(w):
    for suf, rep in (("ations", ""), ("ation", ""), ("ings", ""), ("ing", ""), ("ions", ""), ("ion", ""),
                     ("ies", "y"), ("ied", "y"), ("es", ""), ("ed", ""), ("s", ""), ("ly", ""), ("al", "")):
        if w.endswith(suf) and len(w) - len(suf) >= 3:
            return w[: len(w) - len(suf)] + rep
    return w


def faithfulness(drawio_path, passage_text, allowed):
    """Label words that do not occur in the passage (light stemming; acronyms match their expansion)."""
    raw = passage_text + " " + latex_words(passage_text)
    src_list = tokens(raw)
    src = set(src_list)
    stems = {stem(t) for t in src}
    initials = "".join(t[0] for t in src_list if t[0].isalpha())
    allowed = {a.lower() for a in allowed} | FUNCTION_WORDS

    def ok(w):
        if w in src or w in allowed or len(w) <= 2:
            return True
        if w[0].isdigit():
            return w in src
        if stem(w) in stems:
            return True
        if w.isalpha() and 2 <= len(w) <= 6 and w in initials:     # acronym, e.g. llm <- large language model
            return True
        k = min(len(w), 6)
        return any(t[:k] == w[:k] for t in src if len(t) >= k)
    missing = {}
    for cid, text, plain in label_texts(drawio_path):
        for w in tokens(text):
            if not ok(w):
                missing.setdefault(w, set()).add(cid)
    return missing


def contact_sheet(items, title, out_path):
    from PIL import Image, ImageDraw, ImageFont
    sys.path.insert(0, HERE)
    from drawio_kit import find_font_file
    B = ImageFont.truetype(find_font_file("Helvetica", bold=True), 30)
    R = ImageFont.truetype(find_font_file("Helvetica"), 22)
    W = 1500
    tiles = []
    for png, lab, about in items:
        im = Image.open(png).convert("RGB")
        im = im.resize((W - 60, int(im.height * (W - 60) / im.width)), Image.LANCZOS)
        tiles.append((im, lab, about))
    H = 90 + sum(t[0].height + 110 for t in tiles)
    sheet = Image.new("RGB", (W, H), "#eef1f5")
    dr = ImageDraw.Draw(sheet)
    dr.text((30, 28), title, font=B, fill="#19355d")
    y = 90
    for im, lab, about in tiles:
        dr.text((30, y), lab, font=B, fill="#19355d")
        dr.text((30, y + 40), about[:150], font=R, fill="#333333")
        y += 80
        sheet.paste(im, (30, y)); dr.rectangle([30, y, 30 + im.width, y + im.height], outline="#a4a4a4", width=2)
        y += im.height + 30
    sheet.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("passage_dir"); ap.add_argument("--only", help="comma list of variant keys, e.g. v2")
    ap.add_argument("--note", default=""); ap.add_argument("--no-sheet", action="store_true")
    a = ap.parse_args()
    pdir = os.path.abspath(a.passage_dir)
    cfg = load_cfg(pdir)
    sys.path.insert(0, cfg.get("scripts", HERE))
    from drawio_kit import Diagram, Theme
    from lint_drawio import lint
    from render_preview import render_kit

    spec = importlib.util.spec_from_file_location("build_" + str(abs(hash(pdir))), os.path.join(pdir, "build.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    variants = mod.VARIANTS
    if a.only:
        keep = set(a.only.split(","))
        variants = [v for v in variants if v["key"] in keep]
    theme = Theme.load(cfg["theme"])
    passage = open(cfg["passage"]).read() if os.path.exists(cfg["passage"]) else ""
    allowed = getattr(mod, "ALLOWED_TERMS", [])
    inferred = getattr(mod, "INFERRED", [])
    ignore = getattr(mod, "LINT_IGNORE", [])     # substrings of warnings deliberately accepted

    work = os.path.join(pdir, "_work")
    hist_p = os.path.join(work, "history.json")
    hist = json.load(open(hist_p)) if os.path.exists(hist_p) else []
    it = len(hist) + 1
    snap = os.path.join(work, f"iter_{it:02d}")
    os.makedirs(snap, exist_ok=True)

    results, failed = [], False
    for v in variants:
        base = f"{v['key']}_{v['name']}"
        out = os.path.join(pdir, base + ".drawio")
        png = os.path.join(pdir, base + ".png")
        try:
            d = Diagram(*v["size"], theme=theme, scale=v.get("scale", 1.0))
            v["build"](d)
            d.save(out, quiet=True)
        except Exception:
            failed = True
            print(f"[{v['key']}] BUILD FAILED\n" + traceback.format_exc())
            continue
        render_kit(out, png, scale=cfg.get("preview_scale", 1.5))
        L = lint(out, cfg["theme"], v.get("scale", 1.0))
        kept = [w for w in L["warnings"] if not any(k in w for k in ignore)]
        L["ignored"] = len(L["warnings"]) - len(kept)
        L["warnings"] = kept
        miss = faithfulness(out, passage, allowed)
        shutil.copy(out, snap)
        results.append(dict(v=v, out=out, png=png, lint=L, missing=miss, off_theme=sorted(theme.off_theme)))
        theme.off_theme.clear()
        failed |= bool(L["errors"])

    for r in results:
        v = r["v"]
        print(f"\n[{v['key']}] {v['name']}  ->  {os.path.relpath(r['png'], os.getcwd())}")
        for e in r["lint"]["errors"]:
            print("  ERROR", e)
        for w in r["lint"]["warnings"][:12]:
            print("  WARN ", w)
        if r["lint"].get("ignored"):
            print(f"  ({r['lint']['ignored']} warning(s) matched LINT_IGNORE)")
        if len(r["lint"]["warnings"]) > 12:
            print(f"  ... {len(r['lint']['warnings']) - 12} more warnings")
        if r["missing"]:
            print("  NOT IN PASSAGE:", ", ".join(f"{w} ({', '.join(sorted(ids))[:40]})" for w, ids in sorted(r["missing"].items())))
            print("    -> reword with the passage's terms, add to ALLOWED_TERMS (structural words), or list in INFERRED")
    if inferred:
        print("\nINFERRED (report to the user):")
        for x in inferred:
            print("  -", x)

    if results and not a.no_sheet and not a.only:
        title = os.path.basename(pdir)
        m = re.search(r"^#\s+(.+)$", passage, re.M)
        if m:
            title = m.group(1)
        contact_sheet([(r["png"], f"{r['v']['key']} · {r['v']['name'].replace('_', ' ')}", r["v"].get("about", ""))
                       for r in results], f"Prototypes — {title}", os.path.join(pdir, "prototypes.png"))
        md = [f"# Prototypes — {title}", "", "![prototypes](prototypes.png)", ""]
        for r in results:
            v = r["v"]
            md += [f"## {v['key']} · {v['name'].replace('_', ' ')}", "", f"- file: `{os.path.basename(r['out'])}`",
                   f"- about: {v.get('about', '')}", f"- draft caption: {v.get('caption', '') or '—'}",
                   f"- lint: {len(r['lint']['errors'])} errors, {len(r['lint']['warnings'])} warnings",
                   f"- words not in passage: {', '.join(sorted(r['missing'])) or 'none'}", ""]
        md += ["## Inferred content", ""] + [f"- {x}" for x in inferred or ["none"]]
        open(os.path.join(pdir, "prototypes.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")
        print(f"\nVIEW: {os.path.join(pdir, 'prototypes.png')}")

    hist.append(dict(iter=it, note=a.note, variants={r["v"]["key"]: dict(errors=len(r["lint"]["errors"]),
                     warnings=len(r["lint"]["warnings"]), not_in_passage=len(r["missing"])) for r in results}))
    json.dump(hist, open(hist_p, "w", encoding="utf-8"), indent=1)
    print("\niter  " + "  ".join(f"{k}(err/warn/miss)" for k in sorted(hist[-1]["variants"])) + "  note")
    for h in hist[-6:]:
        print(f"{h['iter']:>4}  " + "  ".join(f"{v['errors']}/{v['warnings']}/{v['not_in_passage']:<9}" for _, v in sorted(h["variants"].items()))
              + f"  {h['note']}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/rebuild.py'] = r'''#!/usr/bin/env python3
"""
rebuild.py — run the edited copy of the original build script and write the enhanced figure.

    python rebuild.py WORK_DIR [--passes all] [--no-check]

WORK_DIR is the <stem>_enhance_work/ folder made by start_enhance.py. It runs build_enhanced.py
with the enhancement theme (work/theme.yaml), writes brief.figure.output (<stem>#ENHANCED.drawio),
optionally runs enhance.py passes on the result, then checks it:
  text origin  -> faithfulness against passage.md (words not in the passage)
  image origin -> compare.py against the source image (only meaningful when fidelity = exact/polish)
and always the academic audit (after_audit.md + after_print_preview.png).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import importlib.util
import json
import os
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.dont_write_bytecode = True
from audit_drawio import audit, load_yaml, print_preview, write_report  # noqa: E402


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work"); ap.add_argument("--passes", default="",
                    help="'all' or comma list for enhance.py, e.g. fonts,sizes,align,fit-page=16")
    ap.add_argument("--no-check", action="store_true")
    a = ap.parse_args()
    work = os.path.abspath(a.work)
    ctx = json.load(open(os.path.join(work, "context.json"), encoding="utf-8"))
    brief = load_yaml(os.path.join(work, "brief.yaml"))
    out = brief["figure"]["output"]
    theme = brief.get("style", {}).get("theme") or ctx["theme"]
    build = os.path.join(work, "build_enhanced.py")
    if not os.path.exists(build):
        sys.exit("no build_enhanced.py — this figure uses the edit route (drawio_doc + enhance.py)")
    if ctx["origin"] == "image-to-drawio":
        cfgp = os.path.join(work, "config.json")
        cfg = json.load(open(cfgp, encoding="utf-8")); cfg.update(out=out, theme=theme, scripts=HERE)
        json.dump(cfg, open(cfgp, "w", encoding="utf-8"), indent=1)
        r = subprocess.run([sys.executable, build], capture_output=True, text=True)
        print(r.stdout.strip()); 
        if r.returncode:
            print(r.stderr); sys.exit("build_enhanced.py failed")
    else:
        from drawio_kit import Diagram, Theme
        spec = importlib.util.spec_from_file_location("build_enhanced", build)
        m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
        key = brief.get("figure", {}).get("variant") or ctx.get("variant_key") or m.VARIANTS[0]["key"]
        v = next(v for v in m.VARIANTS if v["key"] == key)
        size = tuple(brief.get("figure", {}).get("canvas") or v["size"])
        d = Diagram(*size, theme=Theme.load(theme), scale=v.get("scale", 1.0))
        v["build"](d)
        d.save(out)
        globals()["_module"] = m
    # fidelity is measured on the raw rebuild: later passes (fit-page) move content
    if not a.no_check and ctx["origin"] == "image-to-drawio" and ctx.get("source_image") and brief.get("content", {}).get("fidelity", "polish") in ("exact", "polish"):
        from compare import compare
        m = compare(ctx["source_image"], out, os.path.join(work, "compare"), "after", quiet=True)
        print(f"fidelity vs source image: score {m['score']:.3f} (pixel {m['pixel_similarity']:.3f}, edge F1 {m['edge_f1']:.3f})")
    if a.passes:
        flags = []
        for p in (["all"] if a.passes == "all" else a.passes.split(",")):
            k, _, val = p.strip().replace("_", "-").partition("=")
            flags += ["--" + k] + ([val] if val else ["16"] if k == "fit-page" else [])
        subprocess.run([sys.executable, os.path.join(HERE, "enhance.py"), out, "-o", out, "--theme", theme,
                        "--brief", os.path.join(work, "brief.yaml"), "--log", os.path.join(work, "CHANGES.md"),
                        "--update-theme", "--add-roles"] + flags, check=True)
    if a.no_check:
        return
    if ctx["origin"] == "text-to-drawio" and ctx.get("source_text") and os.path.exists(ctx["source_text"]):
        from faithful import faithfulness
        allowed = getattr(globals().get("_module"), "ALLOWED_TERMS", []) + list(brief.get("content", {}).get("allowed_terms", []) or [])
        miss = faithfulness(out, open(ctx["source_text"], encoding="utf-8").read(), allowed)
        print("NOT IN SOURCE TEXT:", ", ".join(sorted(miss)) or "none")
    accepted = list((brief.get("review") or {}).get("accepted_warnings", []) or []) + list(getattr(globals().get("_module"), "LINT_IGNORE", []))
    r = audit(out, brief.get("target"), theme, accepted, reference=ctx["figure"])
    write_report(r, work, "after")
    print_preview(out, work, "after", r["brief"]["width_in"])
    print(f"audit after: {r['score']}/100 ({r['n_issues']} issues) -> {os.path.join(work, 'after_audit.md')}")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/render_preview.py'] = r'''#!/usr/bin/env python3
"""
render_preview.py — rasterize a .drawio file to PNG so it can be compared with the source image.

    python render_preview.py fig.drawio -o preview.png              # page-aligned, 1 px per unit
    python render_preview.py fig.drawio -o preview.png --width 1408 # resize to match source width
    python render_preview.py fig.drawio -o preview.png --engine drawio   # use draw.io desktop CLI

Engines
  kit    (default) built-in PIL renderer. Page-aligned (pixel (0,0) == page origin), so the output
         overlays exactly on the source image. Approximates draw.io: rects/rounded, ellipses, stars,
         triangles, rhombus, text with HTML bold/italic/color/wrap, images (JPEG/PNG + the SVG
         polygons/polylines/circles/paths drawio_kit emits), straight/orthogonal/curved edges,
         arrowheads, dashes, flexArrow, opacity, groups. Unknown shapes fall back to rectangles.
  drawio the real draw.io desktop app (if installed; uses xvfb-run when there is no display).
         Exact rendering but cropped to the diagram bounds — use for a final visual check, not overlay.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import base64
import html
import io
import math
import os
import re
import shutil
import subprocess
import sys
import urllib.parse
import xml.etree.ElementTree as ET
import zlib

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from drawio_kit import catmull_rom, find_font_file, html_lines, render_math  # noqa: E402

from PIL import Image, ImageDraw, ImageFont  # noqa: E402

SS = 2  # supersampling factor for anti-aliasing


# ------------------------------------------------------------------ parsing
def load_model(path):
    root = ET.parse(path).getroot()
    if root.tag == "mxGraphModel":
        return root
    diag = root.find("diagram")
    model = diag.find("mxGraphModel")
    if model is None:  # compressed diagram
        data = base64.b64decode(diag.text.strip())
        xml = urllib.parse.unquote(zlib.decompress(data, -15).decode())
        model = ET.fromstring(xml)
    return model


def parse_style(s):
    tokens, kv = [], {}
    for part in (s or "").split(";"):
        if not part:
            continue
        if "=" in part:
            k, v = part.split("=", 1)
            kv[k] = v
        else:
            tokens.append(part)
    return tokens, kv


def hex2rgba(c, alpha=255):
    if not c or c == "none" or c == "default":
        return None
    c = c.strip()
    if c.startswith("#"):
        c = c[1:]
        if len(c) == 3:
            c = "".join(ch * 2 for ch in c)
        try:
            return (int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16), int(alpha))
        except ValueError:
            return None
    named = {"white": "ffffff", "black": "000000", "red": "ff0000", "gray": "808080", "grey": "808080"}
    return hex2rgba("#" + named[c.lower()], alpha) if c.lower() in named else None


def fnum(kv, k, d):
    try:
        return float(kv.get(k, d))
    except (TypeError, ValueError):
        return d


# ------------------------------------------------------------------ renderer
class Renderer:
    def __init__(self, model, scale=1.0):
        self.model = model
        self.k = scale * SS
        # draw.io also wraps cells in <UserObject>/<object> elements
        self.cells = []
        for o in model.find("root"):
            if o.tag == "mxCell":
                self.cells.append((o, o))
            else:
                c = o.find("mxCell")
                if c is not None:
                    c.set("id", o.get("id")); c.set("value", o.get("label", ""))
                    self.cells.append((c, o))
        self.by_id = {c.get("id"): c for c, _ in self.cells}
        self._abs = {}
        pw = fnum(model.attrib, "pageWidth", 850); ph = fnum(model.attrib, "pageHeight", 1100)
        self.W, self.H = int(math.ceil(pw * self.k)), int(math.ceil(ph * self.k))
        self.math = model.get("math") == "1"
        bg = hex2rgba(model.get("background")) or (255, 255, 255, 255)
        self.img = Image.new("RGB", (self.W, self.H), bg[:3])
        self.draw = ImageDraw.Draw(self.img, "RGBA")

    # ---- geometry
    def abs_geo(self, cid):
        if cid in self._abs:
            return self._abs[cid]
        c = self.by_id.get(cid)
        if c is None or c.find("mxGeometry") is None:
            self._abs[cid] = (0, 0, 0, 0); return self._abs[cid]
        g = c.find("mxGeometry")
        x, y = fnum(g.attrib, "x", 0), fnum(g.attrib, "y", 0)
        w, h = fnum(g.attrib, "width", 0), fnum(g.attrib, "height", 0)
        p = c.get("parent")
        if p and p not in ("0", "1") and p in self.by_id and self.by_id[p].get("vertex") == "1":
            px, py, _, _ = self.abs_geo(p)
            x, y = x + px, y + py
        self._abs[cid] = (x, y, w, h)
        return self._abs[cid]

    def K(self, v):
        return v * self.k

    # ---- helpers
    def color(self, kv, key, default=None, opacity_keys=("opacity",)):
        a = 255
        for ok in opacity_keys:
            if ok in kv:
                a *= fnum(kv, ok, 100) / 100
        return hex2rgba(kv.get(key, default), a)

    def stroke_w(self, kv):
        return max(1, int(round(fnum(kv, "strokeWidth", 1) * self.k)))

    def dashed_line(self, pts, fill, width, pattern):
        on, off = [float(p) * self.k * max(1, width / self.k) ** 0.5 for p in (pattern.split() + ["4"])[:2]]
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            L = math.dist((x1, y1), (x2, y2))
            if L == 0:
                continue
            ux, uy = (x2 - x1) / L, (y2 - y1) / L
            t, draw_on = 0.0, True
            while t < L:
                seg = min(on if draw_on else off, L - t)
                if draw_on:
                    self.draw.line([(x1 + ux * t, y1 + uy * t), (x1 + ux * (t + seg), y1 + uy * (t + seg))],
                                   fill=fill, width=width)
                t += seg; draw_on = not draw_on

    def polyline(self, pts, fill, width, dashed=False, pattern="3 3"):
        if fill is None or len(pts) < 2:
            return
        if dashed:
            self.dashed_line(pts, fill, width, pattern)
        else:
            self.draw.line(pts, fill=fill, width=width, joint="curve")

    # ---- shapes
    def shape_poly(self, kind, x, y, w, h, kv):
        if kind == "star":
            cx, cy, R = x + w / 2, y + h / 2, min(w, h) / 2
            r = R * 0.45
            return [(cx + (R if i % 2 == 0 else r) * math.sin(math.pi * i / 5),
                     cy + 0.06 * h - (R if i % 2 == 0 else r) * math.cos(math.pi * i / 5)) for i in range(10)]
        if kind == "triangle":
            d = kv.get("direction", "east")
            return {"north": [(x, y + h), (x + w / 2, y), (x + w, y + h)],
                    "south": [(x, y), (x + w, y), (x + w / 2, y + h)],
                    "west": [(x + w, y), (x + w, y + h), (x, y + h / 2)]}.get(d, [(x, y), (x + w, y + h / 2), (x, y + h)])
        if kind == "rhombus":
            return [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
        if kind == "hexagon":
            s = w * fnum(kv, "size", 0.25)
            return [(x + s, y), (x + w - s, y), (x + w, y + h / 2), (x + w - s, y + h), (x + s, y + h), (x, y + h / 2)]
        if kind == "parallelogram":
            s = w * fnum(kv, "size", 0.2)
            return [(x + s, y), (x + w, y), (x + w - s, y + h), (x, y + h)]
        if kind == "trapezoid":
            s = w * fnum(kv, "size", 0.2)
            return [(x + s, y), (x + w - s, y), (x + w, y + h), (x, y + h)]
        return None

    def draw_vertex(self, c):
        tokens, kv = parse_style(c.get("style"))
        x, y, w, h = self.abs_geo(c.get("id"))
        X, Y, Wd, Ht = self.K(x), self.K(y), self.K(w), self.K(h)
        shape = kv.get("shape") or next((t for t in tokens if t in ("ellipse", "rhombus", "triangle", "text",
                                                                    "image", "swimlane", "line")), "rect")
        if "mxgraph.basic.star" in (kv.get("shape") or ""):
            shape = "star"
        if "group" in tokens:
            return
        fill = self.color(kv, "fillColor", None, ("opacity", "fillOpacity"))
        stroke = self.color(kv, "strokeColor", "#000000", ("opacity", "strokeOpacity"))
        if "text" in tokens:
            fill = self.color(kv, "fillColor", None, ("opacity", "fillOpacity")) if kv.get("fillColor") not in (None, "none") else None
            stroke = None if kv.get("strokeColor") in (None, "none") else stroke
        sw = self.stroke_w(kv)
        dashed = kv.get("dashed") == "1"
        if shape == "image":
            self.draw_image(kv, X, Y, Wd, Ht)
        elif shape in ("ellipse", "doubleEllipse") or "ellipse" in tokens:
            self.draw.ellipse([X, Y, X + Wd, Y + Ht], fill=fill, outline=None)
            if stroke and dashed:
                ring = [(X + Wd / 2 + Wd / 2 * math.cos(t / 72 * math.tau), Y + Ht / 2 + Ht / 2 * math.sin(t / 72 * math.tau)) for t in range(73)]
                self.polyline(ring, stroke, sw, True, kv.get("dashPattern", "3 3"))
            elif stroke:
                self.draw.ellipse([X, Y, X + Wd, Y + Ht], outline=stroke, width=sw)
        elif shape in ("star", "triangle", "rhombus", "hexagon", "parallelogram", "trapezoid") or \
                self.shape_poly(shape, X, Y, Wd, Ht, kv):
            pts = self.shape_poly(shape, X, Y, Wd, Ht, kv) or [(X, Y), (X + Wd, Y), (X + Wd, Y + Ht), (X, Y + Ht)]
            self.draw.polygon(pts, fill=fill)
            if stroke:
                self.polyline(pts + pts[:1], stroke, sw, dashed, kv.get("dashPattern", "3 3"))
        elif shape == "line":
            self.draw.line([(X, Y + Ht / 2), (X + Wd, Y + Ht / 2)], fill=stroke, width=sw)
        elif "text" not in tokens or fill or stroke:
            r = 0
            if kv.get("rounded") == "1":
                r = (fnum(kv, "arcSize", 16) / 2 * self.k if kv.get("absoluteArcSize") == "1"
                     else min(Wd, Ht) * fnum(kv, "arcSize", 15) / 100)
                r = max(0, min(r, (min(Wd, Ht) - 4) / 2))
            box = [X, Y, X + Wd, Y + Ht]
            g2 = self.color(kv, "gradientColor", None, ("opacity", "fillOpacity"))
            if fill and g2:   # gradient fill (south/north/east/west) drawn as strips inside a rounded mask
                horiz = kv.get("gradientDirection", "south") in ("east", "west")
                flip = kv.get("gradientDirection", "south") in ("north", "west")
                n = int(Wd if horiz else Ht) or 1
                layer = Image.new("RGBA", (max(1, int(Wd)), max(1, int(Ht))), (0, 0, 0, 0))
                ld = ImageDraw.Draw(layer)
                for i in range(n):
                    t = i / max(1, n - 1)
                    t = 1 - t if flip else t
                    col = tuple(int(fill[j] + (g2[j] - fill[j]) * t) for j in range(4))
                    ld.line([(i, 0), (i, Ht)] if horiz else [(0, i), (Wd, i)], fill=col)
                mask = Image.new("L", layer.size, 0)
                ImageDraw.Draw(mask).rounded_rectangle([0, 0, layer.width - 1, layer.height - 1], radius=r, fill=255)
                alpha = Image.composite(layer.split()[3], Image.new("L", layer.size, 0), mask)
                layer.putalpha(alpha)
                self.img.paste(layer, (int(X), int(Y)), layer)
            elif fill:
                self.draw.rounded_rectangle(box, radius=r, fill=fill)
            if stroke and kv.get("strokeColor") != "none":
                if dashed:
                    self.polyline([(X, Y), (X + Wd, Y), (X + Wd, Y + Ht), (X, Y + Ht), (X, Y)], stroke, sw, True,
                                  kv.get("dashPattern", "3 3"))
                else:
                    self.draw.rounded_rectangle(box, radius=r, outline=stroke, width=sw)
        self.draw_label(c.get("value") or "", kv, tokens, X, Y, Wd, Ht)

    # ---- images (raster + the SVG subset drawio_kit emits)
    def draw_image(self, kv, X, Y, Wd, Ht):
        src = kv.get("image", "")
        m = re.match(r"data:image/([\w+.-]+)[,;](?:base64,)?(.*)", src, re.S)
        if not m or Wd < 1 or Ht < 1:
            return
        kind, data = m.group(1), m.group(2)
        try:
            raw = base64.b64decode(data + "=" * (-len(data) % 4))
        except Exception:
            return
        if kind.startswith("svg"):
            self.draw_svg(raw.decode("utf-8", "replace"), X, Y, Wd, Ht, kv)
        else:
            try:
                im = Image.open(io.BytesIO(raw)).convert("RGBA").resize((max(1, int(Wd)), max(1, int(Ht))), Image.LANCZOS)
                self.img.paste(im, (int(X), int(Y)), im)
            except Exception:
                return
        if kv.get("imageBorder"):
            self.draw.rectangle([X, Y, X + Wd, Y + Ht], outline=hex2rgba(kv["imageBorder"]), width=max(1, int(self.k)))

    def draw_svg(self, svg, X, Y, Wd, Ht, kv):
        try:
            root = ET.fromstring(svg)
        except ET.ParseError:
            return
        vb = [float(v) for v in re.split(r"[ ,]+", root.get("viewBox", "").strip())] if root.get("viewBox") else \
            [0, 0, float(re.sub(r"[^\d.]", "", root.get("width", "100"))), float(re.sub(r"[^\d.]", "", root.get("height", "100")))]
        sx, sy = Wd / vb[2], Ht / vb[3]
        op = fnum(kv, "opacity", 100) / 100

        def T(px, py):
            return (X + (px - vb[0]) * sx, Y + (py - vb[1]) * sy)
        for el in root.iter():
            tag = el.tag.split("}")[-1]
            if tag not in ("polygon", "polyline", "circle", "path", "rect", "ellipse", "line"):
                continue
            fill = hex2rgba(el.get("fill", "#000000" if tag != "polyline" else "none"),
                            255 * float(el.get("fill-opacity", 1)) * float(el.get("opacity", 1)) * op)
            stroke = hex2rgba(el.get("stroke", "none"), 255 * float(el.get("stroke-opacity", 1)) * float(el.get("opacity", 1)) * op)
            sw = max(1, int(round(float(el.get("stroke-width", 1)) * sx)))
            dash = el.get("stroke-dasharray")
            polys = []
            if tag in ("polygon", "polyline"):
                nums = [float(v) for v in re.split(r"[ ,]+", el.get("points", "").strip()) if v]
                polys = [([T(nums[i], nums[i + 1]) for i in range(0, len(nums) - 1, 2)], tag == "polygon")]
            elif tag == "line":
                polys = [([T(float(el.get("x1", 0)), float(el.get("y1", 0))), T(float(el.get("x2", 0)), float(el.get("y2", 0)))], False)]
            elif tag == "rect":
                x0, y0 = float(el.get("x", 0)), float(el.get("y", 0)); w0, h0 = float(el.get("width", 0)), float(el.get("height", 0))
                polys = [([T(x0, y0), T(x0 + w0, y0), T(x0 + w0, y0 + h0), T(x0, y0 + h0)], True)]
            elif tag in ("circle", "ellipse"):
                cx, cy = float(el.get("cx", 0)), float(el.get("cy", 0))
                rx = float(el.get("r", el.get("rx", 0))); ry = float(el.get("r", el.get("ry", 0)))
                polys = [([T(cx + rx * math.cos(t / 24 * math.tau), cy + ry * math.sin(t / 24 * math.tau)) for t in range(24)], True)]
            elif tag == "path":
                polys = [(list(map(lambda p: T(*p), sp)), closed) for sp, closed in parse_path(el.get("d", ""))]
            for pts, closed in polys:
                if len(pts) < 2:
                    continue
                if closed and fill:
                    self.draw.polygon(pts, fill=fill)
                if stroke:
                    self.polyline(pts + (pts[:1] if closed else []), stroke, sw, bool(dash), (dash or "3 3").replace(",", " "))

    # ---- labels
    def draw_label(self, value, kv, tokens, X, Y, Wd, Ht):
        if not value:
            return
        if self.math:   # whole-label LaTeX ($$...$$ or \(...\)) -> mathtext image
            plain = html.unescape(re.sub(r"<[^>]+>", "", value)).strip()
            m = re.fullmatch(r"\$\$(.+)\$\$|\\\((.+)\\\)", plain, re.S)
            if m:
                col = kv.get("fontColor", "#000000")
                im = render_math(m.group(1) or m.group(2), fnum(kv, "fontSize", 11) * self.k, col if col.startswith("#") else "#000000")
                if im is not None:
                    al = kv.get("align", "center")
                    px = {"left": X, "right": X + Wd - im.width}.get(al, X + (Wd - im.width) / 2)
                    self.img.paste(im, (int(px), int(Y + (Ht - im.height) / 2)), im)
                    return
        is_html = kv.get("html") == "1"
        lines = html_lines(value) if is_html else html_lines(value, raw_text=True)
        fam = kv.get("fontFamily", "Helvetica")
        size = fnum(kv, "fontSize", 11)
        fstyle = int(fnum(kv, "fontStyle", 0))
        base_color = kv.get("fontColor", "#000000")
        alpha = 255 * fnum(kv, "textOpacity", 100) / 100 * fnum(kv, "opacity", 100) / 100
        spacing = fnum(kv, "spacing", 2)
        pad = {s: (spacing + fnum(kv, "spacing" + s.capitalize(), 0)) * self.k for s in ("left", "right", "top", "bottom")}
        if kv.get("verticalLabelPosition") == "bottom":
            Y, Ht = Y + Ht, size * 2 * self.k
        avail = Wd - pad["left"] - pad["right"]
        # segment fonts
        def seg_font(sg):
            b = bool(fstyle & 1) or sg.get("bold"); i = bool(fstyle & 2) or sg.get("italic")
            sz = (sg.get("size") or size) * self.k
            path = find_font_file(fam, b, i)
            try:
                return ImageFont.truetype(path, max(1, int(round(sz)))) if path else ImageFont.load_default()
            except Exception:
                return ImageFont.load_default()
        # word wrap
        if kv.get("whiteSpace") == "wrap" and avail > 0:
            wrapped = []
            for line in lines:
                cur, cur_w = [], 0.0
                words = []
                for sg in line:
                    for tok in re.split(r"(\s+)", sg["text"]):
                        if tok:
                            words.append(dict(sg, text=tok))
                for wd in words:
                    mm = re.fullmatch(r"(.*?)\\\((.+?)\\\)(.*)", wd["text"])
                    if mm:   # token containing inline math: measure the rendered math
                        mi = render_math(mm.group(2), (wd.get("size") or size) * self.k)
                        ww = (mi.width if mi is not None else 0) + seg_font(wd).getlength(mm.group(1) + mm.group(3))
                    else:
                        ww = seg_font(wd).getlength(wd["text"])
                    if cur and cur_w + ww > avail + 1 and not wd["text"].isspace():
                        wrapped.append(cur); cur, cur_w = [], 0.0
                    if not cur and wd["text"].isspace():
                        continue
                    cur.append(wd); cur_w += ww
                wrapped.append(cur)
            lines = wrapped
        lh = size * 1.2 * self.k
        total_h = lh * len(lines)
        va = kv.get("verticalAlign", "middle")
        top = {"top": Y + pad["top"], "bottom": Y + Ht - pad["bottom"] - total_h}.get(va, Y + (Ht - total_h) / 2)
        align = kv.get("align", "center")
        rot = fnum(kv, "rotation", 0)
        layer, ox, oy = (None, 0, 0)
        if rot:
            layer = Image.new("RGBA", (int(Wd) + 4, int(Ht) + 4), (0, 0, 0, 0))
            ldraw = ImageDraw.Draw(layer)
            ox, oy = X, Y
        else:
            ldraw = self.draw
        for i, line in enumerate(lines):
            # expand emoji into their own segments (no emoji font is assumed: drawn as placeholders)
            line = [dict(sg, text=run, emoji=e) for sg in line for e, run in split_emoji(sg["text"])]
            # inline LaTeX \( ... \) -> math segments rendered with mathtext
            ml = []
            for sg in line:
                parts = re.split(r"\\\((.+?)\\\)", sg["text"]) if not sg.get("emoji") else [sg["text"]]
                for pi, p in enumerate(parts):
                    if p:
                        ml.append(dict(sg, text=p, math=(pi % 2 == 1)))
            line = ml
            fonts = [seg_font(sg) for sg in line]
            mims = [render_math(sg["text"], (sg.get("size") or size) * self.k, sg.get("color") or (base_color if base_color.startswith("#") else "#000000"))
                    if sg.get("math") else None for sg in line]
            widths = [mi.width if mi is not None else f.size * 1.05 if sg.get("emoji") else f.getlength(sg["text"])
                      for f, sg, mi in zip(fonts, line, mims)]
            lw = sum(widths)
            x0 = {"left": X + pad["left"], "right": X + Wd - pad["right"] - lw}.get(align, X + (Wd - lw) / 2)
            yb = top + i * lh + lh / 2
            cx = x0
            for f, sg, wdt, mi in zip(fonts, line, widths, mims):
                col = hex2rgba(sg.get("color") or base_color, alpha) or (0, 0, 0, int(alpha))
                if mi is not None:
                    tgt = layer if rot and layer is not None else self.img
                    tgt.paste(mi, (int(cx - ox), int(yb - oy - mi.height / 2)), mi)
                    cx += wdt
                    continue
                if sg.get("emoji"):   # icon placeholder: soft rounded square the size of the glyph
                    r = f.size * 0.45
                    ldraw.rounded_rectangle([cx - ox + wdt / 2 - r, yb - oy - r, cx - ox + wdt / 2 + r, yb - oy + r],
                                            radius=r * 0.35, fill=(150, 160, 175, int(alpha * 0.45)),
                                            outline=(90, 100, 115, int(alpha * 0.8)), width=max(1, int(self.k)))
                    cx += wdt
                    continue
                ldraw.text((cx - ox, yb - oy), sg["text"], font=f, fill=col, anchor="lm")
                if (fstyle & 4) or sg.get("underline"):
                    ldraw.line([(cx - ox, yb - oy + f.size * 0.55), (cx - ox + wdt, yb - oy + f.size * 0.55)], fill=col, width=max(1, int(self.k)))
                cx += wdt
        if rot and layer is not None:
            layer = layer.rotate(-rot, expand=True, resample=Image.BICUBIC)
            cxp, cyp = X + Wd / 2, Y + Ht / 2
            self.img.paste(layer, (int(cxp - layer.width / 2), int(cyp - layer.height / 2)), layer)

    # ---- edges
    def endpoint(self, cid, toward, frac=None):
        x, y, w, h = self.abs_geo(cid)
        if frac is not None:
            return (x + w * frac[0], y + h * frac[1])
        cx, cy = x + w / 2, y + h / 2
        dx, dy = toward[0] - cx, toward[1] - cy
        if dx == 0 and dy == 0:
            return (cx, cy)
        tokens, kv = parse_style(self.by_id[cid].get("style"))
        if "ellipse" in tokens or "star" in (kv.get("shape") or ""):
            a = math.atan2(dy, dx)
            return (cx + w / 2 * math.cos(a), cy + h / 2 * math.sin(a))
        sx = (w / 2) / abs(dx) if dx else float("inf")
        sy = (h / 2) / abs(dy) if dy else float("inf")
        s = min(sx, sy)
        return (cx + dx * s, cy + dy * s)

    def draw_edge(self, c):
        tokens, kv = parse_style(c.get("style"))
        g = c.find("mxGeometry")
        if g is None:
            return
        sp = g.find("mxPoint[@as='sourcePoint']"); tp = g.find("mxPoint[@as='targetPoint']")
        arr = g.find("Array[@as='points']")
        way = [(fnum(p.attrib, "x", 0), fnum(p.attrib, "y", 0)) for p in arr] if arr is not None else []
        src, tgt = c.get("source"), c.get("target")

        def frac(prefix):
            return (fnum(kv, prefix + "X", 0.5), fnum(kv, prefix + "Y", 0.5)) if prefix + "X" in kv else None
        if src and src in self.by_id:
            sx, sy, sw_, sh_ = self.abs_geo(src)
            ref = way[0] if way else (self.abs_geo(tgt)[0] + self.abs_geo(tgt)[2] / 2, self.abs_geo(tgt)[1] + self.abs_geo(tgt)[3] / 2) if tgt in self.by_id else \
                ((fnum(tp.attrib, "x", 0), fnum(tp.attrib, "y", 0)) if tp is not None else (sx, sy))
            p0 = self.endpoint(src, ref, frac("exit"))
        else:
            p0 = (fnum(sp.attrib, "x", 0), fnum(sp.attrib, "y", 0)) if sp is not None else (0, 0)
        if tgt and tgt in self.by_id:
            ref = way[-1] if way else p0
            p1 = self.endpoint(tgt, ref, frac("entry"))
        else:
            p1 = (fnum(tp.attrib, "x", 0), fnum(tp.attrib, "y", 0)) if tp is not None else p0
        pts = [p0] + way + [p1]
        style = kv.get("edgeStyle", "")
        if kv.get("curved") == "1" and len(pts) > 2:
            pts = drawio_curve(pts)
        elif "orthogonal" in style or "elbow" in style:
            pts = orthogonalize(pts, kv)
        P = [(self.K(x), self.K(y)) for x, y in pts]
        stroke = self.color(kv, "strokeColor", "#000000")
        sw = self.stroke_w(kv)
        if kv.get("shape") == "flexArrow":
            self.draw_flex(P, kv)
            return
        end, start = kv.get("endArrow", "classic"), kv.get("startArrow", "none")
        P2 = list(P)
        if end not in ("none", "") and len(P2) >= 2:
            P2[-1] = self.arrowhead(P2[-2], P2[-1], end, kv.get("endFill", "1") != "0", fnum(kv, "endSize", 6), stroke, sw)
        if start not in ("none", "") and len(P2) >= 2:
            P2[0] = self.arrowhead(P2[1], P2[0], start, kv.get("startFill", "1") != "0", fnum(kv, "startSize", 6), stroke, sw)
        self.polyline(P2, stroke, sw, kv.get("dashed") == "1", kv.get("dashPattern", "3 3"))
        if c.get("value"):
            mid = P[len(P) // 2] if len(P) > 2 else ((P[0][0] + P[-1][0]) / 2, (P[0][1] + P[-1][1]) / 2)
            w = 200 * self.k
            self.draw_label(c.get("value"), dict(kv, align="center", verticalAlign="middle"), tokens,
                            mid[0] - w / 2, mid[1] - 10 * self.k, w, 20 * self.k)

    def arrowhead(self, a, b, kind, filled, size, color, sw):
        L = math.dist(a, b)
        if L == 0 or color is None:
            return b
        ux, uy = (b[0] - a[0]) / L, (b[1] - a[1]) / L
        ln = (size * 1.5) * self.k + sw * 0.5
        hw = (size * 0.6) * self.k + sw * 0.5
        base = (b[0] - ux * ln, b[1] - uy * ln)
        left = (base[0] - uy * hw, base[1] + ux * hw); right = (base[0] + uy * hw, base[1] - ux * hw)
        if kind == "oval":
            r = hw
            self.draw.ellipse([b[0] - r, b[1] - r, b[0] + r, b[1] + r], fill=color if filled else None, outline=color)
            return (b[0] - ux * r, b[1] - uy * r)
        if kind == "diamond":
            mid = (b[0] - ux * ln / 2, b[1] - uy * ln / 2)
            pts = [b, (mid[0] - uy * hw, mid[1] + ux * hw), base, (mid[0] + uy * hw, mid[1] - ux * hw)]
        elif kind == "classic":
            notch = (b[0] - ux * ln * 0.7, b[1] - uy * ln * 0.7)
            pts = [b, left, notch, right]; base = notch
        elif kind in ("open", "dash"):
            self.draw.line([left, b, right], fill=color, width=sw)
            return b
        else:
            pts = [b, left, right]
        self.draw.polygon(pts, fill=color if filled else None, outline=color)
        return base

    def draw_flex(self, P, kv):
        fill = self.color(kv, "fillColor", "#ffffff"); stroke = self.color(kv, "strokeColor", "#000000")
        width = fnum(kv, "width", 10) * self.k
        ew = fnum(kv, "endWidth", 20) * self.k
        es = fnum(kv, "endSize", 7) * self.k * 2
        if len(P) < 2:
            return
        b, a = P[-1], P[-2]
        L = math.dist(a, b) or 1
        ux, uy = (b[0] - a[0]) / L, (b[1] - a[1]) / L
        base = (b[0] - ux * es, b[1] - uy * es)
        body = P[:-1] + [base]
        self.draw.line(body, fill=fill, width=int(width), joint="curve")
        hw = (width + ew) / 2
        self.draw.polygon([b, (base[0] - uy * hw, base[1] + ux * hw), (base[0] + uy * hw, base[1] - ux * hw)], fill=fill, outline=stroke)

    # ---- main
    def render(self):
        for c, _ in self.cells:
            if c.get("vertex") == "1":
                self.draw_vertex(c)
            elif c.get("edge") == "1":
                self.draw_edge(c)
        return self.img.resize((self.W // SS, self.H // SS), Image.LANCZOS)


def drawio_curve(pts, n=14):
    """draw.io 'curved=1': waypoints act as quadratic CONTROL points; the curve passes through the
    midpoints between consecutive waypoints (and the two end points), not through the waypoints."""
    p0, way, p1 = pts[0], pts[1:-1], pts[-1]
    if not way:
        return pts
    anchors = [p0] + [((a[0] + b[0]) / 2, (a[1] + b[1]) / 2) for a, b in zip(way, way[1:])] + [p1]
    out = [p0]
    for i, c in enumerate(way):
        a, b = anchors[i], anchors[i + 1]
        for k in range(1, n + 1):
            t = k / n
            out.append(((1 - t) ** 2 * a[0] + 2 * (1 - t) * t * c[0] + t * t * b[0],
                        (1 - t) ** 2 * a[1] + 2 * (1 - t) * t * c[1] + t * t * b[1]))
    return out


def _is_emoji(ch):
    o = ord(ch)
    return o >= 0x1F000 or 0x2600 <= o <= 0x27BF or 0x2B00 <= o <= 0x2BFF or 0x1F900 <= o


def split_emoji(text):
    """Split into (is_emoji, run) pieces; variation selectors / ZWJ are dropped."""
    out = []
    for ch in text:
        if ch in "\ufe0f\u200d":
            continue
        e = _is_emoji(ch)
        if out and out[-1][0] == e and not e:
            out[-1] = (e, out[-1][1] + ch)
        else:
            out.append((e, ch))
    return out


def orthogonalize(pts, kv):
    """Approximate draw.io orthogonal routing: axis-aligned segments honoring exit/entry sides."""
    def side_dir(prefix):
        if prefix + "X" not in kv:
            return None
        fx, fy = fnum(kv, prefix + "X", 0.5), fnum(kv, prefix + "Y", 0.5)
        return "h" if fx in (0.0, 1.0) and 0 < fy < 1 else "v" if fy in (0.0, 1.0) else None
    p0, p1 = pts[0], pts[-1]
    if len(pts) == 2:
        d0 = side_dir("exit") or ("h" if abs(p1[0] - p0[0]) >= abs(p1[1] - p0[1]) else "v")
        d1 = side_dir("entry") or d0
        if p0[0] == p1[0] or p0[1] == p1[1]:
            return pts
        if d0 == "h" and d1 == "h":
            mx = (p0[0] + p1[0]) / 2
            return [p0, (mx, p0[1]), (mx, p1[1]), p1]
        if d0 == "v" and d1 == "v":
            my = (p0[1] + p1[1]) / 2
            return [p0, (p0[0], my), (p1[0], my), p1]
        return [p0, (p1[0], p0[1]), p1] if d0 == "h" else [p0, (p0[0], p1[1]), p1]
    out = [p0]
    horiz = (side_dir("exit") or "h") == "h"
    for p in pts[1:]:
        q = out[-1]
        if q[0] != p[0] and q[1] != p[1]:
            out.append((p[0], q[1]) if horiz else (q[0], p[1]))
        out.append(p)
        horiz = not horiz
    return out


def parse_path(d):
    """Minimal SVG path parser (M L H V C Q Z, abs + rel) -> list of (points, closed)."""
    toks = re.findall(r"[MLHVCQZmlhvcqz]|-?\d*\.?\d+(?:e-?\d+)?", d)
    paths, cur, pos, start, i, cmd = [], [], (0.0, 0.0), (0.0, 0.0), 0, None

    def num():
        nonlocal i
        v = float(toks[i]); i += 1; return v
    while i < len(toks):
        if re.match(r"[A-Za-z]", toks[i]):
            cmd = toks[i]; i += 1
            if cmd in "Zz":
                if cur: paths.append((cur, True)); cur = []
                pos = start; continue
        rel = cmd.islower()
        C = cmd.upper()
        ox, oy = pos if rel else (0, 0)
        if C == "M":
            if cur: paths.append((cur, False))
            pos = (ox + num(), oy + num()); start = pos; cur = [pos]; cmd = "l" if rel else "L"
        elif C == "L":
            pos = (ox + num(), oy + num()); cur.append(pos)
        elif C == "H":
            pos = ((pos[0] if rel else 0) + num(), pos[1]); cur.append(pos)
        elif C == "V":
            pos = (pos[0], (pos[1] if rel else 0) + num()); cur.append(pos)
        elif C in "CQ":
            n = 3 if C == "C" else 2
            ctrl = [(ox + num(), oy + num()) for _ in range(n)]
            p0 = pos
            for k in range(1, 13):
                t = k / 12
                if C == "C":
                    p = tuple((1 - t) ** 3 * p0[j] + 3 * (1 - t) ** 2 * t * ctrl[0][j] + 3 * (1 - t) * t * t * ctrl[1][j] + t ** 3 * ctrl[2][j] for j in (0, 1))
                else:
                    p = tuple((1 - t) ** 2 * p0[j] + 2 * (1 - t) * t * ctrl[0][j] + t * t * ctrl[1][j] for j in (0, 1))
                cur.append(p)
            pos = ctrl[-1]
        else:
            i += 1
    if cur:
        paths.append((cur, False))
    return paths


def render_kit(path, out, scale=1.0, width=None, dpi=None):
    model = load_model(path)
    if dpi and (scale == 1.0 or scale is None):
        scale = dpi / 72.0
    img = Renderer(model, scale).render()
    if width and img.width != width:
        img = img.resize((width, round(img.height * width / img.width)), Image.LANCZOS)
    import time
    for attempt in range(5):
        try:
            save_kw = {"dpi": (dpi, dpi)} if dpi else {}
            img.save(out, **save_kw)
            break
        except OSError:
            if attempt == 4:
                raise
            time.sleep(0.2)
    return out


def render_drawio_cli(path, out, scale=1.0):
    exe = shutil.which("drawio") or shutil.which("draw.io")
    if not exe and sys.platform.startswith("win"):
        cand = os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "draw.io", "draw.io.exe")
        if os.path.exists(cand):
            exe = cand
    if not exe:
        raise SystemExit("draw.io CLI not found (install draw.io desktop) — use --engine kit")
    cmd = [exe, "--export", "--format", "png", "--scale", str(scale), "--output", out, path, "--no-sandbox"]
    if not os.environ.get("DISPLAY") and shutil.which("xvfb-run"):
        cmd = ["xvfb-run", "-a"] + cmd
    subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("drawio")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--scale", type=float, default=1.0, help="output px per draw.io unit")
    ap.add_argument("--dpi", type=int, default=None, help="output DPI (e.g. 300 for print-ready PNG exports)")
    ap.add_argument("--width", type=int, default=None, help="resize output to this width (match source)")
    ap.add_argument("--engine", choices=["kit", "drawio"], default="kit")
    a = ap.parse_args()
    out = a.out or os.path.splitext(a.drawio)[0] + ".preview.png"
    if a.engine == "drawio":
        scale = (a.dpi / 72.0) if (a.dpi and a.scale == 1.0) else a.scale
        render_drawio_cli(a.drawio, out, scale)
    else:
        render_kit(a.drawio, out, a.scale, a.width, a.dpi)
    print("wrote", out)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/report.py'] = r'''#!/usr/bin/env python3
"""
report.py — final before/after report and exports for an enhanced figure.

    python report.py WORK_DIR [--export pdf,svg,png] [--dpi 300]

Writes WORK_DIR/before_after.png (screen previews + actual-print-size previews, with audit scores),
WORK_DIR/REPORT.md (brief summary, audit table before/after, CHANGES.md, remaining issues), and
exports next to the figure. Export names replace '#' with '_' (LaTeX \\includegraphics and URLs
break on '#'): <stem>_ENHANCED.pdf / .svg / .png. Vector export needs the draw.io desktop CLI;
without it a high-resolution PNG is rendered with the built-in previewer (no MathJax / no vector).
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import json
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.dont_write_bytecode = True
from audit_drawio import audit, load_yaml, print_preview, write_report  # noqa: E402


def drawio_cli():
    exe = shutil.which("drawio") or shutil.which("draw.io")
    if not exe and sys.platform.startswith("win"):
        cand = os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "draw.io", "draw.io.exe")
        if os.path.exists(cand):
            exe = cand
    return exe


def export(src, dst, fmt, scale=1.0):
    exe = drawio_cli()
    if not exe:
        return False
    cmd = [exe, "--export", "--format", fmt, "--crop", "--output", dst, src, "--no-sandbox"]
    if fmt == "png":
        cmd[4:4] = ["--scale", str(scale)]
    if fmt == "svg":
        cmd.insert(4, "--embed-svg-images")
    if not os.environ.get("DISPLAY") and shutil.which("xvfb-run"):
        cmd = ["xvfb-run", "-a"] + cmd
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=240)
        return os.path.exists(dst)
    except Exception:
        return False


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work"); ap.add_argument("--export", default="pdf,svg,png"); ap.add_argument("--dpi", type=int, default=300)
    a = ap.parse_args()
    work = os.path.abspath(a.work)
    ctx = json.load(open(os.path.join(work, "context.json"), encoding="utf-8"))
    brief = load_yaml(os.path.join(work, "brief.yaml"))
    out = brief["figure"]["output"]
    theme = brief.get("style", {}).get("theme") or ctx["theme"]
    from render_preview import render_kit
    render_kit(out, os.path.join(work, "after.png"), scale=1.5)
    accepted = (brief.get("review") or {}).get("accepted_warnings", []) or []
    rb = audit(ctx["figure"], brief.get("target"), theme)
    ra = audit(out, brief.get("target"), theme, accepted, reference=ctx["figure"])
    write_report(rb, work, "before"); write_report(ra, work, "after")
    w_in = ra["brief"]["width_in"]
    print_preview(ctx["figure"], work, "before", w_in); print_preview(out, work, "after", w_in)

    from PIL import Image, ImageDraw, ImageFont
    from drawio_kit import find_font_file
    B = ImageFont.truetype(find_font_file("Helvetica", bold=True), 26); R = ImageFont.truetype(find_font_file("Helvetica"), 20)
    ims = [Image.open(os.path.join(work, p)).convert("RGB") for p in ("before.png", "after.png", "before_print_preview.png", "after_print_preview.png")]
    colw = 900
    big = [im.resize((colw, int(im.height * colw / im.width))) for im in ims[:2]]
    small = [im if im.width <= colw else im.resize((colw, int(im.height * colw / im.width))) for im in ims[2:]]
    scaled = any(im.width > colw for im in ims[2:])
    H = 70 + max(i.height for i in big) + 90 + max(i.height for i in small) + 40
    W = 60 + 2 * colw + 40
    sheet = Image.new("RGB", (W, H), "#eef1f5"); dr = ImageDraw.Draw(sheet)
    for k, (im, lab, r) in enumerate(zip(big, ("BEFORE", "AFTER  (#ENHANCED)"), (rb, ra))):
        x = 30 + k * (colw + 40)
        dr.text((x, 20), f"{lab} — readiness {r['score']}/100, {r['n_issues']} issues", font=B, fill="#19355d")
        sheet.paste(im, (x, 60)); dr.rectangle([x, 60, x + im.width, 60 + im.height], outline="#a4a4a4", width=2)
    y = 60 + max(i.height for i in big) + 30
    for k, im in enumerate(small):
        x = 30 + k * (colw + 40)
        dr.text((x, y), f"print size {w_in} in at 150 dpi" + (" (scaled to fit this sheet)" if scaled else ""), font=R, fill="#333333")
        sheet.paste(im, (x, y + 34)); dr.rectangle([x, y + 34, x + im.width, y + 34 + im.height], outline="#a4a4a4", width=1)
    sheet.save(os.path.join(work, "before_after.png"))

    exported, notes = [], []
    base = os.path.splitext(out)[0].replace("#", "_")
    for fmt in [f.strip() for f in a.export.split(",") if f.strip()]:
        dst = f"{base}.{fmt}"
        if export(out, dst, fmt, scale=a.dpi / 96):
            exported.append(dst)
        elif fmt == "png":
            render_kit(out, dst, width=int(w_in * a.dpi), dpi=a.dpi)
            exported.append(dst); notes.append("PNG rendered with the built-in previewer (approximate; math via matplotlib)")
        else:
            notes.append(f"{fmt.upper()} needs the draw.io desktop CLI (not found) — export from draw.io: File > Export as > {fmt.upper()}")
    changes = open(os.path.join(work, "CHANGES.md"), encoding="utf-8").read() if os.path.exists(os.path.join(work, "CHANGES.md")) else "(none logged)"
    md = [f"# Enhancement report: `{os.path.basename(out)}`", "",
          f"- input: `{ctx['figure']}` ({ctx['origin']})", f"- output: `{out}`",
          f"- target: {brief.get('target', {}).get('venue', '?')} / {w_in} in / min {ra['brief']['min_font_pt']} pt / "
          f"print {brief.get('target', {}).get('print_mode', '?')} / colorblind {brief.get('target', {}).get('colorblind')}",
          f"- purpose: {brief.get('intent', {}).get('purpose', '')} — takeaway: {brief.get('intent', {}).get('takeaway', '')}",
          f"- caption: {brief.get('content', {}).get('caption', '') or '—'}", "",
          "| area | before | after |", "|---|---|---|"]
    for k in ra["issues"]:
        md.append(f"| {k} | {len(rb['issues'][k])} | {len(ra['issues'][k])} |")
    md += [f"| **score** | **{rb['score']}** | **{ra['score']}** |", "", "![before/after](before_after.png)", "",
           "## Changes", "", changes, "", "## Remaining issues", ""]
    md += [f"- [{k}] {x}" for k, v in ra["issues"].items() for x in v] or ["- none"]
    md += ["", "## Exports", ""] + [f"- `{p}`" for p in exported] + [f"- note: {n}" for n in notes]
    open(os.path.join(work, "REPORT.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")
    print(f"before {rb['score']} -> after {ra['score']}; wrote {os.path.join(work, 'REPORT.md')} and before_after.png")
    for p in exported:
        print("exported", p)
    for n in notes:
        print("note:", n)


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['scripts/start_enhance.py'] = r'''#!/usr/bin/env python3
"""
start_enhance.py — gather everything about a figure and prepare the enhancement work folder.

    python start_enhance.py FIG.drawio [--source-image SRC.png] [--text passage.md]
                            [--artifacts DIR] [--theme THEME.yaml] [--suffix "#ENHANCED"]

Finds associated artifacts automatically:
  image-to-drawio : <stem>_drawio_work/ (config.json, build.py, theme, NOTES.md, inspect/, history)
  text-to-drawio  : the passage folder (passage.md, build.py with VARIANTS, CHOSEN.md, prototypes.md,
                    _work/config.json, theme.yaml in the output root)
Creates  <stem>_enhance_work/  next to the figure with:
  context.json / context.md   origin, artifacts, inventory (labels, colors, fonts, sizes, math, rasters),
                              whether the .drawio was hand-edited after it was generated
  brief.yaml                  pre-filled from context; completed from the interview
  questions.md / .json        the interview, filtered to what is still unknown
  before.png, before_audit.md, before_print_preview.png
  theme.yaml                  theme for the enhanced figure (copy of the original or --theme)
  build_enhanced.py (+ config.json)  copy of the original build script when the rebuild route applies
Output figure (later): <stem>#ENHANCED.drawio next to the input.
"""
from __future__ import annotations

import sys as _sys
if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
import argparse
import glob
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.dont_write_bytecode = True
from audit_drawio import audit, load_yaml, print_preview, write_report  # noqa: E402
from drawio_doc import Doc  # noqa: E402
from enhance import enhanced_path  # noqa: E402


def find_image_artifacts(fig, explicit=None):
    stem = os.path.splitext(fig)[0]
    cands = [explicit] if explicit else [stem + "_drawio_work"]
    for d in cands:
        if d and os.path.exists(os.path.join(d, "config.json")):
            cfg = json.load(open(os.path.join(d, "config.json"), encoding="utf-8"))
            if "source" in cfg:
                return dict(dir=d, cfg=cfg, build=os.path.join(d, "build.py"), notes=os.path.join(d, "NOTES.md"))
    return None


def find_text_artifacts(fig, explicit=None):
    d = explicit or os.path.dirname(fig)
    if not os.path.exists(os.path.join(d, "passage.md")):
        return None
    cfgp = os.path.join(d, "_work", "config.json")
    cfg = json.load(open(cfgp, encoding="utf-8")) if os.path.exists(cfgp) else {}
    key = None
    m = re.match(r"(v\d+)_", os.path.basename(fig))
    if m:
        key = m.group(1)
    chosen = os.path.join(d, "CHOSEN.md")
    if not key and os.path.exists(chosen):
        m = re.search(r"prototype: `(v\d+)_", open(chosen, encoding="utf-8").read())
        key = m.group(1) if m else None
    theme = cfg.get("theme")
    if not theme or not os.path.exists(theme):
        ups = [d, os.path.dirname(d), os.path.dirname(os.path.dirname(d))]
        cands = [os.path.join(u, n) for u in ups for n in ("theme.yaml", "theme.yml", "theme.toml", "theme.json")]
        if cfg.get("scripts"):
            cands.append(os.path.join(os.path.dirname(cfg["scripts"]), "assets", "theme.yaml"))
        cands.append(os.path.join(SKILL, "assets", "theme.yaml"))   # same defaults as text-to-drawio
        theme = next((t for t in cands if os.path.exists(t)), None)
    return dict(dir=d, cfg=cfg, key=key, build=os.path.join(d, "build.py"), passage=os.path.join(d, "passage.md"),
                theme=theme, chosen=chosen if os.path.exists(chosen) else None)


def load_build_module(path, scripts_dir):
    sys.path.insert(0, scripts_dir)
    spec = importlib.util.spec_from_file_location("orig_build_" + str(abs(hash(path))), path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def signature(path):
    doc = Doc.load(path)
    sig = set()
    for c in doc.vertices():
        g = c.abs_geo() or [0, 0, 0, 0]
        sig.add((c.plain[:40], round(g[0]), round(g[1]), round(g[2]), round(g[3]), c.get("fillColor", "")))
    return sig


def rebuild_original(origin, art, tmp):
    """Re-run the original build to a temp file (to detect hand edits). Returns path or None."""
    try:
        if origin == "image-to-drawio":
            cfg = dict(art["cfg"]); out = os.path.join(tmp, "orig.drawio"); cfg["out"] = out
            wd = os.path.join(tmp, "imgwork"); os.makedirs(wd, exist_ok=True)
            shutil.copy(art["build"], os.path.join(wd, "build.py"))
            json.dump(cfg, open(os.path.join(wd, "config.json"), "w", encoding="utf-8"))
            subprocess.run([sys.executable, os.path.join(wd, "build.py")], check=True, capture_output=True, timeout=300)
            return out
        if origin == "text-to-drawio" and art.get("key"):
            from drawio_kit import Diagram, Theme
            m = load_build_module(art["build"], art["cfg"].get("scripts", HERE))
            v = next(v for v in m.VARIANTS if v["key"] == art["key"])
            d = Diagram(*v["size"], theme=Theme.load(art["theme"]), scale=v.get("scale", 1.0))
            v["build"](d)
            out = os.path.join(tmp, "orig.drawio"); d.save(out, quiet=True)
            return out
    except Exception as e:
        print("  (could not rebuild the original for comparison:", str(e).splitlines()[-1] if str(e) else e, ")")
    return None


def inventory(doc):
    v = doc.vertices()
    labels = [c for c in v if c.plain]
    sizes, fams, colors = {}, {}, {}
    for c in labels:
        sizes[c.get("fontSize", "11")] = sizes.get(c.get("fontSize", "11"), 0) + 1
        if not c.is_math:
            fams[c.get("fontFamily", "Helvetica")] = fams.get(c.get("fontFamily", "Helvetica"), 0) + 1
    for c in doc.cells:
        for k in ("fillColor", "strokeColor", "fontColor"):
            x = c.get(k)
            if x and x.startswith("#"):
                colors[x.lower()] = colors.get(x.lower(), 0) + 1
    rasters = [c.id for c in v if c.is_image and re.match(r"data:image/(png|jpe?g|webp)", c.get("image", ""))]
    svgs = [c.id for c in v if c.is_image and "svg+xml" in c.get("image", "")]
    maths = [c.plain for c in v if c.is_math]
    uni = [c.plain for c in labels if not c.is_math and re.search("[λθαβγδσμπΣΔ∑∫≤≥×∈]", c.plain)]
    panels = [c.plain for c in labels if re.match(r"^\(?[a-dA-D1-4][\).]\s", c.plain)]
    pw = doc.page[0]
    emph = sorted((c for c in v if c.plain and not c.is_text and c.abs_geo() and c.abs_geo()[2] < 0.35 * pw
                   and c.get("fillColor", "").lower() not in ("", "none", "#ffffff")),
                  key=lambda c: (-float(c.get("fontStyle", 0) == "1" or "<b>" in c.value), -float(c.get("fontSize", 11))))
    cand = []
    for c in emph + [c for c in labels if c not in emph]:
        t = c.plain.replace("\n", " ")
        if 2 < len(t) <= 40 and t not in cand and not c.is_math:
            cand.append(t)
        if len(cand) >= 3:
            break
    return dict(vertices=len(v), edges=len(doc.edges()), groups=sum(c.is_group for c in doc.cells), labels=len(labels),
                label_texts=[c.plain for c in labels][:80], font_sizes=sizes, font_families=fams, colors=colors,
                raster_images=rasters, svg_layers=svgs, math_labels=maths, unicode_math_labels=uni, panel_labels=panels,
                candidate_emphasis=cand + ["Nothing in particular"], page=list(doc.page), math_flag=doc.math)


def applies(q, ctx):
    a = q.get("applies", "always")
    return {"always": True, "image": ctx["origin"] == "image-to-drawio", "text": ctx["origin"] == "text-to-drawio",
            "math": bool(ctx["inventory"]["math_labels"] or ctx["inventory"]["unicode_math_labels"]),
            "photos": bool(ctx["inventory"]["raster_images"]), "inferred": bool(ctx.get("inferred")),
            "multi_panel": bool(ctx["inventory"]["panel_labels"])}.get(a, True)


def set_key(d, dotted, value):
    ks = dotted.split(".")
    for k in ks[:-1]:
        d = d.setdefault(k, {})
    d[ks[-1]] = value


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("drawio"); ap.add_argument("--source-image"); ap.add_argument("--text")
    ap.add_argument("--artifacts", help="work dir of image-to-drawio or passage dir of text-to-drawio")
    ap.add_argument("--theme", help="theme for the enhanced figure (default: the original's theme)")
    ap.add_argument("--suffix", default="#ENHANCED"); ap.add_argument("--work-dir")
    a = ap.parse_args()

    import yaml
    fig = os.path.abspath(a.drawio)
    stem = os.path.splitext(fig)[0]
    name = os.path.basename(stem)
    out = enhanced_path(fig, a.suffix)
    work = os.path.abspath(a.work_dir or stem + "_enhance_work")
    os.makedirs(work, exist_ok=True)

    img = find_image_artifacts(fig, a.artifacts)
    txt = None if img else find_text_artifacts(fig, a.artifacts)
    origin = "image-to-drawio" if img else "text-to-drawio" if txt else ("image-to-drawio" if a.source_image else
                                                                         "text-to-drawio" if a.text else "unknown")
    art = img or txt or {}
    source_image = a.source_image or (img["cfg"]["source"] if img else "")
    source_text = a.text or (txt["passage"] if txt else "")
    theme_src = a.theme or (img and img["cfg"].get("theme")) or (txt and txt.get("theme")) or \
        os.path.join(SKILL, "assets", "themes", "navy_default.yaml")
    theme_dst = os.path.join(work, "theme.yaml")          # always YAML so brief.py can edit it
    if not os.path.exists(theme_dst) or a.theme:
        from drawio_kit import Theme
        yaml.safe_dump(Theme.load(theme_src).to_dict(), open(theme_dst, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)

    doc = Doc.load(fig)
    inv = inventory(doc)
    inferred, caption = [], ""
    build_path = art.get("build") if art and os.path.exists(art.get("build", "")) else ""
    if txt and build_path:
        try:
            m = load_build_module(build_path, txt["cfg"].get("scripts", HERE))
            inferred = [x for x in getattr(m, "INFERRED", []) if not txt.get("key") or not re.match(r"v\d+:", x) or x.startswith(txt["key"] + ":")]
            v = next((v for v in m.VARIANTS if v["key"] == txt.get("key")), None)
            caption = (v or {}).get("caption", "")
        except Exception as e:
            print("  (could not read build.py:", e, ")")

    edited, diff = None, None
    if build_path:
        with tempfile.TemporaryDirectory() as tmp:
            orig = rebuild_original(origin, art, tmp)
            if orig:
                s1, s2 = signature(orig), signature(fig)
                diff = len(s1 ^ s2)
                edited = diff > 0
    route = "rebuild" if build_path and edited is False else "edit"

    ctx = dict(figure=fig, name=name, output=out, origin=origin, work_dir=work, artifacts=art.get("dir", ""),
               source_image=source_image, source_text=source_text, build_script=build_path, variant_key=art.get("key"),
               theme=theme_dst, original_theme=theme_src, edited_after_build=edited, signature_diff=diff,
               recommended_route=route, inferred=inferred, draft_caption=caption, inventory=inv)
    json.dump(ctx, open(os.path.join(work, "context.json"), "w", encoding="utf-8"), indent=1, default=str)

    # brief
    brief = load_yaml(os.path.join(SKILL, "assets", "brief_template.yaml"))
    bp = os.path.join(work, "brief.yaml")
    if os.path.exists(bp):
        brief = load_yaml(bp)
    for k, v in {"figure.name": name, "figure.input": fig, "figure.origin": origin, "figure.source_image": source_image,
                 "figure.source_text": source_text, "figure.build_script": build_path, "figure.output": out,
                 "style.theme": theme_dst, "plan.route": route}.items():
        set_key(brief, k, v)
    if caption and not brief.get("content", {}).get("caption"):
        set_key(brief, "content.caption", caption)
    yaml.safe_dump(brief, open(bp, "w", encoding="utf-8"), sort_keys=False, allow_unicode=True)

    # rebuild route setup
    if build_path:
        dst = os.path.join(work, "build_enhanced.py")
        if not os.path.exists(dst):
            shutil.copy(build_path, dst)
        if origin == "image-to-drawio":
            cfg = dict(img["cfg"]); cfg.update(out=out, theme=theme_dst, scripts=HERE, workdir=work)
            json.dump(cfg, open(os.path.join(work, "config.json"), "w", encoding="utf-8"), indent=1)

    # before: preview + audit
    from render_preview import render_kit
    render_kit(fig, os.path.join(work, "before.png"), scale=1.5)
    r = audit(fig, brief.get("target"), theme_dst)
    write_report(r, work, "before")
    print_preview(fig, work, "before", r["brief"]["width_in"])

    # questions
    bank = load_yaml(os.path.join(SKILL, "assets", "questions.yaml"))
    qs_out, md = [], [f"# Interview for `{name}`", "",
                      "Ask round by round (<= 3-4 questions per message, options as buttons when possible). "
                      "Write answers into brief.yaml at the given key. Skip a question if the user already answered it.", ""]
    for rnd in bank["rounds"]:
        items = [q for q in rnd["questions"] if applies(q, ctx)]
        if not items:
            continue
        md += [f"## {rnd['name']}", ""]
        for q in items:
            opts = inv.get(q["options_from"], []) if q.get("options_from") else q.get("options", [])
            extra = ""
            if q["id"] == "inferred":
                extra = "\n  - inferred items: " + "; ".join(inferred)
            if q["id"] == "caption" and caption:
                extra = f"\n  - draft caption: {caption}"
            md.append(f"- **{q['ask']}**  → `{q['brief_key']}`" + (f"\n  - options: {' | '.join(map(str, opts))}" if opts else "\n  - (free text)") + extra)
            qs_out.append(dict(round=rnd["name"], id=q["id"], ask=q["ask"], options=opts, key=q["brief_key"]))
        md.append("")
    open(os.path.join(work, "questions.md"), "w", encoding="utf-8").write("\n".join(md) + "\n")
    json.dump(qs_out, open(os.path.join(work, "questions.json"), "w", encoding="utf-8"), indent=1)

    cm = [f"# Enhancement context: `{name}`", "",
          f"- origin: **{origin}**; artifacts: `{art.get('dir', '—')}`",
          f"- source image: `{source_image or '—'}`; source text: `{source_text or '—'}`",
          f"- build script: `{build_path or '—'}`" + (f" (variant {art.get('key')})" if art.get("key") else ""),
          f"- hand-edited after generation: {edited if edited is not None else 'unknown'}"
          + (f" ({diff} differing elements)" if diff else ""),
          f"- recommended route: **{route}** — " + ("re-run an edited copy of the build script (build_enhanced.py), then enhance.py"
                                                    if route == "rebuild" else "edit the .drawio directly (drawio_doc API + enhance.py)"),
          f"- theme for enhancement: `{theme_dst}`",
          f"- output: `{out}`", "",
          f"## Inventory", "",
          f"- {inv['vertices']} vertices, {inv['edges']} edges, {inv['groups']} groups, {inv['labels']} labels; page {inv['page']}",
          f"- font families {inv['font_families']}; sizes {inv['font_sizes']}",
          f"- colors {inv['colors']}",
          f"- math labels: {len(inv['math_labels'])} (flag {'on' if inv['math_flag'] else 'off'}); Unicode-math labels: {len(inv['unicode_math_labels'])}",
          f"- raster images: {len(inv['raster_images'])}; SVG layers: {len(inv['svg_layers'])}; panel labels: {inv['panel_labels'] or '—'}",
          f"- inferred content (text origin): {inferred or '—'}", f"- draft caption: {caption or '—'}", "",
          f"## Before", "", f"- audit score **{r['score']}/100** ({r['n_issues']} issues) — see before_audit.md",
          "- previews: before.png, before_print_preview.png (actual print size)"]
    open(os.path.join(work, "context.md"), "w", encoding="utf-8").write("\n".join(cm) + "\n")
    print("\n".join(cm))
    print(f"\nnext: interview the user with {os.path.join(work, 'questions.md')}, fill brief.yaml, then follow SKILL.md (enhance mode, C4)")


if __name__ == "__main__":
    main()
'''

# ================================================================================================
FILES['assets/brief_template.yaml'] = r'''# Figure brief — filled by start_enhance.py from context, completed from the interview.
figure:
  name: ""                 # stem of the input .drawio
  input: ""                # path of the .drawio being enhanced
  origin: ""               # image-to-drawio | text-to-drawio | unknown
  source_image: ""         # image-to-drawio: original raster
  source_text: ""          # text-to-drawio: passage.md
  build_script: ""         # original build.py if any (rebuild route)
  output: ""               # <name>#ENHANCED.drawio
intent:
  purpose: ""
  takeaway: ""
  audience: ""
  emphasis: []
target:
  venue: generic
  column: double           # single | one_and_half | double | custom
  width_in: 7.0
  min_font_pt: 7
  print_mode: color        # color | grayscale_safe | grayscale
  grayscale: false
  colorblind: true
style:
  theme: ""                # theme file used for the enhanced figure
  font: keep
  palette: keep
  panel_labels: none
content:
  fidelity: polish         # exact | polish | redesign   (image origin)
  text_density: keep
  terminology: ""
  inferred: keep           # keep | drop | itemized       (text origin)
  photos: keep
  caption_mode: draft
  caption: ""
  merge_from_prototypes: ""
math:
  notation: keep
  equation_refs: ""
plan:
  route: ""                # rebuild | edit
  steps: []
  approved: false
'''

# ================================================================================================
FILES['assets/build_template_image.py'] = r'''#!/usr/bin/env python3
"""
Build script for __IMAGE_NAME__  ->  __OUT_NAME__
Generated by new_build.py (image mode). Edit the LAYOUT section, then run:
    python __SCRIPTS__/iterate.py __WORKDIR__ --note "what changed"
All coordinates are SOURCE-IMAGE pixels (read them off inspect/source.grid*.png or zoom crops).
Colors / fonts / sizes are THEME ROLES from theme.yaml — never hard-code hex values here.
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = json.load(open(os.path.join(HERE, "config.json")))
sys.path.insert(0, os.environ.get("IMAGE_TO_DRAWIO_SCRIPTS", CFG["scripts"]))
from drawio_kit import Diagram  # noqa: E402


def build():
    d = Diagram.from_source(CFG["source"], theme=CFG["theme"], scale=CFG.get("scale", 1.0),
                            background=CFG.get("background", "background"))
    W, H = d.width, d.height

    # ------------------------------------------------------------------ LAYOUT
    # Order = z-order: background/panels -> fills/regions -> content shapes -> connectors -> text/annotations.
    # Use d.group("name") for things that move together (tile = box + icon + label).
    # Examples (delete when you start):
    #   pid, (cx, cy, cw, ch) = d.panel(20, 20, 400, 300, "**1. Title**", header_fill="primary")
    #   d.rect(60, 90, 120, 40, "steel", fill="background", stroke="text", size="body", id="box_steel")
    #   d.text("**Bold** caption", 220, 280, size="label", color="primary")
    #   d.edge([], src="box_steel", tgt="box_other", arrow="end", route="orthogonal", exit=(1, .5), entry=(0, .5))
    #   d.image_crop((x0, y0, x1, y1))     # photo / complex illustration copied from the source

    return d


if __name__ == "__main__":
    build().save(CFG["out"])
'''

# ================================================================================================
FILES['assets/build_template_text.py'] = r'''"""
Prototypes for passage: __TITLE__
Edit v1 / v2 / v3 below, then run:   python __SCRIPTS__/prototype.py __PASSAGE_DIR__
Coordinates are px on each variant's canvas; colors / sizes are theme roles (theme.yaml).
Every label must come from passage.md. Anything you infer goes in INFERRED (it is reported to the user).
"""
from figlib import *   # NODE, SOFT, DARK, MUTED, GHOST, pill, lane, chain, cycle, layers, frame, bracket,
                       # spread, annotated_equation, rating_dots, card, continuum, legend, axes, table, connect

# Claims shown in a figure that the passage does NOT state (judgments, added categories, guessed ratings).
INFERRED = [
    # "v1: rating 'moderate' cost for X (not stated)",
]
# Structural words used in labels that are not content claims (e.g. "layer", "input", "output").
ALLOWED_TERMS = []
# Lint warnings you have looked at and accept on purpose (substring match), e.g. ["uneven gaps"].
LINT_IGNORE = []


def v1(d):
    """__FAMILY1__"""
    # d.rect(...), chain(d, [...], x, y, w, h), d.edge([], src=..., tgt=...), d.equation(r"...", x, y) ...
    pass


def v2(d):
    """__FAMILY2__"""
    pass


def v3(d):
    """__FAMILY3__"""
    pass


VARIANTS = [
    # key, file-name style, canvas (w, h), function, one-line description, draft caption (from the text)
    dict(key="v1", name="__FAMILY1__", size=(1100, 500), build=v1, about="", caption=""),
    dict(key="v2", name="__FAMILY2__", size=(1100, 500), build=v2, about="", caption=""),
    dict(key="v3", name="__FAMILY3__", size=(1100, 500), build=v3, about="", caption=""),
]
'''

# ================================================================================================
FILES['assets/questions.yaml'] = r'''# Interview bank for enhance-drawio. start_enhance.py filters it (skips what the context already
# answers, adds origin-specific questions) and writes questions.md. Ask in rounds of <= 3-4
# questions, offering the options as buttons when the client supports it; record answers in brief.yaml.
# brief_key: where the answer goes in brief.yaml.  applies: always | image | text | math | photos | inferred
rounds:
  - name: Purpose and placement
    questions:
      - id: purpose
        ask: What is this figure's job in the paper?
        options: [Method / system overview, Conceptual framework, Comparison of approaches, Explaining an equation, Results / evidence, Graphical abstract]
        brief_key: intent.purpose
        applies: always
      - id: domain
        ask: What academic or scientific field is this figure for?
        options: [Computer Science / AI / ML, Physics / Engineering, Biology / Medicine, Chemistry / Materials, Economics / Finance, Social Science / Causal Inference, Mathematics / Statistics, Other]
        brief_key: intent.domain
        applies: always
      - id: venue
        ask: Where will it be published?
        options: [IEEE journal / conference, Elsevier journal, Springer (LNCS), Thesis / dissertation, Other (tell me)]
        brief_key: target.venue
        applies: always
        note: sets width presets and minimum font size from assets/venues.yaml (confirm with author guidelines)
      - id: width
        ask: How wide will it be printed?
        options: [Single column, 1.5 column, Full text width (double column), Custom width]
        brief_key: target.column
        applies: always
      - id: takeaway
        ask: In one sentence, what should a reader take away from this figure?
        free_text: true
        brief_key: intent.takeaway
        applies: always

  - name: Audience and content
    questions:
      - id: audience
        ask: Who is the main audience?
        options: [Specialists in this field, Broad engineering audience, Interdisciplinary / general readers]
        brief_key: intent.audience
        applies: always
      - id: emphasis
        ask: Which part should stand out most?
        options_from: candidate_emphasis      # filled from the figure's labels
        brief_key: intent.emphasis
        applies: always
      - id: density
        ask: How much text should stay inside the figure?
        options: [Keep current labels, Shorten labels and move explanations to the caption, Minimal labels (caption carries the detail)]
        brief_key: content.text_density
        applies: always
      - id: terminology
        ask: Any terms, abbreviations or symbols that must match the paper (or elements to remove/rename)?
        free_text: true
        brief_key: content.terminology
        applies: always

  - name: Visual style
    questions:
      - id: presentation_style
        ask: How would you prefer to present the core concepts or equations?
        options: [(Recommended) Active graphical demonstration (populated matrix computation, 2D continuous coordinate spaces, Bayesian plate diagrams, causal DAGs), High-level structural overview (swimlane workflow, summary cards, pipeline architecture)]
        brief_key: style.presentation_style
        applies: always
      - id: print_mode
        ask: How will it be printed or viewed?
        options: [Color only (online), Must also work in grayscale, Grayscale only]
        brief_key: target.print_mode
        applies: always
      - id: colorblind
        ask: Should colors stay distinguishable for color-blind readers?
        options: [Yes (recommended), Not needed]
        brief_key: target.colorblind
        applies: always
      - id: font
        ask: Which font should labels use to match the paper?
        options: [Times New Roman (Times-based journals), Helvetica / Arial (sans), Latin Modern / Computer Modern (LaTeX look), Keep current]
        brief_key: style.font
        applies: always
      - id: palette
        ask: Which colors?
        options: [Keep current theme, My palette (I'll give hex codes), Color-blind-safe preset (Okabe-Ito), Grayscale preset]
        brief_key: style.palette
        applies: always

  - name: Math, labels and caption
    questions:
      - id: notation
        ask: How should math be typeset? (e.g. bold vectors, upright operators, the paper's macros)
        options: [Keep as is, Match my paper's notation (I'll describe it), Plain-text symbols only]
        brief_key: math.notation
        applies: math
      - id: eq_numbers
        ask: Should the figure reference equation numbers from the paper (e.g. "Eq. (3)")?
        options: ["No", "Yes (give numbers)"]
        brief_key: math.equation_refs
        applies: math
      - id: panels
        ask: Panel labels?
        options: [None, "(a), (b), ...", "a, b, ... (bold)"]
        brief_key: style.panel_labels
        applies: multi_panel
      - id: caption
        ask: Caption?
        options: [Use the draft caption, I'll provide one, Write a new one from the paper text]
        brief_key: content.caption_mode
        applies: always

  - name: Origin-specific
    questions:
      - id: fidelity
        ask: How close should it stay to the original image?
        options: [Keep the exact layout, Keep the structure but polish freely, Free redesign of the same content]
        brief_key: content.fidelity
        applies: image
      - id: photos
        ask: The figure embeds photos / raster crops. What should happen to them?
        options: [Keep, Replace with higher-resolution versions I'll supply, Remove]
        brief_key: content.photos
        applies: photos
      - id: inferred
        ask: The prototype shows content the text does not state (listed below). Keep it?
        options: [Keep all, Drop all, Decide item by item]
        brief_key: content.inferred
        applies: inferred
      - id: merge
        ask: Borrow anything from the other prototypes?
        free_text: true
        brief_key: content.merge_from_prototypes
        applies: text

  - name: Confirm
    questions:
      - id: plan_ok
        ask: Here is the enhancement plan (summarized from your answers). Go ahead?
        options: [Yes, Change something first]
        brief_key: plan.approved
        applies: always
'''

# ================================================================================================
FILES['assets/theme.yaml'] = r'''# Default theme (text-to-drawio + enhance). Build scripts use ROLES (fill="primary", size="body"), never raw values.
# YAML, TOML or JSON all work. Every prototype of every passage shares this file -> consistent figures.
colors:
  primary: "#19355d"      # key shapes, headers, main arrows, emphasis text
  secondary: "#4276b0"    # second category, highlights, tints (with opacity roles)
  neutral: "#a4a4a4"      # context, moderators, frames, muted annotations
  text: "#000000"         # body text, outlines of neutral items
  background: "#ffffff"   # canvas and box fills
  # add semantic roles only when the text needs another category (e.g. highlight: "#5b8c4a")

fonts:
  sans:  { family: "Helvetica" }          # labels; set google: true for e.g. "Roboto" / "Inter"
  serif: { family: "Times New Roman" }    # optional: captions in LaTeX-like papers
default_font: sans

sizes:                 # px at 100 % scale (a 1000-px-wide figure in a 6.5-in column => 1 px ~ 0.47 pt)
  title: 20            # rarely used: academic figures carry their title in the caption
  heading: 15          # card / panel headers
  subheading: 13.5     # column pills, lane titles, layer tabs
  body: 13             # box labels
  label: 11.5          # secondary labels, callouts
  small: 10.5          # annotations, legends, axis ticks (do not go smaller)
  equation: 30         # main displayed equation
  eq_term: 20          # equations inside diagrams
  eq_small: 16         # equations in tables / strips

line_height: 1.2
strokes: { hairline: 0.75, thin: 1, normal: 1.5, thick: 2, heavy: 3 }
opacity: { wash: 8, tint: 18, soft: 35, strong: 60 }
shape:   { corner_radius: 8, shadow: false }
spacing: { padding: 10, gap: 30 }
'''

# ================================================================================================
FILES['assets/theme_image.yaml'] = r'''# Theme for image-to-drawio builds (sizes in SOURCE-IMAGE px). Every color / font / size should come from here.
# Build scripts refer to ROLES (e.g. fill="primary", size="heading", font="serif"), never raw values.
colors:
  primary: "#19355d"      # headers, key shapes, main arrows
  secondary: "#4276b0"    # accents, highlights, secondary series
  neutral: "#a4a4a4"      # borders, muted fills, de-emphasized marks
  text: "#000000"         # body text, outlines, axes
  background: "#ffffff"   # canvas and default fills
  # Add figure-specific roles as needed (keep names semantic, not "green"):
  # highlight: "#5b8c4a"
  # cluster_d: "#c2a27a"

fonts:
  # family must be installed where the .drawio is viewed; web-safe families are safest.
  # Set google: true to load a Google Font in draw.io (e.g. family: "Roboto").
  sans:  { family: "Helvetica" }
  serif: { family: "Times New Roman" }
  mono:  { family: "Courier New" }
default_font: sans

sizes:            # font sizes in SOURCE-IMAGE pixels (scaled with the diagram)
  title: 28
  heading: 20
  subheading: 16
  body: 14
  label: 12
  small: 10
  caption: 18

line_height: 1.2

strokes:          # stroke widths in px
  hairline: 0.75
  thin: 1
  normal: 1.5
  thick: 2.5
  heavy: 4

opacity:          # named opacities (percent) for tints / faded marks
  tint: 25
  soft: 45
  faded: 35

shape:
  corner_radius: 8    # px, used when rounded=True
  shadow: false

spacing:
  padding: 10         # inner padding for labeled boxes
  gap: 12             # default gap for row()/grid() helpers
'''

# ================================================================================================
FILES['assets/venues.yaml'] = r'''# Typical figure sizes. ALWAYS confirm against the venue's current author guidelines; these are
# starting points. width_in = printed figure width; min_font_pt = smallest lettering at print size.
generic:
  columns: {single: 3.5, one_and_half: 5.0, double: 7.0}
  min_font_pt: 7
  font: Helvetica
ieee:
  columns: {single: 3.5, double: 7.16}
  min_font_pt: 8
  font: Times New Roman
  note: IEEE column 3.5 in, text width 7.16 in; check the IEEE graphics guidelines for current minimums
elsevier:
  columns: {single: 3.54, one_and_half: 5.51, double: 7.48}     # 90 / 140 / 190 mm
  min_font_pt: 7
  font: Helvetica
  note: Elsevier artwork sizing 90 / 140 / 190 mm; check the journal's guide for authors
springer_lncs:
  columns: {single: 4.8, double: 4.8}                              # 122 mm text width
  min_font_pt: 8
  font: Times New Roman
thesis:
  columns: {single: 6.0, double: 6.5}
  min_font_pt: 9
  font: Times New Roman
slides:
  columns: {single: 10.0, double: 13.33}
  min_font_pt: 18
  font: Helvetica
'''

# ================================================================================================
FILES['assets/themes/grayscale.yaml'] = r'''# Grayscale roles with well-separated lightness (L* ~ 20 / 45 / 70 / 0 / 100).
colors: {primary: "#262626", secondary: "#6b6b6b", neutral: "#aaaaaa", text: "#000000", background: "#ffffff"}
fonts: {sans: {family: "Helvetica"}, serif: {family: "Times New Roman"}}
default_font: sans
sizes: {title: 20, heading: 15, subheading: 13.5, body: 13, label: 11.5, small: 10.5, equation: 30, eq_term: 20, eq_small: 16}
line_height: 1.2
strokes: {hairline: 0.75, thin: 1, normal: 1.5, thick: 2, heavy: 3}
opacity: {wash: 8, tint: 22, soft: 40, strong: 65}
shape: {corner_radius: 8, shadow: false}
spacing: {padding: 10, gap: 30}
'''

# ================================================================================================
FILES['assets/themes/navy_default.yaml'] = r'''# Theme for text-to-drawio figures. Build scripts use ROLES (fill="primary", size="body"), never raw values.
# YAML, TOML or JSON all work. Every prototype of every passage shares this file -> consistent figures.
colors:
  primary: "#19355d"      # key shapes, headers, main arrows, emphasis text
  secondary: "#4276b0"    # second category, highlights, tints (with opacity roles)
  neutral: "#a4a4a4"      # context, moderators, frames, muted annotations
  text: "#000000"         # body text, outlines of neutral items
  background: "#ffffff"   # canvas and box fills
  # add semantic roles only when the text needs another category (e.g. highlight: "#5b8c4a")

fonts:
  sans:  { family: "Helvetica" }          # labels; set google: true for e.g. "Roboto" / "Inter"
  serif: { family: "Times New Roman" }    # optional: captions in LaTeX-like papers
default_font: sans

sizes:                 # px at 100 % scale (a 1000-px-wide figure in a 6.5-in column => 1 px ~ 0.47 pt)
  title: 20            # rarely used: academic figures carry their title in the caption
  heading: 15          # card / panel headers
  subheading: 13.5     # column pills, lane titles, layer tabs
  body: 13             # box labels
  label: 11.5          # secondary labels, callouts
  small: 10.5          # annotations, legends, axis ticks (do not go smaller)
  equation: 30         # main displayed equation
  eq_term: 20          # equations inside diagrams
  eq_small: 16         # equations in tables / strips

line_height: 1.2
strokes: { hairline: 0.75, thin: 1, normal: 1.5, thick: 2, heavy: 3 }
opacity: { wash: 8, tint: 18, soft: 35, strong: 60 }
shape:   { corner_radius: 8, shadow: false }
spacing: { padding: 10, gap: 30 }
'''

# ================================================================================================
FILES['assets/themes/okabe_ito.yaml'] = r'''# Color-blind-safe roles from the Okabe-Ito palette (Okabe & Ito, 2008).
colors: {primary: "#0072b2", secondary: "#e69f00", neutral: "#999999", text: "#000000", background: "#ffffff",
         accent3: "#009e73", accent4: "#d55e00", accent5: "#cc79a7", accent6: "#56b4e9"}
fonts: {sans: {family: "Helvetica"}, serif: {family: "Times New Roman"}}
default_font: sans
sizes: {title: 20, heading: 15, subheading: 13.5, body: 13, label: 11.5, small: 10.5, equation: 30, eq_term: 20, eq_small: 16}
line_height: 1.2
strokes: {hairline: 0.75, thin: 1, normal: 1.5, thick: 2, heavy: 3}
opacity: {wash: 8, tint: 18, soft: 35, strong: 60}
shape: {corner_radius: 8, shadow: false}
spacing: {padding: 10, gap: 30}
'''

# ================================================================================================
FILES['examples/01_two_panel_simple.py'] = r'''#!/usr/bin/env python3
"""
Tier 1 — simple two-panel comparison (BERT auto-encoding vs GPT autoregressive), 1408x736 source.
Shows: header bands, token boxes in a row with equal gaps, attached connectors, elbow arrow,
partially faded box, mixed bold/regular captions, figure-specific theme roles.
"""
from _common import run

THEME = {  # figure-specific roles layered over the user's theme
    "colors": {"canvas": "#eef1f5", "token": "#d6d6d6", "ghost": "#e9e9e9"},
    "sizes": {"title": 42, "heading": 28, "body": 27, "caption": 25},
    "strokes": {"normal": 4, "thick": 6, "heavy": 9},
}


def layout(d):
    d.background = "canvas"
    d.text("**Transformer Architectures: Two Paths to Language**", 704, 45, size="title", color="primary")
    for side, x0, title, block, caption in (
        ("L", 52, "AUTO-ENCODING (e.g., BERT)", "Encoder",
         "**Bidirectional context.**\nLearns by filling gaps.\n**Best for:** Understanding & Classification."),
        ("R", 750, "AUTOREGRESSIVE (e.g., GPT)", "Decoder",
         "**Unidirectional order.**\nLearns by predicting next word.\n**Best for:** Generation & Creation."),
    ):
        d.panel(x0, 110, 622, 570, title, header_h=70, header_fill="primary", header_color="background",
                title_size="heading", stroke="none", shadow=True, id=f"panel{side}")
        cx = x0 + 311
        d.text(caption, cx + (0 if side == "L" else -2), 592, size="caption", id=f"cap{side}")

    tok = dict(fill="token", stroke="text", sw="normal", rounded=4, size="body")
    # left: input row with [MASK], predicted token above, encoder below
    d.text("**Input:**", 120, 367, size="body")
    xs = d.row_widths(182, 336, [97, 138, 97, 97], 62, gap=13)   # one constant gap, measured widths
    for (x, y, w, h), t in zip(xs, ["[A]", "[MASK]", "[C]", "[D]"]):
        if t == "[MASK]":
            d.rect(x, y, w, h, "**[MASK]**", fill="primary", stroke="secondary", color="background",
                   rounded=4, size="body", id="L_MASK")
        else:
            d.rect(x, y, w, h, t, id=f"L_{t[1]}", **tok)
    mask_cx = xs[1][0] + xs[1][2] / 2
    d.rect(mask_cx - 48, 218, 96, 62, "[B]", id="L_B", **tok)
    d.rect(mask_cx - 85, 455, 170, 62, "**Encoder**", fill="primary", stroke="primary", color="background",
           rounded=4, size="body", id="Enc")
    d.edge([], src="L_MASK", tgt="L_B", color="secondary", sw="thick", exit=(.5, 0), entry=(.5, 1))
    d.edge([], src="Enc", tgt="L_MASK", sw="thick", exit=(.5, 0), entry=(.5, 1))
    d.edge([(108, 486)], tgt="Enc", sw="thick", entry=(0, .5))

    # right: input row with faded last token, decoder -> output -> [D]
    d.text("**Input:**", 822, 249, size="body")
    xs = d.row(885, 218, 4, 97, 62, gap=12)
    d.rect(xs[3][0] + xs[3][2] - 10, 218, 38, 62, fill="ghost", stroke="ghost", rounded=4, opacity="faded")
    for (x, y, w, h), t in zip(xs, ["[A]", "[B]", "[C]", "[D]"]):
        d.rect(x, y, w, h, t, id=f"R_{t[1]}", opacity=75 if t == "[D]" else None, **tok)
    d.rect(978, 336, 197, 62, "**Output:**", fill="primary", stroke="secondary", color="background",
           rounded=4, size="body", id="Out")
    d.rect(978, 455, 197, 62, "**Decoder**", fill="primary", stroke="primary", color="background",
           rounded=4, size="body", id="Dec")
    d.edge([], src="Dec", tgt="Out", sw="thick", exit=(.5, 0), entry=(.5, 1))
    d.edge([(810, 486)], tgt="Dec", sw="thick", entry=(0, .5))
    d.edge([], src="Out", tgt="R_D", color="secondary", sw="thick", route="orthogonal", exit=(1, .5), entry=(.5, 1))

    # divider + direction glyph arrows
    d.line([(710, 110), (710, 715)], color="neutral", sw="thick")
    d.arrow([(330, 676), (395, 676)], color="primary", sw="heavy")
    d.arrow([(395, 706), (330, 706)], color="primary", sw="heavy")
    d.arrow([(1020, 692), (1103, 692)], color="primary", sw="heavy")


if __name__ == "__main__":
    run(layout, (1408, 736), THEME, "01_two_panel_simple")
'''

# ================================================================================================
FILES['examples/02_tile_dashboard.py'] = r'''#!/usr/bin/env python3
"""
Tier 2 — three-column infographic with tiles, icons and mini flow diagrams (NLP / NLU / NLG), 1024x559.
Shows: per-panel tint roles, header+body panels, tiles as groups (box+icon+label) laid out with
fit_row() so gaps are identical, emoji icon placeholders, nested NLI box with branching connectors,
orthogonal routing with explicit waypoints, bullet lists, footer rule with centered label.
Tip: run with --scale 1.5 so small text is comfortable to edit.
"""
from _common import run

THEME = {
    "colors": {
        "canvas": "#f4f7fa", "ink": "#1f2a36", "muted": "#6b7785", "rule": "#c5ced8", "connector": "#8fb3cc",
        "nlp_head": "#dbe9f3", "nlp_body": "#eef5f8", "nlp_line": "#9fb8cc",
        "nlu_head": "#e3e6f3", "nlu_body": "#eef0f7", "nlu_line": "#aab3d3",
        "nlg_head": "#f4ecd8", "nlg_body": "#f8f4ea", "nlg_line": "#cdbd94",
        "tile_green": "#d7ecd6", "tile_green_line": "#8fbf8c", "tile_blue": "#dfe6f6", "tile_blue_line": "#9aa9d4",
        "tile_purple": "#ece2f4", "tile_purple_line": "#b49ac6", "tile_tan": "#f2e4c0", "tile_tan_line": "#c9ad6a",
        "nli_fill": "#f1e8f5", "contra": "#f7dfb5", "contra_line": "#d2a85c", "hyp": "#f5e8c8",
    },
    "sizes": {"title": 21, "heading": 15.5, "body": 13.5, "label": 13, "small": 11, "icon": 30, "icon_sm": 26, "icon_md": 28},
    "strokes": {"thin": 1.2, "normal": 1.5},
}


def tile(d, id, box, col, icon, label, icy, lcy, size="body"):
    x, y, w, h = box
    with d.group(id):
        d.rect(x, y, w, h, fill=f"tile_{col}", stroke=f"tile_{col}_line", sw="normal", rounded=6, id=id + "_box")
        d.icon(icon, x + w / 2, icy, size=d.theme.size("icon"))
        d.text(label, x + w / 2, lcy, size=size, color="ink")


def layout(d):
    d.background = "canvas"
    d.text("**FUNCTIONAL BREAKDOWN OF NATURAL LANGUAGE PROCESSING: NLP, NLU, NLI & NLG**", 512, 35,
           size="title", color="ink")
    cols = [("nlp", 33, 297, "**1. NATURAL LANGUAGE PROCESSING**\n**(NLP)** - Overarching Field", "🧠",
             "Field combining computer science, AI,\nand linguistics to enable interactions\nbetween computers and human language."),
            ("nlu", 364, 297, "**2. NATURAL LANGUAGE UNDERSTANDING**\n**(NLU)** - Intent & Semantics", "📋",
             "Subfield focused on machine understanding\nof text and speech, discerning meaning,\nintent, and relationships."),
            ("nlg", 695, 297, "**3. NATURAL LANGUAGE GENERATION**\n**(NLG)** - Creation & Output", "📝",
             "Process of producing coherent natural\nlanguage text or speech from structured\ndata or conceptual intent.")]
    for key, x, w, head, icon, desc in cols:
        d.rect(x, 65, w, 53, head, fill=f"{key}_head", stroke=f"{key}_line", rounded=6, size="heading", color="ink", id=f"{key}_head")
        d.rect(x, 125, w, 415, fill=f"{key}_body", stroke=f"{key}_line", rounded=6, shadow=True, id=f"{key}_body")
        d.icon(icon, x + 24, 160, size=d.theme.size("icon"))
        d.text(desc, x + 45, 160, anchor="w", size="label", color="ink")
    # connectors between panels (elbow from body up into next header)
    d.arrow([(346, 140), (346, 92), (363, 92)], color="connector", route="straight")
    d.arrow([(678, 140), (678, 92), (694, 92)], color="connector", route="straight")

    # --- NLP panel: two tile rows with equal gaps
    r1 = d.fit_row(43, 320, 198, 3, 102, gap=9)          # symmetric inside the 33..330 panel
    for box, (ic, lb, idn) in zip(r1, [("🗂️", "Information\nExtraction", "ie"), ("🔤", "Machine\nTranslation", "mt"),
                                       ("🔊", "Speech\nRecognition", "sr1")]):
        tile(d, "t_" + idn, box, "green", ic, lb, icy=232, lcy=276)
    r2 = d.row(43, 312, 2, r1[0][2], 108, gap=9)
    for box, (ic, lb, idn) in zip(r2, [("🎙️", "Speech\nRecognition", "sr2"), ("😊", "Sentiment\nAnalysis", "sa")]):
        tile(d, "t_" + idn, box, "green", ic, lb, icy=350, lcy=395)
    with d.group("text_analytics"):
        d.icon("🔍", 274, 350, size=d.theme.size("icon")); d.text("Text\nAnalytics", 274, 392, size="body", color="ink")
    with d.group("raw"):
        d.icon("🗄️", 72, 456, size=d.theme.size("icon_sm")); d.text("Raw\nLanguage\nData", 72, 494, size="body", color="ink")
    d.rect(r2[1][0], 466, 80, 48, "**NLP**\n**System**", fill="tile_green", stroke="tile_green_line", rounded=6, size="body", color="ink", id="nlp_sys")
    with d.group("processed"):
        d.icon("📄", 282, 456, size=d.theme.size("icon_sm")); d.text("Processed\nInformation", 284, 498, size="body", color="ink")
    d.edge([(100, 490)], tgt="nlp_sys", color="ink", entry=(0, .5))
    d.edge([(220, 486), (240, 486), (240, 458), (264, 458)], src="nlp_sys", color="ink", exit=(1, .4))

    # --- NLU panel
    nlu = [("blue", "⚙️", "Semantic Role\nLabeling (SRL)", "srl"), ("purple", "🌳", "Named Entity\nRecognition", "ner"),
           ("purple", "🕸️", "Coreference\nResolution", "cor")]
    for box, (col, ic, lb, idn) in zip(d.fit_row(374, 651, 198, 3, 102, gap=9), nlu):   # same inner margin/gap/height in every panel
        tile(d, "t_" + idn, box, col, ic, lb, icy=230, lcy=274, size="label")
    with d.group("nlu_left"):
        d.rect(374, 308, 111, 222, fill="tile_blue", stroke="tile_blue_line", rounded=6)
        d.icon("🕸️", 398, 345, size=d.theme.size("icon_md")); d.text("Entity\n(NER)", 452, 342, size="body", color="ink")
        d.text("Coreference\nResolution", 428, 396, size="body", color="ink")
        d.icon("🧩", 428, 436, size=d.theme.size("icon_sm")); d.text("Word Sense\nDisambiguation\n(WSD)", 428, 478, size="label", color="ink")
    with d.group("nli"):
        d.rect(492, 308, 159, 222, fill="nli_fill", stroke="tile_purple_line", rounded=6)
        d.text("**NATURAL LANGUAGE**\n**INFERENCE (NLI)**", 572, 332, size="body", color="ink")
        d.text("Premise (P) & Hypothesis (H)", 572, 358, size="small", color="ink")
        d.text("Logical Relations:", 590, 397, size="small", color="ink")
        d.rect(500, 390, 25, 25, "P", fill="tile_blue", stroke="tile_blue_line", rounded=3, size="body", color="ink", id="P")
        d.rect(500, 437, 25, 25, "H", fill="hyp", stroke="tile_tan_line", rounded=3, size="body", color="ink", id="H")
        rel = d.column(538, 408, 3, 97, 22, gap=8)
        for (x, y, w, h), (lb, f, s, idn) in zip(rel, [("Entailment", "tile_green", "tile_green_line", "ent"),
                                                        ("Contradiction", "contra", "contra_line", "con"),
                                                        ("Neutral", "background", "neutral", "neu")]):
            d.rect(x, y, w, h, lb, fill=f, stroke=s, rounded=4, size="label", color="ink", id=idn)
        d.text("*(Examples for examples)*", 572, 506, size="small", color="muted")
    d.edge([], src="P", tgt="H", color="ink", sw="thin", exit=(.5, 1), entry=(.5, 0))
    for t in ("ent", "con", "neu"):  # branch: right out of H, then up/down into each relation
        d.edge([], src="H", tgt=t, color="ink", sw="thin", route="orthogonal", exit=(1, .5), entry=(0, .5))

    # --- NLG panel
    nlg = [("✍️", "Content\nDetermination", "cd"), ("🔀", "Sentence\nPlanning", "sp"), ("🌐", "Surface\nRealization", "sre")]
    boxes = d.fit_row(705, 982, 198, 3, 102, gap=9)
    for box, (ic, lb, idn) in zip(boxes, nlg):
        tile(d, "t_" + idn, box, "tan", ic, lb, icy=266, lcy=219, size="label")
    d.rect(810, 330, 95, 60, "**NLG**\n**Model**", fill="tile_tan", stroke="tile_tan_line", rounded=6, size="heading", color="ink", id="nlg_model")
    for (x, y, w, h), idn in zip(boxes, ("cd", "sp", "sre")):
        cx = x + w / 2
        tx = min(max(cx, 825), 890)
        d.edge([(cx, 300), (cx, 315), (tx, 315), (tx, 330)], src=f"t_{idn}_box", tgt="nlg_model", color="ink",
               exit=(.5, 1), entry=((tx - 810) / 95, 0))
    with d.group("data_intent"):
        d.icon("📄", 742, 354, size=d.theme.size("icon_sm")); d.text("Data/Intent", 742, 386, size="body", color="ink")
    d.edge([(768, 360)], tgt="nlg_model", color="ink", entry=(0, .5))
    d.edge([(927, 360)], src="nlg_model", color="ink", exit=(1, .5))
    with d.group("coherent"):
        d.icon("💬", 946, 352, size=d.theme.size("icon_sm")); d.text("Coherent\nText", 946, 394, size="body", color="ink")
    with d.group("applications"):
        for cy, s in zip((422, 446, 466, 487, 507), ("**APPLICATIONS:**", "•  Document Summarization",
                                                     "•  Dialogue Systems (Chatbots)", "•  Question Answering", "•  Paraphrasing")):
            d.text(s, 712, cy, anchor="w", size="body", color="ink")
    d.icon("💬", 946, 440, size=d.theme.size("icon_md")); d.icon("🔗", 958, 497, size=d.theme.size("icon_sm"))

    # footer rule with centered label
    d.line([(33, 547), (425, 547)], color="rule", sw="thin"); d.line([(600, 547), (990, 547)], color="rule", sw="thin")
    d.text("RESEARCH & EDUCATION", 512, 547, size="small", color="muted")


if __name__ == "__main__":
    run(layout, (1024, 559), THEME, "02_tile_dashboard")
'''

# ================================================================================================
FILES['examples/03_vector_space_plot.py'] = r'''#!/usr/bin/env python3
"""
Tier 3 — conceptual plot / geometry (discrete entities -> continuous vector space), 616x295 source.
Shows: wavy wide arrow (flex_arrow; curve waypoints are CONTROL points), axes with arrowheads,
rotated axis label, invisible anchor so vectors stay attached when points move, dashed cluster
circle, annotation pointers attached to shapes, header bands using a neutral fill at reduced opacity.
All coordinates were measured on the true pixel grid (inspect_image.py zoom / numeric probes).
Tip: run with --scale 2 (the source is small).
"""
from _common import run

THEME = {"sizes": {"heading": 13.5, "body": 15, "subheading": 13, "caption": 12, "label": 12, "note": 10.5, "axis": 8.5},
         "strokes": {"normal": 1.5, "thick": 2.5}}


def layout(d):
    # ---- panel 1: discrete entities
    d.panel(8, 9, 189, 275, "1. Discrete Entities (Input)", header_h=34, header_fill="neutral", header_opacity=55,
            header_color="primary", title_size="heading", stroke="neutral", sw="thick", id="p1")
    for (x, y, w, h), t in zip(d.column(57, 66, 3, 91, 36, gap=23.5), ["steel", "concrete", "banana"]):
        d.rect(x, y, w, h, t, stroke="neutral", sw="thick", size="body", id="e_" + t)
    d.text("**Separate, discrete entities**\n**with no inherent relationship.**", 102, 263, size="caption", color="primary")

    # ---- middle: mapping arrow (anchors pass through the midpoints of consecutive waypoints)
    d.text("**Mapping & Training**\n**Process**", 268, 99, size="subheading", color="primary")
    d.flex_arrow([(186, 149), (220, 124), (260, 140), (310, 160), (347, 140)], width=14, end_width=16, end_size=6,
                 fill="primary")
    d.text("**Assign to vectors &**\n**Train to reflect**\n**relationships**", 265, 190, size="caption", color="primary")

    # ---- panel 2: vector space
    d.panel(343, 9, 264, 275, "2. Continuous Vector Space (Output)", header_h=34, header_fill="neutral",
            header_opacity=55, header_color="primary", title_size="heading", stroke="neutral", sw="thick", id="p2")
    d.arrow([(375, 235), (375, 58)], color="neutral", sw="normal", head="block", head_size=5)   # Dimension 2 axis
    d.arrow([(362, 223), (562, 223)], color="neutral", sw="normal", head="block", head_size=5)  # Dimension 1 axis
    d.text("Dimension 2", 365, 144, size="axis", rotation=-90)
    d.text("Dimension 1", 465, 232, size="axis")
    d.circle(452, 112, 42, fill="none", stroke="neutral", sw="thick", dashed=True, dash="5 4", id="cluster")
    d.circle(375, 223, 1, fill="none", stroke="none", id="origin")            # invisible anchor for the vectors
    for name, (cx, cy) in {"steel": (429, 112), "concrete": (441, 114), "banana": (510, 195)}.items():
        d.dot(cx, cy, 4, fill="primary", stroke="primary", id="pt_" + name)
        d.edge([], src="origin", tgt="pt_" + name, color="primary", sw="thick", head_size=6)
    d.text("steel", 428, 100, size="label"); d.text("concrete", 469, 119, size="label"); d.text("banana", 525, 209, size="label")
    d.text("**Similar entities**\n**have similar vectors**\n**(close together)**", 550, 72, size="note", color="primary")
    d.edge([(523, 93)], tgt="cluster", color="primary", sw="normal", head_size=5)
    d.text("**Dissimilar entity**\n**is far away**", 558, 164, size="note", color="neutral")
    d.edge([(534, 178)], tgt="pt_banana", color="neutral", sw="normal", head_size=5)
    d.text("**Vectors in continuous space reflect**\n**meaningful relationships.**", 477, 263, size="caption", color="primary")


if __name__ == "__main__":
    run(layout, (616, 295), THEME, "03_vector_space_plot")
'''

# ================================================================================================
FILES['examples/04_complex_manifold.py'] = r'''#!/usr/bin/env python3
"""
Tier 4 — complex scientific figure (manifold hypothesis, 3-D surface + 2-D projection), 1408x768.
Shows: pseudo-3-D mesh surface (surface_mesh -> embedded SVG), wireframe cube with back/front edge
z-ordering, seeded scatter points (inside a polygon with keep-out zones, and along a curve),
photos cropped from the source (image_crop), translucent projection beams (polygon), stars,
closed curved loop, dashed straight-line distance attached to two stars, callouts with pointers,
small-caps serif title (raw HTML), justified serif caption with italics, credit line.
Run with --image pointing at the source so the two photo crops can be extracted.
"""
import math

from _common import run

THEME = {
    "colors": {"chunk_b": "#5b8c4a", "cluster_d": "#c2a27a"},
    "sizes": {"title": 30, "heading": 24, "body": 18.5, "label": 17, "subheading": 20, "small": 12.5, "caption": 23.5},
    "opacity": {"tint": 18, "beam": 28, "soft": 35, "chunk_c": 75, "faded": 35},
}


def photo(d, bbox, label):
    if d.source:
        return d.image_crop(bbox, border="text")
    x0, y0, x1, y1 = bbox
    return d.rect(x0, y0, x1 - x0, y1 - y0, f"*[photo: {label}]*", fill="neutral", fill_opacity="soft", size="small")


def layout(d):
    d.text('<span style="font-variant: small-caps">Visualizing the Manifold Hypothesis: Chunk A vs. Chunk B Geometry</span>',
           704, 30, w=1200, size="title", font="serif", raw=True)

    # ================= Panel A =================
    d.rect(35, 57, 665, 538, rounded=12, sw="thick", id="panelA")
    d.text("**A) Ambient High-Dimensional Data Space (D=3,072)**", 50, 85, anchor="w", size="heading")
    for p in ([(330, 110), (330, 410)], [(330, 410), (152, 465)], [(330, 410), (678, 505)]):   # back edges
        d.line(p, color="neutral", sw="thin")
    cube = [(152, 165), (330, 110), (678, 165), (678, 505), (340, 590), (152, 465)]
    keepout = [(40, 120, 218, 192), (470, 105, 660, 182), (572, 203, 684, 352), (500, 368, 692, 432),
               (448, 470, 696, 582), (40, 518, 305, 578), (40, 330, 158, 448), (44, 194, 150, 314)]
    pts = d.scatter_in_polygon(cube, 190, keepout=keepout, min_dist=12, seed=7)
    with d.group("points_behind"):
        for x, y in pts[::2]:
            d.dot(x, y, 4.5, fill="neutral", stroke="text", stroke_opacity=25)
    d.surface_mesh(back=[(172, 345), (240, 240), (330, 185), (430, 200), (510, 255), (548, 320)],
                   front=[(176, 352), (240, 420), (325, 495), (410, 470), (490, 385), (550, 326)],
                   nu=30, nv=14, bulge=38, fill="secondary", fill_opacity="tint", line="primary", id="manifold_surface")
    d.polygon([(143, 203), (143, 307), (282, 286), (280, 276)], fill="secondary", fill_opacity="beam", id="beam_A")
    d.polygon([(581, 215), (581, 343), (500, 309), (500, 301)], fill="chunk_b", fill_opacity="beam", id="beam_B")
    photo(d, (49, 199, 143, 309), "sound concrete")
    photo(d, (581, 211, 676, 346), "alligator cracking")
    loop = [(300, 272), (318, 250), (380, 252), (455, 288), (470, 312), (400, 330), (372, 356), (330, 322), (312, 290), (300, 272)]
    d.line(loop, color="primary", sw=3, route="curved", id="manifold_loop")
    d.star(289, 281, 34, fill="secondary", id="star_A")
    d.star(491, 305, 34, fill="chunk_b", id="star_B")
    d.edge([], src="star_A", tgt="star_B", arrow=None, dashed=True, sw="normal", id="geometric_distance")
    with d.group("points_front"):
        for x, y in pts[1::2]:
            d.dot(x, y, 4.5, fill="neutral", stroke="text", stroke_opacity=25)
    for p in ([(152, 165), (152, 465)], [(152, 165), (330, 110)], [(330, 110), (678, 165)], [(678, 165), (678, 505)],
              [(152, 465), (340, 590)], [(340, 590), (678, 505)]):                                   # front edges
        d.line(p, sw="normal")
    d.text("{secondary|**CHUNK A:**} Sound\nConcrete Pavement\n(3,072-D Vector)", 128, 160)
    d.text("**Latent Manifold** of\nMeaningful Human\nLanguage", 565, 145)
    d.text("{chunk_b|**CHUNK B:**}\nAlligator\nCracking\nDistress\n(3,072-\nD Vector)", 96, 388)
    d.text("**Geometric Distance**\n(Cosine Similarity,\nstraight line - Mismatch)", 592, 402, size="label")
    d.text("**True Semantic**\n**Distance** along Manifold", 170, 548)
    arr = dict(route="curved", sw="normal", head_size=7)
    d.arrow([(205, 150), (275, 160), (281, 265)], **arr)          # middle point = curve CONTROL point
    d.arrow([(458, 162), (412, 170), (408, 215)], **arr)
    d.arrow([(152, 338), (210, 288), (272, 286)], **arr)
    d.arrow([(512, 380), (465, 345), (410, 343)], dashed=True, **arr)
    d.arrow([(180, 502), (182, 462), (240, 462)], **arr)
    d.arrow([(282, 522), (312, 515), (312, 478)], **arr)
    d.callout(455, 475, 235, 100, "**Topological Mismatch:**\nLinear distance metrics often\ncut through empty space\nbetween manifold folds.",
              tip=(444, 402), base_at=497, fill="neutral", fill_opacity="soft", size="label", id="callout_topo")

    # ================= Panel B =================
    d.rect(712, 57, 681, 538, rounded=12, sw="thick", id="panelB")
    d.text("**B) Unfolded Latent Manifold Projection (d=2)**", 727, 85, anchor="w", size="heading")

    def arc_y(x):
        return 270 + (0.0020 if x < 1042 else 0.00247) * (x - 1042) ** 2
    # curved edge: waypoints are control points, so sample densely to pin the curve to the parabola
    d.line([(742, 465)] + [(x, arc_y(x)) for x in range(770, 1340, 30)] + [(1344, 465)], color="secondary", sw=3.5, route="curved")
    with d.group("points_2d"):
        for i, (x, y) in enumerate(d.scatter_along(arc_y, 750, 905, 38, spread=26, min_dist=10.5, seed=1)):
            d.dot(x, y, 5.5, fill="neutral", stroke="text", stroke_opacity=25)
        for i, (x, y) in enumerate(d.scatter_along(arc_y, 905, 950, 6, spread=22, min_dist=10.5, seed=2)):
            d.dot(x, y, 5.5, fill="neutral", stroke="text", stroke_opacity=25, opacity="faded" if i == 5 else None)
        for x, y in d.scatter_along(arc_y, 1040, 1105, 7, spread=14, min_dist=10.5, seed=3):
            d.dot(x, y, 5.5, fill="neutral", stroke="text", stroke_opacity=25)
        for x, y in d.scatter_along(arc_y, 1165, 1312, 34, spread=18, min_dist=10.5, seed=4):
            d.dot(x, y, 5.5, fill="cluster_d", stroke="text", stroke_opacity=25)
    d.star(877, 325, 34, fill="secondary", id="star_A_2d")
    d.star(1067, 276, 34, fill="secondary", opacity="chunk_c", id="star_C_2d")
    d.star(1227, 343, 34, fill="chunk_b", id="star_B_2d")
    d.text("Clustering of related pavement\nengineering concepts", 1032, 157)
    d.arrow([(947, 190), (895, 195), (890, 250)], **arr)
    d.arrow([(1117, 180), (1175, 185), (1178, 228)], **arr)
    d.text("{secondary|**Chunk A:**}", 870, 284, size="subheading")
    d.text("**Chunk C:**\nMinor Joint Spalling", 1068, 236)
    d.text("**Chunk D:**\nModerate\nSurface Scaling", 1262, 242)
    d.text("{chunk_b|**Chunk B:**}\nAlligator Cracking", 1068, 344, size="subheading")
    d.arrow([(842, 402), (1010, 495), (1212, 380)], **arr)
    d.arrow([(848, 412), (1040, 560), (1266, 408)], **arr)
    d.text("**Semantic Continuity**\npreserved by\nManifold Hypothesis", 1043, 520)
    d.callout(1162, 475, 186, 82, "**RAG target:** Retrieve\ncontext based on local\nmanifold proximity.",
              tip=(1228, 384), base_at=1232, size="label", id="callout_rag")

    # ================= credit + caption =================
    d.text("Figure adapted from Mason Smetana\u2019s Doctoral Dissertation (University of Pittsburgh, "
           "Swanson School of Engineering, 2026).", 1393, 607, anchor="e", size="small")
    d.text("Figure 1: Geometric representation of the Manifold Hypothesis applied to two engineering text chunks. "
           "(A) Shows Chunk A and Chunk B in high-dimensional ambient space. While computationally close in D-space, "
           "straight-line distance misinterprets the manifold curvature. (B) Illustrates the same points on an unfolded "
           "2D semantic manifold, where true conceptual *proximity* and *continuity* are revealed, critical for "
           "optimizing semantic retrieval in RAG systems.",
           35, 692, w=1360, h=130, anchor="w", size="caption", font="serif", wrap=True, extra={"align": "left"})


if __name__ == "__main__":
    run(layout, (1408, 768), THEME, "04_complex_manifold")
'''

# ================================================================================================
FILES['examples/_common.py'] = r'''"""Shared CLI for the example builds:  python 0X_example.py [--image SRC] [--out OUT.drawio] [--theme T] [--scale S]"""
import argparse
import os
import sys

SCRIPTS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts")
sys.path.insert(0, SCRIPTS)
from drawio_kit import Diagram, Theme  # noqa: E402


def run(layout, size, theme_overrides=None, name="example", coord_scale=(1.0, 1.0)):
    ap = argparse.ArgumentParser()
    ap.add_argument("--image", help="source image (needed for image_crop; also sets page size)")
    ap.add_argument("--out", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "out", name + ".drawio"))
    ap.add_argument("--theme", default=None, help="theme file (default: assets/theme_image.yaml)")
    ap.add_argument("--scale", type=float, default=1.0)
    a = ap.parse_args()
    theme = Theme.load(a.theme or os.path.join(os.path.dirname(SCRIPTS), "assets", "theme_image.yaml"), overrides=theme_overrides)
    if a.image:
        d = Diagram.from_source(a.image, theme=theme, scale=a.scale, coord_scale=coord_scale)
    else:
        d = Diagram(*size, theme=theme, scale=a.scale, coord_scale=coord_scale)
    layout(d)
    d.save(a.out)
    return a.out
'''

# ================================================================================================
FILES['examples/A_methods_pipeline/build.py'] = r'''"""Prototypes for passage A (methods: RAG pipeline). Run: python ../../scripts/prototype.py ."""
from figlib import *

INFERRED = []
ALLOWED_TERMS = ["layer", "representation", "retrieval", "evaluation", "data", "parameter", "tuning",
                 "practitioner", "online", "answering", "database", "human", "loop", "ingest"]

HL = style(DARK, bold=False)


def v1(d):
    """Swimlane pipeline: offline indexing lane above, online answering lane below, feedback loop."""
    lane(d, 20, 20, 1060, 175, "Offline indexing")
    lane(d, 20, 215, 1060, 170, "Online answering")
    X = [x for x, _, _, _ in d.row(80, 0, 6, 140, 60, gap=30)]
    off = chain(d, ["Inspection reports\n(scanned)", "OCR\nto text", "Semantic\nchunking",
                    "Embedding model\n(3,072-D)", "Vector database"], 80, 78, 140, 60, gap=30,
                shapes=["document", None, None, None, "cylinder"], ids=["o_rep", "o_ocr", "o_chk", "o_emb", "o_db"])
    d.text("*+ route · milepost · year*", X[4] + 64, 154, anchor="e", size="small", color="primary")
    d.rect(X[5], 78, 140, 60, "**Parameter tuning**\nchunk size · k · prompt", id="tune",
           **style(SOFT, dashed=True, size="label"))
    on = chain(d, ["Practitioner\nquestion", "Query\nembedding", "Retrieve top-k\nsimilar chunks", "Cross-encoder\nre-ranking",
                   "Prompt + LLM\ncited answer", "**Engineer review**\naccuracy · usefulness"], 80, 270, 140, 60, gap=30,
               styles=[SOFT] * 5 + [HL], ids=["n_q", "n_emb", "n_topk", "n_rr", "n_llm", "n_rev"])
    d.text("*same embedding model*", X[1] + 70, 346, size="small", color="secondary")
    d.edge([(X[4] + 70, 205), (X[2] + 70, 205)], src="o_db", tgt="n_topk", color="primary", sw="thick",
           exit=(.5, 1), entry=(.5, 0), label="*retrieve*", label_size="small")
    d.edge([], src="n_rev", tgt="tune", color="secondary", sw="thick", dashed=True, exit=(.5, 0), entry=(.5, 1), label="ratings")
    d.edge([(X[5] + 70, 50), (X[2] + 70, 50)], src="tune", tgt="o_chk", color="secondary", sw="thick", dashed=True,
           exit=(.5, 0), entry=(.5, 0))
    d.text("*closes the loop between system output and expert judgment*", (X[2] + X[5]) / 2 + 70, 38, size="small", color="secondary")


def v2(d):
    """Layered architecture: five layers, flow downward, tuning feedback on the right."""
    spec = [("Data\nlayer", ["Scanned inspection\nreports", "OCR\nto text", "Metadata\nroute · milepost · year"]),
            ("Representation\nlayer", ["Semantic\nchunks", "Embeddings\n(3,072-D)", "Vector\ndatabase"]),
            ("Retrieval\nlayer", ["Query\nembedding", "Retrieve top-k\nsimilar chunks", "Cross-encoder\nre-ranking"]),
            ("Generation\nlayer", ["Structured\nprompt", "Large language\nmodel", "Answer with\ncitations"]),
            ("Evaluation\nlayer", ["**Engineer review**\naccuracy · usefulness", "Ratings\nlog", "Parameter\ntuning"])]
    ids, boxes = layers(d, 40, 20, 940, spec, styles={2: SOFT, 3: SOFT})
    d.get(ids[(4, 0)])["style"].update(fillColor=d.theme.color("primary"), fontColor=d.theme.color("background"))
    yb = boxes[4][1] + boxes[4][3] / 2; yr = boxes[1][1] + boxes[1][3] / 2
    d.edge([(980, yb), (1015, yb), (1015, yr), (980, yr)], color="secondary", sw="thick", dashed=True)
    d.text("{secondary|*tune chunk size · k · prompt*}", 1032, (yb + yr) / 2, rotation=-90, size="label")


def v3(d):
    """Numbered cycle around a human-in-the-loop hub; closing arrow = tuning feedback."""
    labels = ["Ingest & OCR\nscanned reports", "Semantic\nchunking", "Embed chunks\n(3,072-D)", "Vector index\n+ metadata",
              "Embed question,\nretrieve top-k", "Cross-encoder\nre-ranking", "LLM answer\nwith citations", "**Engineer**\n**review**"]
    cycle(d, labels, 470, 320, 330, 240, styles=[NODE] * 4 + [SOFT] * 3 + [HL])
    d.text("{secondary|*tune chunk size · k · prompt*}", 245, 58, size="label")
    d.circle(470, 320, 108, "**Human-in-the-loop RAG**\nfor pavement\ninspection reports", fill="primary", stroke="primary",
             color="background", size="subheading")
    legend(d, 220, 648, [(NODE, "Offline indexing"), (SOFT, "Online answering"), (HL, "Expert judgment")])


VARIANTS = [
    dict(key="v1", name="swimlane_pipeline", size=(1100, 400), build=v1,
         about="Offline and online stages as two lanes; retrieval and the tuning loop drawn as routed connectors.",
         caption="Retrieval-augmented generation framework for pavement inspection reports: offline indexing (top) and online answering with engineer review and parameter tuning (bottom)."),
    dict(key="v2", name="layered_architecture", size=(1060, 600), build=v2,
         about="Five system layers stacked top to bottom with their components; tuning feedback on the side.",
         caption="Layered view of the proposed RAG framework, from scanned reports to expert evaluation; engineer ratings tune chunk size, k and the prompt template."),
    dict(key="v3", name="numbered_cycle", size=(940, 680), build=v3,
         about="Eight stages around a human-in-the-loop hub; offline, online and expert stages styled differently.",
         caption="The RAG workflow as a closed loop: engineer review of cited answers feeds back into chunking, retrieval depth and prompting."),
]
'''

# ================================================================================================
FILES['examples/A_methods_pipeline/passage.md'] = r'''# Methods — RAG framework for pavement inspection reports

We propose a retrieval-augmented generation (RAG) framework for answering engineering questions over
pavement inspection reports. In an offline indexing stage, scanned reports are converted to text through
OCR, segmented into semantically coherent chunks, and encoded with a 3,072-dimensional embedding model;
the resulting vectors are stored in a vector database together with metadata such as route, milepost,
and inspection year. At query time, a practitioner's question is embedded with the same model, and the
top-k most similar chunks are retrieved and re-ranked by a cross-encoder. The re-ranked context is
inserted into a structured prompt, and a large language model generates an answer with citations to the
source chunks. Finally, pavement engineers review each answer for factual accuracy and usefulness; their
ratings are logged and used to tune the chunk size, the value of k, and the prompt template, closing the
loop between system output and expert judgment.
'''

# ================================================================================================
FILES['examples/B_conceptual_framework/build.py'] = r'''"""Prototypes for passage B (conceptual framework). Run: python ../../scripts/prototype.py ."""
import math
from figlib import *



def v1(d):
    """Causal columns: drivers -> mechanisms -> indicators -> decisions, moderators beneath."""
    cols = {"drv": 150, "mech": 425, "ind": 700, "dec": 975}
    for k, t in zip(cols, ["External drivers", "Distress mechanisms", "Condition indicators", "M&R decisions"]):
        pill(d, cols[k], 20, t, w=210)
    ys = [80, 150, 220, 290]
    drv = ["Traffic loading", "Temperature cycling", "Moisture infiltration", "Freeze–thaw action"]
    mech = ["**Fatigue (alligator) cracking**\n*load repetitions*", "**Rutting**\n*load + high temperature*",
            "**Thermal cracking**\n*low-temperature contraction*", "**Joint spalling**\n*moisture + freeze–thaw*"]
    for i, (a, b) in enumerate(zip(drv, mech)):
        d.rect(cols["drv"] - 95, ys[i], 190, 50, a, id=f"d{i}", **NODE)
        d.rect(cols["mech"] - 100, ys[i], 200, 50, b, id=f"m{i}", **dict(SOFT, size="label"))
    for a, b in [(0, 0), (0, 1), (1, 1), (1, 2), (2, 3), (3, 3)]:
        d.edge([], src=f"d{a}", tgt=f"m{b}", color="primary", sw="normal", exit=(1, .5), entry=(0, .5))
    d.rect(cols["mech"] - 110, 66, 220, 288, fill="none", stroke="secondary", sw="thin", dashed=True, rounded=10, id="mech_grp")
    d.rect(cols["ind"] - 95, 135, 190, 60, "**Pavement Condition**\n**Index (PCI)**", id="pci", **DARK)
    d.rect(cols["ind"] - 95, 225, 190, 60, "**International Roughness**\n**Index (IRI)**", id="iri", **DARK)
    d.rect(cols["ind"] - 110, 120, 220, 180, fill="none", stroke="primary", sw="thin", dashed=True, rounded=10, id="ind_grp")
    dec = ["Preventive treatment\n*e.g., crack sealing*", "Rehabilitation\n*e.g., overlay*", "Full reconstruction"]
    for i, t in enumerate(dec):
        d.rect(cols["dec"] - 95, 80 + i * 95, 190, 60, t, id=f"r{i}", **NODE)
    d.rect(cols["dec"] - 110, 66, 220, 288, fill="none", stroke="primary", sw="thin", dashed=True, rounded=10, id="dec_grp")
    d.edge([], src="mech_grp", tgt="ind_grp", color="primary", sw="heavy", exit=(1, .5), entry=(0, .5))
    d.text("*quantified by*", (cols["mech"] + 110 + cols["ind"] - 110) / 2, 196, anchor="s", size="small", color="primary")
    d.edge([], src="ind_grp", tgt="dec_grp", color="primary", sw="heavy", exit=(1, .5), entry=(0, .5))
    d.text("*inform*", (cols["ind"] + 110 + cols["dec"] - 110) / 2, 196, anchor="s", size="small", color="primary")
    d.arrow([(1112, 85), (1112, 335)], color="neutral", sw="heavy", head_size=6)
    d.text("*preventive → major rehabilitation*", 1132, 210, rotation=-90, size="small", color="neutral")
    d.rect(55, 400, 500, 64, "**Internal moderators** — material quality · construction practice", id="mod", **MUTED)
    d.edge([(445, 377)], src="mod", tgt="mech_grp", color="neutral", sw="thick", dashed=True, exit=(.78, 0), entry=(.5, 1))
    d.text("*moderate the structural response*", 455, 377, anchor="w", size="small", color="neutral")

def v2(d):
    """Stepped funnel: four stages narrowing top-to-bottom, chips inside, moderators to the side."""
    cx = 440
    stages = [("External drivers", ["Traffic loading", "Temperature cycling", "Moisture infiltration", "Freeze–thaw"], 840, "neutral", 18),
              ("Distress mechanisms", ["Fatigue cracking", "Rutting", "Thermal cracking", "Joint spalling"], 700, "secondary", 18),
              ("Condition indicators", ["PCI", "IRI"], 560, "secondary", 35),
              ("M&R decisions", ["Crack sealing", "Overlay", "Reconstruction"], 440, "primary", 100)]
    y = 20
    for i, (name, chips, w, col, op) in enumerate(stages):
        h = 116
        dark = op == 100
        d.rect(cx - w / 2, y, w, h, fill=col, fill_opacity=op, stroke=col, sw="normal", rounded=14, id=f"st{i}")
        d.text(f"**{name}**", cx, y + 22, size="subheading", color="background" if dark else "primary")
        cw = min(170, (w - 40 - 14 * (len(chips) - 1)) / len(chips))
        if len(chips) == 2:
            cw = 200
        total = cw * len(chips) + 14 * (len(chips) - 1)
        for j, (bx, by, bw, bh) in enumerate(d.row(cx - total / 2, y + 48, len(chips), cw, 46, gap=14)):
            d.rect(bx, by, bw, bh, chips[j], fill="background", stroke="primary" if not dark else "background",
                   rounded=23, size="body", id=f"c{i}{j}")
        if i < 3:
            d.shape("triangle_down", cx - 18, y + h + 6, 36, 18, fill="primary", stroke="primary")
        y += h + 30
    d.text("{background|*preventive → major rehabilitation*}", cx, 20 + 3 * 146 + 102, size="small")
    d.rect(860, 176, 165, 96, "**Moderators**\nmaterial quality\nconstruction practice", fill="neutral", fill_opacity="tint",
           stroke="neutral", dashed=True, rounded=True, size="label", id="mod")
    d.edge([], src="mod", tgt="st1", color="neutral", sw="thick", dashed=True, exit=(0, .5), entry=(1, .5))

def v3(d):
    """Concept map: pavement structure as the hub, labeled relations."""
    hub = (530, 260)
    d.ellipse(hub[0] - 100, hub[1] - 46, 200, 92, "**Pavement**\n**structure**", fill="primary", stroke="primary",
              color="background", size="subheading", id="hub")
    ys = [110, 220, 330, 440]
    for i, t in enumerate(["Traffic\nloading", "Temperature\ncycling", "Moisture\ninfiltration", "Freeze–thaw\naction"]):
        d.rect(60, ys[i] - 28, 150, 56, t, id=f"drv{i}", **NODE)
        d.edge([], src=f"drv{i}", tgt="hub", color="primary", sw="normal")
    d.text("*act on*", 300, 214, size="small", color="primary")
    for i, (t, x) in enumerate([("Material quality", 420), ("Construction\npractice", 640)]):
        d.rect(x - 85, 40, 170, 56, t, id=f"mod{i}", **MUTED)
        d.edge([], src=f"mod{i}", tgt="hub", color="neutral", sw="normal", dashed=True)
    d.text("*moderate*", 530, 150, size="small", color="neutral")
    mech = ["**Fatigue cracking**\n*load repetitions*", "**Rutting**\n*load + high T*",
            "**Thermal cracking**\n*low-T contraction*", "**Joint spalling**\n*moisture + freeze–thaw*"]
    for i, t in enumerate(mech):
        d.rect(840, ys[i] - 28, 180, 56, t, id=f"mech{i}", **dict(SOFT, size="label"))
        d.edge([], src="hub", tgt=f"mech{i}", color="secondary", sw="normal")
    d.text("{secondary|*manifests as*}", 750, 214, size="small")
    d.rect(825, 66, 210, 418, fill="none", stroke="secondary", sw="thin", dashed=True, rounded=12, id="mech_grp")
    d.rect(430, 470, 200, 56, "**PCI · IRI**\ncondition indicators", id="ind", **DARK)
    d.edge([(930, 498)], src="mech_grp", tgt="ind", color="primary", sw="thick", exit=(.5, 1), entry=(1, .5))
    d.text("*quantified by*", 780, 488, anchor="s", size="small", color="primary")
    d.rect(410, 570, 240, 56, "**M&R decisions**\n*crack sealing → reconstruction*", id="dec", **NODE)
    d.edge([], src="ind", tgt="dec", color="primary", sw="thick", exit=(.5, 1), entry=(.5, 0))
    d.text("*inform*", 545, 548, anchor="w", size="small", color="primary")


INFERRED = []
ALLOWED_TERMS = []

VARIANTS = [
    dict(key="v1", name="causal_columns", size=(1150, 500), build=v1,
         about="Drivers -> mechanisms -> indicators -> decisions, with the specific driver-mechanism links and moderators.",
         caption="Conceptual framework of pavement deterioration: external drivers act through distinct distress mechanisms, quantified by PCI and IRI, which inform maintenance and rehabilitation decisions."),
    dict(key="v2", name="stepped_funnel", size=(1040, 610), build=v2,
         about="Four stages narrowing top to bottom; compact overview without individual links.",
         caption="From drivers to decisions: the cumulative effect of distress mechanisms is quantified by condition indicators that inform treatment selection."),
    dict(key="v3", name="concept_map", size=(1060, 640), build=v3,
         about="Pavement structure as the hub with labeled relations (act on, moderate, manifests as, quantified by, inform).",
         caption="Concept map of pavement deterioration: traffic and environmental exposure act on the pavement structure, moderated by material quality and construction practice."),
]
'''

# ================================================================================================
FILES['examples/B_conceptual_framework/passage.md'] = r'''# Conceptual framework — pavement deterioration

Pavement deterioration is commonly understood as the outcome of interacting external drivers and
internal material properties. Repeated traffic loading and environmental exposure (temperature cycling,
moisture infiltration, and freeze-thaw action) act on the pavement structure, whose response is
moderated by material quality and construction practice. These interactions manifest through distinct
distress mechanisms: fatigue (alligator) cracking is driven primarily by load repetitions, rutting by
load combined with high temperatures, thermal cracking by low-temperature contraction, and joint
spalling by moisture and freeze-thaw damage. Agencies quantify the cumulative effect of these distresses
using condition indicators such as the Pavement Condition Index (PCI) and the International Roughness
Index (IRI). These indicators, in turn, inform maintenance and rehabilitation decisions, ranging from
preventive treatments like crack sealing to major rehabilitation such as overlays or full reconstruction.
'''

# ================================================================================================
FILES['examples/C_comparison/build.py'] = r'''"""Prototypes for passage C (comparison of three approaches). Run: python ../../scripts/prototype.py ."""
import math
from figlib import *

METHODS = [("Rule-based", "keyword dictionaries +\nexpert heuristics"),
           ("Classical ML", "e.g., SVM on\nTF-IDF features"),
           ("LLM-based", "prompting or\nfine-tuning")]
# ratings: 3 = most favourable; words come from the text
CRIT = ["Labeled data\nrequired", "Interpretability", "Generalization", "Cost"]
R = [[(3, "none"), (3, "full"), (1, "poor on unseen terms"), (2, "costly to maintain")],
     [(1, "moderate"), (2, "partial (feature weights)"), (2, "within a single agency"), (2, "moderate")],
     [(2, "little task-specific"), (1, "least"), (3, "across agencies & styles"), (1, "highest compute")]]


def v1(d):
    """Comparison matrix with 3-dot ratings and the source wording under each rating."""
    x0, w0, cw, hh, rh = 20, 220, 190, 64, 96
    d.rect(x0, 20, w0, hh, "**Approach**", fill="primary", stroke="background", color="background", size="subheading")
    for j, c in enumerate(CRIT):
        d.rect(x0 + w0 + j * cw, 20, cw, hh, f"**{c}**", fill="primary", stroke="background", color="background", size="subheading")
    for i, (m, desc) in enumerate(METHODS):
        y = 20 + hh + i * rh
        band = "secondary" if i % 2 == 0 else "background"
        d.rect(x0, y, w0 + 4 * cw, rh, fill=band, fill_opacity="wash" if i % 2 == 0 else None, stroke="neutral", sw="thin")
        d.text(f"**{m}**", x0 + 16, y + 30, anchor="w", size="subheading", color="primary")
        d.text(f"*{desc}*", x0 + 16, y + 64, anchor="w", size="label")
        for j, (n, word) in enumerate(R[i]):
            cx = x0 + w0 + j * cw + cw / 2
            rating_dots(d, cx, y + 36, n)
            d.text(word, cx, y + 68, size="small", w=cw - 16, wrap=True)
    yb = 20 + hh + 3 * rh + 18
    rating_dots(d, 40, yb + 12, 3, r=5, gap=14); d.text("most favourable", 66, yb + 12, anchor="w", size="small")
    rating_dots(d, 210, yb + 12, 1, r=5, gap=14); d.text("least favourable", 236, yb + 12, anchor="w", size="small")
    d.rect(400, yb - 4, 580, 34, "**No single approach dominates:** choose by data availability, auditability, and budget.",
           fill="secondary", fill_opacity="tint", stroke="secondary", rounded=17, size="label")

def v2(d):
    """Three cards: how it works, strengths, limitations, best fit; conclusion banner."""
    cards = [("Rule-based", "Keyword dictionaries and expert-defined heuristics.",
              ["No training data needed", "Fully interpretable"], ["Poor generalization to unseen terms", "Costly to maintain"],
              "need for auditability;\nno training data"),
             ("Classical ML", "Support vector machines trained on TF-IDF features.",
              ["Partial interpretability (feature weights)", "Generalizes within a single agency"], ["Needs moderate labeled data", "Vocabulary of a single agency"],
              "moderate labeled data within\na single agency"),
             ("LLM-based", "Prompting or fine-tuning a large language model.",
              ["Little task-specific labeled data", "Best generalization across agencies"], ["Least interpretable", "Highest computational cost"],
              "generalization across agencies;\ndeployment budget available")]
    for i, (name, how, pros, cons, fit) in enumerate(cards):
        x = 20 + i * 333
        w = 313
        with d.group(f"card{i}"):
            d.rect(x, 20, w, 390, fill="background", stroke="neutral", sw="thin", rounded=12, shadow=True)
            d.rect(x, 20, w, 54, f"**{name}**", fill="primary", stroke="primary", color="background", rounded=12, size="heading")
            d.text(f"*{how}*", x + 18, 104, anchor="w", w=w - 36, h=40, wrap=True, size="label")
            d.text("**Strengths**", x + 18, 150, anchor="w", size="body", color="secondary")
            for k, p in enumerate(pros):
                d.text(f"{{secondary|**+**}}  {p}", x + 18, 178 + k * 26, anchor="w", size="label")
            d.text("**Limitations**", x + 18, 250, anchor="w", size="body", color="primary")
            for k, c in enumerate(cons):
                d.text(f"{{primary|**–**}}  {c}", x + 18, 278 + k * 26, anchor="w", size="label")
            d.rect(x + 14, 326, w - 28, 70, f"**Best fit**\n{fit}", fill="secondary", fill_opacity="tint",
                   stroke="none", rounded=8, size="label")
    d.rect(20, 430, 979, 50, "**No single approach dominates** — the choice depends on data availability, the need for auditability, and deployment budget.",
           fill="primary", stroke="primary", color="background", rounded=25, size="body")

def v3(d):
    """Positioning chart: generalization (x) vs interpretability (y); bubble area = computational cost."""
    ox, oy, x1, y1 = 120, 500, 860, 60
    d.rect(ox, y1, x1 - ox, oy - y1, fill="neutral", fill_opacity="wash", stroke="none")
    d.text("{neutral|*ideal*}", x1 - 40, y1 + 80, size="small")
    d.star(x1 - 40, y1 + 48, 26, fill="background", stroke="neutral", dashed=True)
    d.arrow([(ox, oy), (x1 + 20, oy)], color="text", sw="normal")
    d.arrow([(ox, oy), (ox, y1 - 20)], color="text", sw="normal")
    d.text("**Generalization** across agencies and writing styles", (ox + x1) / 2, oy + 48, size="body")
    d.text("**Interpretability**", ox - 62, (oy + y1) / 2, rotation=-90, size="body")
    for t, x in (("poor", ox + 110), ("within a single agency", (ox + x1) / 2), ("across agencies", x1 - 110)):
        d.text(t, x, oy + 18, size="small", color="neutral")
    for t, y in (("least", oy - 70), ("partial", (oy + y1) / 2), ("full", y1 + 70)):
        d.text(t, ox - 30, y, size="small", color="neutral")
    d.line([(ox + 145, y1 + 35), (ox + 440, y1 + 157), (x1 - 65, oy - 115)], color="neutral", sw="thick",
           dashed=True, route="curved")     # middle point = control point (curve bows toward it)
    d.text("{neutral|*trade-off frontier*}", ox + 470, y1 + 150, anchor="w", size="small")
    pts = [("Rule-based", ox + 110, y1 + 70, 16, "no labeled data"), ("Classical ML", (ox + x1) / 2, (oy + y1) / 2, 26, "moderate labeled data"),
           ("LLM-based", x1 - 110, oy - 70, 50, "little labeled data")]
    for i, (name, x, y, r, lab) in enumerate(pts):
        d.circle(x, y, r, fill="secondary" if i < 2 else "primary", fill_opacity=60 if i < 2 else 85, stroke="primary", id=f"b{i}")
        if i == 0:
            d.text(f"**{name}**\n*{lab}*", x, y + r + 8, anchor="n", size="label")
        elif i == 1:
            d.text(f"**{name}**\n*{lab}*", x + r + 14, y + 12, anchor="w", size="label")
        else:
            d.text(f"**{name}**\n*{lab}*", x - r - 14, y, anchor="e", size="label")
    # legend
    lx, ly = 150, 598
    d.text("Bubble size = computational cost", lx, ly, anchor="w", size="small")
    for k, r in enumerate((8, 14, 22)):
        d.circle(lx + 250 + k * 50, ly, r, fill="secondary", fill_opacity=60, stroke="primary")
    d.text("low → highest", lx + 390, ly, anchor="w", size="small", color="neutral")


INFERRED = [
    "v1: Classical ML 'Cost' rating (2 of 3, 'moderate') — the passage does not state ML cost",
    "v1: ordinal ratings (1-3 dots) are an interpretation of the passage's qualitative wording",
    "v2: 'Best fit when' statements are derived from the closing sentence (data availability, auditability, budget)",
    "v3: bubble sizes for rule-based and classical ML (only 'LLM = highest computational cost' is stated)",
    "v3: 'trade-off frontier' and 'ideal' corner are interpretive annotations",
]
ALLOWED_TERMS = ["choose", "most", "none", "approach", "cost", "favourable", "ideal", "frontier", "bubble", "size", "low", "medium", "fit",
                 "strengths", "limitations", "maintenance", "upkeep", "terms", "styles", "tuning"]

VARIANTS = [
    dict(key="v1", name="rating_matrix", size=(1000, 440), build=v1,
         about="Approaches x criteria with 3-dot ratings and the passage's own wording under each rating.",
         caption="Comparison of rule-based, classical machine-learning and LLM-based approaches to distress identification from inspection text."),
    dict(key="v2", name="method_cards", size=(1020, 500), build=v2,
         about="One card per approach: how it works, strengths, limitations, best fit; conclusion banner.",
         caption="Strengths and limitations of three families of methods for automated pavement distress identification."),
    dict(key="v3", name="positioning_chart", size=(940, 630), build=v3,
         about="Generalization vs. interpretability, bubble size = computational cost, trade-off frontier.",
         caption="No approach dominates: interpretability decreases as generalization across agencies increases."),
]
'''

# ================================================================================================
FILES['examples/C_comparison/passage.md'] = r'''# Comparison — methods for distress identification

Three families of methods have been applied to automated pavement distress identification from textual
inspection records. Rule-based systems rely on keyword dictionaries and expert-defined heuristics; they
require no training data and are fully interpretable, but they generalize poorly to unseen terminology
and are costly to maintain. Classical machine-learning models, such as support vector machines trained
on TF-IDF features, require moderate amounts of labeled data, offer partial interpretability through
feature weights, and generalize reasonably within a single agency's vocabulary. Large language model
(LLM) approaches, used either through prompting or fine-tuning, need little task-specific labeled data
and generalize best across agencies and writing styles, but they are the least interpretable and carry
the highest computational cost. These trade-offs suggest that no single approach dominates; the
appropriate choice depends on data availability, the need for auditability, and deployment budget.
'''

# ================================================================================================
FILES['examples/D_equation/build.py'] = r'''"""Prototypes for passage D (equation with explanation). Run: python ../../scripts/prototype.py ."""
import math
from figlib import *
from drawio_kit import measure_math


SIM = r"\mathrm{sim}(\mathbf{q},\mathbf{d})=\frac{\mathbf{q}\cdot\mathbf{d}}{\lVert\mathbf{q}\rVert\,\lVert\mathbf{d}\rVert}"


def v1(d):
    """Annotated equation: each term its own LaTeX label, color-coded, bracketed callouts, definition strip."""
    terms = [(r"s(q,d)", "primary"), (r"=", "text"), (r"\lambda", "secondary"),
             (r"\,\mathrm{sim}(\mathbf{q},\mathbf{d})", "primary"), (r"+", "text"), (r"(1-\lambda)", "secondary"),
             (r"\,\mathrm{BM25}(q,d)", "text")]
    annotated_equation(d, terms, 550, 205,
        notes_above=[(2, "secondary", "**Semantic weight λ** (0 to 1)\nλ = 1 → pure dense retrieval\nλ = 0.7 → best answer accuracy"),
                     (5, "secondary", "**Lexical weight 1 − λ**\nλ = 0 → keyword search")],
        notes_below=[(0, "primary", "**Hybrid relevance score**\nof chunk *d* for query *q*;\ntop-*k* go to the re-ranker"),
                     (3, "primary", "**Cosine similarity** of the\n3,072-D query and chunk\nembeddings (semantic evidence)"),
                     (6, "text", "**BM25 lexical matching**\nscore, normalized to [0, 1]\n(lexical evidence)")])
    d.rect(40, 425, 1020, 60, fill="secondary", fill_opacity="wash", stroke="neutral", sw="thin", rounded=10)
    d.text("**where**", 70, 455, anchor="w", size="label", color="primary")
    d.equation(SIM, 140, 455, anchor="w", size="eq_small", color="primary")
    d.text("is the cosine similarity between the embeddings of the query and the chunk.",
           140 + measure_math(SIM, d.theme.size("eq_small"))[0] + 24, 455, anchor="w", size="body")


def v2(d):
    """Computation graph: how the score is assembled from two evidence paths."""
    d.rect(30, 120, 130, 54, "Query *q*", id="q", **NODE)
    d.rect(30, 290, 130, 54, "Chunk *d*", id="dd", **NODE)
    d.rect(230, 60, 190, 60, "Embedding model\n(3,072-D)", id="emb", **SOFT)
    d.rect(230, 340, 190, 60, "Tokenize +\nBM25 scoring", id="bm", **NODE)
    d.rect(480, 50, 230, 80, "", id="cos", **SOFT)
    d.text("**Cosine similarity**", 595, 66, size="label", color="primary")
    d.equation(r"\frac{\mathbf{q}\cdot\mathbf{d}}{\lVert\mathbf{q}\rVert\,\lVert\mathbf{d}\rVert}", 595, 102, size="eq_small", color="primary")
    d.rect(480, 335, 230, 70, "", id="bmn", **NODE)
    d.text("**Normalized lexical score**", 595, 352, size="label")
    d.equation(r"\mathrm{BM25}(q,d)\in[0,1]", 595, 383, size="eq_small")
    d.circle(800, 90, 28, "", id="m1", fill="background", stroke="secondary", sw="thick")
    d.equation(r"\times\lambda", 800, 90, size="eq_small", color="secondary")
    d.circle(800, 370, 36, "", id="m2", fill="background", stroke="secondary", sw="thick")
    d.equation(r"\times(1-\lambda)", 800, 370, size="label", color="secondary")
    d.circle(890, 230, 30, "", id="plus", fill="primary", stroke="primary")
    d.equation(r"+", 890, 230, size="eq_term", color="background")
    d.rect(960, 195, 140, 70, "", id="s", **DARK)
    d.equation(r"s(q,d)", 1030, 220, size="eq_term", color="background")
    d.text("{background|rank → top-*k*}", 1030, 248, size="small")
    d.rect(620, 200, 170, 60, "", id="lam", fill="secondary", fill_opacity="tint", stroke="secondary", dashed=True, rounded=True)
    d.equation(r"\lambda\in[0,1]", 705, 220, size="eq_small", color="secondary")
    d.text("*0.7 best in our tests*", 705, 246, size="small", color="secondary")
    for a, b, ex, en in [("q", "emb", (1, .5), (0, .5)), ("dd", "emb", (1, .2), (0, .8)), ("q", "bm", (1, .8), (0, .2)),
                         ("dd", "bm", (1, .5), (0, .5)), ("emb", "cos", (1, .5), (0, .5)), ("bm", "bmn", (1, .5), (0, .5)),
                         ("cos", "m1", (1, .5), (0, .5)), ("bmn", "m2", (1, .5), (0, .5)), ("m1", "plus", (1, .5), (.5, 0)),
                         ("m2", "plus", (1, .5), (.5, 1)), ("plus", "s", (1, .5), (0, .5))]:
        d.edge([], src=a, tgt=b, color="primary", sw="thick", exit=ex, entry=en, route="orthogonal" if a in ("m1", "m2") else "straight")
    d.edge([], src="lam", tgt="m1", color="secondary", sw="normal", dashed=True, exit=(.85, 0), entry=(.5, 1))
    d.edge([], src="lam", tgt="m2", color="secondary", sw="normal", dashed=True, exit=(.85, 1), entry=(.5, 0))
    d.text("{secondary|**semantic path**}", 325, 30, size="label")
    d.text("**lexical path**", 325, 430, size="label")

def v3(d):
    """Intuition + glossary: vector angle for cosine similarity, λ continuum, symbol table."""
    # left: geometry
    d.rect(20, 20, 420, 330, fill="secondary", fill_opacity="wash", stroke="neutral", sw="thin", rounded=10)
    d.text("**Dense evidence: angle between embeddings**", 230, 44, size="label", color="primary")
    ox, oy = 90, 300
    d.circle(ox, oy, 1, fill="none", stroke="none", id="o")
    d.circle(380, 120, 1, fill="none", stroke="none", id="qt")
    d.circle(400, 250, 1, fill="none", stroke="none", id="dt")
    d.edge([], src="o", tgt="qt", color="primary", sw="heavy")
    d.edge([], src="o", tgt="dt", color="secondary", sw="heavy")
    d.equation(r"\mathbf{q}", 395, 108, size="eq_term", color="primary")
    d.equation(r"\mathbf{d}", 415, 252, size="eq_term", color="secondary")
    a1, a2 = math.atan2(120 - oy, 380 - ox), math.atan2(250 - oy, 400 - ox)
    arc = [(ox + 110 * math.cos(a1 + (a2 - a1) * t / 20), oy + 110 * math.sin(a1 + (a2 - a1) * t / 20)) for t in range(21)]
    d.polygon(arc, closed=False, stroke="text", fill=None, sw="normal")
    d.equation(r"\theta", ox + 130, oy - 70, size="eq_term")
    d.equation(r"\mathrm{sim}(\mathbf{q},\mathbf{d})=\cos\theta", 230, 330, size="eq_small", color="primary")
    # right: lambda continuum
    d.rect(470, 20, 610, 330, fill="background", stroke="neutral", sw="thin", rounded=10)
    d.text("**Balancing semantic and lexical evidence**", 775, 44, size="label", color="primary")
    d.equation(r"s(q,d)=\lambda\,\mathrm{sim}(\mathbf{q},\mathbf{d})+(1-\lambda)\,\mathrm{BM25}(q,d)", 775, 100, size="eq_term")
    x0, x1, yb = 520, 1030, 210
    d.rect(x0, yb - 8, x1 - x0, 16, fill="neutral", gradient="secondary", gradient_dir="east", stroke="none", rounded=8)
    for v, lab in ((0, "keyword search"), (1, "pure dense retrieval")):
        x = x0 + v * (x1 - x0)
        d.equation(rf"\lambda={v}", x, yb + 34, size="eq_small", color="primary")
        d.text(lab, x, yb + 62, size="small")
    xb = x0 + 0.7 * (x1 - x0)
    d.shape("triangle_down", xb - 11, yb - 36, 22, 18, fill="primary", stroke="primary")
    d.text("**λ = 0.7** best answer accuracy", xb, yb - 50, size="small", color="primary")
    d.text("lexical evidence ← → semantic evidence", 775, yb + 110, size="small", color="neutral")
    # bottom: symbol table
    rows = [(r"s(q,d)", "hybrid relevance score; top-k chunks go to the re-ranker", ""),
            (r"\mathrm{sim}(\mathbf{q},\mathbf{d})", "cosine similarity of 3,072-D query and chunk embeddings", r"[-1,1]"),
            (r"\mathrm{BM25}(q,d)", "lexical (keyword) matching score, normalized", r"[0,1]"),
            (r"\lambda", "weight on semantic vs. lexical evidence", r"[0,1]")]
    ty, rh = 370, 52
    for j, (h, w) in enumerate((("Symbol", 220), ("Meaning", 700), ("Range", 140))):
        xs = [20, 240, 940][j]
        d.rect(xs, ty, w, 36, f"**{h}**", fill="primary", stroke="background", color="background", size="label")
    for i, (sym, mean, rng) in enumerate(rows):
        y = ty + 36 + i * rh
        d.rect(20, y, 1060, rh, fill="secondary" if i % 2 == 0 else "background", fill_opacity="wash" if i % 2 == 0 else None,
               stroke="neutral", sw="thin")
        d.equation(sym, 130, y + rh / 2, size="eq_small", color="primary")
        d.text(mean, 260, y + rh / 2, anchor="w", size="body")
        if rng:
            d.equation(rng, 1010, y + rh / 2, size="eq_small")


INFERRED = [
    "v3: range [-1, 1] for cosine similarity is a mathematical property, not stated in the passage",
    "v3: 'angle between embeddings' (cos theta) is an explanatory interpretation of cosine similarity",
]
LINT_IGNORE = ["uneven gaps"]   # v1 callouts are aligned to their terms, not evenly spaced
ALLOWED_TERMS = ["path", "paths", "rank", "tokenize", "weight", "symbol", "meaning", "range", "angle", "balancing",
                 "kinds", "tests", "normalized", "semantic", "lexical", "evidence", "model"]

VARIANTS = [
    dict(key="v1", name="annotated_equation", size=(1100, 500), build=v1,
         about="Each term is its own LaTeX label, color-coded, with bracketed callouts; definition strip below.",
         caption="Hybrid relevance score: a lambda-weighted combination of embedding cosine similarity and normalized BM25."),
    dict(key="v2", name="computation_graph", size=(1120, 470), build=v2,
         about="How the score is assembled: semantic and lexical paths, weighting by lambda, sum, ranking.",
         caption="Computation of the hybrid score s(q, d) from dense (embedding) and lexical (BM25) evidence."),
    dict(key="v3", name="intuition_glossary", size=(1100, 640), build=v3,
         about="Geometric meaning of cosine similarity, the lambda continuum with the best setting, and a symbol table.",
         caption="Interpreting the hybrid score: lambda moves retrieval between keyword search (0) and pure dense retrieval (1); 0.7 gave the best accuracy."),
]
'''

# ================================================================================================
FILES['examples/D_equation/passage.md'] = r'''# Equation — hybrid relevance score

Each candidate chunk $d$ is scored against the query $q$ with a hybrid relevance function
$$s(q,d) = \lambda\,\mathrm{sim}(\mathbf{q},\mathbf{d}) + (1-\lambda)\,\mathrm{BM25}(q,d),$$
where $\mathrm{sim}(\mathbf{q},\mathbf{d}) = \frac{\mathbf{q}\cdot\mathbf{d}}{\lVert\mathbf{q}\rVert\,\lVert\mathbf{d}\rVert}$
is the cosine similarity between the 3,072-dimensional embeddings of the query and the chunk,
$\mathrm{BM25}(q,d)$ is a lexical matching score normalized to $[0,1]$, and $\lambda \in [0,1]$ controls
the balance between semantic and lexical evidence. Setting $\lambda = 1$ recovers pure dense retrieval,
whereas $\lambda = 0$ reduces to keyword search; in our experiments $\lambda = 0.7$ gave the best answer
accuracy. The top-$k$ chunks ranked by $s(q,d)$ are passed to the re-ranker.
'''

# ================================================================================================
FILES['examples/E_matrix_computation/build.py'] = r'''"""Prototypes for passage E (Matrix Computation & Topic Models: NMF vs. LDA).
Run: python ../../scripts/prototype.py .
"""
import math
import os
import sys

from figlib import *
from drawio_kit import Diagram, measure_math


def v1(d):
    """v1: Explicit populated matrix factorization grid V ≈ W × H with cell dot-product arithmetic."""
    # Outer Panel
    d.rect(20, 15, 1080, 390, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Non-negative Matrix Factorization (NMF): Linear Algebra Computation",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Decomposing document-term frequency matrix V into topic word weights W and document topic weights H (V ≈ W × H)",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # 1. Document-Term Matrix V (5 words × 5 docs)
    vx_top = 165
    d.text("**Document-Term Matrix V**", vx_top, 72, anchor="center", size="label", bold=True, color="v_matrix")
    d.text("Input frequency / TF-IDF matrix (5 × 5)", vx_top, 86, anchor="center", size="small", italic=True, color="neutral")

    doc_labels = ["d_1", "d_2", "d_3", "d_4", "d_5"]
    grid_vx, grid_vy = 105, 115
    cell_w, cell_h = 36, 26

    for j, d_lab in enumerate(doc_labels):
        d.text(d_lab, grid_vx + j * cell_w + cell_w // 2, grid_vy - 10, anchor="center", size="small", bold=True, color="text")

    word_labels = ["neural", "network", "factor", "topic", "corpus"]
    for i, w_lab in enumerate(word_labels):
        d.text(f"w_{i+1}: {w_lab}", grid_vx - 8, grid_vy + i * cell_h + cell_h // 2, anchor="e", size="small", color="text")

    v_data = [
        ["0.72", "0.00", "0.65", "0.00", "0.00"],
        ["0.55", "0.00", "0.58", "0.00", "0.00"],   # Target cell is (1, 2) -> network in d_3
        ["0.00", "0.45", "0.00", "0.62", "0.00"],
        ["0.00", "0.68", "0.15", "0.54", "0.40"],
        ["0.20", "0.35", "0.00", "0.30", "0.50"],
    ]

    target_r, target_c = 1, 2
    for r in range(5):
        for c in range(5):
            cx = grid_vx + c * cell_w
            cy = grid_vy + r * cell_h
            val = v_data[r][c]
            if r == target_r and c == target_c:
                d.rect(cx, cy, cell_w, cell_h, label=f"**{val}**", fill="calc_light", stroke="calc_highlight", sw="thick", size="small", color="calc_highlight")
            elif val == "0.00":
                d.rect(cx, cy, cell_w, cell_h, label=val, fill="sparse_zero", stroke="panel_border", sw="hairline", size="small", color="neutral")
            else:
                d.rect(cx, cy, cell_w, cell_h, label=val, fill="v_light", stroke="v_matrix", sw="thin", size="small", color="v_matrix")

    d.text("Sparse lexical frequencies", vx_top, 258, anchor="center", size="small", italic=True, color="v_matrix")

    # Approx Symbol
    d.text("≈", 305, 175, anchor="center", size="title", bold=True, color="primary")
    d.text("(Eq. 14)", 305, 198, anchor="center", size="small", bold=True, color="neutral")

    # 2. Topic Matrix W (5 words × 2 topics)
    wx_top = 405
    d.text("**Topic Matrix W**", wx_top, 72, anchor="center", size="label", bold=True, color="w_matrix")
    d.text("Word weights for each topic (K = 2)", wx_top, 86, anchor="center", size="small", italic=True, color="neutral")

    grid_wx, grid_wy = 355, 115
    w_col_w = 50
    d.text("t_1 (DL)", grid_wx + w_col_w // 2, grid_wy - 10, anchor="center", size="small", bold=True, color="w_matrix")
    d.text("t_2 (Mining)", grid_wx + w_col_w + w_col_w // 2, grid_wy - 10, anchor="center", size="small", bold=True, color="w_matrix")

    w_data = [
        ["0.82", "0.05"],
        ["0.75", "0.10"],   # Active Row 2 highlighted!
        ["0.04", "0.70"],
        ["0.08", "0.85"],
        ["0.15", "0.52"],
    ]

    for r in range(5):
        for c in range(2):
            cx = grid_wx + c * w_col_w
            cy = grid_wy + r * cell_h
            val = w_data[r][c]
            if r == target_r:
                d.rect(cx, cy, w_col_w, cell_h, label=f"**{val}**", fill="calc_light", stroke="calc_highlight", sw="normal", size="small", color="calc_highlight")
            else:
                d.rect(cx, cy, w_col_w, cell_h, label=val, fill="w_light", stroke="w_matrix", sw="thin", size="small", color="w_matrix")

    d.text("Word-topic distributions", wx_top, 258, anchor="center", size="small", italic=True, color="w_matrix")
    d.rect(grid_wx - 3, grid_wy + target_r * cell_h - 2, 2 * w_col_w + 6, cell_h + 4, fill="none", stroke="calc_highlight", sw="thick", rounded=4)
    d.text("Row w_2", grid_wx - 8, grid_wy + target_r * cell_h + cell_h // 2, anchor="e", size="small", bold=True, color="calc_highlight")

    # Multiply Symbol
    d.text("×", 485, 175, anchor="center", size="title", bold=True, color="primary")

    # 3. Coefficient Matrix H (2 topics × 5 docs)
    hx_top = 615
    d.text("**Coefficient Matrix H**", hx_top, 72, anchor="center", size="label", bold=True, color="h_matrix")
    d.text("Topic weights for each document", hx_top, 86, anchor="center", size="small", italic=True, color="neutral")

    grid_hx, grid_hy = 525, 145
    h_col_w = 36

    for j, d_lab in enumerate(doc_labels):
        d.text(d_lab, grid_hx + j * h_col_w + h_col_w // 2, grid_hy - 10, anchor="center", size="small", bold=True, color="text")

    d.text("t_1", grid_hx - 6, grid_hy + cell_h // 2, anchor="e", size="small", bold=True, color="h_matrix")
    d.text("t_2", grid_hx - 6, grid_hy + cell_h + cell_h // 2, anchor="e", size="small", bold=True, color="h_matrix")

    h_data = [
        ["0.88", "0.05", "0.70", "0.10", "0.05"],
        ["0.12", "0.85", "0.25", "0.80", "0.75"],
    ]

    for r in range(2):
        for c in range(5):
            cx = grid_hx + c * h_col_w
            cy = grid_hy + r * cell_h
            val = h_data[r][c]
            if c == target_c:
                d.rect(cx, cy, h_col_w, cell_h, label=f"**{val}**", fill="calc_light", stroke="calc_highlight", sw="normal", size="small", color="calc_highlight")
            else:
                d.rect(cx, cy, h_col_w, cell_h, label=val, fill="h_light", stroke="h_matrix", sw="thin", size="small", color="h_matrix")

    d.text("Document mixture weights", hx_top, 258, anchor="center", size="small", italic=True, color="h_matrix")
    d.rect(grid_hx + target_c * h_col_w - 2, grid_hy - 3, h_col_w + 4, 2 * cell_h + 6, fill="none", stroke="calc_highlight", sw="thick", rounded=4)
    d.text("Col d_3", grid_hx + target_c * h_col_w + h_col_w // 2, grid_hy + 2 * cell_h + 12, anchor="center", size="small", bold=True, color="calc_highlight")

    # 4. Calculation Trace Box (Right)
    d.rect(740, 68, 345, 215, fill="background", stroke="calc_highlight", sw="normal", rounded=6)
    d.text("**Actual Dot Product Arithmetic**", 755, 86, anchor="w", size="label", bold=True, color="calc_highlight")
    d.text(
        "**Reconstructing Cell V_{2,3} (word 'network' in doc d_3):**\n\n"
        "**V_{i,j} ≈ ∑_{r=1}^K (W_{i,r} × H_{r,j}) = w_{i,:} · h_{:,j}**\n\n"
        "**V_{2,3} ≈ (W_{2,1} × H_{1,3}) + (W_{2,2} × H_{2,3})**\n"
        "     **= (0.75 × 0.70) + (0.10 × 0.25)**\n"
        "     **= 0.525 + 0.025 = 0.55 ≈ 0.58**\n\n"
        "{neutral|• **Additive Non-negativity:** Both W ≥ 0 and H ≥ 0.}\n"
        "{neutral|• Parts-based representation without cancellation.}",
        755, 162, anchor="w", size="small"
    )
    d.arrow([(grid_wx + 2 * w_col_w + 4, grid_wy + target_r * cell_h + cell_h // 2), (735, 135)], color="calc_highlight", sw="thick")
    d.arrow([(grid_hx + target_c * h_col_w + h_col_w // 2, grid_hy - 6), (735, 155)], color="calc_highlight", sw="thick")
    d.rect(755, 244, 315, 26, label="**Reconstructs Target Cell V_{2,3} = 0.58**", fill="calc_light", stroke="calc_highlight", sw="normal", rounded=4, size="small", color="calc_highlight")

    # Bottom Summary Bar
    d.rect(35, 305, 1050, 68, fill="background", stroke="primary", sw="thin", rounded=6)
    d.text(
        "**Linear Algebra Insight:** High-dimensional sparse frequency matrix **V** (M × N) is factorized into topic matrix **W** (M × K word weights)\n"
        "and coefficient matrix **H** (K × N document weights), reducing dimensionality from M to K (K << M) while capturing thematic structure.",
        560, 339, anchor="center", size="small", align="center", color="primary"
    )


def v2(d):
    """v2: Comparative feature spaces: Raw TF-IDF Orthogonal Failure vs. NMF Low-Rank Space vs. LDA Generative Model."""
    d.rect(20, 15, 1100, 390, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Graphical Comparison: Raw TF-IDF Space vs. NMF Low-Rank Space vs. LDA Generative Space",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Demonstrating geometric representations, dimensionality compression, and generative assumptions",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # Sub-panel 1: Raw TF-IDF Space
    d.rect(35, 70, 345, 315, fill="background", stroke="v_matrix", sw="thin", rounded=6)
    d.rect(35, 70, 345, 30, "**1. Raw TF-IDF Space** (Prior Concept)", fill="v_matrix", stroke="v_matrix", color="background", rounded=6, size="label")

    o1_x, o1_y = 135, 230
    d.arrow([(o1_x, o1_y), (o1_x + 130, o1_y)], color="neutral", sw="thin")
    d.arrow([(o1_x, o1_y), (o1_x, o1_y - 95)], color="neutral", sw="thin")
    d.arrow([(o1_x, o1_y), (o1_x - 45, o1_y + 35)], color="neutral", sw="thin")
    d.text("w_1 ('car')", o1_x + 135, o1_y, anchor="w", size="small", color="neutral")
    d.text("w_2 ('auto')", o1_x, o1_y - 100, anchor="s", size="small", color="neutral")
    d.text("w_3 ('neural')", o1_x - 50, o1_y + 35, anchor="e", size="small", color="neutral")

    d.arrow([(o1_x, o1_y), (o1_x + 90, o1_y)], color="v_matrix", sw="thick")
    d.circle(o1_x + 90, o1_y, 4, fill="v_matrix", stroke="v_matrix")
    d.text("Doc A ('car')", o1_x + 90, o1_y + 14, anchor="center", size="small", bold=True, color="v_matrix")

    d.arrow([(o1_x, o1_y), (o1_x, o1_y - 70)], color="secondary", sw="thick")
    d.circle(o1_x, o1_y - 70, 4, fill="secondary", stroke="secondary")
    d.text("Doc B ('auto')", o1_x + 6, o1_y - 75, anchor="w", size="small", bold=True, color="secondary")

    d.line([(o1_x + 18, o1_y), (o1_x + 18, o1_y - 18)], color="calc_highlight", sw="thin")
    d.line([(o1_x, o1_y - 18), (o1_x + 18, o1_y - 18)], color="calc_highlight", sw="thin")
    d.text("θ = 90° (cos θ = 0)", o1_x + 25, o1_y - 30, anchor="w", size="small", bold=True, color="calc_highlight")

    d.rect(45, 315, 320, 56, fill="sparse_zero", stroke="panel_border", sw="thin", rounded=4)
    d.text("• **Orthogonality Failure:** Synonyms have 0 similarity.\n"
           "• **Extreme Sparsity:** Dimensionality = |Vocabulary| (M ~ 10⁴).",
           55, 343, anchor="w", size="small", color="text")

    # Sub-panel 2: NMF Latent Semantic Subspace
    d.rect(395, 70, 345, 315, fill="background", stroke="w_matrix", sw="thin", rounded=6)
    d.rect(395, 70, 345, 30, "**2. NMF Latent Topic Space** (Linear Algebra)", fill="w_matrix", stroke="w_matrix", color="background", rounded=6, size="label")

    o2_x, o2_y = 470, 230
    d.arrow([(o2_x, o2_y), (o2_x + 145, o2_y)], color="w_matrix", sw="thin")
    d.arrow([(o2_x, o2_y), (o2_x, o2_y - 95)], color="h_matrix", sw="thin")
    d.text("t_1 (Automotive)", o2_x + 145, o2_y + 14, anchor="center", size="small", bold=True, color="w_matrix")
    d.text("t_2 (Mining)", o2_x, o2_y - 100, anchor="s", size="small", bold=True, color="h_matrix")

    d.rect(o2_x, o2_y - 95, 170, 95, fill="w_light", fill_opacity="wash", stroke="none")
    d.text("Non-negative Orthant (W, H ≥ 0)", o2_x + 75, o2_y - 80, anchor="center", size="small", italic=True, color="neutral")

    d.arrow([(o2_x, o2_y), (o2_x + 115, o2_y - 25)], color="v_matrix", sw="thick")
    d.arrow([(o2_x, o2_y), (o2_x + 105, o2_y - 36)], color="secondary", sw="thick")
    d.circle(o2_x + 115, o2_y - 25, 4, fill="v_matrix", stroke="v_matrix")
    d.circle(o2_x + 105, o2_y - 36, 4, fill="secondary", stroke="secondary")

    d.text("Doc A ('car')", o2_x + 120, o2_y - 20, anchor="w", size="small", bold=True, color="v_matrix")
    d.text("Doc B ('auto')", o2_x + 110, o2_y - 42, anchor="w", size="small", bold=True, color="secondary")
    d.text("θ ≈ 5° (cos θ ≈ 0.99)", o2_x + 40, o2_y - 30, anchor="center", size="small", bold=True, color="calc_highlight")

    d.rect(405, 315, 320, 56, fill="w_light", fill_opacity="soft", stroke="w_matrix", sw="hairline", rounded=4)
    d.text("• **Semantic Clustering:** Synonyms align along Topic 1.\n"
           "• **Low-Rank Compression:** Reduced to K topics (K << M).",
           415, 343, anchor="w", size="small", color="primary")

    # Sub-panel 3: LDA Generative Bayesian Model
    d.rect(755, 70, 345, 315, fill="background", stroke="h_matrix", sw="thin", rounded=6)
    d.rect(755, 70, 345, 30, "**3. LDA Generative Model** (Hierarchical Bayesian)", fill="h_matrix", stroke="h_matrix", color="background", rounded=6, size="label")

    pl_x, pl_y = 770, 115
    plate_cy = 180

    # Hyperparameter alpha
    d.circle(pl_x + 20, plate_cy, 14, fill="lda_light", stroke="lda_accent", sw="normal")
    d.text("α", pl_x + 20, plate_cy, anchor="center", size="label", bold=True, color="lda_accent")
    d.text("Dirichlet", pl_x + 20, 150, anchor="center", size="small", color="neutral")
    d.arrow([(pl_x + 34, plate_cy), (pl_x + 70, plate_cy)], color="primary", sw="thick")

    # Document Plate D
    d.rect(pl_x + 65, 125, 165, 115, fill="none", stroke="neutral", sw="thin", rounded=6)
    d.text("Document d in Corpus (D)", pl_x + 147, 230, anchor="center", size="small", italic=True, color="neutral")

    # theta_d (topic mixture for document d)
    d.circle(pl_x + 95, plate_cy, 14, fill="w_light", stroke="w_matrix", sw="normal")
    d.text("θ_d", pl_x + 95, plate_cy, anchor="center", size="label", bold=True, color="w_matrix")
    d.arrow([(pl_x + 109, plate_cy), (pl_x + 141, plate_cy)], color="primary", sw="thick")

    # Word Plate N
    d.rect(pl_x + 135, 135, 85, 80, fill="none", stroke="neutral", sw="hairline", rounded=4)
    d.text("Word n (N)", pl_x + 177, 205, anchor="center", size="small", color="neutral")

    # z_dn (topic indicator)
    d.circle(pl_x + 155, plate_cy, 12, fill="w_light", stroke="w_matrix", sw="normal")
    d.text("z", pl_x + 155, plate_cy, anchor="center", size="small", bold=True, color="w_matrix")
    d.arrow([(pl_x + 167, plate_cy), (pl_x + 188, plate_cy)], color="primary", sw="thick")

    # w_dn (observed word token)
    d.circle(pl_x + 200, plate_cy, 12, fill="calc_light", stroke="calc_highlight", sw="thick")
    d.text("w", pl_x + 200, plate_cy, anchor="center", size="small", bold=True, color="calc_highlight")

    # Topic Plate K and beta
    d.circle(pl_x + 285, plate_cy, 14, fill="lda_light", stroke="lda_accent", sw="normal")
    d.text("β", pl_x + 285, plate_cy, anchor="center", size="label", bold=True, color="lda_accent")
    d.text("Dirichlet", pl_x + 285, 150, anchor="center", size="small", color="neutral")
    d.arrow([(pl_x + 271, plate_cy), (pl_x + 215, plate_cy)], color="primary", sw="thick")

    d.rect(765, 315, 320, 56, fill="lda_light", fill_opacity="soft", stroke="lda_accent", sw="hairline", rounded=4)
    d.text("• **Generative Mechanism:** Documents are random topic mixtures.\n"
           "• **Probabilistic Foundations:** Discrete Dirichlet distributions.",
           775, 343, anchor="w", size="small", color="text")


def v3(d):
    """v3: Generative Probabilistic Architecture: Complete LDA Plate Model & Comparison with NMF Factorization."""
    d.rect(20, 15, 1080, 420, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Latent Dirichlet Allocation (LDA) Generative Mechanism vs. NMF Factorization",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Hierarchical Bayesian plate notation and paradigm contrast (Algebraic vs. Generative Probabilistic)",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # Left: Deep Plate Diagram
    d.rect(35, 70, 540, 345, fill="background", stroke="primary", sw="thin", rounded=6)
    d.text("**LDA Hierarchical Bayesian Plate Notation**", 50, 90, anchor="w", size="label", bold=True, color="primary")

    # Priors alpha and beta
    d.circle(80, 195, 16, fill="lda_light", stroke="lda_accent", sw="normal")
    d.text("α", 80, 195, anchor="center", size="heading", bold=True, color="lda_accent")
    d.text("Dirichlet Prior\nover topics", 80, 150, anchor="center", size="small", color="neutral")

    d.arrow([(96, 195), (145, 195)], color="primary", sw="thick")

    # Plate D: Documents
    d.rect(140, 115, 275, 175, fill="none", stroke="primary", sw="thin", rounded=6)
    d.text("Documents d ∈ {1, ..., D}", 277, 280, anchor="center", size="small", bold=True, color="primary")

    # theta_d: topic distribution
    d.circle(180, 195, 16, fill="w_light", stroke="w_matrix", sw="normal")
    d.text("θ_d", 180, 195, anchor="center", size="label", bold=True, color="w_matrix")
    d.text("Topic mixture\nθ_d ~ Dir(α)", 180, 150, anchor="center", size="small", color="neutral")

    d.arrow([(196, 195), (240, 195)], color="primary", sw="thick")

    # Plate N: Words inside Doc d
    d.rect(235, 130, 165, 130, fill="none", stroke="neutral", sw="hairline", rounded=4)
    d.text("Words n ∈ {1, ..., N_d}", 317, 250, anchor="center", size="small", color="neutral")

    # z_dn: latent topic
    d.circle(270, 195, 14, fill="w_light", stroke="w_matrix", sw="normal")
    d.text("z_{dn}", 270, 195, anchor="center", size="small", bold=True, color="w_matrix")
    d.text("Latent topic", 270, 165, anchor="center", size="small", color="neutral")

    d.arrow([(284, 195), (326, 195)], color="primary", sw="thick")

    # w_dn: observed word
    d.circle(340, 195, 14, fill="calc_light", stroke="calc_highlight", sw="heavy")
    d.text("w_{dn}", 340, 195, anchor="center", size="small", bold=True, color="calc_highlight")
    d.text("Observed word", 340, 165, anchor="center", size="small", color="calc_highlight")

    # Topic Plate K and beta
    d.circle(525, 195, 16, fill="lda_light", stroke="lda_accent", sw="normal")
    d.text("β", 525, 195, anchor="center", size="heading", bold=True, color="lda_accent")
    d.text("Dirichlet Prior\nover vocab", 525, 150, anchor="center", size="small", color="neutral")

    d.arrow([(509, 195), (465, 195)], color="primary", sw="thick")

    d.rect(420, 115, 65, 175, fill="none", stroke="secondary", sw="thin", rounded=6)
    d.text("K topics", 452, 280, anchor="center", size="small", bold=True, color="secondary")

    d.circle(452, 195, 14, fill="h_light", stroke="h_matrix", sw="normal")
    d.text("ϕ_k", 452, 195, anchor="center", size="small", bold=True, color="h_matrix")
    d.text("Word dist\nϕ_k ~ Dir(β)", 452, 150, anchor="center", size="small", color="neutral")

    d.arrow([(438, 195), (354, 195)], color="primary", sw="thick")

    # Generative Step Checklist below plate
    d.rect(50, 305, 510, 95, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("**Generative Story:**\n"
           "1. For each topic k ∈ {1, ..., K}: draw word distribution ϕ_k ~ Dir(β)\n"
           "2. For each document d ∈ {1, ..., D}: draw topic mixture θ_d ~ Dir(α)\n"
           "3. For each word n ∈ {1, ..., N_d}: draw topic z_{dn} ~ Cat(θ_d), then word w_{dn} ~ Cat(ϕ_{z_{dn}})",
           60, 350, anchor="w", size="small", color="text")

    # Right: Paradigm Comparison Table
    d.rect(595, 70, 485, 345, fill="background", stroke="primary", sw="thin", rounded=6)
    d.text("**Comparison: Linear Algebra (NMF) vs. Generative Bayesian (LDA)**", 610, 90, anchor="w", size="label", bold=True, color="primary")

    headers = [("Property", 110), ("NMF (Lee & Seung, 1999)", 175), ("LDA (Blei et al., 2003)", 175)]
    ty, rh = 115, 48
    cx = 610
    for h, w in headers:
        d.rect(cx, ty, w, 28, f"**{h}**", fill="primary", stroke="background", color="background", size="small")
        cx += w

    rows = [
        ("Core Mechanism", "Matrix decomposition\nV ≈ W × H", "Hierarchical Bayesian model\nDirichlet-Multinomial mixtures"),
        ("Mathematical Basis", "Linear algebra\nNon-negative constraints (W, H ≥ 0)", "Probability theory\nLatent prior distributions (α, β)"),
        ("Representation", "Word weights W, doc weights H\nParts-based additive reconstruction", "Discrete word distributions ϕ_k\nRandom topic mixtures θ_d"),
        ("Predecessors", "Built upon BoW and TF-IDF\nSolves synonym orthogonality", "Built upon BoW and TF-IDF\nProbabilistic document mining"),
        ("Dimensionality", "Compresses M words to K topics\nDeterministic low-rank factor space", "Uncovers K latent topics\nFull posterior distribution"),
    ]

    for i, (prop, nmf_val, lda_val) in enumerate(rows):
        y = ty + 28 + i * rh
        bg = "secondary" if i % 2 == 0 else "background"
        op = "wash" if i % 2 == 0 else None
        d.rect(610, y, 110, rh, f"**{prop}**", fill=bg, fill_opacity=op, stroke="panel_border", sw="thin", size="small", color="primary")
        d.rect(720, y, 175, rh, nmf_val, fill=bg, fill_opacity=op, stroke="panel_border", sw="thin", size="small", color="text")
        d.rect(895, y, 175, rh, lda_val, fill=bg, fill_opacity=op, stroke="panel_border", sw="thin", size="small", color="text")


INFERRED = [
    "v1: numerical cell values for V, W, and H are concrete pedagogical instantiations",
    "v2: 3D orthogonal axes and 90° angle illustrate the geometric concept of BoW/TF-IDF term independence",
    "v2 & v3: Dirichlet hyperparameters α, β and latent indicators z, θ, ϕ reflect standard LDA formulation (Blei et al., 2003)",
]

LINT_IGNORE = ["near-miss", "uneven gaps"]

ALLOWED_TERMS = [
    "0.00", "0.025", "0.04", "0.05", "0.08", "0.10", "0.12", "0.15", "0.20", "0.25", "0.30", "0.35",
    "0.40", "0.45", "0.50", "0.52", "0.525", "0.54", "0.55", "0.58", "0.62", "0.65", "0.68", "0.70",
    "0.72", "0.75", "0.80", "0.82", "0.85", "0.88", "0.99", "actual", "arithmetic", "both", "capturing",
    "cell", "col", "computation", "dimensional", "dimensionality", "dot", "high", "lexical", "network",
    "product", "reducing", "representation", "row", "sparse", "while", "without", "align", "along",
    "assumptions", "auto", "automotive", "car", "clustering", "comparison", "compression", "cos",
    "demonstrating", "foundations", "geometric", "graphical", "have", "low", "mechanism", "rank",
    "raw", "reduced", "representations", "semantic", "similarity", "space", "basis", "built", "cat",
    "constraints", "contrast", "core", "deterministic", "draw", "mathematical", "multinomial",
    "notation", "observed", "reconstruction", "theory", "additive", "based", "cancellation", "insight",
    "parts", "reconstructing", "reconstructs", "target", "extreme", "failure", "orthant", "orthogonality",
    "sparsity", "synonyms", "compresses", "full", "paradigm", "plate", "posterior", "predecessors",
    "property", "solves", "story", "synonym"
]

VARIANTS = [
    dict(key="v1", name="matrix_factorization_grid", size=(1120, 420), build=v1,
         about="Populated 2D matrix grids for V ≈ W × H with active row × column dot product reconstruction.",
         caption="Non-negative Matrix Factorization: decomposing document-term matrix V into word topic weights W and document topic weights H."),
    dict(key="v2", name="comparative_feature_spaces", size=(1140, 420), build=v2,
         about="Three comparative feature spaces: Raw TF-IDF orthogonal failure vs. NMF low-rank subspace vs. LDA generative model.",
         caption="Comparison of document representations: TF-IDF orthogonal word axes fail on synonyms, whereas NMF projects documents onto shared semantic topic rays and LDA models topic mixtures."),
    dict(key="v3", name="hierarchical_generative_plate", size=(1120, 450), build=v3,
         about="Hierarchical Bayesian plate notation for LDA with hyperparameters and generative story, contrasted with NMF.",
         caption="Latent Dirichlet Allocation: generative probabilistic plate model with Dirichlet priors and discrete word distributions, contrasted with algebraic NMF."),
]
'''

# ================================================================================================
FILES['examples/E_matrix_computation/passage.md'] = r'''# Passage: Non-negative Matrix Factorization and Latent Dirichlet Allocation

For two decades, the standard unsupervised techniques for literature and document mining were Non-negative Matrix Factorization (NMF) (Lee & Seung, 1999) and Latent Dirichlet Allocation (LDA) (Blei et al., 2003). NMF is a linear algebra technique that decomposes a document-term matrix (the input document-term frequency matrix, V) into two smaller non-negative matrices: 

$$V \approx W \times H \quad (14)$$

where W is the topic matrix, containing the word weights for each topic, and H is the coefficient matrix, containing the topic weights for each document. LDA on the other hand is a generative probabilistic hierarchical Bayesian model. It posits that each document d in a corpus is a random mixture over K latent topics, and each topic k is a discrete distribution over vocabulary words.  
Such approaches build upon the prior concepts of BoW and TF-IDF, providing powerful tools for uncovering latent thematic structure in large, unlabeled text corpora. Since they were modeled long before the advent of more advanced neural processes.
'''

# ================================================================================================
FILES['examples/E_matrix_computation/theme.yaml'] = r'''colors:
  primary: "#1f497d"
  secondary: "#2563eb"
  neutral: "#64748b"
  text: "#0f172a"
  background: "#ffffff"
  
  # Panels & Containers
  panel_bg: "#f8fafc"
  panel_border: "#cbd5e1"
  panel_header: "#1e293b"
  
  # Matrix & Thematic Accents
  v_matrix: "#0f766e"        # Teal for Document-Term Matrix V / TF-IDF
  v_light: "#ccfbf1"
  w_matrix: "#1d4ed8"        # Blue for Topic Matrix W (Words x Topics)
  w_light: "#dbeafe"
  h_matrix: "#7c3aed"        # Purple for Coefficient Matrix H (Topics x Docs)
  h_light: "#ede9fe"
  calc_highlight: "#ea580c"  # Orange for active row/col dot product
  calc_light: "#ffedd5"
  sparse_zero: "#f1f5f9"     # Light gray for zero sparse cells
  dist_line: "#475569"
  lda_accent: "#b45309"      # Amber for LDA Bayesian variables
  lda_light: "#fef3c7"

fonts:
  sans:  { family: "Helvetica" }
  serif: { family: "Times New Roman" }
  mono:  { family: "Courier New" }
default_font: sans

sizes:
  title: 17
  heading: 13
  subheading: 11
  body: 10
  label: 9
  small: 8
  caption: 9.5
  eq_large: 16
  eq_term: 14
  eq_small: 12

line_height: 1.25

strokes:
  hairline: 0.75
  thin: 1
  normal: 1.5
  thick: 2
  heavy: 3

opacity:
  wash: 10
  tint: 25
  soft: 45
  faded: 35

shape:
  corner_radius: 6
'''

# ================================================================================================
FILES['examples/F_metric_manifold_geometry/build.py'] = r'''"""Prototypes for passage F (Metric Geometry, Clustering Objectives, and Manifold Projections).
Run: python ../../scripts/prototype.py .
"""
import math
import os
import sys

from figlib import *
from drawio_kit import Diagram, measure_math


def v1(d):
    """v1: Continuous 2D Metric Space: Unit Hypersphere Arc, Cosine-Euclidean Equivalence, and K-means Centroid Attraction."""
    d.rect(20, 15, 1080, 390, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Metric Geometry: Unit Normalization, Cosine Equivalence, and K-means Euclidean Space",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Demonstrating monotonic relation ||u-v||_2^2 = 2 - 2(u·v) and WCSS / SSE centroid minimization (Equations 12 & 13)",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # Left Sub-panel: Unit Normalization & Metric Equivalence
    d.rect(35, 70, 510, 320, fill="background", stroke="primary", sw="thin", rounded=6)
    d.text("**1. Unit Normalization & Cosine-Euclidean Equivalence**", 50, 90, anchor="w", size="label", bold=True, color="primary")

    ox, oy = 115, 235
    d.arrow([(ox, oy), (ox + 195, oy)], color="neutral", sw="thin")
    d.arrow([(ox, oy), (ox, oy - 150)], color="neutral", sw="thin")
    d.text("e_1", ox + 200, oy, anchor="w", size="small", color="neutral")
    d.text("e_2", ox, oy - 155, anchor="s", size="small", color="neutral")

    # Unit circle arc (r = 130)
    r = 130
    arc_pts = []
    for deg in range(0, 95, 10):
        rad = math.radians(deg)
        arc_pts.append((ox + int(r * math.cos(rad)), oy - int(r * math.sin(rad))))
    for i in range(len(arc_pts) - 1):
        d.line([arc_pts[i], arc_pts[i+1]], color="unit_circle", sw="thin")

    d.text("Unit circle ||u||_2 = 1", ox + 105, oy - 130, anchor="w", size="small", color="neutral")

    # Vectors u and v
    deg_u, deg_v = 65, 20
    ux = ox + int(r * math.cos(math.radians(deg_u)))
    uy = oy - int(r * math.sin(math.radians(deg_u)))
    vx = ox + int(r * math.cos(math.radians(deg_v)))
    vy = oy - int(r * math.sin(math.radians(deg_v)))

    d.arrow([(ox, oy), (ux, uy)], color="c1_color", sw="thick")
    d.circle(ux, uy, 4, fill="c1_color", stroke="c1_color")
    d.text("u (||u||_2 = 1)", ux - 6, uy - 10, anchor="e", size="small", bold=True, color="c1_color")

    d.arrow([(ox, oy), (vx, vy)], color="c3_color", sw="thick")
    d.circle(vx, vy, 4, fill="c3_color", stroke="c3_color")
    d.text("v (||v||_2 = 1)", vx + 8, vy + 4, anchor="w", size="small", bold=True, color="c3_color")

    # Angle theta and Chord
    d.text("θ", ox + 40, oy - 22, anchor="center", size="label", bold=True, color="primary")
    d.line([(ox + 30, oy - 5), (ox + 26, oy - 25)], color="primary", sw="thin")

    d.line([(ux, uy), (vx, vy)], color="split_line", sw="normal", dashed=True)
    d.text("||u - v||_2", (ux + vx) // 2 + 16, (uy + vy) // 2 - 6, anchor="center", size="small", bold=True, color="split_line")
    d.text("u · v = cos(θ)", ox + 45, oy + 16, anchor="w", size="small", color="primary")

    # Equivalence Card
    d.rect(50, 275, 480, 100, fill="panel_bg", stroke="primary", sw="thin", rounded=6)
    d.text("**Monotonic Equivalence Formulation:**\n\n"
           "**||u - v||_2² = 2 - 2(u · v) = 2 - 2 cos(θ)**\n\n"
           "• As angular proximity θ → 0, cosine similarity u · v → 1.\n"
           "• Minimizing Euclidean distance directly minimizes cosine distance on normalized embeddings.",
           65, 325, anchor="w", size="small", color="text")

    # Right Sub-panel: K-means in Euclidean Space
    d.rect(565, 70, 520, 320, fill="background", stroke="primary", sw="thin", rounded=6)
    d.text("**2. K-means Euclidean Space & Objective Function**", 580, 90, anchor="w", size="label", bold=True, color="primary")

    # Cluster 1 (C_1)
    c1_x, c1_y = 670, 165
    d.ellipse(c1_x - 65, c1_y - 50, 130, 100, fill="c1_light", fill_opacity="wash", stroke="c1_color", sw="hairline", dashed=True)
    d.text("Cluster C_1", c1_x - 20, c1_y - 38, anchor="center", size="small", bold=True, color="c1_color")
    d.rect(c1_x - 12, c1_y - 12, 24, 24, label="**μ_1**", fill="c1_color", stroke="primary", sw="normal", rounded=True, color="background", size="small")

    pts_c1 = [(c1_x - 40, c1_y - 25), (c1_x + 35, c1_y - 20), (c1_x - 25, c1_y + 30), (c1_x + 30, c1_y + 25)]
    for px, py in pts_c1:
        d.circle(px, py, 3, fill="c1_color", stroke="primary")
        d.line([(c1_x, c1_y), (px, py)], color="dist_line", sw="hairline", dashed=True)
    d.text("(x - μ_1)", c1_x - 24, c1_y - 10, anchor="center", size="small", italic=True, color="dist_line")

    # Cluster 2 (C_2)
    c2_x, c2_y = 940, 165
    d.ellipse(c2_x - 65, c2_y - 50, 130, 100, fill="c2_light", fill_opacity="wash", stroke="c2_color", sw="hairline", dashed=True)
    d.text("Cluster C_2", c2_x + 20, c2_y - 38, anchor="center", size="small", bold=True, color="c2_color")
    d.rect(c2_x - 12, c2_y - 12, 24, 24, label="**μ_2**", fill="c2_color", stroke="primary", sw="normal", rounded=True, color="background", size="small")

    pts_c2 = [(c2_x - 35, c2_y - 25), (c2_x + 40, c2_y - 18), (c2_x - 30, c2_y + 28), (c2_x + 30, c2_y + 30)]
    for px, py in pts_c2:
        d.circle(px, py, 3, fill="c2_color", stroke="primary")
        d.line([(c2_x, c2_y), (px, py)], color="dist_line", sw="hairline", dashed=True)
    d.text("(x - μ_2)", c2_x + 24, c2_y + 12, anchor="center", size="small", italic=True, color="dist_line")

    # WCSS Objective Card
    d.rect(580, 275, 490, 100, fill="panel_bg", stroke="primary", sw="thin", rounded=6)
    d.text("**K-means Minimization Objective (Eq. 12 & 13):**\n\n"
           "**WCSS = SSE = ∑_{i=1}^k ∑_{x ∈ C_i} (x - μ_i)²**\n\n"
           "• Evaluates each embedding **x** against cluster centroid **μ_i**.\n"
           "• Output partition: **C = {c_1, c_2, ..., c_k} = K-means[E(x), k]**.",
           595, 325, anchor="w", size="small", color="text")


def v2(d):
    """v2: Alternative Non-Euclidean Clustering: PDDP Divisive Hyperplane vs. HDBSCAN Density Cores & Noise."""
    d.rect(20, 15, 1080, 390, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Non-Euclidean Clustering: PDDP Divisive Hyperplane vs. HDBSCAN Density Cores",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Eigenvector SVD variance splitting (deterministic binary trees) vs. hierarchical density reachability with outlier noise segregation",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # Left: PDDP
    d.rect(35, 70, 510, 320, fill="background", stroke="eigenvector", sw="thin", rounded=6)
    d.rect(35, 70, 510, 30, "**PDDP: Principal Direction Divisive Clustering** (Boley, 1998)", fill="eigenvector", stroke="eigenvector", color="background", rounded=6, size="label")

    pddp_ox, pddp_oy = 290, 190
    pts_pddp = [
        (-50, -20), (-35, -12), (-20, -5), (-10, -2), (0, 0),
        (15, 5), (30, 10), (48, 18), (-40, -5), (35, 5),
        (-25, -25), (25, 22), (-15, 10), (10, -10)
    ]
    for dx, dy in pts_pddp:
        d.circle(pddp_ox + dx, pddp_oy + dy, 3, fill="neutral", stroke="text")

    # Primary Eigenvector v_1
    d.arrow([(pddp_ox - 75, pddp_oy - 28), (pddp_ox + 75, pddp_oy + 28)], color="eigenvector", sw="thick")
    d.text("v_1 (Primary Eigenvector)", pddp_ox + 40, pddp_oy + 42, anchor="center", size="small", bold=True, color="eigenvector")

    # Splitting Hyperplane
    d.line([(pddp_ox - 22, pddp_oy + 55), (pddp_ox + 22, pddp_oy - 55)], color="split_line", sw="thick", dashed=True)
    d.text("SVD Splitting Hyperplane", pddp_ox + 35, pddp_oy - 50, anchor="w", size="small", bold=True, color="split_line")

    d.rect(50, 280, 480, 95, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("• **Mechanism:** Iteratively splits clusters along primary eigenvector of variance via SVD.\n"
           "• **Tree Structure:** Creates deterministic binary trees (Boley, 1998; Theocharopoulos et al., 2025).\n"
           "• **Geometry:** Divisive hyperplane partitioning rather than spherical Euclidean centroids.",
           65, 325, anchor="w", size="small", color="text")

    # Right: HDBSCAN
    d.rect(565, 70, 520, 320, fill="background", stroke="c3_color", sw="thin", rounded=6)
    d.rect(565, 70, 520, 30, "**HDBSCAN: Hierarchical Density-Based Clustering** (Campello et al., 2013)", fill="c3_color", stroke="c3_color", color="background", rounded=6, size="label")

    # Dense core 1
    hdb_cx, hdb_cy = 680, 185
    d.ellipse(hdb_cx - 50, hdb_cy - 40, 100, 80, fill="c3_light", fill_opacity="soft", stroke="c3_color", sw="thin")
    d.text("Dense Core 1", hdb_cx, hdb_cy, anchor="center", size="small", bold=True, color="c3_color")
    for dx, dy in [(-25, -10), (-10, 15), (15, -15), (20, 10), (0, -5), (-15, -20), (5, 20)]:
        d.circle(hdb_cx + dx, hdb_cy + dy, 3, fill="c3_color", stroke="primary")

    # Dense core 2
    hdb_c2x, hdb_c2y = 860, 175
    d.ellipse(hdb_c2x - 40, hdb_c2y - 30, 80, 60, fill="c1_light", fill_opacity="soft", stroke="c1_color", sw="thin")
    d.text("Dense Core 2", hdb_c2x, hdb_c2y, anchor="center", size="small", bold=True, color="c1_color")
    for dx, dy in [(-15, -5), (10, 8), (-5, 12), (12, -8)]:
        d.circle(hdb_c2x + dx, hdb_c2y + dy, 3, fill="c1_color", stroke="primary")

    # Outliers / Noise
    outliers = [(780, 155), (800, 215), (710, 240), (620, 235), (930, 210)]
    for ox, oy in outliers:
        d.rect(ox - 5, oy - 5, 10, 10, shape="diamond", fill="noise_color", stroke="primary", sw="thin")

    d.rect(940, 140, 130, 24, fill="noise_light", stroke="noise_color", sw="hairline", rounded=4)
    d.rect(950, 147, 10, 10, shape="diamond", fill="noise_color", stroke="primary", sw="thin")
    d.text("**Outlier Noise**", 1010, 152, anchor="center", size="small", bold=True, color="noise_color")

    d.rect(580, 280, 490, 95, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("• **Mechanism:** Operates on cluster density across varying scales.\n"
           "• **Noise Segregation:** Explicitly segregates outliers as unassigned noise (Campello et al., 2013).\n"
           "• **Advantage:** Discovers arbitrarily shaped clusters without enforcing spherical partitions.",
           595, 325, anchor="w", size="small", color="text")


def v3(d):
    """v3: Non-Linear Manifold Reduction: High-D Embedding Manifold vs. t-SNE vs. UMAP 2D Projections."""
    d.rect(20, 15, 1100, 390, fill="panel_bg", stroke="panel_border", sw="thin", rounded=8)
    d.text("Non-Linear Manifold Reduction for Clustering: High-D Space vs. t-SNE vs. UMAP",
           35, 30, anchor="w", size="heading", bold=True, color="primary")
    d.text("Comparing local neighborhood preservation against global Riemannian geometry and topological relationships",
           35, 48, anchor="w", size="label", italic=True, color="neutral")

    # Panel 1: High-D Space
    d.rect(35, 70, 260, 315, fill="background", stroke="primary", sw="thin", rounded=6)
    d.rect(35, 70, 260, 30, "**High-Dimensional Space**", fill="primary", stroke="primary", color="background", rounded=6, size="label")

    d.text("Embedding Matrix E(x) in R^(n × d)", 165, 115, anchor="center", size="small", color="neutral")
    d.circle(165, 185, 45, fill="panel_bg", stroke="primary", sw="thin", dashed=True)
    d.text("Complex Non-linear\nManifold (d >> 2)", 165, 185, anchor="center", size="small", align="center", color="primary")
    d.text("Normalized unit length ||u||_2 = 1", 165, 255, anchor="center", size="small", color="neutral")

    d.rect(45, 290, 240, 80, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("• Dimensionality curse hinders direct clustering.\n"
           "• Embeddings lie on complex non-linear submanifolds.",
           55, 330, anchor="w", size="small", color="text")

    # Projection Arrows
    d.arrow([(295, 160), (335, 130)], color="neutral", sw="thick")
    d.arrow([(295, 210), (335, 240)], color="secondary", sw="thick")
    d.text("t-SNE", 315, 135, anchor="center", size="small", bold=True, color="neutral")
    d.text("UMAP", 315, 235, anchor="center", size="small", bold=True, color="secondary")

    # Panel 2: t-SNE 2D Projection
    d.rect(345, 70, 370, 315, fill="background", stroke="neutral", sw="thin", rounded=6)
    d.rect(345, 70, 370, 30, "**t-SNE 2D Projection** (van der Maaten, 2008)", fill="neutral", stroke="neutral", color="background", rounded=6, size="label")

    d.text("Preserves local neighborhood probabilities in 2D", 530, 115, anchor="center", size="small", color="neutral")

    tsne_cx1, tsne_cy1 = 415, 175
    tsne_cx2, tsne_cy2 = 505, 215
    tsne_cx3, tsne_cy3 = 625, 165

    for cx, cy, col in [(tsne_cx1, tsne_cy1, "c1_color"), (tsne_cx2, tsne_cy2, "c2_color"), (tsne_cx3, tsne_cy3, "c3_color")]:
        d.circle(cx, cy, 18, fill=col, fill_opacity="wash", stroke=col, sw="hairline")
        for dx, dy in [(-8, -4), (6, -6), (-4, 7), (8, 5), (0, 0)]:
            d.circle(cx + dx, cy + dy, 3, fill=col, stroke=col)

    d.line([(tsne_cx1, tsne_cy1), (tsne_cx2, tsne_cy2)], color="split_line", sw="thin", dashed=True)
    d.line([(tsne_cx2, tsne_cy2), (tsne_cx3, tsne_cy3)], color="split_line", sw="thin", dashed=True)
    d.text("Distances distorted", 520, 180, anchor="center", size="small", color="split_line")

    d.rect(355, 290, 350, 80, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("• **Strength:** Exceptional for 2D visual cluster separation.\n"
           "• **Limitation:** Preserves neither global geometry nor distance scales.\n"
           "• **Scale:** Slow scaling to large document corpora.",
           365, 330, anchor="w", size="small", color="text")

    # Panel 3: UMAP 2D Projection
    d.rect(735, 70, 370, 315, fill="background", stroke="secondary", sw="thin", rounded=6)
    d.rect(735, 70, 370, 30, "**UMAP 2D Projection** (McInnes et al., 2018)", fill="secondary", stroke="secondary", color="background", rounded=6, size="label")

    d.text("Riemannian geometry & algebraic topology", 920, 115, anchor="center", size="small", color="neutral")

    umap_cx1, umap_cy1 = 805, 175
    umap_cx2, umap_cy2 = 890, 215
    umap_cx3, umap_cy3 = 1015, 165

    for cx, cy, col in [(umap_cx1, umap_cy1, "c1_color"), (umap_cx2, umap_cy2, "c2_color"), (umap_cx3, umap_cy3, "c3_color")]:
        d.circle(cx, cy, 18, fill=col, fill_opacity="wash", stroke=col, sw="hairline")
        for dx, dy in [(-8, -4), (6, -6), (-4, 7), (8, 5), (0, 0)]:
            d.circle(cx + dx, cy + dy, 3, fill=col, stroke=col)

    d.line([(umap_cx1, umap_cy1), (umap_cx2, umap_cy2)], color="c3_color", sw="normal")
    d.line([(umap_cx2, umap_cy2), (umap_cx3, umap_cy3)], color="c3_color", sw="normal")
    d.text("Global topology preserved", 910, 180, anchor="center", size="small", bold=True, color="c3_color")

    d.rect(745, 290, 350, 80, fill="panel_bg", stroke="panel_border", sw="thin", rounded=4)
    d.text("• **Strength:** Preserves both local clustering and global relationships.\n"
           "• **Theoretical Basis:** Grounded in Riemannian manifold geometry.\n"
           "• **Efficiency:** Scales efficiently to large document corpora.",
           755, 330, anchor="w", size="small", color="text")


INFERRED = [
    "v1: 2D coordinate frame and unit circle arc are visual geometric projections of high-dimensional unit hyperspheres",
    "v2: synthetic 2D point distributions and splitting hyperplanes illustrate SVD divisive variance and density cores",
    "v3: Swiss roll manifold and 3-cluster projections illustrate topological preservation vs distortion properties",
]

LINT_IGNORE = ["near-miss", "uneven gaps"]

ALLOWED_TERMS = [
    "unit", "circle", "chord", "cosine", "equivalence", "formulation", "proximity", "directly",
    "partition", "divisive", "hyperplane", "structure", "creates", "partitioning", "rather",
    "spherical", "reachability", "discovers", "arbitrarily", "shaped", "enforcing", "partitions",
    "strength", "limitation", "efficiency", "theoretical", "submanifolds", "curse", "hinders",
    "direct", "distorted", "separation", "slow", "scaling", "large", "basis", "grounded", "scales",
    "angular", "demonstrating", "equations", "output", "similarity", "advantage", "core", "cores",
    "dense", "mechanism", "splitting", "varying", "without", "comparing", "complex", "dimensional",
    "dimensionality", "document", "high", "lie", "matrix"
]

VARIANTS = [
    dict(key="v1", name="metric_geometry_clusters", size=(1120, 420), build=v1,
         about="Continuous 2D metric space: unit circle arc, cosine-Euclidean chord equivalence, and K-means centroid attraction.",
         caption="Metric space geometry: normalized transformer embeddings exhibit monotonic equivalence between Euclidean chord distance and cosine distance, while K-means minimizes centroid WCSS/SSE."),
    dict(key="v2", name="divisive_density_clustering", size=(1120, 420), build=v2,
         about="Alternative non-Euclidean clustering: PDDP SVD primary eigenvector splitting vs. HDBSCAN density cores and outlier noise.",
         caption="Non-Euclidean clustering methods: PDDP deterministically splits along principal eigenvectors of variance, whereas HDBSCAN clusters along density reachability and segregates outlier noise."),
    dict(key="v3", name="manifold_projection_spaces", size=(1140, 420), build=v3,
         about="Non-linear manifold reduction: High-D space compared to t-SNE and UMAP 2D projections.",
         caption="Manifold reduction for clustering: High-dimensional embeddings projected into 2D via t-SNE (local neighborhoods only) versus UMAP (preserving both local clustering and global Riemannian topology)."),
]
'''

# ================================================================================================
FILES['examples/F_metric_manifold_geometry/passage.md'] = r'''# Passage: Metric Distance, Clustering Objectives, and Manifold Projection

When transformer embeddings are normalized to unit length (|u|_2=1) squared Euclidean distance is monotonically related to cosine distance: |u-v|_2^2=2-2(u⋅v). K-means in particular operates in the Euclidean space, with an objective functions of minimizing within-cluster sum of squares (WCSS), also known as sum of square error (SSE):

$$\text{WCSS} = \text{SSE} = \sum_{i=1}^k \sum_{x \in C_i} (x - \mu_i)^2 \quad (12)$$

where x is each embedding evaluated against the centroid embedding μ for each cluster C_i it belongs to, such that:

$$C = \{c_1, c_2, \dots, c_k\} = \text{K-means}[E(x), k] \quad (13)$$

Other clustering methods are not inherently Euclidean-based, including principal direction divisive clustering (PDDP) and hierarchal density-based spatial clustering of applications with noise (HDBSCAN). PDDP iteratively splits clusters along their primary eigenvector of variance via singular value decomposition, creating deterministic binary trees (Boley, 1998; Theocharopoulos et al., 2025). Conversely, HDBSCAN operates on cluster density, explicitly segregating outliers as unassigned noise (Campello et al., 2013).
Effective clustering frequently benefits from non-linear manifold reduction, which can be achieved through t-distributed Stochastic Neighborhood Embeddings (t-SNE) or Uniform Manifold Approximation and Projection (UMAP). t-SNE preserves local neighborhood probabilities, exceptional for 2D visualization but preserves neither global geometry nor distance scales (van der Maaten & Hinton, 2008). UMAP is grounded in Riemannian geometry and algebraic topology, preserves both local clustering and global inter-cluster relationships while scaling efficiently to large corpora (McInnes et al., 2018).
'''

# ================================================================================================
FILES['examples/F_metric_manifold_geometry/theme.yaml'] = r'''colors:
  primary: "#1f497d"
  secondary: "#2563eb"
  neutral: "#64748b"
  text: "#0f172a"
  background: "#ffffff"
  
  # Panels
  panel_bg: "#f8fafc"
  panel_border: "#cbd5e1"
  panel_header: "#1e293b"
  
  # Geometric Clusters
  c1_color: "#1d4ed8"
  c1_light: "#dbeafe"
  c2_color: "#7c3aed"
  c2_light: "#ede9fe"
  c3_color: "#059669"
  c3_light: "#d1fae5"
  
  # Metrics & Outliers
  eigenvector: "#ea580c"
  split_line: "#b91c1c"
  noise_color: "#dc2626"
  noise_light: "#fee2e2"
  dist_line: "#475569"
  unit_circle: "#94a3b8"

fonts:
  sans:  { family: "Helvetica" }
  serif: { family: "Times New Roman" }
  mono:  { family: "Courier New" }
default_font: sans

sizes:
  title: 17
  heading: 13
  subheading: 11
  body: 10
  label: 9
  small: 8
  caption: 9.5
  eq_large: 16
  eq_term: 14
  eq_small: 12

line_height: 1.25

strokes:
  hairline: 0.75
  thin: 1
  normal: 1.5
  thick: 2
  heavy: 3

opacity:
  wash: 10
  tint: 25
  soft: 45
  faded: 35

shape:
  corner_radius: 6
'''


LIBS = {"drawio_kit", "figlib", "drawio_doc", "faithful"}
TOOLS = sorted(p[len("scripts/"):-3] for p in FILES if p.startswith("scripts/") and p.endswith(".py"))


def _digest():
    h = hashlib.sha1()
    for k in sorted(FILES):
        h.update(k.encode()); h.update(FILES[k].encode("utf-8"))
    return h.hexdigest()[:10]


def unpack(dest, force=False, quiet=False):
    """Write every bundled file under dest/. Existing files are kept unless force=True."""
    dest = os.path.abspath(dest)
    written = 0
    for rel, text in FILES.items():
        p = os.path.join(dest, *rel.split("/"))
        if os.path.exists(p) and not force:
            continue
        os.makedirs(os.path.dirname(p), exist_ok=True)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)
        written += 1
    with open(os.path.join(dest, ".toolkit_version"), "w") as f:
        f.write(f"{VERSION}-{_digest()}\n")
    if not quiet:
        print(f"unpacked {written} file(s) to {dest}  (skill root: scripts/ assets/ examples/)")
    return dest


def home():
    """Unpacked copy used by the dispatcher (created on first use)."""
    tag = f"{VERSION}-{_digest()}"
    cands = [os.environ.get("DRAWIO_TOOLKIT_HOME"),
             os.path.join(os.path.expanduser("~"), ".cache", "drawio_toolkit", tag),
             os.path.join(tempfile.gettempdir(), "drawio_toolkit", tag)]
    for c in cands:
        if not c:
            continue
        marker = os.path.join(c, ".toolkit_version")
        try:
            if os.path.exists(marker) and open(marker).read().strip() == tag:
                return c
            os.makedirs(c, exist_ok=True)
            return unpack(c, force=True, quiet=True)
        except OSError:
            continue
    raise SystemExit("drawio_toolkit: no writable location to unpack to (set DRAWIO_TOOLKIT_HOME)")


def setup():
    """Unpack (if needed) and put scripts/ on sys.path. Returns the scripts directory."""
    s = os.path.join(home(), "scripts")
    if s not in sys.path:
        sys.path.insert(0, s)
    return s


def main(argv):
    if not argv or argv[0] in ("-h", "--help", "help"):
        print(__doc__); return 0
    cmd, rest = argv[0], argv[1:]
    if cmd == "unpack":
        force = "--force" in rest
        dirs = [a for a in rest if not a.startswith("--")]
        unpack(dirs[0] if dirs else "tidyfigs", force=force); return 0
    if cmd == "path":
        print(os.path.join(home(), "scripts")); return 0
    if cmd == "list":
        print("tools:", " ".join(t for t in TOOLS if t not in LIBS))
        print("libraries:", " ".join(sorted(LIBS)))
        print("files:"); [print("  " + k) for k in FILES]; return 0
    tool = cmd[:-3] if cmd.endswith(".py") else cmd
    if tool not in TOOLS:
        print(f"unknown tool '{cmd}'. Tools: {', '.join(TOOLS)}", file=sys.stderr); return 2
    script = os.path.join(home(), "scripts", tool + ".py")
    return subprocess.call([sys.executable, script] + rest)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
